"""online-office-delivery 钉钉（dingtalk_sheet）后端黑盒测试。

正常/复杂路径用受控的 dws 测试替身验证管线逻辑。替身严格按官方 dws v1.0.64
的 --help/schema/reference 核对过的行为实现：
  - 双信封：raw 命令 {success,result,...}、shortcut {ok,outcome,result/error}、
    sheet import create 扁平 {success,taskId,documentUrl,nodeId}、错误 {error} 走 stderr；
  - `sheet import create --file ... (--folder-token|--workspace)` 与 timed_out
    next_command 续查；
  - `+list-sheets` 严格返回 {count,sheets:[{sheetId,title}]}，缺集合即失败；
  - `+read` 严格返回 cells 且 complete/hasMore/truncationReasons；截断失败关闭；
  - csv-get 返回带 "[row=N] " 前缀的 CSV；csv-put 写入；
  - sheet new / delete-sheet（delete 需 --yes）。
替身只用于管线与协议回归，不冒充线上交付；真实钉钉交付受组织 CLI 访问授权
阻塞（见 dingtalk-adapter-review-20261015/authorization-blocker.json），
线上 Jiuwen/WorkBuddy 两边均标 not_run。测试不 import 被测模块，全部子进程黑盒。
"""
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import skill_dir  # noqa: E402

from openpyxl import Workbook  # noqa: E402

SKILL = "online-office-delivery"
DINGTALK_SCRIPT = "dingtalk_delivery.py"
DISPATCH_SCRIPT = "online_delivery.py"

DOUBLE_SOURCE = r'''
import csv
import io
import json
import os
import re
import sys
import uuid

TOP_HELP = "DingTalk Workspace CLI dws\nAvailable Commands:\n  sheet     钉钉表格管理\n  doc       钉钉文档管理\n  auth      认证管理\n"
SHEET_HELP = "\n".join([
    "Available Commands:",
    "  +list-sheets   list sheets strictly",
    "  +read          strict full read",
    "  create         create a spreadsheet doc",
    "  import         import local xlsx/xls",
    "  new            create a worksheet",
    "  delete-sheet   delete a worksheet",
    "  csv-get        read a range as csv",
    "  csv-put        write csv at a cell",
    "  range          range operations",
    "  range read     read a range",
    "  range update   update a range",
])
RANGE_HELP = "Available Commands:\n  read    read a range\n  update  update a range\n"


def state_path():
    return os.environ.get("DOUBLE_STATE") or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "dws-state.json")


def load_state():
    if os.path.exists(state_path()):
        with open(state_path(), encoding="utf-8") as f:
            return json.load(f)
    return {"workbooks": {}, "pending": {}, "timeout_used": False}


def save_state(st):
    with open(state_path(), "w", encoding="utf-8") as f:
        json.dump(st, f, ensure_ascii=False)


def out(obj):
    print(json.dumps(obj, ensure_ascii=False))


def err(obj, message=None):
    if message is not None:
        sys.stderr.write(message + "\n")
    else:
        sys.stderr.write(json.dumps(obj, ensure_ascii=False) + "\n")


def flags(argv):
    fl = {}
    i = 0
    while i < len(argv):
        a = argv[i]
        if a.startswith("--"):
            if "=" in a:
                k, v = a[2:].split("=", 1)
                fl[k] = v
            elif i + 1 < len(argv) and not argv[i + 1].startswith("-"):
                fl[a[2:]] = argv[i + 1]
                i += 1
            else:
                fl[a[2:]] = True
        elif a.startswith("-") and a not in ("-",):
            k = a[1:]
            if i + 1 < len(argv) and not argv[i + 1].startswith("-"):
                fl[k] = argv[i + 1]
                i += 1
            else:
                fl[k] = True
        i += 1
    return fl


def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n - 1


def parse_range(rng):
    rng = (rng or "").strip()
    m = re.match(r"([A-Za-z]+)(\d+):([A-Za-z]+)(\d+)$", rng)
    if m:
        return (col_index(m.group(1)), int(m.group(2)) - 1,
                col_index(m.group(3)), int(m.group(4)) - 1)
    m = re.match(r"([A-Za-z]+)(\d+)$", rng)
    if m:
        c, r = col_index(m.group(1)), int(m.group(2)) - 1
        return c, r, c, r
    return None


def load_file_sheets(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        sheets.append((name, [list(r) for r in ws.iter_rows(values_only=True)]))
    wb.close()
    return sheets


def norm(v):
    return "" if v is None else v


def finish_import(st, pend, task_id):
    sheets = load_file_sheets(pend["file"])
    token = "dt" + uuid.uuid4().hex[:12]
    node = "N" + uuid.uuid4().hex[:12]
    url = "https://alidocs.dingtalk.com/i/nodes/" + token
    sheets_map = {}
    for i, (title, rows) in enumerate(sheets):
        sheets_map[title] = {"sheetId": "sid%d" % i, "rows": rows}
    if os.environ.get("DOUBLE_MUTATE_FILE") and os.path.basename(
            pend["file"]) == os.environ["DOUBLE_MUTATE_FILE"]:
        first = next(iter(sheets_map))
        for row in sheets_map[first]["rows"]:
            if row:
                row[0] = "MUTATED"
                break
    st["workbooks"][node] = {"url": url, "name": pend.get("name"),
                             "sheets": sheets_map}
    st["pending"].pop(task_id, None)
    save_state(st)
    data = {"success": True, "taskId": task_id, "documentUrl": url,
            "documentName": pend.get("name") or "imported",
            "documentType": "1", "nodeId": node}
    if os.environ.get("DOUBLE_OMIT_URL"):
        data.pop("documentUrl", None)
    out(data)
    return 0


def cmd_import_create(argv):
    fl = flags(argv)
    if os.environ.get("DOUBLE_AUTH_FAIL"):
        err({"error": {"category": "internal", "code": 5,
                       "message": "[UNCLASSIFIED] resolve access token: load "
                                  "access token: reading secure data file: "
                                  "The system cannot find the file specified."}})
        return 1
    if os.environ.get("DOUBLE_ORG_DENY"):
        err({"error": {"category": "internal", "code": 5,
                       "message": "当前组织未授权您通过CLI访问个人数据，需要组织"
                                  "主管理员开启后重新登录"}})
        return 1
    f = fl.get("file")
    if not f:
        err(None, "--file is required")
        return 2
    if not fl.get("folder-token") and not fl.get("workspace"):
        err(None, "--folder-token or --workspace is required")
        return 2
    path = os.path.join(os.getcwd(), f) if not os.path.isabs(f) else f
    if not os.path.exists(path):
        err(None, "file not found: %s" % f)
        return 2
    st = load_state()
    task_id = "tk" + uuid.uuid4().hex[:10]
    st["pending"][task_id] = {"file": path, "name": fl.get("name")}
    if os.environ.get("DOUBLE_TIMEOUT_ONCE") and not st["timeout_used"]:
        st["timeout_used"] = True
        save_state(st)
        out({"success": False, "timed_out": True, "taskId": task_id,
             "status": "processing",
             "next_command": "dws sheet import get --task-id %s" % task_id})
        return 0
    return finish_import(st, st["pending"][task_id], task_id)


def cmd_import_get(argv):
    fl = flags(argv)
    st = load_state()
    pend = st["pending"].get(fl.get("task-id"))
    if not pend:
        err({"error": {"category": "internal", "code": 5,
                       "message": "unknown task"}})
        return 1
    return finish_import(st, pend, fl.get("task-id"))


def get_wb(fl):
    st = load_state()
    wb = st["workbooks"].get(fl.get("node"))
    return st, wb


def sheet_list(wb, fl):
    sheets = []
    for i, (t, s) in enumerate(wb["sheets"].items()):
        if fl.get("title") is not None and t != fl["title"]:
            continue
        if fl.get("sheet-id") and s["sheetId"] != fl["sheet-id"]:
            continue
        sheets.append({"sheetId": s["sheetId"], "title": t, "index": i,
                       "rowCount": len(s["rows"]),
                       "columnCount": max((len(r) for r in s["rows"]),
                                          default=0)})
    return sheets


def cmd_list_sheets(argv):
    fl = flags(argv)
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"type": "api", "subtype": "not_found",
                       "message": "spreadsheet not found"}})
        return 1
    sheets = sheet_list(wb, fl)
    out({"ok": True, "outcome": "success",
         "result": {"count": len(sheets), "sheets": sheets}})
    return 0


def _annotated_csv(values, first_row):
    lines = []
    for i, row in enumerate(values):
        buf = io.StringIO()
        csv.writer(buf).writerow([norm(v) for v in row])
        lines.append("[row=%d] %s" % (first_row + i,
                                      buf.getvalue().rstrip("\r\n")))
    return "\n".join(lines)


def _resolve(wb, fl):
    sid = fl.get("sheet-id")
    for t, s in wb["sheets"].items():
        if s["sheetId"] == sid or t == sid:
            return t, s
    return None, None


def cmd_read(argv):
    fl = flags(argv)
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"type": "api", "subtype": "not_found",
                       "message": "spreadsheet not found"}})
        return 1
    t, s = _resolve(wb, fl)
    if not s:
        out({"ok": False, "outcome": "failure",
             "error": {"type": "api", "subtype": "not_found",
                       "message": "sheet not found"}})
        return 1
    rng = parse_range(fl.get("range") or "A1")
    c0, r0, c1, r1 = rng
    cells = []
    for r in range(r0, r1 + 1):
        row = s["rows"][r] if r < len(s["rows"]) else []
        cells.append([{"value": norm(row[c]) if c < len(row) else ""}
                      for c in range(c0, c1 + 1)])
    out({"success": True, "cells": cells,
         "colIndices": [chr(65 + c) for c in range(c0, c1 + 1)],
         "rowIndices": list(range(r0 + 1, r1 + 2)),
         "complete": True, "hasMore": False, "truncationReasons": [],
         "resolvedRange": fl.get("range"),
         "returnedRange": fl.get("range")})
    return 0


def cmd_csv_get(argv):
    fl = flags(argv)
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "spreadsheet not found"}})
        return 1
    t, s = _resolve(wb, fl)
    if not s:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "sheet not found"}})
        return 1
    rng = parse_range(fl.get("range") or "A1")
    c0, r0, c1, r1 = rng
    values = []
    for r in range(r0, r1 + 1):
        row = s["rows"][r] if r < len(s["rows"]) else []
        values.append([norm(row[c]) if c < len(row) else ""
                       for c in range(c0, c1 + 1)])
    out({"success": True, "csv": _annotated_csv(values, r0 + 1),
         "colIndices": [chr(65 + c) for c in range(c0, c1 + 1)],
         "rowIndices": list(range(r0 + 1, r1 + 2)),
         "hasMore": False, "truncationReasons": [],
         "resolvedRange": fl.get("range"),
         "returnedRange": fl.get("range")})
    return 0


def cmd_sheet_new(argv):
    fl = flags(argv)
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "spreadsheet not found"}})
        return 1
    name = fl.get("name")
    if not name:
        err(None, "--name is required")
        return 2
    sid = "sid" + uuid.uuid4().hex[:8]
    wb["sheets"][name] = {"sheetId": sid, "rows": []}
    save_state(st)
    out({"success": True, "sheetId": sid, "title": name})
    return 0


def cmd_csv_put(argv):
    fl = flags(argv)
    if os.environ.get("DOUBLE_DENY_WRITE"):
        out({"ok": False, "outcome": "failure",
             "error": {"message": "write denied by policy"}})
        return 1
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "spreadsheet not found"}})
        return 1
    t, s = _resolve(wb, fl)
    if not s:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "sheet not found"}})
        return 1
    m = re.match(r"([A-Za-z]+)(\d+)$", fl.get("start-cell") or "A1")
    col = col_index(m.group(1))
    row = int(m.group(2)) - 1
    rows = [r for r in csv.reader((fl.get("csv") or "").splitlines())]
    for dr, cells in enumerate(rows):
        while len(s["rows"]) <= row + dr:
            s["rows"].append([])
        rr = s["rows"][row + dr]
        while len(rr) <= col + len(cells) - 1:
            rr.append("")
        for dc, v in enumerate(cells):
            rr[col + dc] = v
    save_state(st)
    out({"success": True, "written": len(rows)})
    return 0


def cmd_delete_sheet(argv):
    fl = flags(argv)
    if fl.get("yes") is not True:
        err(None, "high-risk command: --yes is required")
        return 10
    st, wb = get_wb(fl)
    if not wb:
        out({"ok": False, "outcome": "failure",
             "error": {"message": "spreadsheet not found"}})
        return 1
    for t, s in list(wb["sheets"].items()):
        if s["sheetId"] == fl.get("sheet-id"):
            del wb["sheets"][t]
            save_state(st)
            out({"success": True, "deleted": t})
            return 0
    out({"ok": False, "outcome": "failure",
         "error": {"message": "sheet not found"}})
    return 1


def cmd_auth_status(argv):
    if os.environ.get("DOUBLE_AUTH_STATUS_FALSE"):
        out({"success": True, "authenticated": False, "message": "未登录"})
        return 0
    out({"success": True, "authenticated": True, "message": "ok"})
    return 0


def main(argv):
    if "--version" in argv or argv[0] in ("-v", "--version"):
        sys.stdout.write("dws version v1.0.64-double\n")
        return 0
    if not argv or argv[0] in ("-h", "--help"):
        sys.stdout.write(TOP_HELP)
        return 0
    if argv[0] == "auth" and len(argv) > 1 and argv[1] == "status":
        return cmd_auth_status(argv[2:])
    if argv[0] == "sheet":
        if len(argv) < 2 or argv[1] in ("-h", "--help"):
            sys.stdout.write(SHEET_HELP + "\n")
            return 0
        if argv[1] == "range" and (len(argv) < 3 or argv[2] in ("-h", "--help")):
            sys.stdout.write(RANGE_HELP + "\n")
            return 0
        if argv[1] == "import" and len(argv) > 2 and argv[2] == "create":
            return cmd_import_create(argv[3:])
        if argv[1] == "import" and len(argv) > 2 and argv[2] == "get":
            return cmd_import_get(argv[3:])
        handlers = {
            "+list-sheets": cmd_list_sheets,
            "+read": cmd_read,
            "csv-get": cmd_csv_get,
            "csv-put": cmd_csv_put,
            "new": cmd_sheet_new,
            "delete-sheet": cmd_delete_sheet,
        }
        if argv[1] in handlers:
            return handlers[argv[1]](argv[2:])
    sys.stderr.write("unknown command: %r\n" % (argv[:2],))
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
'''


def write_double(base):
    d = base / "bin"
    d.mkdir(parents=True, exist_ok=True)
    py = d / "dws_double.py"
    py.write_text(DOUBLE_SOURCE.strip() + "\n", encoding="utf-8")
    if os.name == "nt":
        cmd = d / "dws.cmd"
        cmd.write_text('@echo off\r\n"%s" "%s" %%*\r\n' % (sys.executable, py),
                       encoding="ascii")
    else:
        sh = d / "dws"
        sh.write_text("#!/usr/bin/env python3\n" + DOUBLE_SOURCE.strip() + "\n",
                      encoding="utf-8")
        os.chmod(str(sh), 0o755)
    return str(d)


def make_env(state_path, double_dir=None, **knobs):
    env = dict(os.environ)
    for k in list(env):
        if k.upper().startswith("DOUBLE_"):
            del env[k]
    env["PYTHONUTF8"] = "1"
    env["DOUBLE_STATE"] = str(state_path)
    if double_dir is None:
        empty = os.path.dirname(str(state_path)) or "."
        env["PATH"] = empty
    else:
        env["PATH"] = double_dir + os.pathsep + env.get("PATH", "")
    for k, v in knobs.items():
        env["DOUBLE_" + k] = str(v)
    return env


def run_backend(script, plan_path, outdir, env, cli=None, extra=None,
                cwd=None, timeout=180):
    path = os.path.join(skill_dir(SKILL), "scripts", script)
    outdir = str(outdir)
    os.makedirs(outdir, exist_ok=True)
    cmd = [sys.executable, path, "--plan", str(plan_path),
           "--outdir", outdir]
    if cli:
        cmd += ["--cli", cli]
    cmd += extra or []
    return subprocess.run(cmd, capture_output=True, text=True,
                          encoding="utf-8", cwd=cwd or outdir, timeout=timeout,
                          env=env)


def write_plan(base, target=None, files=None, name="plan.json", **extra):
    plan = {"target": target or {"type": "dingtalk_sheet",
                                 "folder_token": "FOLDER_TOKEN_TEST"},
            "files": files or []}
    plan.update(extra)
    p = base / name
    p.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    return p


def read_report(outdir):
    return json.loads((outdir / "delivery-report.json").read_text(
        encoding="utf-8"))


def read_state(state_path):
    with open(state_path, encoding="utf-8") as f:
        return json.load(f)


def build_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["序号", "姓名", "编号", "金额", "备注"])
    ws.append([1, "张三", "001", 1999.5, "正常"])
    ws.append([2, "李四", "010", 12345.67, ""])
    ws2 = wb.create_sheet("汇总")
    ws2.append(["合计", 14345.17])
    wb.save(str(path))


def test_blocked_without_cli(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    env = make_env(tmp_path / "st.json", double_dir=None)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert rep["missing"]["clis"] == ["dws"]
    assert "results" not in rep
    assert "https://" not in json.dumps(rep)
    assert rep["backend"] == "dingtalk_sheet"


def test_help_gate_blocks_unverified_command(tmp_path):
    # 替身的 --help 固定；通过只保留已知命令验证门禁：改用隐藏方式不可行，
    # 因此这里验证门禁逻辑在真实 dws 帮助下的正向通过（见 delivered 用例），
    # 异常方向用「PATH 中放入一个不含命令的假 dws」验证 blocked。
    fake = tmp_path / "fakebin"
    fake.mkdir()
    stub = fake / ("dws.cmd" if os.name == "nt" else "dws")
    stub.write_text("@echo off\r\necho not a dws\r\n" if os.name == "nt"
                    else "#!/bin/sh\necho not a dws\n", encoding="ascii")
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    env = make_env(tmp_path / "st.json", str(fake))
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert "commands_from_help" in rep["missing"]


def test_input_errors(tmp_path):
    bad_json = tmp_path / "bad.json"
    bad_json.write_text("{not json", encoding="utf-8")
    outdir = tmp_path / "out1"
    r = run_backend(DINGTALK_SCRIPT, bad_json, outdir,
                    make_env(tmp_path / "s1.json"))
    assert r.returncode == 2
    assert read_report(outdir)["status"] == "input_error"

    build_xlsx(tmp_path / "data.xlsx")
    empty = write_plan(tmp_path, files=[], name="plan-empty.json")
    outdir2 = tmp_path / "out2"
    r = run_backend(DINGTALK_SCRIPT, empty, outdir2, make_env(tmp_path / "s2.json"))
    assert r.returncode == 2
    assert any("空" in e for e in read_report(outdir2)["errors"])

    # folder_token/workspace 缺失：官方要求至少一个，拒绝猜测默认位置。
    no_folder = write_plan(tmp_path, target={"type": "dingtalk_sheet"},
                           files=[{"local": str(tmp_path / "data.xlsx")}],
                           name="plan-nofolder.json")
    outdir3 = tmp_path / "out3"
    r = run_backend(DINGTALK_SCRIPT, no_folder, outdir3,
                    make_env(tmp_path / "s3.json"))
    assert r.returncode == 2
    assert any("folder_token" in e for e in read_report(outdir3)["errors"])

    # .csv 不受支持（钉钉 import 不支持 csv；本地回读基准亦不适用）。
    csvf = tmp_path / "names.csv"
    csvf.write_text("a,b\n1,2\n", encoding="utf-8")
    plan_csv = write_plan(tmp_path, files=[{"local": str(csvf)}],
                          name="plan-csv.json")
    outdir4 = tmp_path / "out4"
    r = run_backend(DINGTALK_SCRIPT, plan_csv, outdir4,
                    make_env(tmp_path / "s4.json"))
    assert r.returncode == 2
    assert any("扩展名" in e for e in read_report(outdir4)["errors"])

    missing = write_plan(tmp_path, files=[{"local": str(tmp_path / "nope.xlsx")}],
                         name="plan-missing.json")
    outdir5 = tmp_path / "out5"
    r = run_backend(DINGTALK_SCRIPT, missing, outdir5,
                    make_env(tmp_path / "s5.json"))
    assert r.returncode == 2
    assert any("不存在" in e for e in read_report(outdir5)["errors"])


def test_delivered_xlsx_full_pipeline(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[
        {"local": str(tmp_path / "data.xlsx"), "name": "Q3钉钉表"}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "delivered"
    entry = rep["results"][0]
    assert entry["status"] == "delivered"
    assert entry["url"].startswith("https://alidocs.dingtalk.com/i/nodes/")
    assert entry["node_id"]
    assert entry["readback"]["verified"] is True
    assert entry["readback"]["sheets_local"] == ["人员", "汇总"]
    sampled = {s["sheet"]: s for s in entry["readback"]["sampled"]}
    assert sampled["人员"]["range"] == "A1:E3"
    assert sampled["人员"]["mismatch_count"] == 0
    assert sampled["人员"]["source"] == "read_strict"
    pc = entry["permission_check"]
    assert pc["verified"] is True and pc["readback_matched"] is True
    assert pc["cleanup"]["ok"] is True
    imports = [e for e in rep["cli_invocation_log"]
               if "import" in e["argv"] and "create" in e["argv"]]
    assert imports
    file_arg = imports[0]["argv"][imports[0]["argv"].index("--file") + 1]
    assert file_arg == "data.xlsx"
    assert imports[0]["cwd"].lower() == str(tmp_path).lower()
    deletes = [e for e in rep["cli_invocation_log"]
               if "delete-sheet" in e["argv"]]
    assert deletes and "--yes" in deletes[0]["argv"]
    assert rep["probe"]["syntax_verification"]["all_required_verified"] is True
    st = read_state(tmp_path / "st.json")
    node = entry["node_id"]
    wbk = st["workbooks"][node]
    assert "__delivery_check__" not in wbk["sheets"]


def test_dispatch_via_online_delivery_entry(tmp_path):
    # 明确入口：online_delivery.py 对 target.type=dingtalk_sheet 分派到本后端。
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_backend(DISPATCH_SCRIPT, plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["backend"] == "dingtalk_sheet"
    assert rep["status"] == "delivered"


def test_import_timeout_follows_next_command(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, TIMEOUT_ONCE=1)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    entry = read_report(outdir)["results"][0]
    assert entry["status"] == "delivered"
    assert entry["import"]["attempts"] == 2
    assert entry["import"]["followed_next_command"] is True


def test_readback_mismatch_fails(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, MUTATE_FILE="data.xlsx")
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "failed"
    entry = rep["results"][0]
    assert entry["readback"]["verified"] is False
    m = entry["readback"]["sampled"][0]["mismatches"][0]
    assert m["online"] == "MUTATED"


def test_missing_url_not_fabricated(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, OMIT_URL=1)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    entry = read_report(outdir)["results"][0]
    assert "url" not in entry
    assert any("documentUrl" in x for x in entry["reasons"])
    assert "https://alidocs" not in json.dumps(read_report(outdir))


def test_permission_denied_fails(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, DENY_WRITE=1)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    entry = read_report(outdir)["results"][0]
    assert entry["readback"]["verified"] is True
    assert entry["permission_check"]["verified"] is False
    assert "金丝雀写入失败" in entry["permission_check"]["error"]


def test_unauthorized_blocks(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    for knob in ("AUTH_FAIL", "ORG_DENY", "AUTH_STATUS_FALSE"):
        env = make_env(tmp_path / ("st-%s.json" % knob), double_dir, **{knob: 1})
        o = tmp_path / ("out-%s" % knob)
        r = run_backend(DINGTALK_SCRIPT, plan, o, env)
        assert r.returncode == 3, r.stderr + r.stdout
        rep = read_report(o)
        assert rep["status"] == "blocked"
        assert "cli_authorization" in rep["missing"]
        assert "alidocs" not in json.dumps(rep)


def test_org_denied_reports_admin_entry(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, ORG_DENY=1)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    missing = rep["missing"]["cli_authorization"]
    assert missing["admin_url"].startswith("https://open-dev.dingtalk.com")
    assert "组织" in missing["reason"]


def test_keep_check_sheet_leaves_marker(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}],
                      keep_check_sheet=True)
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    entry = read_report(outdir)["results"][0]
    assert entry["status"] == "delivered"
    assert entry["permission_check"]["cleanup"]["attempted"] is False
    st = read_state(tmp_path / "st.json")
    wbk = st["workbooks"][entry["node_id"]]
    check = wbk["sheets"]["__delivery_check__"]
    assert check["rows"][0][0].startswith("delivery-check-")


def test_profile_passthrough(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_backend(DINGTALK_SCRIPT, plan, outdir, env,
                    extra=["--profile", "corp:user"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    business = [e for e in rep["cli_invocation_log"]
                if "-f" in e["argv"] or "--profile" in e["argv"]]
    assert business
    for e in business:
        argv = e["argv"]
        if "--help" in argv or "--version" in argv:
            continue
        assert "--profile" in argv and "corp:user" in argv
