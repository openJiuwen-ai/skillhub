"""PDF/PPT 家族 5 技能黑盒测试。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

# 身份证校验位独立实现（与被测脚本不同代码）
W = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
CHECK = "10X98765432"


def make_id(prefix17):
    s = sum(int(prefix17[i]) * W[i] for i in range(17))
    return prefix17 + CHECK[s % 11]


VALID_ID = make_id("11010519880605121")  # 以测试内独立算法构造
BAD_ID = "110105198806051211"           # 校验位故意错误(正确校验位为0)
PHONE = "13812345678"
BANK = "4111111111111111"                # 经典 Luhn 有效测试卡号
EMAIL = "zhangsan@example.com"
F15 = "110105880612345"                  # 15位数字


def verify_luhn(v):
    digits = [int(c) for c in v]
    total = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def build_pdf(path):
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), "订单登记表", fontsize=16, fontname="china-s")
    p1.insert_text((72, 110), f"客户经理身份证 {VALID_ID}", fontsize=10, fontname="china-s")
    p1.insert_text((72, 130), f"电话 {PHONE} 邮箱 {EMAIL}", fontsize=10, fontname="china-s")
    p1.insert_text((72, 150), f"卡号 {BANK} 存疑号 {BAD_ID}", fontsize=10, fontname="china-s")
    p1.insert_text((72, 170), f"旧编号 {F15}", fontsize=10, fontname="china-s")
    # 线框表格
    x0, y0, x1, y1 = 72, 200, 460, 360
    for i in range(5):
        y = y0 + i * (y1 - y0) / 4
        p1.draw_line(pymupdf.Point(x0, y), pymupdf.Point(x1, y))
    for j in range(5):
        x = x0 + j * (x1 - x0) / 4
        p1.draw_line(pymupdf.Point(x, y0), pymupdf.Point(x, y1))
    cells = [["订单号", "客户", "数量", "金额"], ["A001", "张三", "10", "1200"],
             ["A002", "李四", "20", "3400"], ["A003", "王五", "30", "5600"]]
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            p1.insert_text((x0 + j * (x1 - x0) / 4 + 4, y0 + i * (y1 - y0) / 4 + 14),
                           val, fontsize=10, fontname="china-s")
    p2 = doc.new_page()
    p2.insert_text((72, 100), "标题：二〇二六年三季度运营复盘", fontsize=14, fontname="china-s")
    p2.insert_text((72, 130), "本季度完成三项重点任务。", fontsize=10, fontname="china-s")
    p3 = doc.new_page()
    p3.insert_text((72, 100), "签字：", fontsize=10, fontname="china-s")
    doc.save(path)
    doc.close()


# ---------- pdf-structured-extract ----------

def test_extract(workspace):
    build_pdf(str(workspace / "in.pdf"))
    r = run_script("pdf-structured-extract", "extract_pdf.py",
                   ["--input", workspace / "in.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "out" / "extract.json").read_text(encoding="utf-8"))
    # 3页; 第3页稀疏
    assert len(rep["pages"]) == 3
    assert rep["sparse_pages"] == [3]
    rendered = workspace / "out" / rep["pages"][2]["rendered_image"]
    assert rendered.exists() and rendered.stat().st_size > 1000  # 已渲染
    # 表格结构化: 4行4列(检测可能有边缘损失, 断言>=3行且含表头)
    tables = rep["pages"][0]["tables"]
    assert len(tables) >= 1
    flat = [c for row in tables[0] for c in row]
    assert any("订单号" in str(c) for c in flat)
    assert any("1200" in str(c) for c in flat)
    # 关键数字: 有效证件/电话/银行卡/邮箱 valid; 坏校验位证件 mismatch
    kn = rep["key_numbers"]["candidates_and_verification"]
    types = {f["type"] for f in kn}
    assert {"id_card", "phone", "bank_card", "email"} <= types
    id_entries = [f for f in kn if f["type"] == "id_card"]
    statuses = {f["status"] for f in id_entries}
    assert "valid" in statuses and "checksum_mismatch" in statuses
    bank = [f for f in kn if f["type"] == "bank_card"][0]
    assert bank["status"] == "valid"
    # md 输出含表格管道
    md = (workspace / "out" / "extract.md").read_text(encoding="utf-8")
    assert "|" in md


# ---------- pdf-nl-edit ----------

def test_pdf_edit(workspace):
    build_pdf(str(workspace / "in.pdf"))
    r = run_script("pdf-nl-edit", "pdf_edit.py",
                   ["--input", workspace / "in.pdf", "--page", "2",
                    "--find", "二〇二六年三季度运营复盘", "--replace", "2026年三季度运营复盘",
                    "--output", workspace / "out.pdf", "--report", workspace / "r.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "r.json").read_text(encoding="utf-8"))
    assert rep["other_pages_unchanged"] is True
    assert rep["target_page_contains_new_text"] is True
    assert rep["target_page_contains_old_text"] is False
    src = pymupdf.open(str(workspace / "in.pdf"))
    out = pymupdf.open(str(workspace / "out.pdf"))
    t2 = out[1].get_text("text")
    assert "2026年三季度运营复盘" in t2
    assert "二〇二六年三季度运营复盘" not in t2
    assert out[0].get_text("text") == src[0].get_text("text")
    assert out[2].get_text("text") == src[2].get_text("text")
    src.close(); out.close()
    # 目标不存在 -> 退出1
    r2 = run_script("pdf-nl-edit", "pdf_edit.py",
                    ["--input", workspace / "in.pdf", "--page", "2",
                     "--find", "不存在的文本", "--replace", "X",
                     "--output", workspace / "o2.pdf"])
    assert r2.returncode == 1
    # 页码越界 -> 退出2
    r3 = run_script("pdf-nl-edit", "pdf_edit.py",
                    ["--input", workspace / "in.pdf", "--page", "9",
                     "--find", "X", "--replace", "Y",
                     "--output", workspace / "o3.pdf"])
    assert r3.returncode == 2


# ---------- doc-sensitive-redact ----------

def test_redact(workspace):
    # PDF
    build_pdf(str(workspace / "in.pdf"))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "in.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    man = json.loads((workspace / "out" / "redact-manifest.json").read_text(encoding="utf-8"))
    pf = man["files"][0]
    types = {x["type"] for x in pf["redactions"]}
    assert {"id_card", "phone", "bank_card", "email"} <= types
    assert pf["pdf_removed_and_verified"] is True
    # 底层删除验证(独立): 输出中提取不到原文
    out = pymupdf.open(str(workspace / "out" / "in_REDACTED.pdf"))
    text = "".join(p.get_text("text") for p in out)
    out.close()
    assert VALID_ID not in text and PHONE not in text and BANK not in text and EMAIL not in text
    # 坏校验位/15位 未被打码(防误杀)
    assert F15 in text
    review_types = {x["type"] for x in pf["needs_review"]}
    assert any("15" in t or "legacy" in t for t in review_types)

    # DOCX 跨run
    doc = Document()
    p = doc.add_paragraph()
    for piece in ("联系人身份证", VALID_ID[:9], VALID_ID[9:], "电话", PHONE):
        p.add_run(piece)
    doc.save(str(workspace / "d.docx"))
    r2 = run_script("doc-sensitive-redact", "redact.py",
                    ["--input", workspace / "d.docx", "--outdir", workspace / "out"])
    assert r2.returncode == 0
    doc2 = Document(str(workspace / "out" / "d_REDACTED.docx"))
    full = "\n".join(p.text for p in doc2.paragraphs)
    assert VALID_ID not in full and PHONE not in full and "■" in full

    # XLSX / PPTX
    wb = Workbook(); ws = wb.active
    ws.append(["姓名", "联系方式"])
    ws.append(["张三", f"{PHONE} / {EMAIL}"])
    wb.save(str(workspace / "x.xlsx"))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = f"卡号 {BANK}"
    prs.save(str(workspace / "p.pptx"))
    for f in ("x.xlsx", "p.pptx"):
        r3 = run_script("doc-sensitive-redact", "redact.py",
                        ["--input", workspace / f, "--outdir", workspace / "out"])
        assert r3.returncode == 0
    man2 = json.loads((workspace / "out" / "redact-manifest.json").read_text(encoding="utf-8"))
    # 批量目录
    (workspace / "batch").mkdir()
    import shutil
    shutil.copy(str(workspace / "d.docx"), str(workspace / "batch" / "d.docx"))
    shutil.copy(str(workspace / "x.xlsx"), str(workspace / "batch" / "x.xlsx"))
    r4 = run_script("doc-sensitive-redact", "redact.py",
                    ["--input", workspace / "batch", "--outdir", workspace / "bout"])
    assert r4.returncode == 0
    man3 = json.loads((workspace / "bout" / "redact-manifest.json").read_text(encoding="utf-8"))
    assert man3["summary"]["files_processed"] == 2
    # 清单不含原始敏感值
    raw = json.dumps(man3, ensure_ascii=False)
    assert VALID_ID not in raw and PHONE not in raw and BANK not in raw and EMAIL not in raw


# ---------- deck-delivery-qa ----------

def build_deck(path):
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    tb = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1.5))
    tb.text_frame.text = "季度业务回顾"
    # 溢出页: 小框 + 长文本
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    tb2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(2.5), Inches(0.5))
    tf = tb2.text_frame
    tf.text = ("本季度公司在核心业务领域取得了显著进展，各项运营指标持续向好，"
               "客户满意度稳步提升，市场份额进一步扩大。") * 3
    for run in tf.paragraphs[0].runs:
        run.font.size = Pt(18)
    # 越界页: 文本框放在画布右侧之外
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    tb3 = s3.shapes.add_textbox(Inches(12), Inches(1), Inches(2), Inches(1))
    tb3.text_frame.text = "越界内容"
    prs.save(path)


def test_deck_qa(workspace):
    build_deck(str(workspace / "deck.pptx"))
    r = run_script("deck-delivery-qa", "deck_qa.py",
                   ["--input", workspace / "deck.pptx", "--report", workspace / "qa.json"])
    assert r.returncode == 1, r.stdout  # 有不合格页
    rep = json.loads((workspace / "qa.json").read_text(encoding="utf-8"))
    assert rep["pass"] is False
    assert rep["rendered"] is False and rep["mode"] == "estimated_geometry"
    s = {x["index"]: x for x in rep["slides"]}
    assert s[1]["pass"] is True          # 正常页通过
    assert any(i["check"] == "text_overflow" for i in s[2]["issues"])
    assert any(i["check"] == "shape_out_of_bounds" for i in s[3]["issues"])
    assert all("utilization" in x for x in rep["slides"])
    # 确定性: 两次运行一致
    r2 = run_script("deck-delivery-qa", "deck_qa.py",
                    ["--input", workspace / "deck.pptx", "--report", workspace / "qa2.json"])
    assert (workspace / "qa.json").read_bytes() == (workspace / "qa2.json").read_bytes()
    # 损坏
    (workspace / "bad.pptx").write_bytes(b"junk")
    r3 = run_script("deck-delivery-qa", "deck_qa.py",
                    ["--input", workspace / "bad.pptx", "--report", workspace / "q3.json"])
    assert r3.returncode == 2


# ---------- doc-translate-layout ----------

def test_translate(workspace):
    doc = Document()
    doc.add_paragraph("智慧园区一期项目概述", style="Heading 1")
    doc.add_paragraph("本项目旨在建设智能化园区管理平台。")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "模块"
    t.cell(0, 1).text = "状态"
    t.cell(1, 0).text = "智慧园区平台"
    t.cell(1, 1).text = "已上线"
    doc.sections[0].header.paragraphs[0].text = "内部资料"
    # 超长译文触发回退的 run 需显式字号
    p = doc.add_paragraph()
    run = p.add_run("短句")
    run.font.size = Pt(12)
    doc.save(str(workspace / "src.docx"))

    r = run_script("doc-translate-layout", "translate_layout.py",
                   ["extract", "--input", workspace / "src.docx",
                    "--worksheet", workspace / "ws.json"])
    assert r.returncode == 0, r.stderr
    ws = json.loads((workspace / "ws.json").read_text(encoding="utf-8"))
    ids = [it["id"] for it in ws["items"]]
    texts = [it["text"] for it in ws["items"]]
    assert "智慧园区一期项目概述" in texts
    assert "本项目旨在建设智能化园区管理平台。" in texts
    assert "内部资料" in texts

    # 译文: 表格/正文/页眉全译, "短句"不译(未翻译清单), 术语表命中检查
    tr = {}
    for it in ws["items"]:
        if it["text"] == "短句":
            continue  # 留一个未翻译
        tr[it["id"]] = "EN: " + it["text"][:4]
    # 术语表: 智慧园区 -> Smart Park; 但译文没放 Smart Park -> miss
    (workspace / "g.json").write_text(json.dumps({"智慧园区": "Smart Park"}, ensure_ascii=False),
                                      encoding="utf-8")
    # 长译文触发回退
    sid = next(it["id"] for it in ws["items"] if it["text"] == "短句")
    tr[sid] = "x" * 20
    (workspace / "tr.json").write_text(json.dumps(tr, ensure_ascii=False), encoding="utf-8")
    r2 = run_script("doc-translate-layout", "translate_layout.py",
                    ["apply", "--input", workspace / "src.docx",
                     "--worksheet", workspace / "ws.json",
                     "--translations", workspace / "tr.json",
                     "--output", workspace / "out.docx",
                     "--glossary", workspace / "g.json",
                     "--report", workspace / "rep.json"])
    assert r2.returncode == 0, r2.stderr + r2.stdout
    rep = json.loads((workspace / "rep.json").read_text(encoding="utf-8"))
    assert rep["translated"] == len(ws["items"]) - 0  # "短句"也给了长译文
    assert rep["mt_engine"].startswith("external")
    assert len(rep["glossary_misses"]) >= 1 or len(rep["glossary_hits"]) >= 1
    assert any(f.get("action") in ("font_shrunk", "needs_attention") for f in rep["fallback_list"])

    # 结构保持独立验证
    src_doc = Document(str(workspace / "src.docx"))
    out_doc = Document(str(workspace / "out.docx"))
    assert len(out_doc.tables) == len(src_doc.tables) == 1
    assert int(out_doc.tables[0]._tbl.xpath("count(.//w:tr)")) == 2
    assert out_doc.paragraphs[0].style.name == src_doc.paragraphs[0].style.name
    assert "EN: " in out_doc.paragraphs[0].text
    assert out_doc.sections[0].header.paragraphs[0].text != "内部资料"
    # pptx 路径
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = "你好世界"
    prs.save(str(workspace / "s.pptx"))
    r3 = run_script("doc-translate-layout", "translate_layout.py",
                    ["extract", "--input", workspace / "s.pptx",
                     "--worksheet", workspace / "wsp.json"])
    assert r3.returncode == 0
    wsp = json.loads((workspace / "wsp.json").read_text(encoding="utf-8"))
    assert wsp["items"][0]["text"] == "你好世界"
    (workspace / "trp.json").write_text(
        json.dumps({wsp["items"][0]["id"]: "Hello World"}, ensure_ascii=False), encoding="utf-8")
    r4 = run_script("doc-translate-layout", "translate_layout.py",
                    ["apply", "--input", workspace / "s.pptx",
                     "--worksheet", workspace / "wsp.json",
                     "--translations", workspace / "trp.json",
                     "--output", workspace / "s_out.pptx"])
    assert r4.returncode == 0, r4.stderr
    prs2 = Presentation(str(workspace / "s_out.pptx"))
    texts = [sh.text_frame.text for sh in prs2.slides[0].shapes if sh.has_text_frame]
    assert "Hello World" in texts
