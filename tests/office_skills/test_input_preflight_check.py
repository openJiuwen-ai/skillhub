"""input-preflight-check 单项标准黑盒测试。

独立验证（normal / complex / abnormal 三类）：
- 真实 md/docx/xlsx 输入、真实子进程执行（不 import 被测实现）、断言落盘产物
- 关键回归（benchmarks/preflight-complex/before-review 冻结缺陷）：
  1) 不可能的日历日期（2026-02-30）与无法解析的值（NOT_A_DATE）不得通过
     必需日期范围检查——必须按真实日历解析并统计不可解析/缺失值后阻断；
  2) 损坏 docx 仅因文件名命中模式不得判 ready/pass；
  3) 空清单/空目录/重复id/重复表头/公式缺缓存/空requirements不得静默通过。
- 预期值由本测试手工独立推导（行号/计数），不 mock 被测实现。
"""
import datetime
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

import docx  # noqa: E402
from openpyxl import Workbook  # noqa: E402

SKILL = "input-preflight-check"
SCRIPT = "preflight.py"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def run(*args):
    return run_script(SKILL, SCRIPT, list(args))


def write_json(path, obj):
    path.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")


# ---------- 正常 ----------

def test_normal_checklist_all_ready_pass(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "营业执照副本.md").write_text(
        "统一社会信用代码 91330100MA27X0000X；经营范围：软件开发。", encoding="utf-8")
    (mat / "审计报告_2023.md").write_text("审计意见：标准无保留意见。", encoding="utf-8")
    cfg = {"items": [
        {"id": "M1", "name": "营业执照", "required": True,
         "patterns": ["营业执照", "统一社会信用代码"], "fuzzy": ["执照"]},
        {"id": "M2", "name": "近三年审计报告", "required": True,
         "patterns": ["审计报告", "2023"]},
    ]}
    write_json(workspace / "c.json", cfg)
    r = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    states = {i["id"]: i["state"] for i in rep["items"]}
    assert states == {"M1": "ready", "M2": "ready"}
    assert rep["verdict"] == "pass"
    assert rep["reasons"] == []
    assert rep["files_scanned"] == 2
    # 确定性：同输入两次运行结果一致
    r2 = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
             "--report", workspace / "r2.json")
    assert r2.returncode == 0
    assert read_json(workspace / "r2.json") == rep


def test_normal_data_clean_pass(workspace):
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["日期", "销售额(元)", "地区"])
    ws.append(["2026-01-05", 100, "华东"])
    ws.append([datetime.date(2026, 3, 2), 200, "华北"])  # 真实日期单元格
    ws.append(["2026.12.31", 300, "华南"])  # 点分日期字符串
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "销售额(元)", "地区"],
           "units": {"销售额(元)": "元"},
           "numeric_columns": ["销售额(元)"],
           "date_column": "日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    write_json(workspace / "req.json", req)
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "pass"
    assert rep["reasons"] == []
    assert rep["rows"] == 3
    assert rep["sheet"] == "销售"
    by = {c["check"]: c for c in rep["checks"]}
    assert by["duplicate_headers"]["ok"] is True
    assert by["required_columns"]["ok"] is True
    assert by["unit"]["ok"] is True
    assert by["numeric"]["ok"] is True and by["numeric"]["unparseable_count"] == 0
    d = by["date_range"]
    assert d["ok"] is True and d["parsed_rows"] == 3
    assert d["unparseable_count"] == 0 and d["missing_count"] == 0
    assert d["out_of_range_count"] == 0
    assert by["min_rows"]["ok"] is True


# ---------- 复杂 ----------

def test_complex_corrupt_required_docx_not_ready_blocks(workspace):
    """冻结缺陷回归：损坏合同.docx 仅因文件名命中不得 ready/pass。"""
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "说明.md").write_text("材料说明：参见各附件。", encoding="utf-8")
    (mat / "合同.docx").write_bytes(b"broken document")  # 损坏 docx
    d = docx.Document()
    d.add_paragraph("协议条款：本文件由甲乙双方签署。")
    d.save(str(mat / "协议.docx"))  # 可读 docx，内容不含"合同"
    cfg = {"items": [
        {"id": "C1", "name": "合同", "required": True, "patterns": ["合同"]},
        {"id": "C2", "name": "协议", "required": True, "patterns": ["协议条款"]},
    ]}
    write_json(workspace / "c.json", cfg)
    r = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    states = {i["id"]: i["state"] for i in rep["items"]}
    assert states["C1"] == "needs_confirm"  # 损坏文件不得判 ready
    assert states["C2"] == "ready"
    c1 = [i for i in rep["items"] if i["id"] == "C1"][0]
    assert c1["matched_files"] == []
    assert c1["unreadable_files"] and c1["unreadable_files"][0]["file"] == "合同.docx"
    assert c1["unreadable_files"][0]["error"]
    assert rep["verdict"] == "block"  # 必需材料仅命中损坏文件 → 阻断
    assert any("无法读取" in x for x in rep["reasons"])
    # 同一损坏文件若为非必需 → 不阻断，但必须 confirm_needed（不得静默 pass）
    cfg["items"][0]["required"] = False
    write_json(workspace / "c2.json", cfg)
    r2 = run("checklist", "--checklist", workspace / "c2.json", "--dir", mat,
             "--report", workspace / "r2.json")
    assert r2.returncode == 0
    rep2 = read_json(workspace / "r2.json")
    assert rep2["verdict"] == "confirm_needed"
    assert rep2["items"][0]["state"] == "needs_confirm"


def test_complex_impossible_unparseable_missing_dates_block(workspace):
    """冻结缺陷回归：2026-02-30 / NOT_A_DATE / 空值 / 越界不得通过日期范围检查。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "金额"])
    ws.append(["2026-02-30", 100])  # 行2：不存在的日历日期
    ws.append(["NOT_A_DATE", 100])  # 行3：无法解析
    ws.append(["", 300])            # 行4：日期缺失（行非空）
    ws.append(["2027-01-01", 100])  # 行5：越界
    ws.append(["2026-06-15", 100])  # 行6：范围内有效
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "金额"], "date_column": "日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    write_json(workspace / "req.json", req)
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    d = [c for c in rep["checks"] if c["check"] == "date_range"][0]
    assert d["ok"] is False
    assert d["parsed_rows"] == 2  # 行5(越界)与行6可解析；不可解析/缺失不计入
    assert d["unparseable_count"] == 2
    assert {u["row"] for u in d["unparseable"]} == {2, 3}
    assert {u["value"] for u in d["unparseable"]} == {"2026-02-30", "NOT_A_DATE"}
    assert d["missing_count"] == 1 and d["missing_rows"] == [4]
    assert d["out_of_range_count"] == 1
    assert d["out_of_range"][0]["row"] == 5
    assert any("日期不可解析" in x for x in rep["reasons"])
    assert any("日期缺失" in x for x in rep["reasons"])
    assert any("日期超出范围" in x for x in rep["reasons"])


def test_complex_duplicate_headers_and_numeric_garbage_block(workspace):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "金额", "金额"])  # 重复表头 → 列引用歧义
    ws.append(["2026-01-10", 100, 200])
    ws.append(["2026-02-11", "NOT_A_NUMBER", 300])  # 行3：数值不可解析
    ws.append(["2026-03-12", None, 400])            # 行4：数值缺失
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "金额"], "numeric_columns": ["金额"],
           "date_column": "日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    write_json(workspace / "req.json", req)
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    by = {c["check"]: c for c in rep["checks"]}
    dh = by["duplicate_headers"]
    assert dh["ok"] is False and dh["duplicates"] == ["金额"]
    num = by["numeric"]
    assert num["ok"] is False
    assert num["unparseable_count"] == 1
    assert num["unparseable"][0]["row"] == 3
    assert num["unparseable"][0]["value"] == "NOT_A_NUMBER"
    assert num["missing_count"] == 1 and num["missing_rows"] == [4]
    assert any("表头重复" in x for x in rep["reasons"])
    assert any("数值不可解析" in x for x in rep["reasons"])
    assert any("数值缺失" in x for x in rep["reasons"])


def test_complex_formula_without_cache_blocks(workspace):
    """公式无缓存值（未读取到的范围）不得静默通过。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "金额(元)", "计算"])
    ws.append(["2026-01-05", 100, "=B2*2"])  # C2 公式，openpyxl 不写缓存
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "金额(元)"], "units": {"金额(元)": "元"},
           "numeric_columns": ["金额(元)"], "date_column": "日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    write_json(workspace / "req.json", req)
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    fc = [c for c in rep["checks"] if c["check"] == "formula_cache"][0]
    assert fc["ok"] is False
    assert fc["count"] == 1
    assert fc["formula_cells_without_cache"] == ["C2"]
    assert any("公式无缓存值" in x for x in rep["reasons"])
    # 其余检查本身均通过（阻断仅因公式缺缓存）
    for c in rep["checks"]:
        if c["check"] != "formula_cache":
            assert c["ok"] is True


# ---------- 异常 ----------

def test_abnormal_nonexistent_dir_exit2(workspace):
    cfg = {"items": [{"id": "M1", "name": "合同", "required": True,
                      "patterns": ["合同"]}]}
    write_json(workspace / "c.json", cfg)
    r = run("checklist", "--checklist", workspace / "c.json",
            "--dir", workspace / "no_such_dir",
            "--report", workspace / "r.json")
    assert r.returncode == 2
    assert "not found" in r.stderr


def test_abnormal_empty_checklist_blocks(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "合同.docx").write_bytes(b"broken document")
    write_json(workspace / "c.json", {"items": []})
    r = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    assert rep["items"] == []
    assert any("清单为空" in x for x in rep["reasons"])


def test_abnormal_empty_dir_not_silent_pass(workspace):
    mat = workspace / "mats"
    mat.mkdir()  # 空目录
    # 仅可选条目：不得静默 pass，至少 confirm_needed
    cfg = {"items": [{"id": "O1", "name": "可选材料", "required": False,
                      "patterns": ["可选"]}]}
    write_json(workspace / "c1.json", cfg)
    r = run("checklist", "--checklist", workspace / "c1.json", "--dir", mat,
            "--report", workspace / "r1.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r1.json")
    assert rep["verdict"] == "confirm_needed"
    assert rep["items"][0]["state"] == "missing"
    assert any("材料目录为空" in x for x in rep["reasons"])
    # 含必需条目：缺失 → block
    cfg["items"][0]["required"] = True
    write_json(workspace / "c2.json", cfg)
    r2 = run("checklist", "--checklist", workspace / "c2.json", "--dir", mat,
             "--report", workspace / "r2.json")
    assert r2.returncode == 0
    assert read_json(workspace / "r2.json")["verdict"] == "block"


def test_abnormal_duplicate_and_missing_ids_block(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "营业执照.md").write_text("统一社会信用代码 91xxx。", encoding="utf-8")
    cfg = {"items": [
        {"id": "M1", "name": "营业执照", "required": True,
         "patterns": ["营业执照"]},
        {"id": "M1", "name": "营业执照副本", "required": True,
         "patterns": ["营业执照"]},
        {"name": "无id条目", "required": False, "patterns": ["不存在的模式"]},
    ]}
    write_json(workspace / "c.json", cfg)
    r = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    assert any("id重复" in x for x in rep["reasons"])
    assert any("缺少id" in x for x in rep["reasons"])
    # 两个 M1 都能命中可用文件（阻断仅因清单自身问题）
    assert sum(1 for i in rep["items"] if i["state"] == "ready") == 2


def test_abnormal_no_pattern_items_block(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "合同.md").write_text("合同内容。", encoding="utf-8")
    cfg = {"items": [{"id": "M1", "name": "合同", "required": True}]}
    write_json(workspace / "c.json", cfg)
    r = run("checklist", "--checklist", workspace / "c.json", "--dir", mat,
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    assert any("未配置patterns" in x for x in rep["reasons"])


def test_abnormal_corrupt_data_file_exit2(workspace):
    (workspace / "d.xlsx").write_bytes(b"not an xlsx at all")
    write_json(workspace / "req.json", {"columns": ["日期"]})
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 2
    assert "cannot read data" in r.stderr


def test_abnormal_empty_requirements_block(workspace):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "金额"])
    ws.append(["2026-01-05", 100])
    wb.save(str(workspace / "d.xlsx"))
    write_json(workspace / "req.json", {})
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    assert any("requirements为空" in x for x in rep["reasons"])


def test_abnormal_date_column_missing_blocks(workspace):
    wb = Workbook()
    ws = wb.active
    ws.append(["日期", "金额"])
    ws.append(["2026-01-05", 100])
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "金额"], "date_column": "签署日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    write_json(workspace / "req.json", req)
    r = run("data", "--data", workspace / "d.xlsx",
            "--requirements", workspace / "req.json",
            "--report", workspace / "r.json")
    assert r.returncode == 0
    rep = read_json(workspace / "r.json")
    assert rep["verdict"] == "block"
    d = [c for c in rep["checks"] if c["check"] == "date_range"][0]
    assert d["ok"] is False
    assert d["error"] == "date_column_not_in_header"
    assert any("日期列缺失" in x for x in rep["reasons"])
