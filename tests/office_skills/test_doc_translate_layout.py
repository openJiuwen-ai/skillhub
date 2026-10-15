"""doc-translate-layout 单项标准黑盒测试。

独立验证（normal / complex / abnormal 三类）：
- 真实 docx/pptx 输入、真实子进程执行（不 import 被测实现）、断言落盘产物
- 关键回归（benchmarks/translation-complex 冻结缺陷）：
  1) 多 run Word 段落回填必须完整替换，译文后不得残留原文；
  2) PPT 同页不同 shape 的段落 id 必须唯一，两个框不得被回填同一译文；
  3) 工作表与源文件不一致必须拒绝（非零退出、不产出），不得把旧译文覆盖到新文档。
- 预期值由本测试内独立回读推导，不 mock 被测实现。
"""
import hashlib
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

import docx  # noqa: E402
from docx.shared import Pt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

SKILL = "doc-translate-layout"
SCRIPT = "translate_layout.py"


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(*args):
    return run_script(SKILL, SCRIPT, list(args))


def build_normal_docx(path):
    """标题 + 表格(2 格) + 多 run 段落 + 显式页眉。"""
    d = docx.Document()
    d.add_heading("季度预算汇总", level=1)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "项目名称"
    t.rows[0].cells[1].text = "预算金额"
    p = d.add_paragraph()
    p.add_run("已批准").bold = True
    p.add_run("年度预算")
    sec = d.sections[0]
    sec.header.is_linked_to_previous = False
    sec.header.paragraphs[0].text = "内部资料 勿外传"
    d.save(path)


def build_pptx(path, boxes):
    """boxes: [(slide_idx, text)]，每项在独立文本框。"""
    prs = Presentation()
    slides = {}
    for slide_idx, text in boxes:
        if slide_idx not in slides:
            slides[slide_idx] = prs.slides.add_slide(prs.slide_layouts[6])
        slides[slide_idx].shapes.add_textbox(
            Inches(1), Inches(1 + 2 * len(slides[slide_idx].shapes)),
            Inches(4), Inches(1)).text = text
    prs.save(path)


def pptx_box_texts(path):
    return [sh.text for sl in Presentation(str(path)).slides
            for sh in sl.shapes if sh.has_text_frame]


# ---------- 正常 ----------

def test_normal_docx_translate_keep_structure(workspace):
    src = workspace / "normal.docx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    gl = workspace / "terms.json"
    out = workspace / "translated.docx"
    rep = workspace / "report.json"
    build_normal_docx(src)
    src_hash = sha256(src)

    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    worksheet = read_json(ws)
    assert worksheet["item_count"] == 5
    assert len(worksheet["source_sha256"]) == 64
    ids = [it["id"] for it in worksheet["items"]]
    texts = [it["text"] for it in worksheet["items"]]
    assert texts[0] == "季度预算汇总" and texts[3] == "已批准年度预算"
    assert texts[4] == "内部资料 勿外传"
    assert worksheet["items"][1]["location"] == "table"

    write_json(tr, {ids[0]: "Quarterly Budget Overview",
                    ids[1]: "Item Name", ids[2]: "Budget Amount",
                    ids[3]: "Approved"})
    write_json(gl, {"季度": "Quarterly", "汇总": "Summary"})

    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out, "--glossary", gl, "--report", rep)
    assert r.returncode == 0, r.stderr + r.stdout

    d = docx.Document(str(out))
    # 正文段落：标题与多 run 段
    assert d.paragraphs[0].text == "Quarterly Budget Overview"
    assert d.paragraphs[0].style.name == "Heading 1"
    assert d.paragraphs[1].text == "Approved"
    assert d.paragraphs[1].runs[0].bold is True  # 首 run 格式保留
    # 表格结构与译文
    assert len(d.tables) == 1 and len(d.tables[0].rows[0].cells) == 2
    assert d.tables[0].rows[0].cells[0].text == "Item Name"
    assert d.tables[0].rows[0].cells[1].text == "Budget Amount"
    # 未翻译的页眉保留原文
    assert d.sections[0].header.paragraphs[0].text == "内部资料 勿外传"
    # 源文件未被修改
    assert sha256(src) == src_hash

    report = read_json(rep)
    assert report["translated"] == 4
    assert report["untranslated_ids"] == [ids[4]]
    assert report["glossary_hits"] == [{"id": ids[0], "term": "季度"}]
    assert report["glossary_misses"] == [{"id": ids[0], "term": "汇总",
                                          "expected": "Summary"}]
    assert report["format_degradations"] == [
        {"id": ids[3], "runs_with_text_before": 2,
         "note": "mixed run formatting collapsed into first run's format"}]
    # 中→英译文普遍超过 2x 阈值且无显式字号 -> 全部记 needs_attention
    fallback = {f["id"]: f["action"] for f in report["fallback_list"]}
    assert set(fallback) == {ids[0], ids[1], ids[2]}
    assert all(a == "needs_attention" for a in fallback.values())
    assert report["source_binding"]["matched"] is True


def test_normal_pptx_translate_two_slides(workspace):
    src = workspace / "normal.pptx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "translated.pptx"
    build_pptx(src, [(0, "预算"), (1, "时间表")])

    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    items = read_json(ws)["items"]
    assert len(items) == 2
    ids = [it["id"] for it in items]
    assert len(set(ids)) == 2
    assert ids[0].startswith("s01-") and ids[1].startswith("s02-")
    assert "sp" in ids[0]  # id 含 shape 标识

    write_json(tr, {ids[0]: "Budget", ids[1]: "Timeline"})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 0, r.stderr + r.stdout
    assert pptx_box_texts(out) == ["Budget", "Timeline"]


# ---------- 复杂 ----------

def test_complex_docx_multirun_complete_replacement(workspace):
    """冻结缺陷 1：多 run 回填只替换首 run，译文后残留原文。"""
    src = workspace / "runs.docx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "translated.docx"
    rep = workspace / "report.json"
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("已批准").bold = True
    p.add_run("预算")
    d.save(src)

    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    items = read_json(ws)["items"]
    assert items[0]["text"] == "已批准预算"
    write_json(tr, {items[0]["id"]: "Approved budget"})

    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out, "--report", rep)
    assert r.returncode == 0, r.stderr + r.stdout
    para = docx.Document(str(out)).paragraphs[0]
    assert para.text == "Approved budget"  # 完整替换，无“预算”残留
    assert "预算" not in para.text and "已批准" not in para.text
    assert para.runs[0].bold is True
    degradations = read_json(rep)["format_degradations"]
    assert degradations and degradations[0]["runs_with_text_before"] == 2


def test_complex_pptx_same_slide_shapes_unique_ids(workspace):
    """冻结缺陷 2：同页不同 shape 段落 id 重复，两框被回填同一译文。"""
    src = workspace / "shapes.pptx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "translated.pptx"
    build_pptx(src, [(0, "预算"), (0, "时间表")])  # 同页两个文本框

    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    items = read_json(ws)["items"]
    ids = [it["id"] for it in items]
    assert len(ids) == 2 and len(set(ids)) == 2  # id 唯一
    assert all(i.startswith("s01-") for i in ids)

    write_json(tr, {it["id"]: ("Budget" if it["text"] == "预算" else "Timeline")
                    for it in items})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 0, r.stderr + r.stdout
    texts = pptx_box_texts(out)
    assert texts == ["Budget", "Timeline"]  # 各自正确，不互相覆盖
    assert texts[0] != texts[1]


def test_complex_overflow_font_shrink_and_flag(workspace):
    """超长译文：docx 显式字号降 1pt；pptx 记 needs_attention。"""
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("短句").font.size = Pt(12)
    src = workspace / "short.docx"
    d.save(src)
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "out.docx"
    rep = workspace / "rep.json"
    long_text = "A translation that is far longer than the source text."
    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    write_json(tr, {read_json(ws)["items"][0]["id"]: long_text})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out, "--report", rep)
    assert r.returncode == 0, r.stderr + r.stdout
    fallback = read_json(rep)["fallback_list"]
    assert fallback == [{"id": "p-0000", "src_len": 2, "tgt_len": len(long_text),
                         "action": "font_shrunk"}]
    new_size = docx.Document(str(out)).paragraphs[0].runs[0].font.size
    assert new_size is not None and new_size.pt == 11.0  # 12pt -> 11pt

    psrc = workspace / "short.pptx"
    pws = workspace / "pws.json"
    ptr = workspace / "ptr.json"
    pout = workspace / "pout.pptx"
    prep = workspace / "prep.json"
    build_pptx(psrc, [(0, "Hi")])
    r = run("extract", "--input", psrc, "--worksheet", pws)
    assert r.returncode == 0, r.stderr
    write_json(ptr, {read_json(pws)["items"][0]["id"]: long_text})
    r = run("apply", "--input", psrc, "--worksheet", pws, "--translations", ptr,
            "--output", pout, "--report", prep)
    assert r.returncode == 0, r.stderr + r.stdout
    assert read_json(prep)["fallback_list"][0]["action"] == "needs_attention"
    assert pptx_box_texts(pout) == [long_text]


def test_complex_verify_roundtrip_and_drift(workspace):
    """verify 回读：正常产物通过；文本/结构被改动必须检出。"""
    d = docx.Document()
    d.add_paragraph("第一段落")
    d.add_paragraph("第二段落")
    src = workspace / "src.docx"
    d.save(src)
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "out.docx"
    vr = workspace / "verify.json"
    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    write_json(tr, {"p-0000": "A translation far longer than the source.",
                    "p-0001": "Para two"})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 0, r.stderr + r.stdout

    r = run("verify", "--input", src, "--worksheet", ws, "--output", out,
            "--translations", tr, "--report", vr)
    assert r.returncode == 0, r.stderr + r.stdout
    result = read_json(vr)
    assert result["structure_preserved"] is True
    assert result["structure_problems"] == []
    assert result["text_failures"] == []
    assert len(result["long_text_risks"]) == 1
    assert result["long_text_risks"][0]["id"] == "p-0000"
    assert result["verified"] is True

    # 文本被改动 -> verify 拒绝
    tampered = docx.Document(str(out))
    tampered.paragraphs[0].text = "WRONG TEXT"
    tampered.save(str(out))
    r = run("verify", "--input", src, "--worksheet", ws, "--output", out,
            "--translations", tr)
    assert r.returncode == 2
    assert "text" in r.stderr.lower() or "mismatch" in r.stderr.lower()

    # 结构被改动（新增段落）-> verify 拒绝
    out2 = workspace / "out2.docx"
    d2 = docx.Document(str(out))
    d2.paragraphs[0].text = "A translation far longer than the source."
    d2.add_paragraph("EXTRA PARAGRAPH")
    d2.save(str(out2))
    vr2 = workspace / "verify2.json"
    r = run("verify", "--input", src, "--worksheet", ws, "--output", out2,
            "--translations", tr, "--report", vr2)
    assert r.returncode == 2
    result2 = read_json(vr2)
    assert result2["structure_preserved"] is False
    assert result2["structure_problems"]
    assert result2["verified"] is False


# ---------- 异常 ----------

def test_abnormal_stale_worksheet_rejected(workspace):
    """冻结缺陷 3：源文件已变更仍退出 0 覆盖旧译文。"""
    a = workspace / "a.docx"
    b = workspace / "b.docx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "stale.docx"
    d = docx.Document()
    d.add_paragraph("已批准预算")
    d.save(a)
    d2 = docx.Document()
    d2.add_paragraph("另一份完全不同的内容")
    d2.save(b)
    a_hash, b_hash = sha256(a), sha256(b)

    r = run("extract", "--input", a, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    write_json(tr, {read_json(ws)["items"][0]["id"]: "Approved budget"})

    r = run("apply", "--input", b, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 2  # 拒绝，不再退出 0
    assert not out.exists()  # 不产出被旧译文覆盖的文件
    assert "binding" in r.stderr or "mismatch" in r.stderr
    assert sha256(a) == a_hash and sha256(b) == b_hash  # 双方源文件不动


def test_abnormal_worksheet_without_binding_rejected(workspace):
    src = workspace / "a.docx"
    d = docx.Document()
    d.add_paragraph("内容")
    d.save(src)
    ws = workspace / "nobinding.json"
    tr = workspace / "tr.json"
    out = workspace / "out.docx"
    write_json(ws, {"items": [{"id": "p-0000", "text": "内容",
                               "style": "", "location": "body"}]})
    write_json(tr, {"p-0000": "Content"})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 2
    assert "source_sha256" in r.stderr or "binding" in r.stderr
    assert not out.exists()


def test_abnormal_unknown_translation_ids_rejected(workspace):
    src = workspace / "a.docx"
    ws = workspace / "ws.json"
    tr = workspace / "tr.json"
    out = workspace / "out.docx"
    d = docx.Document()
    d.add_paragraph("内容")
    d.save(src)
    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    write_json(tr, {"p-0000": "Content", "p-9999": "ghost"})
    r = run("apply", "--input", src, "--worksheet", ws, "--translations", tr,
            "--output", out)
    assert r.returncode == 2
    assert "not in worksheet" in r.stderr
    assert not out.exists()


def test_abnormal_tampered_worksheet_texts_rejected(workspace):
    """绑定哈希虽匹配但工作表原文被改动：结构核对必须拒绝。"""
    src = workspace / "a.docx"
    ws = workspace / "ws.json"
    tampered = workspace / "ws-tampered.json"
    tr = workspace / "tr.json"
    out = workspace / "out.docx"
    d = docx.Document()
    d.add_paragraph("原始内容")
    d.save(src)
    r = run("extract", "--input", src, "--worksheet", ws)
    assert r.returncode == 0, r.stderr
    data = read_json(ws)
    data["items"][0]["text"] += "被篡改"
    write_json(tampered, data)
    write_json(tr, {"p-0000": "Content"})
    r = run("apply", "--input", src, "--worksheet", tampered,
            "--translations", tr, "--output", out)
    assert r.returncode == 2
    assert "differs from worksheet" in r.stderr or "does not match" in r.stderr
    assert not out.exists()


def test_abnormal_corrupt_and_missing_inputs(workspace):
    bad = workspace / "bad.docx"
    bad.write_bytes(b"this is not a real docx file")
    r = run("extract", "--input", bad, "--worksheet", workspace / "ws.json")
    assert r.returncode == 2
    assert "cannot read" in r.stderr

    r = run("extract", "--input", workspace / "missing.docx",
            "--worksheet", workspace / "ws.json")
    assert r.returncode == 2

    src = workspace / "a.docx"
    d = docx.Document()
    d.add_paragraph("内容")
    d.save(src)
    r = run("apply", "--input", src, "--worksheet", workspace / "missing.json",
            "--translations", workspace / "tr.json",
            "--output", workspace / "out.docx")
    assert r.returncode == 2

    r = run("apply", "--input", src, "--worksheet", workspace / "ws2.json",
            "--translations", workspace / "missing-tr.json",
            "--output", workspace / "out.docx")
    assert r.returncode == 2

    # 损坏的 pptx 同样拒绝
    badpptx = workspace / "bad.pptx"
    badpptx.write_bytes(b"not a pptx")
    r = run("extract", "--input", badpptx, "--worksheet", workspace / "ws3.json")
    assert r.returncode == 2
