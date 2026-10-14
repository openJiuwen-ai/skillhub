"""word-template-editor 验收: 用独立构建的模板与手工推导的预期值做黑盒断言。"""
import json
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402


def build_template(path):
    """构建含跨run占位/多级编号/表格/页眉页脚的模板。"""
    doc = Document()
    # 基础样式段
    p1 = doc.add_paragraph()
    r1 = p1.add_run("项目名称：")
    r1.font.size = Pt(12)
    r1.font.bold = True
    # 占位符被拆到3个run: {{ | project | }}
    p1.add_run("{{")
    p1.add_run("project")
    p1.add_run("}}")
    # 多级编号列表(使用内置列表样式 -> 产生 numbering 引用)
    doc.add_paragraph("背景说明一", style="List Number")
    doc.add_paragraph("背景说明二", style="List Number")
    doc.add_paragraph("背景说明三", style="List Bullet")
    # 表格内占位
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "金额"
    table.cell(0, 1).text = "{{amount}}"
    table.cell(1, 0).text = "日期"
    table.cell(1, 1).text = "{{  date  }}"
    # 未提供数据的占位
    doc.add_paragraph("备注：{{remark}}")
    # 页眉页脚占位
    sec = doc.sections[0]
    sec.header.paragraphs[0].add_run("机密——{{level}}")
    sec.footer.paragraphs[0].add_run("第 1 页")
    doc.save(path)
    return path


def test_normal_fill(workspace):
    tpl = build_template(str(workspace / "tpl.docx"))
    data = {"project": "智慧园区一期", "amount": "1280000", "date": "2026-10-14",
            "level": "内部"}
    (workspace / "data.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    r = run_script("word-template-editor", "fill_template.py",
                   ["--template", tpl, "--data", str(workspace / "data.json"),
                    "--output", str(workspace / "out.docx"),
                    "--report", str(workspace / "report.json")])
    assert r.returncode == 0, r.stderr + r.stdout
    report = json.loads((workspace / "report.json").read_text(encoding="utf-8"))

    # 断言: 匹配了预期的4个键; remark 未提供
    assert set(report["matched"]) == {"project", "amount", "date", "level"}
    assert "remark" in report["unmatched"]
    assert report["replacements"] == 4
    assert report["residual_placeholders"] == ["remark"]  # 未提供数据的占位原样保留
    assert report["numbering_xml_unchanged"] is True
    assert report["opc_check"]["ok"] is True

    # 独立验证: 重新打开输出, 检查文本与字符格式
    doc = Document(str(workspace / "out.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert any("项目名称：智慧园区一期" in t for t in texts), texts
    # 首个run格式保留: "项目名称："仍为粗体12pt, 替换文本继承占位首个run格式
    para = next(p for p in doc.paragraphs if "智慧园区一期" in p.text)
    runs = [r for r in para.runs if r.text]
    assert runs[0].font.bold is True and runs[0].font.size == Pt(12)
    tbl = doc.tables[0]
    assert tbl.cell(0, 1).text == "1280000"
    assert tbl.cell(1, 1).text == "2026-10-14"
    assert "内部" in doc.sections[0].header.paragraphs[0].text
    # 编号段落样式仍在
    assert any("背景说明一" == p.text for p in doc.paragraphs)


def test_complex_cross_run_and_style_names(workspace):
    """复杂: 占位完全跨run、样式名不变、numbering part字节不变。"""
    doc = Document()
    p = doc.add_paragraph(style="Heading 1")
    for piece in ("关于", "{{", "title", "}}", "的通知"):
        p.add_run(piece)
    doc.add_paragraph("正文占位 {{body}}")
    tpl = str(workspace / "t2.docx")
    doc.save(tpl)
    with zipfile.ZipFile(tpl) as zf:
        numbering_before = zf.read("word/numbering.xml") if "word/numbering.xml" in zf.namelist() else None
    data = {"title": "年度审计", "body": "见附件"}
    (workspace / "d.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    r = run_script("word-template-editor", "fill_template.py",
                   ["--template", tpl, "--data", str(workspace / "d.json"),
                    "--output", str(workspace / "o2.docx")])
    assert r.returncode == 0, r.stderr
    doc2 = Document(str(workspace / "o2.docx"))
    hp = doc2.paragraphs[0]
    assert hp.text == "关于年度审计的通知", hp.text
    assert hp.style.name == "Heading 1"
    with zipfile.ZipFile(str(workspace / "o2.docx")) as zf:
        numbering_after = zf.read("word/numbering.xml") if "word/numbering.xml" in zf.namelist() else None
    assert numbering_before == numbering_after


def test_abnormal_missing_data_and_corrupt(workspace):
    tpl = build_template(str(workspace / "tpl.docx"))
    # 数据缺失: 未匹配清单, 退出仍为0
    (workspace / "empty.json").write_text("{}", encoding="utf-8")
    r = run_script("word-template-editor", "fill_template.py",
                   ["--template", tpl, "--data", str(workspace / "empty.json"),
                    "--output", str(workspace / "o3.docx"),
                    "--report", str(workspace / "r3.json")])
    assert r.returncode == 0
    rep = json.loads((workspace / "r3.json").read_text(encoding="utf-8"))
    assert set(rep["unmatched"]) == {"project", "amount", "date", "level", "remark"}
    assert rep["replacements"] == 0
    # 损坏文件: 非零退出
    bad = workspace / "bad.docx"
    bad.write_bytes(b"not a zip at all")
    r2 = run_script("word-template-editor", "fill_template.py",
                    ["--template", str(bad), "--data", str(workspace / "empty.json"),
                     "--output", str(workspace / "o4.docx")])
    assert r2.returncode == 2
    # 数据JSON损坏
    (workspace / "bad.json").write_text("{invalid", encoding="utf-8")
    r3 = run_script("word-template-editor", "fill_template.py",
                    ["--template", tpl, "--data", str(workspace / "bad.json"),
                     "--output", str(workspace / "o5.docx")])
    assert r3.returncode == 2
