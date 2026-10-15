"""proposal-bid-generation 标准黑盒测试。

独立预期值，不 mock 实现、不 import 被测模块；通过子进程调用 CLI 契约，
对产物（xlsx/json/docx）做独立回读断言。覆盖正常/复杂/异常三类。
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def run_bid(ws, tender_path, quals_path, outdir, extra=None):
    args = ["--tender", tender_path, "--quals", quals_path, "--outdir", outdir]
    return run_script("proposal-bid-generation", "bid_matrix.py", args + (extra or []))


def read_matrix(outdir):
    with open(os.path.join(outdir, "response-matrix.json"), encoding="utf-8") as f:
        return json.load(f)


def draft_text(outdir):
    doc = Document(os.path.join(outdir, "bid-draft.docx"))
    return "\n".join(p.text for p in doc.paragraphs)


# ---------- 正常 ----------

NORMAL_TENDER = """1. 投标人须具备有效的ISO9001质量管理体系认证
2. 投标人近三年须具有智慧园区类项目业绩
3. 投标人须提供售后服务承诺书
4. 投标人注册资本不低于1000万元
5. 须盖章
6. 投标人须具备有效期内的安全生产许可证
"""


def test_normal_matrix_and_draft(workspace):
    (workspace / "tender.txt").write_text(NORMAL_TENDER, encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()
    doc = Document()
    doc.add_paragraph("质量管理体系认证证书 ISO9001，证书编号C-2026-118，有效期至2027年6月。")
    doc.save(str(quals / "资质-ISO9001.docx"))
    (quals / "业绩-智慧园区.md").write_text(
        "智慧园区一期项目合同，金额800万元，2024年验收通过。", encoding="utf-8")
    (quals / "售后-服务承诺.txt").write_text(
        "本公司郑重承诺提供为期三年的售后服务，7×24小时响应。", encoding="utf-8")
    (quals / "安全生产许可证.pdf").write_bytes(b"%PDF-1.4 not really parsed")
    r = run_bid(workspace, workspace / "tender.txt", quals, workspace / "out")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    # 6 条款全保留（含短条款"须盖章"）；3 条内容级证据 responded；
    # pdf 仅文件名候选 → needs_confirmation；注册资本/盖章无证据 → gap
    assert rep["clauses"] == 6
    assert rep["responded"] == 3
    assert rep["needs_confirmation"] == 1
    assert rep["gaps"] == 2
    assert rep["parse_mode"] == "numbered"
    assert rep["quals_files"] == 4
    assert rep["quals_unsupported"] == 1
    m = read_matrix(workspace / "out")
    by_no = {c["no"]: c for c in m["clauses"]}
    assert by_no["R01"]["status"] == "responded"
    assert by_no["R02"]["status"] == "responded"
    assert by_no["R03"]["status"] == "responded"
    assert by_no["R04"]["status"] == "gap"
    assert by_no["R05"]["text"] == "须盖章"
    assert by_no["R05"]["status"] == "gap"
    assert by_no["R06"]["status"] == "needs_confirmation"
    assert by_no["R06"]["reason"] == "candidate_only"
    # 证据可回溯：文件真实存在且带原文摘录
    ev = by_no["R01"]["evidence"][0]
    assert (quals / ev["file"]).exists()
    assert "ISO9001" in ev["sentence"]
    # xlsx 行数 = 条款数 + 表头；responded 行证据文件存在
    wb = load_workbook(str(workspace / "out" / "response-matrix.xlsx"))
    ws = wb["响应矩阵"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 6
    for row in rows:
        if row[2] == "responded":
            for f in str(row[3]).split("；"):
                assert (quals / f).exists(), f
    # 草稿：不自动承诺满足；缺口与待确认都有标记；要求人工复核
    full = draft_text(workspace / "out")
    assert "响应：满足" not in full
    assert "【待补充-缺口】" in full
    assert "【待确认-" in full
    assert "请复核" in full
    assert "ISO9001" in full


def test_fallback_paragraph_mode_and_empty_quals(workspace):
    tender = """本公司就智慧园区运维服务进行公开招标。
投标人应具备有效的质量管理体系认证。
近三年具有同类项目实施业绩。
"""
    (workspace / "tender.txt").write_text(tender, encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()  # 空目录：全部缺口但矩阵完整
    r = run_bid(workspace, workspace / "tender.txt", quals, workspace / "out")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["parse_mode"] == "paragraph_fallback"
    assert rep["clauses"] == 3  # 无编号 → 每个非空段落一条（虚拟编号）
    assert rep["responded"] == 0
    assert rep["gaps"] == 3
    m = read_matrix(workspace / "out")
    nos = [c["no"] for c in m["clauses"]]
    assert nos == ["R01", "R02", "R03"]
    assert all(c["status"] == "gap" for c in m["clauses"])
    gaps = json.loads((workspace / "out" / "gaps.json").read_text(encoding="utf-8"))
    assert len(gaps["gaps"]) == 3


# ---------- 复杂 ----------

def test_frozen_negative_evidence_and_short_clause(workspace):
    """主控冻结用例（benchmarks/bid-complex/before-review）复现。

    资质材料明确写“未取得ISO9001……不得宣称已经取得”，不得据此 responded；
    第 4 条短强制条款“须盖章”必须完整保留。
    """
    (workspace / "tender.txt").write_text(
        "1. 投标人必须具备有效ISO9001质量体系认证\n"
        "2. 投标人必须提供近三年同类项目合同\n"
        "3. 投标人必须提供售后服务承诺书\n"
        "4. 须盖章\n", encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()
    (quals / "ISO9001说明.txt").write_text(
        "本公司目前未取得ISO9001质量体系认证。不得宣称已经取得。", encoding="utf-8")
    r = run_bid(workspace, workspace / "tender.txt", quals, workspace / "out")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["clauses"] == 4  # 短条款不再被长度过滤丢弃
    assert rep["responded"] == 0  # 否定证据不得判 responded
    m = read_matrix(workspace / "out")
    by_no = {c["no"]: c for c in m["clauses"]}
    assert by_no["R04"]["text"] == "须盖章"
    assert by_no["R04"]["status"] == "gap"
    r01 = by_no["R01"]
    assert r01["status"] == "gap"
    assert r01["reason"] == "negative_evidence"
    negs = [w for w in r01["warnings"] if w["kind"] == "negative"]
    assert negs and negs[0]["file"] == "ISO9001说明.txt"
    assert "未取得" in negs[0]["sentence"]
    gaps = json.loads((workspace / "out" / "gaps.json").read_text(encoding="utf-8"))
    g01 = next(g for g in gaps["gaps"] if g["no"] == "R01")
    assert g01["reason"] == "negative_evidence"
    assert g01["negative_evidence"][0]["file"] == "ISO9001说明.txt"
    # 草稿不得自动承诺资格满足，且必须提示相反陈述
    full = draft_text(workspace / "out")
    assert "响应：满足" not in full
    assert "本公司目前未取得ISO9001质量体系认证" in full
    assert "不得据此宣称满足" in full
    assert "【待补充-缺口】" in full


def test_docx_table_continuation_sources(workspace):
    """docx 招标文件：第X条+续行、短条款、表格条款与出处，首行表头跳过。"""
    doc = Document()
    doc.add_paragraph("智慧园区运维服务项目招标公告")  # 前言（非条款）
    doc.add_paragraph("第一条 投标人须具备ISO9001质量体系认证")
    doc.add_paragraph("且提供证书复印件。")  # 续行并入上一条款
    doc.add_paragraph("第二条 须盖章")
    t = doc.add_table(rows=3, cols=2)
    t.rows[0].cells[0].text = "序号"
    t.rows[0].cells[1].text = "要求"
    t.rows[1].cells[0].text = "1"
    t.rows[1].cells[1].text = "投标人须提供售后服务承诺书"
    t.rows[2].cells[0].text = "2"
    t.rows[2].cells[1].text = "中标后30日内完成系统部署"
    doc.save(str(workspace / "tender.docx"))
    quals = workspace / "quals"
    quals.mkdir()
    (quals / "资质-ISO9001.md").write_text(
        "本公司取得ISO9001质量管理体系认证，证书编号C-118。", encoding="utf-8")
    r = run_bid(workspace, workspace / "tender.docx", quals, workspace / "out")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["parse_mode"] == "numbered+table"
    assert rep["clauses"] == 4  # 2 编号条款（含续行）+ 2 表格条款；表头不算条款
    assert rep["preamble_skipped"] == 1
    assert rep["continued_clauses"] == 1
    assert rep["table_clauses"] == 2
    m = read_matrix(workspace / "out")
    by_no = {c["no"]: c for c in m["clauses"]}
    r01 = by_no["R01"]
    assert "且提供证书复印件。" in r01["text"]
    assert r01["continued"] is True
    assert r01["source"] == "docx:段落2"
    assert r01["status"] == "responded"
    assert r01["evidence"][0]["file"] == "资质-ISO9001.md"
    assert by_no["R02"]["text"] == "须盖章"
    assert by_no["R02"]["source"] == "docx:段落4"
    assert by_no["R03"]["source"] == "docx:表1行2"
    assert "售后服务承诺书" in by_no["R03"]["text"]
    assert by_no["R04"]["source"] == "docx:表1行3"
    texts = [c["text"] for c in m["clauses"]]
    assert not any(t == "序号；要求" for t in texts)  # 表头未混入条款


def test_expired_subject_and_conflicting_evidence(workspace):
    """过期证据、主体不符/未核验、--company 核验、证据冲突。"""
    (workspace / "tender.txt").write_text(
        "1. 投标人须具备有效ISO9001质量体系认证\n", encoding="utf-8")

    def one_qual(name, content):
        q = workspace / name
        q.mkdir()
        (q / "证据.md").write_text(content, encoding="utf-8")
        return q

    # 过期证据 → gap(expired_or_invalid)，草稿提示过期/无效
    qe = one_qual("q-expired", "本公司原ISO9001证书已于2023年12月31日过期。\n")
    r = run_bid(workspace, workspace / "tender.txt", qe, workspace / "o1")
    assert r.returncode == 0
    m = read_matrix(workspace / "o1")
    assert m["clauses"][0]["status"] == "gap"
    assert m["clauses"][0]["reason"] == "expired_or_invalid"
    assert draft_text(workspace / "o1").count("过期/无效") >= 1

    # 第三方主体、未指定 --company → needs_confirmation(subject_unverified)
    qo = one_qual("q-other", "华宇科技有限公司具备ISO9001质量管理体系认证。\n")
    r = run_bid(workspace, workspace / "tender.txt", qo, workspace / "o2")
    assert r.returncode == 0
    m = read_matrix(workspace / "o2")
    c = m["clauses"][0]
    assert c["status"] == "needs_confirmation"
    assert c["reason"] == "subject_unverified"

    # --company 与证据主体一致 → responded
    r = run_bid(workspace, workspace / "tender.txt", qo, workspace / "o3",
                extra=["--company", "华宇科技有限公司"])
    assert r.returncode == 0
    m = read_matrix(workspace / "o3")
    assert m["clauses"][0]["status"] == "responded"

    # --company 与证据主体不一致 → needs_confirmation(subject_mismatch)
    r = run_bid(workspace, workspace / "tender.txt", qo, workspace / "o4",
                extra=["--company", "中天建设集团有限公司"])
    assert r.returncode == 0
    m = read_matrix(workspace / "o4")
    c = m["clauses"][0]
    assert c["status"] == "needs_confirmation"
    assert c["reason"] == "subject_mismatch"

    # 肯定证据与否定证据并存 → needs_confirmation(conflicting_evidence)
    qc = workspace / "q-conflict"
    qc.mkdir()
    (qc / "肯定.md").write_text("本公司取得ISO9001质量管理体系认证，证书编号C-118。\n",
                                encoding="utf-8")
    (qc / "否定.md").write_text("另查：本公司未取得ISO9001认证，正在申办中。\n",
                                encoding="utf-8")
    r = run_bid(workspace, workspace / "tender.txt", qc, workspace / "o5")
    assert r.returncode == 0
    m = read_matrix(workspace / "o5")
    c = m["clauses"][0]
    assert c["status"] == "needs_confirmation"
    assert c["reason"] == "conflicting_evidence"
    kinds = {w["kind"] for w in c["warnings"]}
    assert "negative" in kinds


# ---------- 异常 ----------

def test_missing_tender_exit2(workspace):
    quals = workspace / "quals"
    quals.mkdir()
    r = run_bid(workspace, workspace / "nope.txt", quals, workspace / "out")
    assert r.returncode == 2
    assert "招标文件不存在" in r.stderr


def test_corrupt_tender_docx_exit2(workspace):
    (workspace / "tender.docx").write_bytes(b"PK\x03\x04 not a real docx body")
    quals = workspace / "quals"
    quals.mkdir()
    r = run_bid(workspace, workspace / "tender.docx", quals, workspace / "out")
    assert r.returncode == 2
    assert "cannot parse tender" in r.stderr


def test_empty_tender_exit2(workspace):
    (workspace / "tender.txt").write_text("   \n\n  \n", encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()
    r = run_bid(workspace, workspace / "tender.txt", quals, workspace / "out")
    assert r.returncode == 2
    assert "cannot parse tender" in r.stderr


def test_missing_quals_dir_exit2(workspace):
    (workspace / "tender.txt").write_text(
        "1. 投标人须具备有效ISO9001质量体系认证\n", encoding="utf-8")
    r = run_bid(workspace, workspace / "tender.txt",
                workspace / "no-such-dir", workspace / "out")
    assert r.returncode == 2
    assert "资质材料目录不存在" in r.stderr


def test_unreadable_qual_filename_only_needs_confirmation(workspace):
    """内容不可读的资质文件仅文件名命中 → needs_confirmation，绝不 responded。"""
    (workspace / "tender.txt").write_text(
        "1. 投标人须具备有效ISO9001质量体系认证\n", encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()
    (quals / "ISO9001证书扫描件.pdf").write_bytes(b"%PDF-1.4 dummy bytes")
    (quals / "损坏-ISO9001资质.docx").write_bytes(b"PK\x03\x04 broken")
    r = run_bid(workspace, workspace / "tender.txt", quals, workspace / "out")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["quals_files"] == 2
    assert rep["quals_unreadable"] == 1  # 损坏 docx
    assert rep["quals_unsupported"] == 1  # pdf 不做内容提取
    m = read_matrix(workspace / "out")
    c = m["clauses"][0]
    assert c["status"] == "needs_confirmation"
    assert c["reason"] == "candidate_only"
    cand_files = [x["file"] for x in c["candidates"]]
    assert "ISO9001证书扫描件.pdf" in cand_files
    assert "损坏-ISO9001资质.docx" in cand_files
    gaps = json.loads((workspace / "out" / "gaps.json").read_text(encoding="utf-8"))
    assert gaps["gaps"] == []  # 未证实 ≠ 已缺口：单独列 needs_confirmation
    assert gaps["needs_confirmation"][0]["no"] == "R01"
    full = draft_text(workspace / "out")
    assert "【待确认-" in full
    assert "响应：满足" not in full
