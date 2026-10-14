"""sheet-consolidate-reconcile 黑盒测试：金标准为手工推导的匹配统计。"""
import json
import os
import sys
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from openpyxl import Workbook  # noqa: E402


def build_sources(ws):
    # 华东区: 编码/销售额(万元), P001-P004 + 合计行
    wb = Workbook()
    s1 = wb.active
    s1.title = "明细"
    s1.append(["序号", "编码", "销售额(万元)"])
    rows = [(1, "P001", 100), (2, "P002", 200), (3, "P003", 300), (4, "P004", 400)]
    for r in rows:
        s1.append(list(r))
    s1.append(["", "合计", "1000"])
    wb.save(str(ws / "east.xlsx"))

    # 华南区: 产品编码/销售额, P002-P004, P003 值不同 + P004 重复主键
    wb2 = Workbook()
    s2 = wb2.active
    s2.append(["产品编码", "销售额"])
    for r in [("P002", 200), ("P003", 330), ("P004", 400), ("P004", 410)]:
        s2.append(list(r))
    wb2.save(str(ws / "south.xlsx"))

    # 华北区: 产品编码/销售额, P001+P003+P004+P005
    wb3 = Workbook()
    s3 = wb3.active
    s3.append(["产品编码", "销售额"])
    for r in [("P001", 100), ("P003", 300), ("P004", 400), ("P005", 500)]:
        s3.append(list(r))
    wb3.save(str(ws / "north.xlsx"))


CFG = {
    "primary_key": "产品编码",
    "metrics": ["销售额"],
    "sources": [
        {"file": "east.xlsx", "sheet": "明细", "source_name": "华东区",
         "column_map": {"产品编码": "编码", "销售额": "销售额(万元)"}},
        {"file": "south.xlsx", "source_name": "华南区", "column_map": {}},
        {"file": "north.xlsx", "source_name": "华北区", "column_map": {}},
    ],
}


def test_reconcile(workspace):
    build_sources(workspace)
    (workspace / "cfg.json").write_text(json.dumps(CFG, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfg.json", "--output", workspace / "merged.xlsx",
                    "--report", workspace / "rep.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "rep.json").read_text(encoding="utf-8"))
    ms = rep["match_stats"]
    # 手工金标准:
    # 唯一主键 P001..P005 = 5
    assert ms["n_unique_pk"] == 5
    # 三源齐全: P003,P004 = 2
    assert ms["n_present_all"] == 2
    # 匹配率 = 2/5 = 0.4
    assert ms["match_rate"] == 0.4
    # 单源: P005(仅华北) = 1 (P001 在华东+华北, P002 在华东+华南)
    assert ms["n_single_source"] == 1
    # 差异主键: P003(300 vs 330), P004(400 vs 400,410 重复多值) = 2
    assert ms["n_mismatch_pk"] == 2
    # 重复主键: 华南 P004 x2
    assert len(rep["duplicates"]) == 1
    assert rep["duplicates"][0]["pk"] == "P004" and rep["duplicates"][0]["count"] == 2
    # 合计行剔除: 华东 合计
    assert len(rep["skipped_total_rows"]) == 1
    assert rep["skipped_total_rows"][0]["source"] == "华东区"
    # 差异明细含 P003 的各源值
    m_p3 = [m for m in rep["mismatches"] if m["pk"] == "P003"][0]
    assert m_p3["values"]["华东区"] == [300.0]
    assert m_p3["values"]["华南区"] == [330.0]
    # 华北缺失主键 = P002
    assert sorted(rep["source_missing"]["华北区"]) == ["P002"]

    # 合并表独立验证
    from openpyxl import load_workbook
    wb = load_workbook(str(workspace / "merged.xlsx"))
    ws = wb["合并汇总"]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert header[0] == "产品编码"
    assert "销售额@华东区" in header and "销售额@华南区" in header and "销售额@华北区" in header
    data = {r[0]: r for r in rows[1:]}
    assert data["P001"][1] == 100 and data["P001"][3] == 100
    assert data["P005"][3] == 500
    # 源明细行数 = 表头1 + 华东4 + 华南4 + 华北4 = 13
    ws2 = wb["源明细"]
    detail_rows = list(ws2.iter_rows(values_only=True))
    assert len(detail_rows) == 1 + 4 + 4 + 4


def test_failed_source_continues(workspace):
    build_sources(workspace)
    cfg = dict(CFG)
    cfg["sources"] = CFG["sources"] + [
        {"file": "broken.xlsx", "source_name": "西域区", "column_map": {}}]
    (workspace / "cfg2.json").write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfg2.json", "--output", workspace / "m2.xlsx",
                    "--report", workspace / "r2.json"])
    # broken.xlsx 不存在 -> 该源 failed, 其余继续
    assert r.returncode == 0, r.stderr
    rep = json.loads((workspace / "r2.json").read_text(encoding="utf-8"))
    assert "西域区" in rep["failed_sources"]
    assert rep["match_stats"]["n_unique_pk"] == 5


def test_all_sources_fail(workspace):
    (workspace / "cfg3.json").write_text(json.dumps({
        "primary_key": "X", "metrics": ["Y"],
        "sources": [{"file": "nope1.xlsx"}, {"file": "nope2.xlsx"}]},
        ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfg3.json", "--output", workspace / "m3.xlsx",
                    "--report", workspace / "r3.json"])
    assert r.returncode == 2


def _patch_cached_formula(path, formula_body, cached_value):
    """给 openpyxl 生成的公式单元格补写缓存值 <v>，模拟 Excel 保存后的带缓存公式。"""
    zin = zipfile.ZipFile(path)
    items = {n: zin.read(n) for n in zin.namelist()}
    zin.close()
    xml = items["xl/worksheets/sheet1.xml"].decode("utf-8")
    patched = xml.replace(f"<f>{formula_body}</f>",
                          f"<f>{formula_body}</f><v>{cached_value}</v>")
    assert patched != xml, "formula <f> element not found; openpyxl output format changed?"
    items["xl/worksheets/sheet1.xml"] = patched.encode("utf-8")
    zout = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    for n, data in items.items():
        zout.writestr(n, data)
    zout.close()


def test_formula_cache_handling(workspace):
    """复杂：公式单元格——带缓存值参与对账，无缓存公式显式待核、不猜不补0。"""
    # 华南区: P001 销售额为公式且无缓存值(openpyxl 不写缓存), P002=500
    wb = Workbook()
    s = wb.active
    s.append(["产品编码", "销售额"])
    s.append(["P001", "=B3*2"])
    s.append(["P002", 500])
    wb.save(str(workspace / "south.xlsx"))
    # 华东区: P001 同公式但带缓存值1000(模拟Excel保存), P002=500
    wb2 = Workbook()
    s2 = wb2.active
    s2.append(["产品编码", "销售额"])
    s2.append(["P001", "=B3*2"])
    s2.append(["P002", 500])
    wb2.save(str(workspace / "east.xlsx"))
    _patch_cached_formula(str(workspace / "east.xlsx"), "B3*2", 1000)

    (workspace / "cfgf.json").write_text(json.dumps({
        "primary_key": "产品编码", "metrics": ["销售额"],
        "sources": [
            {"file": "east.xlsx", "source_name": "华东区", "column_map": {}},
            {"file": "south.xlsx", "source_name": "华南区", "column_map": {}},
        ]}, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfgf.json", "--output", workspace / "mf.xlsx",
                    "--report", workspace / "rf.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "rf.json").read_text(encoding="utf-8"))

    # 带缓存公式: 值被采用并计数
    fh = rep["formula_handling"]
    assert fh["cached_cells_per_source"] == {"华东区": 1}
    # 无缓存公式: 逐条列明(来源/主键/指标/行号/公式文本), 不静默当缺失
    assert len(fh["no_cache"]) == 1
    nc = fh["no_cache"][0]
    assert nc["source"] == "华南区" and nc["pk"] == "P001" and nc["metric"] == "销售额"
    assert nc["row"] == 2 and nc["formula"] == "=B3*2"
    # 无缓存公式进入 pending_values, 不参与数值比较 => P001 不算差异
    pends = [p for p in rep["pending_values"] if p["pk"] == "P001"]
    assert len(pends) == 1 and pends[0]["formula"] == "=B3*2" and pends[0]["source"] == "华南区"
    assert all(m["pk"] != "P001" for m in rep["mismatches"])
    # P001 行存在于两源, 不是源缺失
    assert not rep["source_missing"]
    # 匹配统计: 唯一主键2, 两源齐全2, 无差异主键
    ms = rep["match_stats"]
    assert ms["n_unique_pk"] == 2 and ms["n_present_all"] == 2
    assert ms["n_mismatch_pk"] == 0 and ms["match_rate"] == 1.0

    # 合并表独立验证
    from openpyxl import load_workbook
    wbm = load_workbook(str(workspace / "mf.xlsx"))
    ws = wbm["合并汇总"]
    rows = {r0[0]: r0 for r0 in ws.iter_rows(min_row=2, values_only=True)}
    # header: 产品编码, 销售额@华东区, 销售额@华南区, 数据一致, 缺失来源
    assert rows["P001"][1] == 1000            # 缓存值被采用
    assert rows["P001"][2] == "(公式未算)"     # 无缓存公式显式标注, 不是空也不是0
    assert rows["P001"][3] == "待核"
    assert rows["P002"][1] == 500 and rows["P002"][2] == 500
    assert rows["P002"][3] == "是"
    # 源明细: 华南 P001 显示公式文本, 华东 P001 显示缓存值
    det = list(wbm["源明细"].iter_rows(min_row=2, values_only=True))
    p1_south = [d for d in det if d[0] == "华南区" and d[1] == "P001"][0]
    p1_east = [d for d in det if d[0] == "华东区" and d[1] == "P001"][0]
    assert p1_south[3] == "(公式未算)=B3*2"
    assert p1_east[3] == 1000


def test_unparseable_value_pending(workspace):
    """复杂/异常：文本值不数值化——列入pending带原值，不算差异，不补0，不出现'None'。"""
    wb = Workbook()
    s = wb.active
    s.append(["产品编码", "销售额"])
    s.append(["P001", "待定"])
    s.append(["P002", 200])
    wb.save(str(workspace / "a.xlsx"))
    wb2 = Workbook()
    s2 = wb2.active
    s2.append(["产品编码", "销售额"])
    s2.append(["P001", 100])
    s2.append(["P002", 200])
    wb2.save(str(workspace / "b.xlsx"))
    (workspace / "cfgu.json").write_text(json.dumps({
        "primary_key": "产品编码", "metrics": ["销售额"],
        "sources": [
            {"file": "a.xlsx", "source_name": "A源", "column_map": {}},
            {"file": "b.xlsx", "source_name": "B源", "column_map": {}},
        ]}, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfgu.json", "--output", workspace / "mu.xlsx",
                    "--report", workspace / "ru.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "ru.json").read_text(encoding="utf-8"))
    pends = [p for p in rep["pending_values"] if p["pk"] == "P001"]
    assert len(pends) == 1
    assert pends[0]["source"] == "A源" and pends[0]["raw"] == "待定"
    assert pends[0].get("row") == 2
    # 不出现字符串化的 None
    assert all(p.get("raw") != "None" for p in rep["pending_values"])
    # P001 只有 B源有数值 => 不构成差异, 不猜测
    assert all(m["pk"] != "P001" for m in rep["mismatches"])
    assert rep["match_stats"]["n_mismatch_pk"] == 0
    # 合并表: P001 A源列为空(非0), 标记待核; P002 一致为是
    from openpyxl import load_workbook
    wbm = load_workbook(str(workspace / "mu.xlsx"))
    rows = {r0[0]: r0 for r0 in wbm["合并汇总"].iter_rows(min_row=2, values_only=True)}
    assert rows["P001"][1] in ("", None)
    assert rows["P001"][2] == 100
    assert rows["P001"][3] == "待核"
    assert rows["P002"][3] == "是"


def test_nonfinite_text_pending(workspace):
    """异常：文本 'nan'/'inf' 虽可被 float() 解析，但非有限值不得当数值——按原值待核、
    不参与对账、不静默变 None、报告不得输出非标准 JSON 的 NaN/Infinity 字面量。"""
    wb = Workbook()
    s = wb.active
    s.append(["产品编码", "销售额"])
    s.append(["P001", "nan"])
    s.append(["P002", "inf"])
    s.append(["P003", "abc"])
    wb.save(str(workspace / "a.xlsx"))
    wb2 = Workbook()
    s2 = wb2.active
    s2.append(["产品编码", "销售额"])
    for r in [("P001", 1), ("P002", 2), ("P003", 3)]:
        s2.append(list(r))
    wb2.save(str(workspace / "b.xlsx"))
    (workspace / "cfgn.json").write_text(json.dumps({
        "primary_key": "产品编码", "metrics": ["销售额"],
        "sources": [
            {"file": "a.xlsx", "source_name": "A源", "column_map": {}},
            {"file": "b.xlsx", "source_name": "B源", "column_map": {}},
        ]}, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfgn.json", "--output", workspace / "mn.xlsx",
                    "--report", workspace / "rn.json"])
    assert r.returncode == 0, r.stderr + r.stdout

    raw_text = (workspace / "rn.json").read_text(encoding="utf-8")
    assert "NaN" not in raw_text and "Infinity" not in raw_text

    rep = json.loads(raw_text)
    by_pk = {p["pk"]: p for p in rep["pending_values"] if p["source"] == "A源"}
    assert by_pk["P001"]["raw"] == "nan" and by_pk["P002"]["raw"] == "inf"
    assert by_pk["P003"]["raw"] == "abc"
    # 非有限文本不构成数值差异, 也不当缺失
    assert rep["mismatches"] == []
    assert rep["match_stats"]["n_mismatch_pk"] == 0
    assert not rep["source_missing"]

    from openpyxl import load_workbook
    wbm = load_workbook(str(workspace / "mn.xlsx"))
    rows = {r0[0]: r0 for r0 in wbm["合并汇总"].iter_rows(min_row=2, values_only=True)}
    for pk in ("P001", "P002", "P003"):
        assert rows[pk][1] in ("", None)   # A源非有限文本不落成 None 数值, 留空
        assert rows[pk][3] == "待核"
