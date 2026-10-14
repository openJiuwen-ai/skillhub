"""chinese-official-doc 测试：MIT复用脚本可运行 + GB/T 9704 排版与机检 + 确定性。"""
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script, skill_dir  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

DRAFT = """标题：关于开展2026年度内部审计工作的通知
主送：各部门、各分子公司：

一、审计范围
2026年度内部审计覆盖全部一级部门，重点关注资金使用与采购流程。

二、工作要求
各部门须在收到通知后五个工作日内完成自查，并将自查报告报送审计部。

落款：九文科技（集团）有限公司审计部
日期：2026年10月14日
"""


def test_layout_and_check(workspace):
    (workspace / "draft.txt").write_text(DRAFT, encoding="utf-8")
    r = run_script("chinese-official-doc", "gbt9704_layout.py",
                   ["--draft", workspace / "draft.txt", "--output", workspace / "gongwen.docx",
                    "--report", workspace / "layout.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "layout.json").read_text(encoding="utf-8"))
    assert rep["draft_elements"]["title"] is True
    assert rep["draft_elements"]["recipient"] is True
    assert rep["draft_elements"]["body"] >= 4
    assert rep["elements_complete"] is True, rep["missing"]
    assert rep["missing"] == []

    # 独立断言（不依赖脚本check逻辑）
    doc = Document(str(workspace / "gongwen.docx"))
    sec = doc.sections[0]
    assert abs(sec.top_margin.mm - 37) < 1 and abs(sec.left_margin.mm - 28) < 1
    paras = [p for p in doc.paragraphs if p.text.strip()]
    title = paras[0]
    assert title.text.startswith("关于开展")
    assert title.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.size.pt == 22
    recipient = paras[1]
    assert recipient.text.startswith("各部门")
    assert (recipient.paragraph_format.first_line_indent or 0) == 0
    body_first = next(p for p in paras if p.text.startswith("一、"))
    assert body_first.runs[0].font.size.pt == 16
    body_indented = next(p for p in paras if p.text.startswith("2026年度"))
    assert body_indented.paragraph_format.first_line_indent.pt >= 30  # 2字符 x 16pt
    right = [p for p in paras if p.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.RIGHT]
    assert any("九文科技" in p.text for p in right)
    assert any("2026年10月14日" in p.text for p in right)

    # check 模式独立可用
    r2 = run_script("chinese-official-doc", "gbt9704_layout.py",
                    ["--check", workspace / "gongwen.docx", "--report", workspace / "chk.json"])
    assert r2.returncode == 0
    chk2 = json.loads((workspace / "chk.json").read_text(encoding="utf-8"))
    assert chk2["elements_complete"] is True, chk2["missing"]


def test_layout_determinism(workspace):
    """同一输入两次排版，输出 docx 主 XML 一致（确定性）。"""
    import zipfile
    (workspace / "draft.txt").write_text(DRAFT, encoding="utf-8")
    run_script("chinese-official-doc", "gbt9704_layout.py",
               ["--draft", workspace / "draft.txt", "--output", workspace / "a.docx"])
    run_script("chinese-official-doc", "gbt9704_layout.py",
               ["--draft", workspace / "draft.txt", "--output", workspace / "b.docx"])
    with zipfile.ZipFile(str(workspace / "a.docx")) as za, \
            zipfile.ZipFile(str(workspace / "b.docx")) as zb:
        assert za.read("word/document.xml") == zb.read("word/document.xml")


def test_upstream_prose_lint_runs_and_deterministic(workspace):
    """MIT 复用的 prose_lint.py 可运行且两次命中一致；只报告不改写。"""
    text = ("需要指出的是，本方案将全面提升管理水平。值得注意的是，相关情况如下："
            "我们将持续发力，赋能业务，实现质的飞跃。")
    (workspace / "s.txt").write_text(text, encoding="utf-8")
    sdir = skill_dir("chinese-official-doc")
    env = dict(os.environ, PYTHONUTF8="1")
    outs = []
    for i in (1, 2):
        p = subprocess.run([sys.executable, os.path.join(sdir, "scripts", "prose_lint.py"),
                            "--json", str(workspace / "s.txt")],
                           capture_output=True, text=True, encoding="utf-8", env=env)
        assert p.returncode in (0, 1), p.stderr
        data = json.loads(p.stdout)
        findings = data if isinstance(data, list) else data.get("findings", [])
        outs.append(sorted((f["label"], f["line"]) for f in findings))
    assert outs[0] == outs[1]
    labels = {l for l, _ in outs[0]}
    assert len(labels) >= 1  # 命中套话/口语类规则


def test_upstream_scripts_help(workspace):
    """上游4个脚本均可在本环境运行（--help 退出0）。"""
    sdir = skill_dir("chinese-official-doc")
    env = dict(os.environ, PYTHONUTF8="1")
    for s in ("draft_length.py", "edit_scope.py", "material_diff.py", "prose_lint.py"):
        p = subprocess.run([sys.executable, os.path.join(sdir, "scripts", s), "--help"],
                           capture_output=True, text=True, encoding="utf-8", env=env)
        assert p.returncode == 0, (s, p.stderr)


def test_abnormal(workspace):
    (workspace / "bad.txt").write_text("", encoding="utf-8")
    r = run_script("chinese-official-doc", "gbt9704_layout.py",
                   ["--draft", workspace / "bad.txt", "--output", workspace / "x.docx"])
    assert r.returncode == 2
    r2 = run_script("chinese-official-doc", "gbt9704_layout.py",
                    ["--draft", workspace / "noexist.txt", "--output", workspace / "y.docx"])
    assert r2.returncode == 2


COMPLEX_DRAFT = """# 关于印发《集团差旅费管理办法（试行）》的通知
主送：集团总部各部门、各子公司：

一、总体要求
为进一步规范差旅费管理，降低行政运行成本，现将《集团差旅费管理办法（试行）》印发给你们，请遵照执行。

（一）适用范围
本办法适用于集团总部及全资、控股子公司。

（二）开支口径
交通费、住宿费在标准以内凭票报销，超支部分自理。

二、执行要求
各单位应当自本办法印发之日起组织专题学习，并将执行中的问题及时反馈财务部。

落款：九文科技（集团）有限公司
日期：2026年10月14日
"""


def _east_asia(paragraph):
    rPr = paragraph.runs[0]._element.find(qn("w:rPr"))
    rf = rPr.find(qn("w:rFonts"))
    return rf.get(qn("w:eastAsia"))


def test_layout_complex_structure(workspace):
    """复杂结构：markdown标题、黑体一级/楷体二级标题、落款右空2字。"""
    (workspace / "complex.txt").write_text(COMPLEX_DRAFT, encoding="utf-8")
    r = run_script("chinese-official-doc", "gbt9704_layout.py",
                   ["--draft", workspace / "complex.txt", "--output", workspace / "c.docx",
                    "--report", workspace / "c.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "c.json").read_text(encoding="utf-8"))
    assert rep["draft_elements"]["title"] is True
    assert rep["elements_complete"] is True, rep["missing"]

    doc = Document(str(workspace / "c.docx"))
    paras = [p for p in doc.paragraphs if p.text.strip()]
    title = paras[0]
    assert title.text.startswith("关于印发")
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert _east_asia(title) == "方正小标宋简体"
    assert title.runs[0].font.size.pt == 22

    h1 = next(p for p in paras if p.text.startswith("一、总体要求"))
    assert _east_asia(h1) == "黑体"
    assert h1.runs[0].font.size.pt == 16
    assert (h1.paragraph_format.first_line_indent or 0) == 0

    h2 = next(p for p in paras if p.text.startswith("（一）适用范围"))
    assert _east_asia(h2) == "楷体_GB2312"
    assert h2.runs[0].font.size.pt == 16

    body = next(p for p in paras if p.text.startswith("为进一步规范"))
    assert _east_asia(body) == "仿宋_GB2312"
    assert abs(body.paragraph_format.first_line_indent.pt - 32) < 2  # 2字符 x 16pt

    signoff = next(p for p in paras if p.text.startswith("九文科技"))
    assert signoff.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert abs(signoff.paragraph_format.right_indent.pt - 32) < 2  # 右空2字
    date_para = next(p for p in paras if p.text.startswith("2026年10月14日"))
    assert date_para.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert abs(date_para.paragraph_format.right_indent.pt - 32) < 2

    # 行距：标题固定31磅，正文固定28磅
    assert abs(title.paragraph_format.line_spacing.pt - 31) < 0.1
    assert abs(body.paragraph_format.line_spacing.pt - 28) < 0.1


def test_layout_missing_elements_reported_not_fabricated(workspace):
    """主送/落款/日期缺失时列为待补（missing），不编造发文要素。"""
    partial = "标题：关于调整值班安排的说明\n\n各部门自下周起按新安排执行值班。\n"
    (workspace / "partial.txt").write_text(partial, encoding="utf-8")
    r = run_script("chinese-official-doc", "gbt9704_layout.py",
                   ["--draft", workspace / "partial.txt", "--output", workspace / "p.docx",
                    "--report", workspace / "p.json"])
    assert r.returncode == 0, r.stderr + r.stdout  # 有标题+正文即可排版
    rep = json.loads((workspace / "p.json").read_text(encoding="utf-8"))
    assert rep["draft_elements"]["recipient"] is False
    assert rep["draft_elements"]["signoff"] is False
    assert rep["draft_elements"]["date"] is False
    assert rep["elements_complete"] is False
    assert "signoff_right_aligned" in rep["missing"]

    doc = Document(str(workspace / "p.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "发字" not in text and "〔" not in text  # 不编造发文字号
    assert "印发" not in text or "印发" in partial  # 不新增草稿外要素


def test_prose_lint_report_only_no_ai_probability_wording(workspace):
    """风险报告只报不改写；输出不得含'AI概率'类表述。"""
    text = ("需要指出的是，本方案将全面提升管理水平。值得注意的是，"
            "相关情况如下：大概就是这样，可能吧。")
    src = workspace / "s2.txt"
    src.write_text(text, encoding="utf-8")
    before = src.read_bytes()
    sdir = skill_dir("chinese-official-doc")
    env = dict(os.environ, PYTHONUTF8="1")
    p = subprocess.run([sys.executable, os.path.join(sdir, "scripts", "prose_lint.py"),
                        "--json", str(src)],
                       capture_output=True, text=True, encoding="utf-8", env=env)
    assert p.returncode in (0, 1), p.stderr
    assert "概率" not in p.stdout, p.stdout
    assert not any(("AI" in (f.get("label") or "")) for f in json.loads(p.stdout))
    assert len(json.loads(p.stdout)) >= 1  # 命中套话/口语规则
    assert src.read_bytes() == before  # 只报告，不改写原稿
