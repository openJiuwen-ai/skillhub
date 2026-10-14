"""业务流程家族 7 技能黑盒测试。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


# ---------- meeting-minutes-actions ----------

TRANSCRIPT = """[10:01:12] 主持人: 今天讨论Q3复盘与下阶段安排。
[10:03:40] 张伟: 建议加强客户回访，我负责本月完成客户回访计划。
[10:06:15] 李娜: 数据平台二期由李娜牵头，下周五前提交方案。
[10:09:02] 王强: 会议决定通过Q3预算追加12万元。
[10:11:30] 张伟: 供应商对接口径由采购部跟进。
[10:14:00] 李娜: 上期的档案整理任务已完成。
"""

PREV_ACTIONS = [
    ["A-001", "李娜", "档案整理", "9月底", "进行中"],
    ["A-002", "王强", "机房巡检", "9月底", "进行中"],
    ["A-003", "张伟", "季度报告初稿", "8月中", "已完成"],
]


def build_prev_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "责任人", "任务", "截止", "状态"])
    for row in PREV_ACTIONS:
        ws.append(row)
    wb.save(path)


def test_minutes(workspace):
    (workspace / "t.txt").write_text(TRANSCRIPT, encoding="utf-8")
    build_prev_xlsx(str(workspace / "prev.xlsx"))
    r = run_script("meeting-minutes-actions", "minutes_actions.py",
                   ["--transcript", workspace / "t.txt",
                    "--prev-actions", workspace / "prev.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    # 行动项: 3条新行动(回访/方案/口径) + 1条"已完成"陈述(档案整理含完成动词也算) -> 新行动>=3
    assert rep["n_actions"] >= 4  # 3承接 + 新增若干
    # xlsx 与纪要待办一致
    wb = load_workbook(str(workspace / "out" / "actions.xlsx"))
    ws = wb["待办追踪"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    doc = Document(str(workspace / "out" / "minutes.docx"))
    action_table_rows = len(doc.tables[0].rows) - 1
    assert len(rows) == action_table_rows == rep["n_actions"]
    # 承接: A-001(档案整理)本期提及且完成 -> 已完成; A-002(机房巡检)未提及 -> 进行中+否
    by_task = {row[2]: row for row in rows}
    a1 = next(row for row in rows if "档案整理" in str(row[2]))
    assert a1[4] == "已完成" and a1[7] == "是"
    a2 = next(row for row in rows if "机房巡检" in str(row[2]))
    assert a2[4] == "进行中" and a2[7] == "否"  # 未提及不推断未完成
    a3 = next(row for row in rows if "季度报告" in str(row[2]))
    assert a3[4] == "已完成"  # 上期已完成保持
    # 决议被识别
    assert rep["decisions"] >= 1
    # ASR 存疑: 12万元 数字被列出
    issues = json.loads((workspace / "out" / "asr-issues.json").read_text(encoding="utf-8"))
    assert any("12万" in i["value"] for i in issues["issues"])
    # 确定性: 两次运行输出一致
    r2 = run_script("meeting-minutes-actions", "minutes_actions.py",
                    ["--transcript", workspace / "t.txt",
                     "--prev-actions", workspace / "prev.xlsx",
                     "--outdir", workspace / "out2"])
    wb2 = load_workbook(str(workspace / "out2" / "actions.xlsx"))
    ws2 = wb2["待办追踪"]
    assert list(ws.iter_rows(values_only=True)) == list(ws2.iter_rows(values_only=True))


# ---------- evidence-digest ----------

def test_evidence(workspace):
    mat = workspace / "materials"
    mat.mkdir()
    (mat / "报告.md").write_text(
        "评价结论：2026年Q3销售额为1200万元。\n客户满意度为92%。\n访谈提到交付及时。", encoding="utf-8")
    doc = Document()
    doc.add_paragraph("访谈记录：受访人认为服务响应良好。")
    doc.add_paragraph("补充数据：Q3销售额为1150万元。")  # 与报告冲突
    doc.save(str(mat / "访谈.docx"))
    wb = Workbook(); ws = wb.active
    ws.append(["指标", "数值"])
    ws.append(["客户满意度", "93%"])  # 与报告 92% 冲突
    ws.append(["投诉量", "5件"])
    wb.save(str(mat / "指标.xlsx"))
    (mat / "截图.png").write_bytes(b"\x89PNG\r\n\x1a\nfakepng")
    (workspace / "q.md").write_text("- Q3销售额是多少？\n- 客户满意度如何？\n- 员工流失率是多少？", encoding="utf-8")

    r = run_script("evidence-digest", "evidence_ledger.py",
                   ["--materials", mat, "--questions", workspace / "q.md",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "out" / "evidence-ledger.json").read_text(encoding="utf-8"))
    s = rep["stats"]
    # 每条证据可回溯: source_file + location 均非空
    assert all(e["source_file"] and e["location"] for e in rep["ledger"])
    # 图片 needs_vision
    assert len(rep["image_sources_needing_vision"]) == 1
    # 冲突: 销售额(1200 vs 1150万元)、客户满意度(92% vs 93%)
    conflict_metrics = " ".join(c["metric"] for c in rep["conflicts"])
    assert "销售额" in conflict_metrics or "满意度" in conflict_metrics
    # 无证据问题: 员工流失率
    oq = json.loads((workspace / "out" / "open-questions.json").read_text(encoding="utf-8"))
    assert any("流失率" in q["question"] for q in oq["open_questions"])


# ---------- input-preflight-check ----------

def test_preflight(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    (mat / "营业执照副本.md").write_text("统一社会信用代码 91330100MA27X0000X；经营范围：软件开发。", encoding="utf-8")
    (mat / "审计报告_2024.md").write_text("审计意见：标准无保留意见。", encoding="utf-8")
    cfg = {"items": [
        {"id": "M1", "name": "营业执照", "required": True,
         "patterns": ["营业执照", "统一社会信用代码"], "fuzzy": ["执照"]},
        {"id": "M2", "name": "近三年审计报告", "required": True, "patterns": ["审计报告", "2023"]},
        {"id": "M3", "name": "税务证明", "required": False, "patterns": ["完税证明"]},
    ]}
    (workspace / "c.json").write_text(json.dumps(cfg, ensure_ascii=False), encoding="utf-8")
    r = run_script("input-preflight-check", "preflight.py",
                   ["checklist", "--checklist", workspace / "c.json", "--dir", mat,
                    "--report", workspace / "r.json"])
    assert r.returncode == 0
    rep = json.loads((workspace / "r.json").read_text(encoding="utf-8"))
    states = {i["id"]: i["state"] for i in rep["items"]}
    assert states["M1"] == "ready"
    assert states["M2"] == "missing"  # 审计报告有但缺2023模式
    assert states["M3"] == "missing"  # 非必需
    assert rep["verdict"] == "block"  # 必需项缺失阻断
    # 名称相近陷阱: 只有fuzzy命中 -> needs_confirm
    (mat / "执照复印件占位.txt").write_text("占位文件，无信用代码", encoding="utf-8")
    (mat / "营业执照副本.md").unlink()
    r2 = run_script("input-preflight-check", "preflight.py",
                    ["checklist", "--checklist", workspace / "c.json", "--dir", mat,
                     "--report", workspace / "r2.json"])
    rep2 = json.loads((workspace / "r2.json").read_text(encoding="utf-8"))
    states2 = {i["id"]: i["state"] for i in rep2["items"]}
    assert states2["M1"] == "needs_confirm"

    # data 模式: 缺列阻断 + 日期越界
    wb = Workbook(); ws = wb.active
    ws.append(["日期", "销售额(元)", "地区"])
    ws.append(["2026-01-05", 100, "华东"])
    ws.append(["2026-02-10", 200, "华北"])
    ws.append(["2027-01-01", 50, "华南"])  # 越界
    wb.save(str(workspace / "d.xlsx"))
    req = {"columns": ["日期", "销售额(元)", "地区", "利润"],
           "units": {"销售额(元)": "元"}, "date_column": "日期",
           "date_range": ["2026-01-01", "2026-12-31"], "min_rows": 1}
    (workspace / "req.json").write_text(json.dumps(req, ensure_ascii=False), encoding="utf-8")
    r3 = run_script("input-preflight-check", "preflight.py",
                    ["data", "--data", workspace / "d.xlsx", "--requirements", workspace / "req.json",
                     "--report", workspace / "r3.json"])
    rep3 = json.loads((workspace / "r3.json").read_text(encoding="utf-8"))
    assert rep3["verdict"] == "block"
    assert "利润" in str(rep3["reasons"])  # 缺列
    assert any(c["check"] == "date_range" and not c["ok"] for c in rep3["checks"])


# ---------- tax-compliance ----------

def test_tax(workspace):
    mat = workspace / "mats"
    mat.mkdir()
    wb = Workbook(); ws = wb.active
    ws.append(["发票号", "金额", "税额"])
    ws.append(["INV001", 10000, 1300])
    ws.append(["INV002", 20000, 2600])
    wb.save(str(mat / "发票台账.xlsx"))
    doc = Document()
    doc.add_paragraph("技术服务合同：甲方委托乙方开发数据平台，合同金额50万元。")
    doc.save(str(mat / "合同-技术服务.docx"))
    (mat / "流水.csv").write_text("日期,金额\n2026-07-05,30000\n2026-08-01,12000\n", encoding="utf-8")
    biz = {"taxpayer_type": "一般纳税人", "activities": ["货物销售", "跨境服务"], "period": "2026Q3"}
    (workspace / "b.json").write_text(json.dumps(biz, ensure_ascii=False), encoding="utf-8")
    r = run_script("tax-compliance", "tax_review.py",
                   ["--materials", mat, "--business", workspace / "b.json",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "out" / "rule-analysis.json").read_text(encoding="utf-8"))
    # 材料归集
    assert rep["materials_summary"]["contracts"] == 1
    assert rep["materials_summary"]["invoice_files"] == 1
    assert rep["materials_summary"]["flow_files"] == 1
    # 规则逐条有来源
    rules = {x["id"]: x for x in rep["rules"]}
    assert "VAT-G-13" in rules  # 一般纳税人+货物销售
    assert "VAT-S-3" not in rules  # 非小规模
    assert "VAT-CROSS-FREE" in rules  # 跨境服务
    for x in rep["rules"]:
        assert x["source"] and "条例" in x["source"] or "法" in x["source"] or "号" in x["source"] or "办法" in x["source"] or "细则" in x["source"]
        assert x["scope"]
    # 有合同材料 -> 跨境规则 applicable；缺项不存在
    assert rules["VAT-CROSS-FREE"]["status"] == "applicable"
    # 免责声明
    assert "不构成专业税务意见" in rep["disclaimer"]
    md = (workspace / "out" / "rule-analysis.md").read_text(encoding="utf-8")
    assert "不构成专业税务意见" in md
    # 小规模分支
    biz2 = {"taxpayer_type": "小规模纳税人", "activities": ["服务"], "period": "2026Q3"}
    (workspace / "b2.json").write_text(json.dumps(biz2, ensure_ascii=False), encoding="utf-8")
    r2 = run_script("tax-compliance", "tax_review.py",
                    ["--materials", mat, "--business", workspace / "b2.json",
                     "--outdir", workspace / "out2"])
    rep2 = json.loads((workspace / "out2" / "rule-analysis.json").read_text(encoding="utf-8"))
    ids2 = {x["id"] for x in rep2["rules"]}
    assert "VAT-S-3" in ids2 and "VAT-G-13" not in ids2
    # 无合同场景 -> 跨境 needs_confirmation
    mat2 = workspace / "mats2"
    mat2.mkdir()
    (mat2 / "发票台账.xlsx").write_bytes((mat / "发票台账.xlsx").read_bytes())
    r3 = run_script("tax-compliance", "tax_review.py",
                    ["--materials", mat2, "--business", workspace / "b.json",
                     "--outdir", workspace / "out3"])
    rep3 = json.loads((workspace / "out3" / "rule-analysis.json").read_text(encoding="utf-8"))
    rules3 = {x["id"]: x for x in rep3["rules"]}
    assert rules3["VAT-CROSS-FREE"]["status"] == "needs_confirmation"
    conf = json.loads((workspace / "out3" / "confirmations.json").read_text(encoding="utf-8"))
    assert any(c["rule_id"] == "VAT-CROSS-FREE" for c in conf["confirmations"])


# ---------- proposal-bid-generation ----------

def test_bid(workspace):
    tender = """一、供应商须具备有效的质量管理体系认证（ISO9001）。
二、供应商近三年须具有智慧园区类项目业绩。
三、项目团队须配备具有PMP证书的项目经理。
四、投标人注册资本不低于1000万元。
五、须提供本地化驻场服务承诺。
六、投标报价不得超过预算500万元。
七、须提供数据安全等级保护证明。
八、中标后30日内完成系统部署。
"""
    (workspace / "tender.md").write_text(tender, encoding="utf-8")
    quals = workspace / "quals"
    quals.mkdir()
    (quals / "资质-ISO9001质量管理体系认证.docx")
    doc = Document(); doc.add_paragraph("质量管理体系认证证书 ISO9001，证书编号有效。")
    doc.save(str(quals / "资质-ISO9001质量管理体系认证.docx"))
    (quals / "业绩-智慧园区项目合同.md").write_text(
        "智慧园区一期项目合同，金额800万元，2024年验收。", encoding="utf-8")
    (quals / "团队-项目经理PMP名单.xlsx")
    wb = Workbook(); ws = wb.active
    ws.append(["姓名", "证书"])
    ws.append(["张伟", "PMP"])
    wb.save(str(quals / "团队-项目经理PMP名单.xlsx"))
    r = run_script("proposal-bid-generation", "bid_matrix.py",
                   ["--tender", workspace / "tender.md", "--quals", quals,
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "out" / "report.json").read_text(encoding="utf-8"))
    assert rep["clauses"] == 8  # 8条款全解析
    assert rep["responded"] >= 2  # ISO9001+智慧园区业绩 (+团队PMP)
    gaps = json.loads((workspace / "out" / "gaps.json").read_text(encoding="utf-8"))
    # 缺口条款不编造
    for g in gaps["gaps"]:
        assert "注册资本" in g["text"] or "驻场" in g["text"] or "预算" in g["text"] \
            or "部署" in g["text"] or "等保" in g["text"] or "数据安全" in g["text"]
    # 证据文件真实存在
    wbm = load_workbook(str(workspace / "out" / "response-matrix.xlsx"))
    wsm = wbm["响应矩阵"]
    for row in wbm["响应矩阵"].iter_rows(min_row=2, values_only=True):
        if row[2] == "responded":
            for ev in str(row[3]).split("；"):
                assert (quals / ev).exists(), ev
    doc = Document(str(workspace / "out" / "bid-draft.docx"))
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "【待补充-缺口】" in full
    assert "ISO9001" in full


# ---------- hr-material-analysis ----------

def test_hr(workspace):
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发工资", "实发工资", "身份证号"])
    ws.append(["E001", "张伟", "研发部", 20000, 18000, "110105198806051210"])
    ws.append(["E002", "李娜", "市场部", 15000, 13500, "110105199001012345"])
    ws.append(["E003", "王强", "研发部", 18000, 16000, "110105198505056789"])
    ws.append(["E003", "王强", "研发部", 18000, 16000, "110105198505056789"])  # 重复工号
    ws.append(["合计", "", "", 53000, 47500, ""])
    wb.save(str(workspace / "payroll.xlsx"))
    wb2 = Workbook(); ws2 = wb2.active
    ws2.append(["工号", "姓名", "应出勤天数", "实出勤天数"])
    ws2.append(["E001", "张伟", 22, 20])   # 缺勤2
    ws2.append(["E002", "李娜", 22, 22])
    # E003 无考勤 -> 缺考勤记录
    wb2.save(str(workspace / "attendance.xlsx"))
    wb3 = Workbook(); ws3 = wb3.active
    ws3.append(["工号", "姓名", "部门"])
    ws3.append(["E001", "张伟", "研发部"])
    ws3.append(["E002", "李娜", "市场部"])  # 花名册无E003
    wb3.save(str(workspace / "roster.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--attendance", workspace / "attendance.xlsx",
                    "--roster", workspace / "roster.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    # 3个唯一工号; 缺勤合计=2; 异常: E003缺考勤(或重复) -> >=1
    assert rep["people"] == 3
    assert rep["total_absence"] == 2.0
    assert rep["duplicate_groups"] == 1
    # 独立验证: 比对表
    wbm = load_workbook(str(workspace / "out" / "compare-report.xlsx"))
    rows = list(wbm["人员比对"].iter_rows(min_row=2, values_only=True))
    by_emp = {row[0]: row for row in rows}
    assert by_emp["E001"][6] == 2  # 缺勤天数
    assert "缺考勤记录" in by_emp["E003"][9]
    # 汇总: 研发2人(应发38000), 市场1人(15000), 未分配0 —— 合计行剔除
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    srows = {row[0]: row for row in wbs["部门汇总"].iter_rows(min_row=2, values_only=True)}
    assert srows["研发部"][1] == 2 and srows["研发部"][2] == 38000
    assert srows["市场部"][2] == 15000
    # 敏感字段脱敏: 输出无明文身份证
    import re as _re
    all_out_text = json.dumps([list(r) for r in rows], ensure_ascii=False)
    assert not _re.search(r"\d{17}[\dXx]", all_out_text)
    mask = json.loads((workspace / "out" / "mask-manifest.json").read_text(encoding="utf-8"))
    assert mask["masked"] is True
    # 简报docx存在且含计算值
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "3 人" in text and "2.0" in text


# ---------- online-office-delivery ----------

def test_online_delivery(workspace):
    plan = {"target": {"type": "lark_sheet", "position": "共享"},
            "files": [{"local": "m.xlsx", "name": "表"}], "readback": True}
    (workspace / "plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    wb = Workbook(); wb.active.append(["a"])
    wb.save(str(workspace / "m.xlsx"))
    r = run_script("online-office-delivery", "online_delivery.py",
                   ["--plan", workspace / "plan.json", "--outdir", workspace / "out"])
    assert r.returncode == 3  # 无授权环境 -> blocked
    rep = json.loads((workspace / "out" / "delivery-result.json").read_text(encoding="utf-8"))
    assert rep["status"] == "blocked"
    assert rep["probe"]["clis_found"] == []
    assert "lark-cli" in rep["missing"]["clis"]
    assert "不伪造" in rep["note"] or "冒充" in rep["note"]
    # 损坏计划
    (workspace / "bad.json").write_text("{", encoding="utf-8")
    r2 = run_script("online-office-delivery", "online_delivery.py",
                    ["--plan", workspace / "bad.json", "--outdir", workspace / "o2"])
    assert r2.returncode == 2
