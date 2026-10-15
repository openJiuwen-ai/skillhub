"""hr-material-analysis 黑盒测试：金标准手工推导。

正常（三表金标准复算/脱敏）、复杂（工号冲突不静默挑行、不可解析/非有限/无缓存
公式不按0计入、姓名不一致/缺记录标注、一致重复去重、仅薪资表不做缺勤结论、
考勤缺列不可计算）、异常（缺工号列报错、损坏文件、未提供任何表）。
冻结探针对应：benchmarks/hr-complex/before-review/（duplicate/bad-header）。
"""
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

ID_RE = re.compile(r"\d{17}[\dXx]")


def read_compare(outdir):
    wb = load_workbook(os.path.join(outdir, "compare-report.xlsx"))
    rows = list(wb["人员比对"].iter_rows(values_only=True))
    return {r[0]: r for r in rows[1:]}, wb


def test_normal_three_tables_golden(workspace):
    """三表金标准：3人/缺勤2+0.5=2.5/部门以花名册为准/合计行剔除/敏感脱敏。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发", "身份证号"])
    ws.append(["E001", "张伟", "研发部", 20000, 18000, "110105198806051210"])
    ws.append(["E002", "李娜", "市场部", 15000, 13500, "110105199001012345"])
    ws.append(["E004", "赵敏", "研发部", 18000, 16000, "110105198505056789"])
    ws.append(["合计", "", "", 53000, 47500, ""])
    wb.save(str(workspace / "payroll.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "应出勤", "实出勤"])
    ws.append(["E001", "张伟", 22, 20])
    ws.append(["E002", "李娜", 22, 22])
    ws.append(["E004", "赵敏", 22, 21.5])
    wb.save(str(workspace / "attendance.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门"])
    ws.append(["E001", "张伟", "研发部"])
    ws.append(["E002", "李娜", "市场部"])
    ws.append(["E004", "赵敏", "销售部"])  # 花名册部门优先于薪资表
    wb.save(str(workspace / "roster.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--attendance", workspace / "attendance.xlsx",
                    "--roster", workspace / "roster.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["people"] == 3 and rep["people_in_summary"] == 3
    assert rep["anomalies"] == 0 and rep["conflict_groups"] == 0
    assert rep["total_absence"] == 2.5  # 2 + 0 + 0.5，独立复算
    assert rep["pay_sum"] == 53000.0
    by_emp, _ = read_compare(workspace / "out")
    # 独立复算：缺勤 = 应出勤 - 实出勤（手工金标准）
    assert by_emp["E001"][6] == 2 and by_emp["E004"][6] == 0.5
    assert by_emp["E002"][6] == 0
    for row in by_emp.values():
        if isinstance(row[4], (int, float)) and isinstance(row[5], (int, float)):
            assert row[6] == round(row[4] - row[5], 2)
    assert by_emp["E004"][3] == "销售部"  # 花名册部门优先
    # 部门汇总：合计行剔除，人数/应发独立核对
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    srows = {row[0]: row for row in wbs["部门汇总"].iter_rows(min_row=2, values_only=True)}
    assert set(srows) == {"研发部", "市场部", "销售部"}
    assert srows["研发部"][1] == 1 and srows["研发部"][2] == 20000
    assert srows["市场部"][2] == 15000 and srows["销售部"][2] == 18000
    # 敏感脱敏：输出无明文身份证，脱敏列为 ■
    all_text = json.dumps([list(v) for v in by_emp.values()], ensure_ascii=False)
    assert not ID_RE.search(all_text)
    hdrs = list(load_workbook(str(workspace / "out" / "compare-report.xlsx"))
                ["人员比对"].iter_rows(max_row=1, values_only=True))[0]
    assert "身份证号(脱敏)" in hdrs
    assert by_emp["E001"][hdrs.index("身份证号(脱敏)")] == "■" * 12
    for f in ["hr-report.json", "mask-manifest.json"]:
        assert not ID_RE.search((workspace / "out" / f).read_text(encoding="utf-8"))
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "3 人" in text and "2.50" in text
    assert "未发现表内工号重复" in text
    assert not ID_RE.search(text)
    repj = json.loads((workspace / "out" / "hr-report.json").read_text(encoding="utf-8"))
    assert repj["attendance_provided"] is True
    assert repj["total_absence"] == 2.5 and repj["conflicts"] == []
    assert len(repj["skipped_total_rows"]) == 1  # 合计行留痕
    assert repj["skipped_total_rows"][0]["row"] == 5
    assert repj["dropped_rows"] == []


def test_duplicate_conflict_not_silently_picked(workspace):
    """冻结探针 duplicate：同工号两行 100/900 不得静默挑最后一行进比对与汇总。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "虚构甲", "测试部", 100, 90])
    ws.append(["E001", "虚构甲", "测试部", 900, 800])
    wb.save(str(workspace / "duplicate.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "duplicate.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["people"] == 1 and rep["people_in_summary"] == 0
    assert rep["people_excluded_conflict"] == 1
    assert rep["conflict_groups"] == 1 and rep["duplicate_groups"] == 1
    assert rep["total_absence"] is None and rep["attendance_provided"] is False
    by_emp, _ = read_compare(workspace / "out")
    assert set(by_emp) == {"E001"}
    assert by_emp["E001"][7] is None  # 应发不挑行（既非900也非100）
    assert by_emp["E001"][8] is None
    assert by_emp["E001"][9] == "工号冲突"
    assert "缺考勤" not in (by_emp["E001"][9] or "")  # 未提供考勤表不得硬称缺考勤
    detail = by_emp["E001"][11]
    assert "行2、3" in detail and "100/900" in detail and "90/800" in detail
    # 重复工号 sheet：来源文件+行号+字段值
    wbm = load_workbook(str(workspace / "out" / "compare-report.xlsx"))
    drows = list(wbm["重复工号"].iter_rows(min_row=2, values_only=True))
    assert len(drows) == 1
    assert drows[0][0] == "duplicate.xlsx" and drows[0][1] == "薪资表"
    assert drows[0][2] == "E001" and drows[0][3] == "2、3"
    assert "内容冲突(未计入汇总)" in drows[0][4]
    assert "应发=100/900" in drows[0][5]
    # 汇总排除冲突人员：部门汇总无数据行
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    assert len(list(wbs["部门汇总"].iter_rows(values_only=True))) == 1
    note = list(wbs["汇总说明"].iter_rows(values_only=True))
    assert any("1 人因工号冲突未计入" in str(n) for n in note)
    # 简报：声明排除范围；未提供考勤表不得写缺勤结论
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "未计入汇总" in text and "未提供考勤表" in text
    assert not re.search(r"缺勤\s*[\d.]+\s*天", text)
    repj = json.loads((workspace / "out" / "hr-report.json").read_text(encoding="utf-8"))
    c = repj["conflicts"][0]
    assert c["rows"] == [2, 3] and c["field_values"]["应发"] == [100, 900]
    assert c["field_values"]["实发"] == [90, 800]
    assert repj["total_absence"] is None and repj["pay_sum"] == 0.0
    assert repj["people_absence_checked"] == 0


def test_unparseable_nonfinite_uncached_not_counted_as_zero(workspace):
    """金额/天数为空、不可解析、非有限数、无缓存公式：显式数据质量，不按0计入。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "甲", "一部", 100, 90])
    ws.append(["E002", "乙", "一部", "N/A", 90])      # 应发不可解析
    ws.append(["E003", "丙", "一部", "NaN", 90])       # 应发非有限数
    ws.append(["E004", "丁", "一部", "=B2*2", 90])     # 应发公式无缓存值
    ws.append(["E005", "戊", "一部", "1,000", 900])    # 千分位可解析 → 1000
    wb.save(str(workspace / "payroll.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "应出勤", "实出勤"])
    ws.append(["E001", "甲", 22, 20])
    ws.append(["E002", "乙", "abc", 20])               # 应出勤不可解析
    ws.append(["E003", "丙", 22, 22])
    ws.append(["E004", "丁", 22, "inf"])               # 实出勤非有限数
    ws.append(["E005", "戊", 22, 20])
    wb.save(str(workspace / "attendance.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--attendance", workspace / "attendance.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["people"] == 5 and rep["anomalies"] == 0
    # 缺勤仅可核查 3 人（E001/E003/E005）：2 + 0 + 2 = 4.0
    assert rep["total_absence"] == 4.0
    # 应发合计仅计入 E001(100) + E005(1000) = 1100，其余不按0静默计入
    assert rep["pay_sum"] == 1100.0
    assert rep["data_quality_issues"] == 5
    by_emp, _ = read_compare(workspace / "out")
    assert by_emp["E001"][7] == 100 and by_emp["E001"][6] == 2
    assert by_emp["E002"][7] == "N/A"        # 原样可见，不消失
    assert by_emp["E002"][6] is None
    assert by_emp["E003"][7] == "非有限数(NaN)"
    assert by_emp["E004"][7] is None         # 无缓存值无内容可显示
    assert by_emp["E005"][7] == 1000 and by_emp["E005"][6] == 2
    dq2 = by_emp["E002"][10]
    assert "应发不可解析" in dq2 and "应出勤不可解析" in dq2
    dq4 = by_emp["E004"][10]
    assert "应发公式无缓存值" in dq4 and "实出勤非有限数" in dq4
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    srows = {row[0]: row for row in wbs["部门汇总"].iter_rows(min_row=2, values_only=True)}
    assert srows["一部"][1] == 5 and srows["一部"][2] == 1100
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "数据质量问题 5 项" in text
    repj = json.loads((workspace / "out" / "hr-report.json").read_text(encoding="utf-8"))
    assert repj["data_quality_issue_count"] == 5
    assert repj["people_absence_checked"] == 3
    d = {x["工号"]: x for x in repj["people_detail"]}
    assert any("应发不可解析" in s for s in d["E002"]["data_quality"])
    assert d["E002"]["absence"] is None and d["E001"]["absence"] == 2


def test_name_mismatch_and_missing_records(workspace):
    """姓名跨表不一致/缺考勤/缺薪资显式标注（提供考勤表时）。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "张伟", "一部", 100, 90])
    ws.append(["E002", "李娜", "二部", 200, 180])
    wb.save(str(workspace / "payroll.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "应出勤", "实出勤"])
    ws.append(["E001", "张伟三", 22, 20])   # 姓名不一致
    ws.append(["E003", "王强", 22, 22])     # 缺薪资记录
    wb.save(str(workspace / "attendance.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门"])
    ws.append(["E001", "张伟", "一部"])
    ws.append(["E002", "李娜", "二部"])
    wb.save(str(workspace / "roster.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--attendance", workspace / "attendance.xlsx",
                    "--roster", workspace / "roster.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["people"] == 3 and rep["anomalies"] == 3
    assert rep["total_absence"] == 2.0
    by_emp, _ = read_compare(workspace / "out")
    assert "姓名不一致" in by_emp["E001"][9]
    assert "缺考勤记录" in by_emp["E002"][9]
    assert "缺薪资记录" in by_emp["E003"][9]
    repj = json.loads((workspace / "out" / "hr-report.json").read_text(encoding="utf-8"))
    d = {x["工号"]: x for x in repj["people_detail"]}
    assert d["E001"]["names"] == {"payroll": "张伟", "attendance": "张伟三",
                                  "roster": "张伟"}
    assert by_emp["E001"][1] == "张伟" and by_emp["E001"][2] == "张伟三"


def test_consistent_duplicate_dedup_listed(workspace):
    """同工号内容完全一致的重复行：去重计入一次并单列，不重复计薪。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "甲", "一部", 100, 90])
    ws.append(["E001", "甲", "一部", 100, 90])
    ws.append(["E002", "乙", "二部", 500, 450])
    wb.save(str(workspace / "payroll.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["people"] == 2 and rep["people_in_summary"] == 2
    assert rep["duplicate_groups"] == 1 and rep["conflict_groups"] == 0
    assert rep["pay_sum"] == 600.0  # E001 只计一次 100，不重复计 200
    by_emp, _ = read_compare(workspace / "out")
    assert by_emp["E001"][7] == 100 and by_emp["E001"][9] == "工号重复(一致)"
    wbm = load_workbook(str(workspace / "out" / "compare-report.xlsx"))
    drows = list(wbm["重复工号"].iter_rows(min_row=2, values_only=True))
    assert drows[0][4] == "内容一致(已去重计入)"
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    srows = {row[0]: row for row in wbs["部门汇总"].iter_rows(min_row=2, values_only=True)}
    assert srows["一部"][1] == 1 and srows["一部"][2] == 100
    doc = Document(str(workspace / "out" / "draft.docx"))
    assert "已去重计入" in "\n".join(p.text for p in doc.paragraphs)


def test_payroll_only_no_attendance_conclusions(workspace):
    """仅提供薪资表：不硬称缺考勤、无考勤不得写已核查缺勤0天。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "张伟", "研发部", 100, 90])
    wb.save(str(workspace / "payroll.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["attendance_provided"] is False
    assert rep["total_absence"] is None
    assert rep["people_in_summary"] == 1 and rep["pay_sum"] == 100.0
    by_emp, _ = read_compare(workspace / "out")
    assert by_emp["E001"][9] is None  # 无缺考勤记录等任何异常
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "未提供考勤表" in text and "未核查" in text and "满勤" in text
    assert not re.search(r"缺勤\s*[\d.]+\s*天", text)
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    note = list(wbs["汇总说明"].iter_rows(values_only=True))
    assert any("未核查" in str(n) for n in note)


def test_attendance_without_required_columns(workspace):
    """考勤表缺应出勤/实出勤列：缺勤不可计算，不得写0天。"""
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    ws.append(["E001", "张伟", "一部", 100, 90])
    wb.save(str(workspace / "payroll.xlsx"))
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "打卡天数"])
    ws.append(["E001", "张伟", 20])
    wb.save(str(workspace / "attendance.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "payroll.xlsx",
                    "--attendance", workspace / "attendance.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    repj = json.loads((workspace / "out" / "hr-report.json").read_text(encoding="utf-8"))
    assert repj["attendance_provided"] is True
    assert repj["absence_computable"] is False
    assert repj["total_absence"] is None and repj["people_absence_checked"] == 0
    by_emp, _ = read_compare(workspace / "out")
    assert by_emp["E001"][6] is None and by_emp["E001"][9] is None
    doc = Document(str(workspace / "out" / "draft.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "缺少应出勤/实出勤列" in text and "不可计算" in text
    wbs = load_workbook(str(workspace / "out" / "roster-summary.xlsx"))
    note = list(wbs["汇总说明"].iter_rows(values_only=True))
    assert any("不可计算" in str(n) for n in note)


def test_missing_id_column_error(workspace):
    """冻结探针 bad-header：缺工号表头必须报错，不得0人静默成功。"""
    wb = Workbook(); ws = wb.active
    ws.append(["姓名", "部门", "应发"])
    ws.append(["虚构甲", "测试部", 100])
    wb.save(str(workspace / "bad-header.xlsx"))
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "bad-header.xlsx",
                    "--outdir", workspace / "out"])
    assert r.returncode == 2
    assert "工号" in r.stderr and "0人" in r.stderr.replace(" ", "")
    assert not (workspace / "out" / "compare-report.xlsx").exists()
    assert not (workspace / "out" / "draft.docx").exists()


def test_corrupt_file_and_no_args(workspace):
    """异常：损坏文件/未提供任何表报错；空数据表（表头合法）为真实0人。"""
    (workspace / "bad.xlsx").write_bytes(b"this is not an xlsx file")
    r = run_script("hr-material-analysis", "hr_compare.py",
                   ["--payroll", workspace / "bad.xlsx",
                    "--outdir", workspace / "o1"])
    assert r.returncode == 2 and "ERROR" in r.stderr
    r2 = run_script("hr-material-analysis", "hr_compare.py",
                    ["--outdir", workspace / "o2"])
    assert r2.returncode == 2 and "ERROR" in r2.stderr
    wb = Workbook(); ws = wb.active
    ws.append(["工号", "姓名", "部门", "应发", "实发"])
    wb.save(str(workspace / "empty.xlsx"))
    r3 = run_script("hr-material-analysis", "hr_compare.py",
                    ["--payroll", workspace / "empty.xlsx",
                     "--outdir", workspace / "o3"])
    assert r3.returncode == 0, r3.stderr + r3.stdout
    rep = json.loads(r3.stdout)
    assert rep["people"] == 0 and rep["anomalies"] == 0
    repj = json.loads((workspace / "o3" / "hr-report.json").read_text(encoding="utf-8"))
    assert repj["people_total"] == 0 and repj["dropped_rows"] == []
