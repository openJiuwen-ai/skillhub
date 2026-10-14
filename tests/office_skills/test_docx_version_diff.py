"""docx-version-diff 黑盒测试：改字/表格单元格/分节/相同文件四类场景。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402


def build_v1(path):
    doc = Document()
    doc.add_paragraph("第一段：背景说明")
    doc.add_paragraph("第二段：预算为一百万元")
    doc.add_paragraph("第三段：实施安排")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "项目"
    t.cell(0, 1).text = "金额"
    t.cell(1, 0).text = "咨询费"
    t.cell(1, 1).text = "300000"
    doc.sections[0].header.paragraphs[0].text = "V1 页眉"
    doc.save(path)


def test_modified_one_word(workspace):
    build_v1(str(workspace / "v1.docx"))
    doc = Document(str(workspace / "v1.docx"))
    doc.paragraphs[1].text = "第二段：预算为二百万元"  # 只改一处
    doc.save(str(workspace / "v2.docx"))
    r = run_script("docx-version-diff", "docx_diff.py",
                   ["--old", workspace / "v1.docx", "--new", workspace / "v2.docx",
                    "--report", workspace / "d.json"])
    assert r.returncode == 0, r.stderr
    rep = json.loads((workspace / "d.json").read_text(encoding="utf-8"))
    assert rep["summary"]["paragraph_modified"] == 1
    assert rep["summary"]["paragraph_added"] == 0
    assert rep["summary"]["paragraph_removed"] == 0
    ch = rep["changes"][0]
    assert ch["type"] == "paragraph" and ch["change"] == "modified"
    assert "一百万元" in ch["before"] and "二百万元" in ch["after"]


def test_table_cell_change(workspace):
    build_v1(str(workspace / "v1.docx"))
    doc = Document(str(workspace / "v1.docx"))
    doc.tables[0].cell(1, 1).text = "350000"
    doc.save(str(workspace / "v2.docx"))
    r = run_script("docx-version-diff", "docx_diff.py",
                   ["--old", workspace / "v1.docx", "--new", workspace / "v2.docx",
                    "--report", workspace / "d.json"])
    assert r.returncode == 0
    rep = json.loads((workspace / "d.json").read_text(encoding="utf-8"))
    assert rep["summary"]["table_cell_modified"] == 1
    ch = [c for c in rep["changes"] if c["type"] == "table_cell"][0]
    assert ch["location"] == "table 0 row 1 col 1"
    assert ch["before"] == "300000" and ch["after"] == "350000"


def test_new_section_and_header(workspace):
    build_v1(str(workspace / "v1.docx"))
    doc = Document(str(workspace / "v1.docx"))
    doc.add_section()
    doc.sections[0].header.paragraphs[0].text = "V2 页眉"
    doc.save(str(workspace / "v2.docx"))
    r = run_script("docx-version-diff", "docx_diff.py",
                   ["--old", workspace / "v1.docx", "--new", workspace / "v2.docx",
                    "--report", workspace / "d.json"])
    assert r.returncode == 0
    rep = json.loads((workspace / "d.json").read_text(encoding="utf-8"))
    assert rep["summary"]["section_changed"] >= 1  # 节数变化
    assert rep["summary"]["header_footer_changed"] >= 1  # 页眉变化
    assert any(c["type"] == "section" for c in rep["changes"])
    assert any(c["type"] == "header" and "V2" in (c["after"] or "") for c in rep["changes"])


def test_identical_and_corrupt(workspace):
    build_v1(str(workspace / "v1.docx"))
    import shutil
    shutil.copy(str(workspace / "v1.docx"), str(workspace / "v1copy.docx"))
    r = run_script("docx-version-diff", "docx_diff.py",
                   ["--old", workspace / "v1.docx", "--new", workspace / "v1copy.docx",
                    "--report", workspace / "d.json"])
    assert r.returncode == 0
    rep = json.loads((workspace / "d.json").read_text(encoding="utf-8"))
    assert rep["identical"] is True and rep["changes"] == []
    # 损坏输入
    (workspace / "bad.docx").write_bytes(b"junk")
    r2 = run_script("docx-version-diff", "docx_diff.py",
                    ["--old", workspace / "bad.docx", "--new", workspace / "v1.docx",
                     "--report", workspace / "d2.json"])
    assert r2.returncode == 2
