"""meeting-minutes-actions 标准黑盒测试（正常/复杂/异常）。

独立验证本技能真实行为：转写稿+上期待办 -> 纪要docx + 待办追踪xlsx + asr-issues.json。
复杂场景数据取自冻结基准 benchmarks/meeting-complex（prompt.txt/expected.json，
2026-10-14 冻结），逐条核对 expected.json 的 9 项检查。
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

SKILL = "meeting-minutes-actions"
SCRIPT = "minutes_actions.py"


def run(ws, transcript, prev=None, out="out"):
    args = ["--transcript", ws / transcript, "--outdir", ws / out]
    if prev is not None:
        args[2:2] = ["--prev-actions", ws / prev]
    return run_script(SKILL, SCRIPT, args)


def write_prev(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["编号", "责任人", "任务", "截止", "状态"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def load_rows(path):
    ws = load_workbook(str(path))["待办追踪"]
    return [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)]


def docx_table(path):
    doc = Document(str(path))
    return [[c.text for c in r.cells] for r in doc.tables[0].rows]


def paragraphs_before_first_table(doc):
    out = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            break
        if child.tag == qn("w:p"):
            out.append(Paragraph(child, doc).text)
    return out


# ---------------------------------------------------------------- 正常场景

NORMAL_TRANSCRIPT = """[14:00] 主持人：先过一下上周待办。
[14:02] 王敏：客户名单核对已完成，报表还没发。
[14:03] 王敏：会议确认下周三发布。
[14:05] 李强：新官网文案由李强牵头，下周五前提交初稿。
[14:08] 主持人：周报请大家本周五前提交。
[14:10] 赵蕾：发票整理我来负责，尽快完成。
[14:12] 主持人：A03库存盘点这周没有讨论，维持原记录。
"""

NORMAL_PREV = [
    ("A01", "王敏", "客户名单核对", "2026-10-10", "进行中"),
    ("A02", "李强", "供应商对账", "2026-10-11", "受阻"),
    ("A03", "赵蕾", "库存盘点", "2026-10-12", "暂缓"),
]


def test_normal_two_period_carry_over(workspace):
    (workspace / "t.txt").write_text(NORMAL_TRANSCRIPT, encoding="utf-8")
    write_prev(str(workspace / "prev.xlsx"), NORMAL_PREV)
    r = run(workspace, "t.txt", "prev.xlsx")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["n_actions"] == 6 and rep["carried"] == 3 and rep["new"] == 3
    assert rep["decisions"] == 1 and rep["conflicts"] == 0

    rows = load_rows(workspace / "out" / "actions.xlsx")
    by_id = {r[0]: r for r in rows[1:]}
    # A01 有效提及且明确说完成 -> 已完成，保留来源与"报表还没发"的原始表述
    assert by_id["A01"][4] == "已完成" and by_id["A01"][7] == "是"
    assert by_id["A01"][6] == "L2@14:02" and "报表还没发" in by_id["A01"][8]
    # A02/A03 未提及 -> 状态原样保留，不得推断改写
    assert by_id["A02"][4] == "受阻" and by_id["A02"][7] == "否"
    assert by_id["A03"][4] == "暂缓" and by_id["A03"][7] == "否"
    assert by_id["A03"][6] == "L7@14:12"  # 明确未讨论语句可作来源证据
    # 新增项编号顺延（A01-A03 -> A04 起）
    assert [r[0] for r in rows[1:]] == ["A01", "A02", "A03", "A04", "A05", "A06"]
    # 责任人：句中模式(李强牵头)、自述承诺(我来负责->赵蕾)、
    # 主持人安排但无人认领 -> 待确认（不能拿主持人代填）
    assert by_id["A04"][1] == "李强" and by_id["A04"][3] == "下周五"
    assert by_id["A05"][1] == "待确认" and by_id["A05"][3] == "本周五"
    assert by_id["A06"][1] == "赵蕾" and by_id["A06"][3] == "尽快"
    # 历史完成汇报句（14:02）不重复识别为新增任务
    assert not any("客户名单核对已完成" in r[2] for r in rows[1:] if r[5] == "本期新增")
    # docx 行动项表与 xlsx 完全一致（含来源时间戳列）
    assert docx_table(workspace / "out" / "minutes.docx") == rows
    # 两次运行输出一致（确定性）
    r2 = run(workspace, "t.txt", "prev.xlsx", out="out2")
    assert r2.returncode == 0
    assert load_rows(workspace / "out2" / "actions.xlsx") == rows
    assert docx_table(workspace / "out2" / "minutes.docx") == rows


# ---------------------------------------------------------------- 复杂场景（冻结基准）

# 数据与 benchmarks/meeting-complex/prompt.txt 一致（发送前冻结的验收输入）
BENCH_TRANSCRIPT = """[09:00] 主持人：今天只确认培训反馈和本周材料。
[09:01] 陈青：A02整理培训反馈已完成，附件尚未提供。
[09:02] 林舟：我负责在2026年10月16日前提交新版供应商询价表。这不是A01核对供应商清单。
[09:03] 赵宁：预算汇总表写12万元，口头汇报是15万元；暂未核实，不要定稿。
[09:04] 主持人：安排一个人核对预算，责任人与截止时间下次再定。
[09:05] 林舟：演示材料我会提交，但截止日期还没定。
[09:06] 主持人：本期没有讨论A01、A03、A04，维持原记录并注明未提及。
"""

BENCH_PREV = [
    ("A01", "林舟", "核对供应商清单", "2026-10-13", "受阻"),
    ("A02", "陈青", "整理培训反馈", "2026-10-13", "进行中"),
    ("A03", "赵宁", "迁移历史资料", "2026-10-20", "暂缓"),
    ("A04", "林舟", "旧系统采购", "2026-10-12", "已取消"),
]


def test_complex_benchmark_frozen_checks(workspace):
    (workspace / "t.txt").write_text(BENCH_TRANSCRIPT, encoding="utf-8")
    write_prev(str(workspace / "prev.xlsx"), BENCH_PREV)
    r = run(workspace, "t.txt", "prev.xlsx")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["n_actions"] == 7 and rep["carried"] == 4 and rep["new"] == 3

    rows = load_rows(workspace / "out" / "actions.xlsx")
    by_id = {r[0]: r for r in rows[1:]}
    assert [r[0] for r in rows[1:]] == ["A01", "A02", "A03", "A04", "A05", "A06", "A07"]

    # 检查1: A01 保持受阻、未被提及，不与新增询价表混同
    a01 = by_id["A01"]
    assert a01[4] == "受阻" and a01[7] == "否"
    assert a01[2] == "核对供应商清单" and a01[1] == "林舟"  # 原记录未被改写
    a05 = by_id["A05"]
    assert a05[5] == "本期新增" and "询价表" in a05[2]
    assert a05[0] != a01[0] and "这不是A01" in a05[2]  # 否定句原文保留，两项独立
    # "这不是A01"未给 A01 带来任何进度/提及
    assert a01[6] != "L3@09:02" and a01[4] != "进行中"

    # 检查2: A02 据汇报记已完成；附件情况按原话保留，不宣称已独立核验
    a02 = by_id["A02"]
    assert a02[4] == "已完成" and a02[7] == "是" and a02[6] == "L2@09:01"
    assert "附件尚未提供" in a02[8] and "未独立核验" in a02[8]

    # 检查3/4: A03 暂缓、A04 已取消，均未被提及
    assert by_id["A03"][4] == "暂缓" and by_id["A03"][7] == "否"
    assert by_id["A04"][4] == "已取消" and by_id["A04"][7] == "否"
    # 受阻/暂缓/取消不得无依据变为进行中（本期无任何行处于进行中）
    assert not any(r[4] == "进行中" for r in rows[1:])

    # 检查5: 新询价表 责任人林舟、截止2026-10-16、来源09:02
    assert a05[1] == "林舟" and "2026年10月16日" in a05[3]
    assert a05[6] == "L3@09:02"
    # 检查6: 预算核对 责任人/截止均未知、来源09:04，且不是主持人代填
    a06 = by_id["A06"]
    assert a06[1] == "待确认" and a06[3] == "待确认"
    assert a06[6] == "L5@09:04" and "核对预算" in a06[2]
    # 检查7: 演示材料 责任人林舟、截止未知、来源09:05
    a07 = by_id["A07"]
    assert a07[1] == "林舟" and a07[3] == "待确认" and a07[6] == "L6@09:05"
    assert "演示材料" in a07[2]
    # 全表没有以"主持人"充当责任人
    assert all(r[1] != "主持人" for r in rows[1:])

    # 检查8: 12万元 vs 15万元 冲突完整保留（双方数值都在、不选择、来源09:03）
    issues = json.loads((workspace / "out" / "asr-issues.json").read_text(encoding="utf-8"))
    conflict = [i for i in issues["issues"] if i["type"] == "conflict"]
    assert len(conflict) == 1
    assert set(conflict[0]["values"]) == {"12万元", "15万元"}
    assert "09:03" in conflict[0]["where"]
    assert conflict[0]["note"].find("不选择") >= 0
    nums = [i["value"] for i in issues["issues"] if i["type"] == "number"]
    assert {"12万元", "15万元"} <= set(nums)

    # 历史完成项不重复识别为新增任务（09:01 汇报句不在新增行中）
    news = [r for r in rows[1:] if r[5] == "本期新增"]
    assert len(news) == 3
    assert not any("A02整理培训反馈已完成" in r[2] for r in news)

    # 检查9: 两份文件行动项编号及内容一致（含来源时间戳列）
    table = docx_table(workspace / "out" / "minutes.docx")
    assert table == rows
    # 纪要待确认事项段包含冲突双方数值与来源
    doc = Document(str(workspace / "out" / "minutes.docx"))
    pending_text = "\n".join(p.text for p in doc.paragraphs)
    assert "12万元" in pending_text and "15万元" in pending_text and "09:03" in pending_text


def test_complex_layout_readable_and_summary_surfaces_updates(workspace):
    """版式修复独立验证（对应 native-qa/meeting-minutes-layout-review.json）：
    横向页面、按内容列宽、合适字号、跨页重复表头、摘要前置冲突/完成更新；
    同时确认字段、七个行动项及原状态未被改动。"""
    (workspace / "t.txt").write_text(BENCH_TRANSCRIPT, encoding="utf-8")
    write_prev(str(workspace / "prev.xlsx"), BENCH_PREV)
    r = run(workspace, "t.txt", "prev.xlsx")
    assert r.returncode == 0, r.stderr + r.stdout

    doc = Document(str(workspace / "out" / "minutes.docx"))
    sec = doc.sections[0]
    assert sec.orientation == WD_ORIENT.LANDSCAPE
    assert sec.page_width > sec.page_height  # 横向页面

    table = doc.tables[0]
    assert len(table.columns) == 9  # 9 列字段保留
    # 跨页重复表头（w:tblHeader）
    trPr = table.rows[0]._tr.find(qn("w:trPr"))
    assert trPr is not None and trPr.find(qn("w:tblHeader")) is not None
    # 按内容分配列宽：任务列明显宽于 编号/状态 等窄列
    task_w = table.rows[1].cells[2].width
    assert task_w is not None and task_w > table.rows[1].cells[4].width
    assert task_w > table.rows[1].cells[0].width
    # 表格字号统一为 9pt（可读且小于正文默认）
    sizes = {run.font.size.pt for row in table.rows for c in row.cells
             for p in c.paragraphs for run in p.runs if run.font.size}
    assert sizes == {9.0}

    # 摘要位于行动项表之前：完成更新与预算冲突不再只埋在末尾
    pre = "\n".join(paragraphs_before_first_table(doc))
    assert "本期进展摘要" in pre and "[A02]" in pre and "已完成" in pre
    assert "未决事项摘要" in pre
    assert "12万元" in pre and "15万元" in pre and "09:03" in pre

    # 内容未被版式改动影响：仍是 7 项，原状态保留
    rows = load_rows(workspace / "out" / "actions.xlsx")
    by_id = {r[0]: r for r in rows[1:]}
    assert len(rows) - 1 == 7
    assert by_id["A01"][4] == "受阻"
    assert by_id["A03"][4] == "暂缓"
    assert by_id["A04"][4] == "已取消"
    assert docx_table(workspace / "out" / "minutes.docx") == rows


# ---------------------------------------------------------------- 异常场景

def test_exceptional_missing_transcript_exit2(workspace):
    r = run_script(SKILL, SCRIPT,
                   ["--transcript", workspace / "nope.txt", "--outdir", workspace / "out"])
    assert r.returncode == 2
    assert "ERROR" in r.stderr
    assert not (workspace / "out" / "actions.xlsx").exists()


def test_exceptional_empty_transcript(workspace):
    (workspace / "t.txt").write_text("", encoding="utf-8")
    r = run(workspace, "t.txt")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["n_actions"] == 0 and rep["decisions"] == 0
    rows = load_rows(workspace / "out" / "actions.xlsx")
    assert len(rows) == 1  # 仅表头
    doc = Document(str(workspace / "out" / "minutes.docx"))
    paras = [p.text for p in doc.paragraphs]
    assert "（未识别到明确决议语句）" in paras and "（无待确认事项）" in paras
    assert docx_table(workspace / "out" / "minutes.docx") == rows


def test_exceptional_corrupt_prev_actions_warning(workspace):
    (workspace / "t.txt").write_text("[10:00] 张三：由张三负责提交周报。\n", encoding="utf-8")
    (workspace / "prev.xlsx").write_bytes(b"this is not an xlsx file")
    r = run(workspace, "t.txt", "prev.xlsx")
    assert r.returncode == 0, r.stderr + r.stdout
    assert "WARNING" in r.stderr
    rep = json.loads(r.stdout)
    assert rep["carried"] == 0 and rep["new"] == 1
    rows = load_rows(workspace / "out" / "actions.xlsx")
    assert rows[1][0] == "A-001"  # 无有效上期 -> 默认编号起步
    assert rows[1][1] == "张三"


def test_exceptional_no_prev_no_timestamps(workspace):
    # 无上期待办、转写无时间戳/说话人：仍可解析，按行号定位并标注"无时间戳"
    (workspace / "t.txt").write_text(
        "由王芳负责提交季度报表。\n会议决定通过新方案。\n", encoding="utf-8")
    r = run(workspace, "t.txt")
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads(r.stdout)
    assert rep["new"] == 1 and rep["decisions"] == 1
    rows = load_rows(workspace / "out" / "actions.xlsx")
    assert rows[1][0] == "A-001" and rows[1][1] == "王芳"
    assert rows[1][6] == "L1@无时间戳"
    assert rows[1][3] == "待确认"  # 无截止信息不编造
