"""docx-review-redline 黑盒测试：批注线程 + 修订 + 定位修复验证。"""
import json
import os
import re
import sys
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402


def build_draft(path):
    doc = Document()
    doc.add_paragraph("第一段介绍背景。")
    # AAA OLD ZZZ: 修正 harvey 中间子串错序用例
    doc.add_paragraph("AAA OLD ZZZ")
    # 表格混排
    doc.add_paragraph("表格前的段落，包含 OLD 这个词。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "数值 OLD 备用"
    table.cell(1, 0).text = "备注"
    table.cell(1, 1).text = "正常内容"
    # 跨 run 文本
    p = doc.add_paragraph()
    for piece in ("跨run目标", "文本", "尾部"):
        p.add_run(piece)
    doc.save(path)


def xml_of(docx_path, part):
    with zipfile.ZipFile(docx_path) as zf:
        return zf.read(part).decode("utf-8")


def test_replace_with_comment(workspace):
    build_draft(str(workspace / "draft.docx"))
    plan = {"operations": [
        {"op": "replace", "find": "OLD", "replace": "NEW",
         "comment": "术语统一为NEW"},
        {"op": "replace", "find": "跨run目标文本", "replace": "跨run已改文本",
         "comment": "跨run替换"},
    ]}
    (workspace / "plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    r = run_script("docx-review-redline", "apply_review.py",
                   ["--input", workspace / "draft.docx", "--plan", workspace / "plan.json",
                    "--output", workspace / "out_REVIEWED.docx",
                    "--report", workspace / "r.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    report = json.loads((workspace / "r.json").read_text(encoding="utf-8"))
    assert report["applied"] == 2, report
    assert report["skipped"] == []

    doc_xml = xml_of(str(workspace / "out_REVIEWED.docx"), "word/document.xml")
    # 修订结构: w:del 含 OLD, w:ins 含 NEW
    assert re.search(r"<w:del [^>]*>.*?OLD.*?</w:del>", doc_xml, re.S)
    assert re.search(r"<w:ins [^>]*>.*?NEW.*?</w:ins>", doc_xml, re.S)
    # 批注锚点
    assert "commentRangeStart" in doc_xml and "commentRangeEnd" in doc_xml
    assert "commentReference" in doc_xml
    # comments.xml 存在且含意见文本
    comments_xml = xml_of(str(workspace / "out_REVIEWED.docx"), "word/comments.xml")
    assert "术语统一为NEW" in comments_xml
    assert "跨run已改" in comments_xml or "跨run替换" in comments_xml
    # commentsExtended 存在
    names = zipfile.ZipFile(str(workspace / "out_REVIEWED.docx")).namelist()
    assert "word/commentsExtended.xml" in names
    # [Content_Types].xml 注册了 comments
    ct = xml_of(str(workspace / "out_REVIEWED.docx"), "[Content_Types].xml")
    assert "comments+xml" in ct
    # 关系注册
    rels = xml_of(str(workspace / "out_REVIEWED.docx"), "word/_rels/document.xml.rels")
    assert "comments.xml" in rels


def test_accept_semantics_and_order(workspace):
    """接受全部修订后的文本 = 原文去掉del内容替换为ins内容；验证段落顺序定位正确。"""
    build_draft(str(workspace / "draft.docx"))
    plan = {"operations": [
        {"op": "replace", "find": "OLD", "replace": "NEW", "comment": "c1",
         "scope": "all"},
    ]}
    (workspace / "plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    r = run_script("docx-review-redline", "apply_review.py",
                   ["--input", workspace / "draft.docx", "--plan", workspace / "plan.json",
                    "--output", workspace / "o.docx"])
    assert r.returncode == 0, r.stderr
    report = json.loads(r.stdout)
    # OLD 出现3处: AAA OLD ZZZ段 / 表格前段落 / 表格cell(0,1)
    assert report["applied"] == 3, report

    from docx.oxml.ns import qn
    doc = Document(str(workspace / "o.docx"))

    def accepted_text(p):
        out = []
        for node in p._p.iter():
            if node.tag == qn("w:t") and node.getparent().getparent().tag != qn("w:del"):
                out.append(node.text or "")
        return "".join(out)

    paras = doc.paragraphs
    assert accepted_text(paras[1]) == "AAA NEW ZZZ"
    assert "包含 NEW 这个词" in accepted_text(paras[2])
    # 表格cell文本(接受修订后)
    cell = doc.tables[0].cell(0, 1)
    cell_text = accepted_text(cell.paragraphs[0])
    assert "NEW 备用" in cell_text and "OLD" not in cell_text


def test_comment_only_and_skipped(workspace):
    build_draft(str(workspace / "draft.docx"))
    plan = {"operations": [
        {"op": "comment", "find": "正常内容", "comment": "请补充说明"},
        {"op": "replace", "find": "不存在的文本XX", "replace": "Y", "comment": "z"},
    ]}
    (workspace / "plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    r = run_script("docx-review-redline", "apply_review.py",
                   ["--input", workspace / "draft.docx", "--plan", workspace / "plan.json",
                    "--output", workspace / "o2.docx", "--report", workspace / "r2.json"])
    assert r.returncode == 0, r.stderr
    report = json.loads((workspace / "r2.json").read_text(encoding="utf-8"))
    assert report["applied"] == 1
    assert len(report["skipped"]) == 1
    assert "not found" in report["skipped"][0]["reason"]
    doc_xml = xml_of(str(workspace / "o2.docx"), "word/document.xml")
    assert "commentRangeStart" in doc_xml
    comments = xml_of(str(workspace / "o2.docx"), "word/comments.xml")
    assert "请补充说明" in comments
    # 无修订标记(comment-only不产生ins/del)
    assert "<w:del " not in doc_xml


def test_abnormal_corrupt(workspace):
    (workspace / "bad.docx").write_bytes(b"garbage")
    (workspace / "p.json").write_text('{"operations":[]}', encoding="utf-8")
    r = run_script("docx-review-redline", "apply_review.py",
                   ["--input", workspace / "bad.docx", "--plan", workspace / "p.json",
                    "--output", workspace / "o.docx"])
    assert r.returncode == 2


def test_empty_plan(workspace):
    build_draft(str(workspace / "draft.docx"))
    (workspace / "p.json").write_text('{"operations":[]}', encoding="utf-8")
    r = run_script("docx-review-redline", "apply_review.py",
                   ["--input", workspace / "draft.docx", "--plan", workspace / "p.json",
                    "--output", workspace / "o.docx"])
    assert r.returncode == 0, r.stderr
    report = json.loads(r.stdout)
    assert report["operations_total"] == 0 and report["applied"] == 0
