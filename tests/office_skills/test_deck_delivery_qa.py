"""deck-delivery-qa 黑盒测试：确定性几何 lint，预期值手工推导。

覆盖正常/复杂/异常三类：合规页、混合字号溢出、段落继承字号溢出、形状越界、
0 页空演示、损坏文件、利用率独立复算与确定性。
"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

EMU_PER_PT = 12700


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _textbox(slide, text, left_in, top_in, w_in, h_in, font_pt=None,
             para_font_pt=None):
    tb = slide.shapes.add_textbox(Inches(left_in), Inches(top_in),
                                  Inches(w_in), Inches(h_in))
    tf = tb.text_frame
    if para_font_pt is not None:
        tf.paragraphs[0].font.size = Pt(para_font_pt)
    run = tf.paragraphs[0].add_run()
    run.text = text
    if font_pt is not None:
        run.font.size = Pt(font_pt)
    return tb


def _run(tmp, name, pptx_path):
    rep = tmp / f"{name}.json"
    r = run_script("deck-delivery-qa", "deck_qa.py",
                   ["--input", pptx_path, "--report", rep])
    data = json.loads(rep.read_text(encoding="utf-8"))
    return r, data


def test_normal_deck_passes(workspace):
    """正常：短标题在充裕文本框内——合规、退出 0、估算模式诚实标注。"""
    prs = Presentation()
    _textbox(_blank_slide(prs), "季度业务回顾", 1, 1, 6, 1.5, font_pt=18)
    path = workspace / "normal.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "normal", path)
    assert r.returncode == 0, r.stdout + r.stderr
    assert rep["pass"] is True
    assert rep["mode"] == "estimated_geometry"
    assert rep["rendered"] is False                      # 估算不等于渲染
    assert rep["slide_count"] == 1
    assert rep["deck_issues"] == []
    assert rep["slides"][0]["pass"] is True
    assert rep["slides"][0]["issues"] == []
    assert isinstance(rep["slides"][0]["utilization"], float)


def test_mixed_font_run_overflow(workspace):
    """复杂：10pt 前缀 + 20 个 72pt 汉字（5in×1in 框）——逐 run 计宽、按段内最大
    字号计行高，必须报溢出。只取"第一个显式字号 run"（10pt）会漏报。"""
    prs = Presentation()
    s = _blank_slide(prs)
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    p = tb.text_frame.paragraphs[0]
    r1 = p.add_run(); r1.text = "a"; r1.font.size = Pt(10)
    r2 = p.add_run(); r2.text = "长" * 20; r2.font.size = Pt(72)
    path = workspace / "mixed.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "mixed", path)
    assert r.returncode == 1, r.stdout
    assert rep["pass"] is False
    issues = rep["slides"][0]["issues"]
    over = [i for i in issues if i["check"] == "text_overflow"]
    assert len(over) == 1
    # 手工推导：宽=10*0.55 + 20*72 = 1445.5pt；可用宽=5*72-21.6=338.4pt
    #        → 5 行；高=5*72*1.25=450pt；可用高=1*72-7.2=64.8pt
    assert "72pt" in over[0]["detail"]      # 段内最大字号被采用，非 10pt
    assert "450pt" in over[0]["detail"]
    assert "65pt" in over[0]["detail"] or "64pt" in over[0]["detail"]


def test_paragraph_inherited_font_overflow(workspace):
    """复杂：run 无显式字号、由段落 defRPr(=para.font.size) 继承 72pt——必须按 72pt
    计算并报溢出；把 run.font.size=None 当默认 18pt 会漏报。"""
    prs = Presentation()
    _textbox(_blank_slide(prs), "长" * 20, 1, 1, 5, 1, para_font_pt=72)
    path = workspace / "para.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "para", path)
    assert r.returncode == 1, r.stdout
    assert rep["pass"] is False
    over = [i for i in rep["slides"][0]["issues"] if i["check"] == "text_overflow"]
    assert len(over) == 1
    assert "72pt" in over[0]["detail"]
    assert "450pt" in over[0]["detail"]


def test_shape_out_of_bounds(workspace):
    """异常：形状超出画布——记录 shape_out_of_bounds 并判定不合格。"""
    prs = Presentation()  # 默认 10in×7.5in
    s = _blank_slide(prs)
    _textbox(s, "越界内容", 12, 1, 2, 1, font_pt=18)   # right=14in > 10in
    path = workspace / "oob.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "oob", path)
    assert r.returncode == 1, r.stdout
    assert rep["pass"] is False
    checks = [i["check"] for i in rep["slides"][0]["issues"]]
    assert "shape_out_of_bounds" in checks


def test_utilization_independent_recompute(workspace):
    """复杂：利用率数值独立复算——2in×1in 矩形 / 10in×7.5in 画布 = 0.0267。"""
    prs = Presentation()
    s = _blank_slide(prs)  # 10×7.5in 默认
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
                       Inches(2), Inches(1))
    path = workspace / "util.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "util", path)
    # 空文本框形状不构成占位符问题，页面应通过
    assert r.returncode == 0, r.stdout
    expected = round((2 * 1) / (10 * 7.5), 4)
    assert rep["slides"][0]["utilization"] == expected == 0.0267


def test_empty_deck_not_deliverable(workspace):
    """异常：合法 PPTX 容器但 0 页——不得声称可交付，退出非零并给出 empty_deck。"""
    prs = Presentation()
    path = workspace / "empty.pptx"
    prs.save(str(path))

    r, rep = _run(workspace, "empty", path)
    assert r.returncode == 1, r.stdout
    assert rep["pass"] is False
    assert rep["slide_count"] == 0
    assert rep["slides"] == []
    assert [i["check"] for i in rep["deck_issues"]] == ["empty_deck"]


def test_corrupt_file_exit_2(workspace):
    """异常：损坏的 pptx——退出码 2，不产出伪报告。"""
    bad = workspace / "bad.pptx"
    bad.write_bytes(b"not a zip")
    r = run_script("deck-delivery-qa", "deck_qa.py",
                   ["--input", bad, "--report", workspace / "bad.json"])
    assert r.returncode == 2
    assert "cannot open pptx" in r.stderr


def test_deterministic_and_render_note(workspace):
    """复杂：同输入两次运行报告逐字节一致；render_note 不硬写"无引擎"，rendered=false。"""
    prs = Presentation()
    _textbox(_blank_slide(prs), "长" * 20, 1, 1, 5, 1, font_pt=72)
    path = workspace / "det.pptx"
    prs.save(str(path))

    r1 = run_script("deck-delivery-qa", "deck_qa.py",
                    ["--input", path, "--report", workspace / "d1.json"])
    r2 = run_script("deck-delivery-qa", "deck_qa.py",
                    ["--input", path, "--report", workspace / "d2.json"])
    assert r1.returncode == r2.returncode == 1
    assert (workspace / "d1.json").read_bytes() == (workspace / "d2.json").read_bytes()
    rep = json.loads((workspace / "d1.json").read_text(encoding="utf-8"))
    assert rep["rendered"] is False
    assert "渲染" in rep["render_note"]
