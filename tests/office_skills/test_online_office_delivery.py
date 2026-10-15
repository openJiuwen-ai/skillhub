"""online-office-delivery 黑盒测试。

正常/复杂路径用受控的 lark-cli 测试替身验证真实管线。替身严格按官方文档
核对过的行为实现（content-comparison 审计材料中的 lark-cli 语义）：
+shortcut 命令、{ok, identity, data} 信封、+workbook-import 的 ready/job_status/
timed_out/next_command 续查、+csv-get 返回带 "[row=N] " 前缀的 annotated_csv、
--file 仅接受 cwd 相对路径、显式 --as user/bot 身份、结构化错误信封
（missing_scopes/console_url）、+sheet-delete 为 high-risk 需 --yes。替身只用于
管线逻辑验证，不冒充真实在线交付（真实在线交付见 online-live-review-20261015）。
异常路径（无 CLI/计划非法/语法无法核实）直接验证真实行为。测试不 import
被测模块，全部经子进程黑盒执行。
"""
import json
import os
import subprocess
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import skill_dir  # noqa: E402

from openpyxl import Workbook  # noqa: E402

SKILL = "online-office-delivery"
SCRIPT = "online_delivery.py"

DOUBLE_SOURCE = r'''
import csv
import io
import json
import os
import re
import sys
import uuid

HELP = "lark-cli (documented-behavior double for pipeline tests)\nSubcommands: sheets, slides, drive, api\n"
SHEETS_HELP_LINES = [
    "sheets - spreadsheet operations",
    "Shortcuts:",
    "  +workbook-import   import local .xlsx/.csv as a new spreadsheet",
    "  +workbook-info     list sheets/rows/cols of a workbook",
    "  +csv-get           read values of a range",
    "  +csv-put           write csv values starting at a cell",
    "  +sheet-create      create an empty sheet",
    "  +sheet-delete      delete a sheet (high-risk, requires --yes)",
    "  +sheet-list        list sheets",
]

_IDENTITY = "bot"


def current_identity():
    return os.environ.get("DOUBLE_IDENTITY") or _IDENTITY


def note_identity(fl):
    global _IDENTITY
    if fl.get("as"):
        _IDENTITY = fl["as"]


def state_path():
    return os.environ.get("DOUBLE_STATE") or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "double-state.json")


def load_state():
    if os.path.exists(state_path()):
        with open(state_path(), encoding="utf-8") as f:
            return json.load(f)
    return {"workbooks": {}, "pending": {}, "timeout_used": False}


def save_state(st):
    with open(state_path(), "w", encoding="utf-8") as f:
        json.dump(st, f, ensure_ascii=False)


def out(obj):
    if isinstance(obj, dict) and "identity" not in obj:
        obj = dict(obj)
        obj["identity"] = current_identity()
    print(json.dumps(obj, ensure_ascii=False))


def flags(argv):
    fl = {}
    i = 0
    while i < len(argv):
        a = argv[i]
        if a.startswith("--"):
            if "=" in a:
                k, v = a[2:].split("=", 1)
                fl[k] = v
            elif i + 1 < len(argv) and not argv[i + 1].startswith("--"):
                fl[a[2:]] = argv[i + 1]
                i += 1
            else:
                fl[a[2:]] = True
        i += 1
    return fl


def parse_range(rng):
    def col(s):
        n = 0
        for ch in s:
            n = n * 26 + (ord(ch) - 64)
        return n - 1
    rng = (rng or "").strip()
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)$", rng)
    if m:
        return (col(m.group(1)), int(m.group(2)) - 1,
                col(m.group(3)), int(m.group(4)) - 1)
    m = re.match(r"([A-Z]+)(\d+)$", rng)
    if m:
        c, r = col(m.group(1)), int(m.group(2)) - 1
        return c, r, c, r
    return None


def load_file_sheets(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = [r for r in csv.reader(f)]
        return [(os.path.splitext(os.path.basename(path))[0] or "csv", rows)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for name in wb.sheetnames:
        ws = wb[name]
        sheets.append((name, [list(r) for r in ws.iter_rows(values_only=True)]))
    wb.close()
    return sheets


def finish_import(st, ticket, pend):
    sheets = load_file_sheets(pend["file"])
    token = "shtdbl" + uuid.uuid4().hex[:10]
    url = "https://lark-double.example/sheets/" + token
    wbk = {"token": token, "sheets": {}}
    for i, (title, rows) in enumerate(sheets):
        wbk["sheets"][title] = {"sheet_id": "sid%d" % i, "rows": rows}
    if os.environ.get("DOUBLE_MUTATE_FILE") and os.path.basename(
            pend["file"]) == os.environ["DOUBLE_MUTATE_FILE"]:
        first = next(iter(wbk["sheets"]))
        wbk["sheets"][first]["rows"][0][0] = "MUTATED"
    st["workbooks"][url] = wbk
    if pend.get("name"):
        wbk["name"] = pend["name"]
    st["pending"].pop(ticket, None)
    save_state(st)
    data = {"token": token, "ticket": ticket, "ready": True, "job_status": 0}
    if not os.environ.get("DOUBLE_OMIT_URL"):
        data["url"] = url
    out({"ok": True, "data": data})
    return 0


def cmd_import(argv):
    fl = flags(argv)
    note_identity(fl)
    if os.environ.get("DOUBLE_AUTH_FAIL"):
        out({"ok": False, "error": {"code": 401,
                                    "message": "unauthorized: invalid token"}})
        return 1
    if os.environ.get("DOUBLE_SCOPE_FAIL"):
        sys.stderr.write(json.dumps({
            "ok": False, "identity": current_identity(),
            "error": {"type": "authorization", "subtype": "missing_scope",
                      "code": 99991679,
                      "message": "permission denied: missing scope",
                      "hint": "grant the scope then retry",
                      "missing_scopes": ["sheets:spreadsheet:create"],
                      "console_url": "https://open.example/app/auth?q=sheets"}},
            ensure_ascii=False) + "\n")
        return 1
    if fl.get("ticket"):
        st = load_state()
        pend = st["pending"].get(fl["ticket"])
        if not pend:
            out({"ok": False, "error": {"message": "unknown ticket"}})
            return 1
        return finish_import(st, fl["ticket"], pend)
    f = fl.get("file")
    if not f:
        sys.stderr.write("--file is required\n")
        return 2
    if os.path.isabs(f) or f.startswith("/") or f.startswith("\\\\") or (
            len(f) > 1 and f[1] == ":"):
        sys.stderr.write("unsafe file path: absolute paths are rejected\n")
        return 2
    path = os.path.join(os.getcwd(), f)
    if not os.path.exists(path):
        sys.stderr.write("file not found: %s\n" % f)
        return 2
    st = load_state()
    ticket = "tk" + uuid.uuid4().hex[:10]
    st["pending"][ticket] = {"file": path, "name": fl.get("name"),
                             "folder": fl.get("folder-token")}
    if os.environ.get("DOUBLE_IMPORT_TIMEOUT_ONCE") and not st["timeout_used"]:
        st["timeout_used"] = True
        save_state(st)
        next_cmd = '"%s" "%s" sheets +workbook-import --ticket %s' % (
            sys.executable, os.path.abspath(__file__), ticket)
        out({"ok": True, "data": {"ticket": ticket, "timed_out": True,
                                  "next_command": next_cmd}})
        return 0
    return finish_import(st, ticket, st["pending"][ticket])


def get_workbook(fl):
    st = load_state()
    wbk = st["workbooks"].get(fl.get("url"))
    if not wbk:
        out({"ok": False, "error": {"message": "spreadsheet not found"}})
        return None, None
    return st, wbk


def find_sheet(wbk, fl):
    if fl.get("sheet-id"):
        for t, s in wbk["sheets"].items():
            if s["sheet_id"] == fl["sheet-id"]:
                return t, s
        return None, None
    s = wbk["sheets"].get(fl.get("sheet-name"))
    return (fl.get("sheet-name"), s) if s else (None, None)


def cmd_info(argv):
    fl = flags(argv)
    note_identity(fl)
    st, wbk = get_workbook(fl)
    if not wbk:
        return 1
    sheets = [{"sheet_id": s["sheet_id"], "title": t,
               "row_count": len(s["rows"]),
               "col_count": max((len(r) for r in s["rows"]), default=0)}
              for t, s in wbk["sheets"].items()]
    out({"ok": True, "data": {"token": wbk["token"], "sheets": sheets}})
    return 0


def _annotated_csv(values, first_row):
    lines = []
    for i, row in enumerate(values):
        buf = io.StringIO()
        csv.writer(buf).writerow(["" if v is None else v for v in row])
        lines.append("[row=%d] %s" % (first_row + i, buf.getvalue().rstrip("\r\n")))
    return "\n".join(lines)


def cmd_csv_get(argv):
    fl = flags(argv)
    note_identity(fl)
    st, wbk = get_workbook(fl)
    if not wbk:
        return 1
    t, s = find_sheet(wbk, fl)
    if not s:
        out({"ok": False, "error": {"message": "sheet not found"}})
        return 1
    rng = parse_range(fl.get("range") or "A1:A1")
    if not rng:
        out({"ok": False, "error": {"message": "bad range"}})
        return 1
    c0, r0, c1, r1 = rng
    values = []
    for r in range(r0, r1 + 1):
        row = s["rows"][r] if r < len(s["rows"]) else []
        values.append([row[c] if c < len(row) else ""
                       for c in range(c0, c1 + 1)])
    data = {"sheet_id": s["sheet_id"],
            "col_indices": [chr(65 + c) for c in range(c0, c1 + 1)],
            "row_indices": list(range(r0 + 1, r1 + 2)),
            "annotated_csv": _annotated_csv(values, r0 + 1)}
    out({"ok": True, "data": data})
    return 0


def cmd_csv_put(argv):
    fl = flags(argv)
    note_identity(fl)
    if os.environ.get("DOUBLE_DENY_WRITE"):
        out({"ok": False, "error": {"message": "write denied by policy"}})
        return 1
    st, wbk = get_workbook(fl)
    if not wbk:
        return 1
    t, s = find_sheet(wbk, fl)
    if not s:
        out({"ok": False, "error": {"message": "sheet not found"}})
        return 1
    m = re.match(r"([A-Z]+)(\d+)$", fl.get("start-cell") or "A1")
    if not m:
        out({"ok": False, "error": {"message": "bad start-cell"}})
        return 1
    col = 0
    for ch in m.group(1):
        col = col * 26 + (ord(ch) - 64)
    col -= 1
    row = int(m.group(2)) - 1
    rows = [r for r in csv.reader((fl.get("csv") or "").splitlines())]
    for dr, cells in enumerate(rows):
        while len(s["rows"]) <= row + dr:
            s["rows"].append([])
        rr = s["rows"][row + dr]
        while len(rr) <= col + len(cells) - 1:
            rr.append("")
        for dc, v in enumerate(cells):
            rr[col + dc] = v
    save_state(st)
    out({"ok": True, "data": {"written": len(rows)}})
    return 0


def cmd_sheet_create(argv):
    fl = flags(argv)
    note_identity(fl)
    st, wbk = get_workbook(fl)
    if not wbk:
        return 1
    title = fl.get("title") or ""
    if not title:
        sys.stderr.write("--title is required\n")
        return 2
    sid = "sid" + uuid.uuid4().hex[:8]
    wbk["sheets"][title] = {"sheet_id": sid, "rows": []}
    save_state(st)
    out({"ok": True, "data": {"sheet_id": sid, "title": title}})
    return 0


def cmd_sheet_delete(argv):
    fl = flags(argv)
    note_identity(fl)
    if fl.get("yes") is not True:
        sys.stderr.write(
            "high-risk shortcut: confirmation flag --yes is required\n")
        return 10
    st, wbk = get_workbook(fl)
    if not wbk:
        return 1
    for t, s in list(wbk["sheets"].items()):
        if s["sheet_id"] == fl.get("sheet-id"):
            del wbk["sheets"][t]
            save_state(st)
            out({"ok": True, "data": {"deleted": t}})
            return 0
    out({"ok": False, "error": {"message": "sheet not found"}})
    return 1


COMMANDS = {
    "+workbook-import": cmd_import,
    "+workbook-info": cmd_info,
    "+csv-get": cmd_csv_get,
    "+csv-put": cmd_csv_put,
    "+sheet-create": cmd_sheet_create,
    "+sheet-delete": cmd_sheet_delete,
}


def main(argv):
    if not argv or argv[0] in ("-h", "--help"):
        sys.stdout.write(HELP)
        return 0
    if argv[0] in ("-v", "--version"):
        sys.stdout.write("lark-cli version 1.0.98-double\n")
        return 0
    if argv[0] == "sheets" and (len(argv) < 2 or argv[1] in ("-h", "--help")):
        hidden = os.environ.get("DOUBLE_HIDE_CMD")
        lines = [ln for ln in SHEETS_HELP_LINES
                 if not (hidden and hidden in ln)]
        sys.stdout.write("\n".join(lines) + "\n")
        return 0
    if len(argv) >= 2 and argv[0] == "sheets" and argv[1] in COMMANDS:
        return COMMANDS[argv[1]](argv[2:])
    sys.stderr.write("unknown command: %r\n" % (argv[:2],))
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
'''


def write_double(base):
    d = base / "bin"
    d.mkdir(parents=True, exist_ok=True)
    py = d / "lark_cli_double.py"
    py.write_text(DOUBLE_SOURCE.strip() + "\n", encoding="utf-8")
    if os.name == "nt":
        cmd = d / "lark-cli.cmd"
        cmd.write_text('@echo off\r\n"%s" "%s" %%*\r\n' % (sys.executable, py),
                       encoding="ascii")
    else:
        sh = d / "lark-cli"
        sh.write_text("#!/usr/bin/env python3\n" + DOUBLE_SOURCE.strip() + "\n",
                      encoding="utf-8")
        os.chmod(str(sh), 0o755)
    return str(d)


def make_env(state_path, double_dir=None, **knobs):
    env = dict(os.environ)
    for k in list(env):
        if k.upper().startswith(("LARK_", "FEISHU_")) or k.upper().startswith(
                "DOUBLE_"):
            del env[k]
    env["PYTHONUTF8"] = "1"
    env["DOUBLE_STATE"] = str(state_path)
    if double_dir is None:
        empty = os.path.dirname(str(state_path)) or "."
        env["PATH"] = empty
    else:
        env["PATH"] = double_dir + os.pathsep + env.get("PATH", "")
    for k, v in knobs.items():
        env["DOUBLE_" + k] = str(v)
    return env


def run_delivery(plan_path, outdir, env, cwd=None, timeout=180):
    script = os.path.join(skill_dir(SKILL), "scripts", SCRIPT)
    outdir = str(outdir)
    os.makedirs(outdir, exist_ok=True)
    return subprocess.run(
        [sys.executable, script, "--plan", str(plan_path),
         "--outdir", outdir],
        capture_output=True, text=True, encoding="utf-8",
        cwd=cwd or outdir, timeout=timeout, env=env)


def write_plan(base, target_type="lark_sheet", files=None, name="plan.json",
               **extra):
    plan = {"target": {"type": target_type}, "files": files or []}
    plan.update(extra)
    p = base / name
    p.write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    return p


def read_report(outdir):
    return json.loads((outdir / "delivery-report.json").read_text(
        encoding="utf-8"))


def read_state(state_path):
    with open(state_path, encoding="utf-8") as f:
        return json.load(f)


def build_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["name", "score", "note"])
    ws.append(["alice", 95, ""])
    ws.append(["bob", 87, "x"])
    ws.append(["carol", 72, None])
    ws2 = wb.create_sheet("Summary")
    ws2.append(["total", 254])
    wb.save(str(path))


def test_blocked_without_cli(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    env = make_env(tmp_path / "st.json", double_dir=None)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert rep["missing"]["clis"] == ["lark-cli"]
    assert "results" not in rep
    assert "https://" not in json.dumps(rep)
    assert rep["probe"]["env_vars_present"] == []
    assert json.loads(r.stdout)["status"] == "blocked"


def test_help_gate_blocks_unverified_command(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, HIDE_CMD="+csv-put")
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert rep["missing"]["commands_from_help"] == ["+csv-put"]
    assert "results" not in rep
    argvs = [e["argv"] for e in rep["cli_invocation_log"]]
    assert len(argvs) == 3
    assert argvs[0][1:] == ["--help"]
    assert argvs[1][1:] == ["sheets", "--help"]
    assert argvs[2][1:] == ["--version"]


def test_input_errors(tmp_path):
    bad_json = tmp_path / "bad.json"
    bad_json.write_text("{not json", encoding="utf-8")
    outdir = tmp_path / "out1"
    r = run_delivery(bad_json, outdir, make_env(tmp_path / "s1.json"))
    assert r.returncode == 2
    rep = read_report(outdir)
    assert rep["status"] == "input_error"
    assert rep["errors"] and "无法读取" in rep["errors"][0]

    build_xlsx(tmp_path / "data.xlsx")
    empty = write_plan(tmp_path, files=[], name="plan-empty.json")
    outdir2 = tmp_path / "out2"
    r = run_delivery(empty, outdir2, make_env(tmp_path / "s2.json"))
    assert r.returncode == 2
    rep = read_report(outdir2)
    assert rep["status"] == "input_error"
    assert any("空" in e for e in rep["errors"])

    plan3 = write_plan(tmp_path, name="plan-missing.json",
                       files=[{"local": str(tmp_path / "nope.xlsx")}])
    outdir3 = tmp_path / "out3"
    r = run_delivery(plan3, outdir3, make_env(tmp_path / "s3.json"))
    assert r.returncode == 2
    rep = read_report(outdir3)
    assert any("不存在" in e for e in rep["errors"])


def test_unsupported_target_and_extension(tmp_path):
    docx = tmp_path / "report.docx"
    docx.write_bytes(b"PK\x03\x04stub")
    plan = write_plan(tmp_path, target_type="lark_doc",
                      files=[{"local": str(docx)}])
    outdir = tmp_path / "out1"
    r = run_delivery(plan, outdir, make_env(tmp_path / "s1.json"))
    assert r.returncode == 2
    rep = read_report(outdir)
    assert any("不支持的目标类型" in e for e in rep["errors"])

    xls = tmp_path / "legacy.xls"
    xls.write_bytes(b"stub")
    plan2 = write_plan(tmp_path, files=[{"local": str(xls)}])
    outdir2 = tmp_path / "out2"
    r = run_delivery(plan2, outdir2, make_env(tmp_path / "s2.json"))
    assert r.returncode == 2
    rep = read_report(outdir2)
    assert any(".xls" in e and "扩展名" in e for e in rep["errors"])


def test_delivered_xlsx_full_pipeline(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[
        {"local": str(tmp_path / "data.xlsx"), "name": "Q3-report"}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "delivered"
    entry = rep["results"][0]
    assert entry["status"] == "delivered"
    assert entry["url"].startswith("https://lark-double.example/sheets/")
    assert entry["import"]["ready"] is True
    assert entry["import"]["job_status"] == 0
    assert entry["readback"]["verified"] is True
    assert entry["readback"]["sheets_local"] == ["Data", "Summary"]
    sampled = {s["sheet"]: s for s in entry["readback"]["sampled"]}
    assert sampled["Data"]["range"] == "A1:C4"
    assert sampled["Data"]["mismatch_count"] == 0
    assert sampled["Summary"]["range"] == "A1:B1"
    assert sampled["Summary"]["cells"] == 2
    pc = entry["permission_check"]
    assert pc["verified"] is True and pc["readback_matched"] is True
    assert pc["cleanup"]["ok"] is True
    imports = [e for e in rep["cli_invocation_log"]
               if "+workbook-import" in e["argv"]]
    assert imports
    file_arg = imports[0]["argv"][imports[0]["argv"].index("--file") + 1]
    assert file_arg == "data.xlsx"
    assert imports[0]["cwd"].lower() == str(tmp_path).lower()
    deletes = [e for e in rep["cli_invocation_log"]
               if "+sheet-delete" in e["argv"]]
    assert deletes and "--yes" in deletes[0]["argv"]
    assert rep["probe"]["syntax_verification"]["all_required_verified"] is True
    st = read_state(tmp_path / "st.json")
    assert list(st["workbooks"]) == [entry["url"]]
    wbk = st["workbooks"][entry["url"]]
    assert "__delivery_check__" not in wbk["sheets"]
    assert wbk["sheets"]["Data"]["rows"][1][1] == 95
    assert wbk["sheets"]["Summary"]["rows"][0][1] == 254


def test_delivered_csv_with_readback_warning(tmp_path):
    csvf = tmp_path / "names.csv"
    csvf.write_text("name,score\nalice,95\nbob,87\n", encoding="utf-8")
    plan = write_plan(tmp_path, files=[{"local": str(csvf)}], readback=False)
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "delivered"
    assert any("readback=false" in w for w in rep["warnings"])
    entry = rep["results"][0]
    assert entry["readback"]["verified"] is True
    s = entry["readback"]["sampled"][0]
    assert s["range"] == "A1:B3"
    assert s["mismatch_count"] == 0


def test_import_timeout_follows_next_command(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, IMPORT_TIMEOUT_ONCE=1)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    entry = rep["results"][0]
    assert entry["status"] == "delivered"
    assert entry["import"]["attempts"] == 2
    assert entry["import"]["followed_next_command"] is True
    first = [e for e in rep["cli_invocation_log"]
             if "+workbook-import" in e["argv"]][0]
    assert "timed_out" in first["stdout_excerpt"]


def test_readback_mismatch_fails(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, MUTATE_FILE="data.xlsx")
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "failed"
    entry = rep["results"][0]
    assert entry["status"] == "failed"
    assert entry["readback"]["verified"] is False
    m = entry["readback"]["sampled"][0]["mismatches"][0]
    assert m["cell"] == "A1"
    assert m["local"] == "name"
    assert m["online"] == "MUTATED"
    assert json.loads(r.stdout)["status"] == "failed"


def test_missing_url_not_fabricated(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, OMIT_URL=1)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "failed"
    entry = rep["results"][0]
    assert "url" not in entry
    assert any("未给出 URL" in x for x in entry["reasons"])
    assert "https://" not in json.dumps(rep)


def test_permission_denied_fails(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, DENY_WRITE=1)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "failed"
    entry = rep["results"][0]
    assert entry["readback"]["verified"] is True
    pc = entry["permission_check"]
    assert pc["verified"] is False
    assert "金丝雀写入失败" in pc["error"]


def test_unauthorized_cli_blocks(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, AUTH_FAIL=1)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert "cli_authorization" in rep["missing"]
    assert "https://" not in json.dumps(rep)


def test_keep_check_sheet_leaves_marker(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}],
                      keep_check_sheet=True)
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    entry = rep["results"][0]
    assert entry["status"] == "delivered"
    assert entry["permission_check"]["cleanup"]["attempted"] is False
    st = read_state(tmp_path / "st.json")
    wbk = st["workbooks"][entry["url"]]
    check = wbk["sheets"]["__delivery_check__"]
    assert check["rows"][0][0].startswith("delivery-check-")


def test_multi_file_partial_failure(tmp_path):
    build_xlsx(tmp_path / "good.xlsx")
    bad = tmp_path / "bad.csv"
    bad.write_text("name,score\nalice,95\n", encoding="utf-8")
    plan = write_plan(tmp_path, files=[
        {"local": str(tmp_path / "good.xlsx")},
        {"local": str(bad)}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, MUTATE_FILE="bad.csv")
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "failed"
    good, bad_entry = rep["results"]
    assert good["status"] == "delivered"
    assert good["url"].startswith("https://lark-double.example/sheets/")
    assert bad_entry["status"] == "failed"
    m = bad_entry["readback"]["sampled"][0]["mismatches"][0]
    assert m["cell"] == "A1" and m["online"] == "MUTATED"


def build_special_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["序号", "姓名", "编号", "金额", "备注"])
    ws.append([1, "张三", "001", 1999.5, "正常"])
    ws.append([2, "李四", "010", 12345.67, "含 空格"])
    wb.save(str(path))


def test_row_prefix_and_leading_zero_readback(tmp_path):
    # +csv-get 默认返回带 [row=N] 前缀的 annotated_csv：脚本必须去前缀后逐格
    # 比对，否则首列会因前缀判不一致。文本前导零 "001" 必须按字面保真，"001"!=1。
    build_special_xlsx(tmp_path / "special.xlsx")
    plan = write_plan(tmp_path,
                      files=[{"local": str(tmp_path / "special.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    entry = rep["results"][0]
    assert entry["status"] == "delivered"
    assert entry["readback"]["verified"] is True
    assert entry["readback"]["sampled"][0]["mismatch_count"] == 0
    csvlog = [e for e in rep["cli_invocation_log"] if "+csv-get" in e["argv"]]
    assert csvlog and "annotated_csv" in csvlog[0]["stdout_excerpt"]
    assert "--sheet-id" in csvlog[0]["argv"]


def test_identity_passthrough_and_recorded(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}],
                      identity="user")
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["identity_requested"] == "user"
    entry = rep["results"][0]
    assert entry["import"]["identity"] == "user"
    assert entry["import"]["identity_requested"] == "user"
    sheets_calls = [e for e in rep["cli_invocation_log"]
                    if e["argv"][1:2] == ["sheets"]
                    and "--help" not in e["argv"]]
    assert sheets_calls
    assert all("--as" in e["argv"] and "user" in e["argv"]
               for e in sheets_calls)


def test_identity_mismatch_fails(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}],
                      identity="user")
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, IDENTITY="bot")
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 4, r.stderr + r.stdout
    rep = read_report(outdir)
    entry = rep["results"][0]
    assert entry["status"] == "failed"
    assert any("身份不符" in x for x in entry["reasons"])


def test_missing_scope_surfaces_scopes_and_console_url(tmp_path):
    build_xlsx(tmp_path / "data.xlsx")
    plan = write_plan(tmp_path, files=[{"local": str(tmp_path / "data.xlsx")}])
    outdir = tmp_path / "out"
    double_dir = write_double(tmp_path)
    env = make_env(tmp_path / "st.json", double_dir, SCOPE_FAIL=1)
    r = run_delivery(plan, outdir, env)
    assert r.returncode == 3, r.stderr + r.stdout
    rep = read_report(outdir)
    assert rep["status"] == "blocked"
    assert rep["missing"]["missing_scopes"] == ["sheets:spreadsheet:create"]
    assert rep["missing"]["console_url"].startswith("https://")
    assert rep["missing"]["error"]["subtype"] == "missing_scope"
