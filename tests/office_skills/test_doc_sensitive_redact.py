"""doc-sensitive-redact 标准黑盒测试。

独立验证真实行为（正常/复杂/异常三类），不复用被测实现代码：
识别与打码用自己的校验位/Luhn 独立实现构造期望值；
隐藏部件（删除修订/批注/数值单元格/备注）用底层 ZIP 逐部件扫描验证，
对应冻结证据 benchmarks/redaction-hidden 的验收规则。
"""
import io
import json
import os
import shutil
import sys
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

import pymupdf  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

W = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
CHECK = "10X98765432"


def make_id(prefix17):
    s = sum(int(prefix17[i]) * W[i] for i in range(17))
    return prefix17 + CHECK[s % 11]


def make_bank(prefix15):
    total = 0
    for i, ch in enumerate(reversed(prefix15)):
        d = int(ch)
        if i % 2 == 0:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return prefix15 + str((10 - total % 10) % 10)


VALID_ID = make_id("31010119900101123")
BAD_ID = VALID_ID[:-1] + ("X" if VALID_ID[-1] != "X" else "0")
PHONE = "13900001111"
BANK = make_bank("622202020011223")
F15 = "310101900101123"
EMAIL = "lianxi@chuangyi.test"
H_EMAIL = "kejian@chuangyi.test"
DEL_EMAIL = "shanchu@chuangyi.test"
CMT_EMAIL = "pizhu@chuangyi.test"
SHEET_CMT_EMAIL = "biaopizhu@chuangyi.test"
NOTES_EMAIL = "beizhu@chuangyi.test"
NUM_PHONE = 15912340001
STR_EMAIL = "zifuchuan@chuangyi.test"
PDF_META_EMAIL = "zuozhe@chuangyi.test"
PDF_XMP_EMAIL = "xmp@chuangyi.test"
PDF_ANNOT_EMAIL = "zhushiquan@chuangyi.test"
PDF_ATTACH_EMAIL = "fujian@chuangyi.test"
PDF_TOC_EMAIL = "shuqian@chuangyi.test"
BIN_EMAIL = "ercun@chuangyi.test"
PDF_LINK_EMAIL = "lianjie@chuangyi.test"
PDF_NORMAL_URI = "https://example.com/public"
ATTACH_DOC_EMAIL = "fujianneirong@chuangyi.test"
ATTACH_PUBLIC = "Attachment public unchanged text."
FORM_EMAIL = "biaodan@chuangyi.test"


def build_plain_docx(path):
    doc = Document()
    doc.add_paragraph("公开内容，保持原样。")
    doc.add_paragraph(f"联系人 {EMAIL}，电话 {PHONE}")
    doc.add_paragraph(f"证件 {VALID_ID}，卡号 {BANK}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = f"表内邮箱 {EMAIL}"
    doc.sections[0].header.paragraphs[0].add_run(f"页眉 {H_EMAIL}")
    doc.save(str(path))


def build_hidden_docx(path):
    doc = Document()
    p = doc.add_paragraph("Public unchanged text. ")
    r = p.add_run(EMAIL)
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "50")
    dr = OxmlElement("w:r")
    dt = OxmlElement("w:delText")
    dt.text = " " + DEL_EMAIL
    dr.append(dt)
    deleted.append(dr)
    p._p.append(deleted)
    doc.add_comment(runs=[r], text="批注 " + CMT_EMAIL, author="Synthetic")
    doc.save(str(path))


def build_hidden_xlsx(path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Synthetic phone"
    ws["B1"] = NUM_PHONE
    ws["A2"] = "Public unchanged"
    ws["A2"].comment = Comment(SHEET_CMT_EMAIL, "Synthetic")
    ws["C1"] = f"文本 {STR_EMAIL}"
    wb.save(str(path))


def build_hidden_pptx(path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "Public unchanged"
    s.notes_slide.notes_text_frame.text = NOTES_EMAIL
    prs.save(str(path))


def build_review_docx(path):
    doc = Document()
    doc.add_paragraph(f"存疑证件 {BAD_ID} 与旧编号 {F15}")
    doc.save(str(path))


def zip_scan(path, needles):
    """底层逐部件扫描（与冻结控制器探针同思路，独立于被测实现）。"""
    found = {}
    with zipfile.ZipFile(str(path)) as z:
        for n in z.namelist():
            data = z.read(n)
            hits = [x for x in needles if x.encode("utf-8") in data]
            if hits:
                found[n] = hits
    return found


def load_manifest(outdir):
    return json.loads((outdir / "redact-manifest.json").read_text(encoding="utf-8"))


# ---------- 正常 ----------

def test_redact_normal_docx_visible_content(workspace):
    build_plain_docx(workspace / "plain.docx")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "plain.docx", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    out = workspace / "out" / "plain_REDACTED.docx"
    assert out.exists()
    doc = Document(str(out))
    body = "\n".join(p.text for p in doc.paragraphs)
    header = doc.sections[0].header.paragraphs[0].text
    table_text = doc.tables[0].cell(0, 0).text
    for surface in (body, header, table_text):
        assert EMAIL not in surface and PHONE not in surface
        assert VALID_ID not in surface and BANK not in surface
        assert "■" in surface or surface == ""
    assert "公开内容，保持原样。" in body
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    counts = {x["type"]: x["count"] for x in pf["redactions"]}
    assert counts.get("email", 0) >= 2
    assert counts.get("phone") == 1
    assert counts.get("id_card") == 1
    assert counts.get("bank_card") == 1
    raw = json.dumps(man, ensure_ascii=False)
    for needle in (EMAIL, H_EMAIL, PHONE, VALID_ID, BANK):
        assert needle not in raw


def test_redact_normal_xlsx_pptx_strings(workspace):
    wb = Workbook()
    ws = wb.active
    ws.append(["姓名", "联系方式"])
    ws.append(["张三", f"{PHONE} / {EMAIL}"])
    wb.save(str(workspace / "x.xlsx"))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = f"卡号 {BANK}"
    prs.save(str(workspace / "p.pptx"))
    for name in ("x.xlsx", "p.pptx"):
        r = run_script("doc-sensitive-redact", "redact.py",
                       ["--input", workspace / name, "--outdir", workspace / "out"])
        assert r.returncode == 0, r.stderr + r.stdout
    assert (workspace / "out" / "x_REDACTED.xlsx").exists()
    assert (workspace / "out" / "p_REDACTED.pptx").exists()
    wb2 = load_workbook(str(workspace / "out" / "x_REDACTED.xlsx"))
    v = wb2.active["B2"].value
    assert isinstance(v, str) and PHONE not in v and EMAIL not in v and "■" in v
    wb2.close()
    prs2 = Presentation(str(workspace / "out" / "p_REDACTED.pptx"))
    all_text = "\n".join(sh.text_frame.text for sh in prs2.slides[0].shapes
                         if sh.has_text_frame)
    assert BANK not in all_text and "■" in all_text


# ---------- 复杂：隐藏部件底层验证（冻结缺陷场景） ----------

def test_redact_hidden_parts_lowlevel(workspace):
    inputs = workspace / "inputs"
    inputs.mkdir()
    build_hidden_docx(inputs / "hidden.docx")
    build_hidden_xlsx(inputs / "hidden.xlsx")
    build_hidden_pptx(inputs / "hidden.pptx")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", inputs, "--outdir", workspace / "redacted"])
    assert r.returncode == 0, r.stderr + r.stdout
    docx_needles = [EMAIL, DEL_EMAIL, CMT_EMAIL]
    xlsx_needles = [str(NUM_PHONE), SHEET_CMT_EMAIL, STR_EMAIL]
    pptx_needles = [NOTES_EMAIL]
    leaks = zip_scan(workspace / "redacted" / "hidden_REDACTED.docx", docx_needles)
    assert leaks == {}, leaks
    leaks = zip_scan(workspace / "redacted" / "hidden_REDACTED.xlsx", xlsx_needles)
    assert leaks == {}, leaks
    leaks = zip_scan(workspace / "redacted" / "hidden_REDACTED.pptx", pptx_needles)
    assert leaks == {}, leaks

    man = load_manifest(workspace / "redacted")
    by_ext = {os.path.splitext(os.path.basename(f["file"]))[1]: f for f in man["files"]}
    docx_counts = {x["type"]: x["count"] for x in by_ext[".docx"]["redactions"]}
    assert docx_counts.get("email") == 3
    locs = {loc for x in by_ext[".docx"]["redactions"] if x["type"] == "email"
            for loc in x["locations"]}
    assert any("word/comments.xml" in loc for loc in locs)
    xlsx_counts = {x["type"]: x["count"] for x in by_ext[".xlsx"]["redactions"]}
    assert xlsx_counts.get("phone") == 1
    assert xlsx_counts.get("email") == 2
    pptx_counts = {x["type"]: x["count"] for x in by_ext[".pptx"]["redactions"]}
    assert pptx_counts.get("email") == 1
    assert any("notesSlide" in loc for x in by_ext[".pptx"]["redactions"]
               for loc in x["locations"])

    doc = Document(str(workspace / "redacted" / "hidden_REDACTED.docx"))
    assert doc.paragraphs[0].text.startswith("Public unchanged text.")
    wb = load_workbook(str(workspace / "redacted" / "hidden_REDACTED.xlsx"))
    assert wb.active["A2"].value == "Public unchanged"
    assert wb.active["B1"].value == "■" * 11
    assert SHEET_CMT_EMAIL not in (wb.active["A2"].comment.text or "")
    wb.close()
    prs = Presentation(str(workspace / "redacted" / "hidden_REDACTED.pptx"))
    assert prs.slides[0].notes_slide.notes_text_frame.text == "■" * 12
    raw = json.dumps(man, ensure_ascii=False)
    for needle in docx_needles + xlsx_needles + pptx_needles:
        assert needle not in raw


def test_redact_review_items_not_masked(workspace):
    build_review_docx(workspace / "review.docx")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "review.docx", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    doc = Document(str(workspace / "out" / "review_REDACTED.docx"))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert BAD_ID in text and F15 in text
    man = load_manifest(workspace / "out")
    review_types = {x["type"] for x in man["files"][0]["needs_review"]}
    assert "id_card_invalid_checksum" in review_types
    assert "legacy_15_digit" in review_types
    assert man["files"][0]["redactions"] == []


# ---------- 复杂：嵌套目录、重名、类型过滤 ----------

def test_redact_nested_dirs_unique_outputs(workspace):
    root = workspace / "batch"
    (root / "a").mkdir(parents=True)
    (root / "b").mkdir(parents=True)
    d1 = Document()
    d1.add_paragraph("甲文件 " + EMAIL)
    d1.save(str(root / "a" / "report.docx"))
    d2 = Document()
    d2.add_paragraph("乙文件 " + H_EMAIL)
    d2.save(str(root / "b" / "report.docx"))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", root, "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    oa = workspace / "out" / "a__report_REDACTED.docx"
    ob = workspace / "out" / "b__report_REDACTED.docx"
    assert oa.exists() and ob.exists() and oa.read_bytes() != ob.read_bytes()
    ta = "\n".join(p.text for p in Document(str(oa)).paragraphs)
    tb = "\n".join(p.text for p in Document(str(ob)).paragraphs)
    assert EMAIL not in ta and H_EMAIL not in ta
    assert H_EMAIL not in tb and EMAIL not in tb
    assert "甲文件" in ta and "乙文件" in tb


def test_redact_outdir_inside_input_not_rescanned(workspace):
    root = workspace / "src"
    root.mkdir()
    build_plain_docx(root / "doc.docx")
    (root / "out").mkdir()
    shutil.copy(str(root / "doc.docx"), str(root / "out" / "old_REDACTED.docx"))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", root, "--outdir", root / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    man = load_manifest(root / "out")
    assert man["summary"]["files_processed"] == 1
    assert (root / "out" / "doc_REDACTED.docx").exists()
    assert (root / "out" / "old_REDACTED.docx").exists()


def test_redact_types_filter(workspace):
    doc = Document()
    doc.add_paragraph(f"邮箱 {EMAIL} 电话 {PHONE} 卡号 {BANK}")
    doc.save(str(workspace / "f.docx"))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "f.docx", "--outdir", workspace / "out",
                    "--types", "email"])
    assert r.returncode == 0, r.stderr + r.stdout
    text = "\n".join(p.text for p in Document(str(workspace / "out" / "f_REDACTED.docx")).paragraphs)
    assert EMAIL not in text
    assert PHONE in text and BANK in text
    man = load_manifest(workspace / "out")
    types = {x["type"] for x in man["files"][0]["redactions"]}
    assert types == {"email"}


def test_redact_output_name_collision_blocks(workspace):
    root = workspace / "coll"
    (root / "x" / "y").mkdir(parents=True)
    (root / "x__y").mkdir(parents=True)
    d1 = Document()
    d1.add_paragraph(EMAIL)
    d1.save(str(root / "x" / "y" / "z.docx"))
    d2 = Document()
    d2.add_paragraph(H_EMAIL)
    d2.save(str(root / "x__y" / "z.docx"))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", root, "--outdir", workspace / "out"])
    assert r.returncode == 2
    man = load_manifest(workspace / "out")
    failed = [f for f in man["files"] if f["failed"]]
    ok = [f for f in man["files"] if not f["failed"]]
    assert len(failed) == 1 and "output_name_collision" in failed[0]["error"]
    assert len(ok) == 1
    assert ok[0]["output"] and os.path.exists(ok[0]["output"])
    assert failed[0]["output"] is None


# ---------- 复杂：PDF 隐藏面（元数据/XMP/批注/附件/书签） ----------

def build_pdf_hidden(path):
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 72), f"body {EMAIL}", fontname="helv")
    p1.insert_text((72, 110), f"phone {PHONE}", fontname="helv")
    doc.set_metadata({"title": "Report", "author": PDF_META_EMAIL})
    xmp = (
        '<?xpacket?><x:xmpmeta xmlns:x="adobe:ns:meta/">'
        '<rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
        '<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/">'
        f"<dc:creator>{PDF_XMP_EMAIL}</dc:creator>"
        "</rdf:Description></rdf:RDF></x:xmpmeta>"
    )
    doc.set_xml_metadata(xmp)
    doc.embfile_add("attach.txt", f"see {PDF_ATTACH_EMAIL}".encode(), filename="attach.txt")
    a = p1.add_text_annot((200, 200), f"note {PDF_ANNOT_EMAIL}")
    a.update()
    doc.set_toc([[1, f"chapter by {PDF_TOC_EMAIL}", 1]])
    doc.save(str(path))
    doc.close()


def test_redact_pdf_hidden_surfaces(workspace):
    build_pdf_hidden(workspace / "in.pdf")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "in.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    out = workspace / "out" / "in_REDACTED.pdf"
    d2 = pymupdf.open(str(out))
    text = "".join(pg.get_text("text") for pg in d2)
    assert EMAIL not in text and PHONE not in text
    assert PDF_META_EMAIL not in (d2.metadata.get("author") or "")
    assert PDF_XMP_EMAIL not in (d2.get_xml_metadata() or "")
    assert d2.embfile_count() == 0
    assert PDF_TOC_EMAIL not in json.dumps(d2.get_toc(), ensure_ascii=False)
    annot_contents = []
    for pg in d2:
        for an in (pg.annots() or []):
            annot_contents.append(an.info.get("content") or "")
    assert all(PDF_ANNOT_EMAIL not in c for c in annot_contents)
    d2.close()
    raw = out.read_bytes()
    for needle in (EMAIL, PHONE, PDF_META_EMAIL, PDF_XMP_EMAIL,
                   PDF_ATTACH_EMAIL, PDF_ANNOT_EMAIL, PDF_TOC_EMAIL):
        assert needle.encode() not in raw, needle
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["pdf_removed_and_verified"] is True
    assert pf["removed_embedded_files"] == 1
    counts = {x["type"]: x["count"] for x in pf["redactions"]}
    assert counts.get("email") >= 5 and counts.get("phone") == 1


# ---------- 复杂：PDF 链接目标 / 附件批注 / 表单与图像限定 ----------
# 对应冻结证据 benchmarks/redaction-pdf-sidechannels（mailto 链接目标与
# 压缩 DOCX 附件批注残留 + link-only 独立案例的验收规则）

def build_pdf_links(path):
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 72), "Public unchanged body", fontname="helv")
    p.insert_link({"kind": pymupdf.LINK_URI,
                   "from": pymupdf.Rect(72, 100, 260, 118),
                   "uri": "mailto:" + PDF_LINK_EMAIL})
    p.insert_link({"kind": pymupdf.LINK_URI,
                   "from": pymupdf.Rect(72, 130, 260, 148),
                   "uri": PDF_NORMAL_URI})
    doc.save(str(path))
    doc.close()


def test_redact_pdf_link_targets_cleaned(workspace):
    build_pdf_links(workspace / "links.pdf")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "links.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    out = workspace / "out" / "links_REDACTED.pdf"
    d = pymupdf.open(str(out))
    uris = [l.get("uri") or "" for pg in d for l in pg.get_links()]
    assert uris and all(PDF_LINK_EMAIL not in u for u in uris), uris
    assert PDF_NORMAL_URI in uris
    assert "Public unchanged body" in d[0].get_text()
    d.close()
    assert PDF_LINK_EMAIL.encode() not in out.read_bytes()
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["failed"] is False
    assert pf["pdf_removed_and_verified"] is True
    counts = {x["type"]: x["count"] for x in pf["redactions"]}
    assert counts.get("email", 0) >= 1
    locs = [loc for x in pf["redactions"] if x["type"] == "email"
            for loc in x["locations"]]
    assert any(loc.startswith("link:") for loc in locs), locs
    assert PDF_LINK_EMAIL not in json.dumps(man, ensure_ascii=False)


def build_pdf_attachment(path):
    buf = io.BytesIO()
    docx = Document()
    docx.add_paragraph(ATTACH_PUBLIC)
    docx.add_paragraph("contact " + ATTACH_DOC_EMAIL)
    docx.save(buf)
    pdf = pymupdf.open()
    p = pdf.new_page()
    p.insert_text((72, 72), "Public unchanged body", fontname="helv")
    p.add_file_annot((72, 200), buf.getvalue(), "notes.docx",
                     desc="Supporting material")
    pdf.save(str(path))
    pdf.close()


def test_redact_pdf_attachment_annot_cleaned(workspace):
    build_pdf_attachment(workspace / "att.pdf")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "att.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    out = workspace / "out" / "att_REDACTED.pdf"
    d = pymupdf.open(str(out))
    assert "Public unchanged body" in d[0].get_text()
    blobs = []
    for pg in d:
        for a in (pg.annots() or []):
            if a.type[0] == pymupdf.PDF_ANNOT_FILE_ATTACHMENT:
                blobs.append(a.get_file())
    assert len(blobs) == 1
    blob = blobs[0]
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        names = z.namelist()
        leaks = [n for n in names if ATTACH_DOC_EMAIL.encode() in z.read(n)]
        doc_xml = z.read("word/document.xml").decode("utf-8")
    assert leaks == [], leaks
    assert ATTACH_PUBLIC in doc_xml
    d.close()
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["failed"] is False
    assert pf["pdf_removed_and_verified"] is True
    counts = {x["type"]: x["count"] for x in pf["redactions"]}
    assert counts.get("email", 0) >= 1
    locs = [loc for x in pf["redactions"] if x["type"] == "email"
            for loc in x["locations"]]
    assert any("file_attachment" in loc for loc in locs), locs
    assert ATTACH_DOC_EMAIL not in json.dumps(man, ensure_ascii=False)


def test_redact_pdf_form_and_image_qualification(workspace):
    pdf = pymupdf.open()
    p = pdf.new_page()
    p.insert_text((72, 72), "Public body text", fontname="helv")
    w = pymupdf.Widget()
    w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w.field_name = "contact"
    w.field_value = "form " + FORM_EMAIL
    w.rect = pymupdf.Rect(72, 150, 300, 170)
    p.add_widget(w)
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 16, 16))
    p.insert_image(pymupdf.Rect(72, 200, 136, 264), pixmap=pix)
    pdf.save(str(workspace / "formimg.pdf"))
    pdf.close()
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "formimg.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 0, r.stderr + r.stdout
    out = workspace / "out" / "formimg_REDACTED.pdf"
    d = pymupdf.open(str(out))
    vals = [wd.field_value or "" for wd in d[0].widgets()]
    assert vals and all(FORM_EMAIL not in v for v in vals), vals
    d.close()
    assert FORM_EMAIL.encode() not in out.read_bytes()
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["failed"] is False
    assert pf["pdf_removed_and_verified"] is False
    unv = pf["unverified_surfaces"]
    assert "form_fields" in unv and "image_pages_not_ocr_scanned" in unv
    review_types = {x["type"] for x in pf["needs_review"]}
    assert "form_fields_present_needs_manual_review" in review_types
    assert "images_present_not_ocr_scanned" in review_types
    assert FORM_EMAIL not in json.dumps(man, ensure_ascii=False)


# ---------- 异常 ----------

def test_redact_pdf_uncleanable_attachment_blocked(workspace):
    pdf = pymupdf.open()
    p = pdf.new_page()
    p.insert_text((72, 72), "Public body", fontname="helv")
    blob = bytes([0x81, 0x88, 0xFE]) + b"contact " + \
        ATTACH_DOC_EMAIL.encode() + b" ok" + bytes([0x80, 0x99, 0xFF])
    p.add_file_annot((72, 200), blob, "blob.bin", desc="binary")
    pdf.save(str(workspace / "bin.pdf"))
    pdf.close()
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "bin.pdf", "--outdir", workspace / "out"])
    assert r.returncode == 2
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["failed"] is True
    assert "UncleanableAttachment" in pf["error"], pf["error"]
    assert pf["output"] is None
    assert not (workspace / "out" / "bin_REDACTED.pdf").exists()
    assert ATTACH_DOC_EMAIL not in json.dumps(man, ensure_ascii=False)


def test_redact_mixed_batch_partial_failure(workspace):
    root = workspace / "mix"
    root.mkdir()
    build_plain_docx(root / "good.docx")
    (root / "broken.docx").write_bytes(b"not a real zip archive")
    build_hidden_xlsx(root / "good.xlsx")
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", root, "--outdir", workspace / "out"])
    assert r.returncode == 2
    man = load_manifest(workspace / "out")
    assert man["summary"]["files_failed"] == 1
    assert man["summary"]["files_processed"] == 2
    by_name = {os.path.basename(f["file"]): f for f in man["files"]}
    assert by_name["broken.docx"]["failed"] is True
    assert by_name["broken.docx"]["error"]
    assert by_name["broken.docx"]["output"] is None
    assert by_name["good.docx"]["failed"] is False
    assert os.path.exists(by_name["good.docx"]["output"])
    assert os.path.exists(by_name["good.xlsx"]["output"])
    assert not (workspace / "out" / "broken_REDACTED.docx").exists()


def test_redact_unredactable_binary_blocked(workspace):
    src = workspace / "blob.docx"
    build_plain_docx(src)
    tmp = workspace / "rebuilt.docx"
    with zipfile.ZipFile(str(src)) as zin, \
            zipfile.ZipFile(str(tmp), "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            zout.writestr(item, zin.read(item))
        zout.writestr("word/embeddings/blob.bin",
                      ("binary copy " + BIN_EMAIL).encode("utf-8"))
    os.replace(str(tmp), str(src))
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", src, "--outdir", workspace / "out"])
    assert r.returncode == 2
    man = load_manifest(workspace / "out")
    pf = man["files"][0]
    assert pf["failed"] is True
    assert "VerificationLeak" in pf["error"] or "verification_leak" in pf["error"]
    assert pf["output"] is None
    assert not (workspace / "out" / "blob_REDACTED.docx").exists()
    raw = json.dumps(man, ensure_ascii=False)
    assert BIN_EMAIL not in raw


def test_redact_input_errors(workspace):
    r = run_script("doc-sensitive-redact", "redact.py",
                   ["--input", workspace / "missing", "--outdir", workspace / "out"])
    assert r.returncode == 2
    f = workspace / "note.txt"
    f.write_text("hello", encoding="utf-8")
    r2 = run_script("doc-sensitive-redact", "redact.py",
                    ["--input", f, "--outdir", workspace / "out"])
    assert r2.returncode == 2
    empty = workspace / "empty"
    empty.mkdir()
    r3 = run_script("doc-sensitive-redact", "redact.py",
                    ["--input", empty, "--outdir", workspace / "out"])
    assert r3.returncode == 2
    r4 = run_script("doc-sensitive-redact", "redact.py",
                    ["--input", workspace / "x.docx", "--outdir", workspace / "out",
                     "--types", "email,ssn"])
    assert r4.returncode == 2
