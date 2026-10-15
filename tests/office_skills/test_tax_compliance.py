"""tax-compliance 黑盒测试：规则溯源/有效期门控/事实门控/金额数据错误，金标准手工推导。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from docx import Document  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402


def build_invoice_xlsx(path):
    """发票台账（含千分位字符串金额与整数）。

    金标准（手工推导）:
      金额合计 = 1000.50 + 2000 + 1234.56 = 4235.06 元
      税额合计 = 130.07 + 260 + 160.49 = 550.56 元
    """
    wb = Workbook()
    ws = wb.active
    ws.append(["发票号码", "金额", "税额"])
    ws.append(["F-001", 1000.50, 130.07])
    ws.append(["F-002", "2,000", 260])
    ws.append(["F-003", "1,234.56", 160.49])
    wb.save(str(path))


def build_contract_docx(path):
    doc = Document()
    doc.add_paragraph("技术开发合同")
    doc.add_paragraph("甲方委托乙方提供软件开发服务，合同金额以发票为准。")
    doc.save(str(path))


def build_flow_csv(path):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        f.write("日期,金额\n2026-07-05,3000\n2026-07-20,1500.25\n2026-08-01,999.99\n")


def write_biz(workspace, biz):
    p = workspace / "biz.json"
    p.write_text(json.dumps(biz, ensure_ascii=False), encoding="utf-8")
    return p


def run(workspace, materials, business, out="out"):
    return run_script("tax-compliance", "tax_review.py",
                      ["--materials", materials, "--business", business,
                       "--outdir", workspace / out])


def load_rep(workspace, out="out"):
    return json.loads((workspace / out / "rule-analysis.json").read_text(encoding="utf-8"))


def test_normal_full_facts(workspace):
    m = workspace / "materials"
    m.mkdir()
    build_invoice_xlsx(m / "发票.xlsx")
    build_contract_docx(m / "合同.docx")
    build_flow_csv(m / "流水.csv")
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"], "period": "2026Q3",
                                "jurisdiction": "上海市", "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = load_rep(workspace)
    assert rep["materials_status"] == "ready"
    assert rep["data_errors"] == []
    inv = next(x for x in rep["materials"] if x["class"] == "发票台账")["stats"]
    assert inv["amount_total_exact"] == "4235.06"
    assert inv["tax_total_exact"] == "550.56"
    assert inv["unit"] == "元（人民币，材料口径）"
    assert inv["invoice_rows"] == 3 and inv["error_cells"] == 0
    assert rep["business_facts"]["period_start"] == "2026-07-01"
    assert rep["business_facts"]["period_end"] == "2026-09-30"
    flow = next(x for x in rep["materials"] if x["class"] == "收支流水")["stats"]
    assert flow["flow_rows"] == 3
    status = {a["id"]: a["status"] for a in rep["rules"]}
    assert status["VAT-G-RATES-2026"] == "applicable"
    assert status["CIT-25"] == "applicable"
    assert status["INV-COMPLIANCE"] == "applicable"
    assert status["VAT-G-RATES-HIST"] == "not_applicable_for_period"
    vat = next(a for a in rep["rules"] if a["id"] == "VAT-G-RATES-2026")
    assert vat["source"]["regulation"] == "中华人民共和国增值税法"
    assert "第十条" in vat["source"]["article"]
    assert vat["source"]["url_verified"] is True
    assert vat["source"]["effective_from"] == "2026-01-01"
    assert vat["source"]["checked_on"] == "2026-10-15"
    assert rep["confirmations"] == []
    conf = json.loads((workspace / "out" / "confirmations.json").read_text(encoding="utf-8"))
    assert conf["confirmations"] == []
    wb = load_workbook(str(workspace / "out" / "materials-ledger.xlsx"))
    ws = wb["材料归集"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("文件", "分类", "关键统计")
    assert len(rows) == 4
    assert "数据错误" not in wb.sheetnames


def test_historical_period_routing(workspace):
    """2025Q4 业务期间：现行增值税法体系不适用，历史暂行条例体系适用。"""
    m = workspace / "m"
    m.mkdir()
    build_invoice_xlsx(m / "发票.xlsx")
    biz = write_biz(workspace, {"taxpayer_type": "小规模纳税人",
                                "activities": ["货物销售", "跨境服务", "货物出口"],
                                "period": "2025Q4", "jurisdiction": "上海市",
                                "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = load_rep(workspace)
    status = {a["id"]: a["status"] for a in rep["rules"]}
    assert status["VAT-S-LEVY-2026"] == "not_applicable_for_period"
    assert status["VAT-S-THRESHOLD-2026"] == "not_applicable_for_period"
    assert status["VAT-S-LEVY-HIST"] == "applicable"
    assert status["VAT-CROSS-FREE-HIST"] == "needs_confirmation"
    assert status["EXPORT-REFUND-HIST"] == "needs_confirmation"
    assert status["EXPORT-REFUND-2026"] == "not_applicable_for_period"
    hist = next(a for a in rep["rules"] if a["id"] == "VAT-S-LEVY-HIST")
    assert hist["source"]["regulation"] == "中华人民共和国增值税暂行条例（2017年修订）"
    assert hist["source"]["effective_to"] == "2025-12-31"
    cross = next(a for a in rep["rules"] if a["id"] == "VAT-CROSS-FREE-HIST")
    assert any("cross_border_service_details" in f for f in cross["missing_facts"])
    exp = next(a for a in rep["rules"] if a["id"] == "EXPORT-REFUND-HIST")
    assert any("export_documents" in f for f in exp["missing_facts"])


def test_preference_rules_gated(workspace):
    """2026Q3 小规模：法定3%征收率可确认；起征点/减按1%优惠条件未报齐须待确认。"""
    m = workspace / "m"
    m.mkdir()
    build_invoice_xlsx(m / "发票.xlsx")
    biz = write_biz(workspace, {"taxpayer_type": "小规模纳税人",
                                "activities": ["货物销售"], "period": "2026Q3",
                                "jurisdiction": "上海市", "entity_type": "个体工商户"})
    r = run(workspace, m, biz)
    assert r.returncode == 0, r.stderr + r.stdout
    rep = load_rep(workspace)
    status = {a["id"]: a["status"] for a in rep["rules"]}
    assert status["VAT-S-LEVY-2026"] == "applicable"
    assert status["CIT-25"] == "needs_confirmation"  # 个体工商户非企业主体
    thr = next(a for a in rep["rules"] if a["id"] == "VAT-S-THRESHOLD-2026")
    assert thr["status"] == "needs_confirmation"
    assert any("filing_period_type" in f for f in thr["missing_facts"])
    assert any("sales_computation_basis" in f for f in thr["missing_facts"])
    assert thr.get("conditions")
    lv1 = next(a for a in rep["rules"] if a["id"] == "VAT-S-LEVY1-2026")
    assert lv1["status"] == "needs_confirmation"
    assert any("sales_by_category" in f for f in lv1["missing_facts"])
    assert any("preference_election" in f for f in lv1["missing_facts"])
    # 报齐结构化事实后优惠规则方可确认（依据为业务申报事实）
    biz2 = write_biz(workspace, {
        "taxpayer_type": "小规模纳税人", "activities": ["货物销售"],
        "period": "2026Q3", "jurisdiction": "上海市", "entity_type": "个体工商户",
        "filing_period_type": "季", "sales_computation_basis": "不含税销售额",
        "sales_by_category": {"不动产类": 0, "其他": 100000},
        "preference_election": {"small_scale_levy_1": "享受"}})
    r2 = run(workspace, m, biz2, out="out2")
    assert r2.returncode == 0, r2.stderr + r2.stdout
    rep2 = load_rep(workspace, "out2")
    status2 = {a["id"]: a["status"] for a in rep2["rules"]}
    assert status2["VAT-S-THRESHOLD-2026"] == "applicable"
    assert status2["VAT-S-LEVY1-2026"] == "applicable"


def test_cross_border_and_export_2026(workspace):
    """跨境/出口：无交易条件与单证事实 → 待确认；报齐后可确认；来源为2026年新规。"""
    m = workspace / "m"
    m.mkdir()
    build_contract_docx(m / "合同.docx")
    base = {"taxpayer_type": "一般纳税人", "activities": ["跨境服务", "货物出口"],
            "period": "2026Q3", "jurisdiction": "上海市", "entity_type": "居民企业"}
    r = run(workspace, m, write_biz(workspace, base))
    assert r.returncode == 0, r.stderr + r.stdout
    rep = load_rep(workspace)
    status = {a["id"]: a["status"] for a in rep["rules"]}
    assert status["VAT-CROSS-ZERO-2026"] == "needs_confirmation"
    assert status["VAT-CROSS-EXEMPT-2026"] == "needs_confirmation"
    assert status["EXPORT-REFUND-2026"] == "needs_confirmation"
    assert status["INV-COMPLIANCE"] == "needs_confirmation"  # 无发票台账
    exp = next(a for a in rep["rules"] if a["id"] == "EXPORT-REFUND-2026")
    assert "2026年第5号" in exp["source"]["document_number"]
    assert "2012年第24号" in (exp["source"].get("supersedes") or "")
    assert exp["source"]["url_verified"] is True
    assert any("2012年第24号" not in f for f in [exp["source"]["regulation"]])
    full = dict(base, cross_border_service_details={
        "service": "软件服务", "buyer": "境外单位", "consumed_overseas": True},
        export_documents={"customs_declaration": True, "export_invoice": True,
                          "forex_receipt": True, "retention_documents": True})
    r2 = run(workspace, m, write_biz(workspace, full), out="out2")
    assert r2.returncode == 0, r2.stderr + r2.stdout
    rep2 = load_rep(workspace, "out2")
    status2 = {a["id"]: a["status"] for a in rep2["rules"]}
    assert status2["VAT-CROSS-ZERO-2026"] == "applicable"
    assert status2["EXPORT-REFUND-2026"] == "applicable"


def test_missing_context_and_invalid_amount(workspace):
    """缺期间/管辖地不得输出 applicable；NOT_A_NUMBER 必须报数据错误而非静默按0。"""
    m = workspace / "m"
    m.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["发票号码", "金额", "税额"])
    ws.append(["SYNTHETIC-001", "NOT_A_NUMBER", 13])
    wb.save(str(m / "虚构发票.xlsx"))
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"]})
    r = run(workspace, m, biz)
    assert r.returncode == 1  # 数据错误不得静默成功
    rep = load_rep(workspace)
    assert len(rep["data_errors"]) == 1
    e = rep["data_errors"][0]
    assert e["file"] == "虚构发票.xlsx" and e["row"] == 2 and e["column"] == "金额"
    assert e["value"] == "NOT_A_NUMBER" and "无法解析" in e["reason"]
    inv = next(x for x in rep["materials"] if x["class"] == "发票台账")["stats"]
    assert inv["error_cells"] == 1
    assert rep["business_facts"]["period_start"] is None
    for a in rep["rules"]:
        assert a["status"] != "applicable"
        assert a["status"] == "needs_confirmation"
        assert any("period" in f for f in a["missing_facts"])
        assert any("jurisdiction" in f for f in a["missing_facts"])
    wb2 = load_workbook(str(workspace / "out" / "materials-ledger.xlsx"))
    assert "数据错误" in wb2.sheetnames


def test_empty_materials_not_ready(workspace):
    m = workspace / "empty"
    m.mkdir()
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"], "period": "2026Q3",
                                "jurisdiction": "上海市", "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 1
    rep = load_rep(workspace)
    assert rep["materials_status"] == "not_ready"
    assert rep["materials"] == []
    for a in rep["rules"]:
        assert a["status"] != "applicable"
        if a["status"] == "needs_confirmation":
            assert any("申报材料" in f for f in a["missing_facts"])


def test_uncached_formula_amount(workspace):
    m = workspace / "m"
    m.mkdir()
    wb = Workbook()
    ws = wb.active
    ws.append(["发票号码", "金额", "税额"])
    ws.cell(row=2, column=1, value="F-100")
    ws.cell(row=2, column=2, value="=1+1")  # openpyxl写入的公式无缓存值
    ws.cell(row=2, column=3, value=26)
    wb.save(str(m / "公式发票.xlsx"))
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"], "period": "2026Q3",
                                "jurisdiction": "上海市", "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 1
    rep = load_rep(workspace)
    assert len(rep["data_errors"]) == 1
    assert "无缓存值" in rep["data_errors"][0]["reason"]
    inv = next(x for x in rep["materials"] if x["class"] == "发票台账")["stats"]
    assert inv["error_cells"] == 1 and inv["amount_total_exact"] == "0.00"


def test_corrupt_files_and_usage_errors(workspace):
    m = workspace / "m"
    m.mkdir()
    (m / "bad.xlsx").write_bytes(b"junk")
    (m / "bad.docx").write_bytes(b"junk")
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"], "period": "2026Q3",
                                "jurisdiction": "上海市", "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 1
    rep = load_rep(workspace)
    assert rep["materials_status"] == "not_ready"
    assert all(x["class"] == "读取失败" for x in rep["materials"])
    assert len(rep["data_errors"]) == 2
    assert all("读取失败" in e["reason"] for e in rep["data_errors"])
    # 坏business.json → 2
    (workspace / "bad.json").write_text("{not json", encoding="utf-8")
    r2 = run(workspace, m, workspace / "bad.json")
    assert r2.returncode == 2
    # 材料目录不存在 → 2
    r3 = run(workspace, workspace / "nope", biz)
    assert r3.returncode == 2


def test_invalid_period_never_applicable(workspace):
    m = workspace / "m"
    m.mkdir()
    build_invoice_xlsx(m / "发票.xlsx")
    biz = write_biz(workspace, {"taxpayer_type": "一般纳税人",
                                "activities": ["货物销售"], "period": "2026-02-30",
                                "jurisdiction": "上海市", "entity_type": "居民企业"})
    r = run(workspace, m, biz)
    assert r.returncode == 0
    rep = load_rep(workspace)
    assert rep["business_facts"]["period_start"] is None
    assert rep["business_facts"]["period_raw"] == "2026-02-30"
    for a in rep["rules"]:
        assert a["status"] == "needs_confirmation"
        assert any("业务期间" in f for f in a["missing_facts"])

