"""sheet-consolidate-reconcile 测试：三源对账金标准。"""
import json
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
from helpers import run_script  # noqa: E402

from openpyxl import Workbook, load_workbook  # noqa: E402


def build_sources(ws):
    # A: 华东区 编码/销售额(万元)  P001-P004 + 合计行
    wa = Workbook(); s = wa.active
    s.append(["编码", "销售额(万元)"])
    s.append(["P001", 100])
    s.append(["P002", 200])
    s.append(["P003", 300])
    s.append(["P004", 400])
    s.append(["合计", 1000])
    wa.save(str(ws / "east.xlsx"))
    # B: 华南区 产品编码/销售额  P002-P004, P003不同值, P004重复
    wb_ = Workbook(); s = wb_.active
    s.append(["产品编码", "销售额"])
    s.append(["P002", 200])
    s.append(["P003", 350])   # 与华东不一致
    s.append(["P004", 400])
    s.append(["P004", 420])   # 重复主键
    wb_.save(str(ws / "south.xlsx"))
    # C: 华北区 产品编码/销售额  P001, P003, P004, P005
    wc = Workbook(); s = wc.active
    s.append(["产品编码", "销售额"])
    s.append(["P001", 100])
    s.append(["P003", 300])
    s.append(["P004", 400])
    s.append(["P005", 500])
    wc.save(str(ws / "north.xlsx"))


CFG = {
    "primary_key": "产品编码",
    "metrics": ["销售额"],
    "sources": [
        {"file": "east.xlsx", "source_name": "华东区",
         "column_map": {"产品编码": "编码", "销售额": "销售额(万元)"}},
        {"file": "south.xlsx", "source_name": "华南区", "column_map": {}},
        {"file": "north.xlsx", "source_name": "华北区", "column_map": {}},
    ],
}


def test_consolidate_golden(workspace):
    build_sources(workspace)
    (workspace / "cfg.json").write_text(json.dumps(CFG, ensure_ascii=False), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfg.json", "--output", workspace / "merged.xlsx",
                    "--report", workspace / "rep.json"])
    assert r.returncode == 0, r.stderr + r.stdout
    rep = json.loads((workspace / "rep.json").read_text(encoding="utf-8"))
    ms = rep["match_stats"]
    # 手工金标准: 唯一主键 P001-P005 = 5; 三源齐全 = P003,P004 = 2
    assert ms["n_unique_pk"] == 5
    assert ms["n_present_all"] == 2
    # 单源: P005(仅华北) = 1; P001 是双源(华东+华北), P002 双源(华东+华南)
    assert ms["n_single_source"] == 1
    # 不一致主键: P003(300 vs 350) 与 P004(400/400/420 重复导致多值) = 2
    assert ms["n_mismatch_pk"] == 2
    assert ms["match_rate"] == 0.4
    # 重复主键: 华南 P004
    assert len(rep["duplicates"]) == 1
    assert rep["duplicates"][0]["pk"] == "P004"
    assert rep["duplicates"][0]["source"] == "华南区"
    assert rep["duplicates"][0]["count"] == 2
    # 合计行剔除: 华东 合计
    assert len(rep["skipped_total_rows"]) == 1
    assert rep["skipped_total_rows"][0]["source"] == "华东区"
    # 源缺失: 华南缺 P001/P005; 华北缺 P002; 华东缺 P005
    sm = rep["source_missing"]
    assert set(sm["华南区"]) == {"P001", "P005"}
    assert set(sm["华北区"]) == {"P002"}
    assert set(sm["华东区"]) == {"P005"}
    # 差异明细含 P003 的两个值
    m303 = [m for m in rep["mismatches"] if m["pk"] == "P003"][0]
    assert 300 in m303["values"]["华东区"] and 350 in m303["values"]["华南区"]

    # 合并表可打开且行数正确
    wb = load_workbook(str(workspace / "merged.xlsx"))
    ws = wb["合并汇总"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 5
    hdr = [c.value for c in ws[1]]
    assert "销售额@华东区" in hdr and "销售额@华北区" in hdr  # 口径来源显式标注
    det = wb["源明细"]
    det_rows = list(det.iter_rows(min_row=2, values_only=True))
    assert len(det_rows) == 4 + 4 + 4  # 华东4明细(合计已剔) + 华南4 + 华北4
    dup = wb["重复主键"]
    dup_rows = list(dup.iter_rows(min_row=2, values_only=True))
    assert len(dup_rows) == 1 and dup_rows[0][1] == "P004"


def test_determinism(workspace):
    build_sources(workspace)
    (workspace / "cfg.json").write_text(json.dumps(CFG, ensure_ascii=False), encoding="utf-8")
    outs = []
    for i in (1, 2):
        run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "cfg.json", "--output", workspace / f"m{i}.xlsx",
                    "--report", workspace / f"r{i}.json"])
        rep = json.loads((workspace / f"r{i}.json").read_text(encoding="utf-8"))
        outs.append(json.dumps(rep["match_stats"]))
    assert outs[0] == outs[1]


def test_abnormal(workspace):
    build_sources(workspace)
    cfg = dict(CFG)
    cfg["sources"] = [{"file": "nope.xlsx", "source_name": "X", "column_map": {}}]
    (workspace / "bad.json").write_text(json.dumps(cfg), encoding="utf-8")
    r = run_script("sheet-consolidate-reconcile", "consolidate.py",
                   ["--config", workspace / "bad.json", "--output", workspace / "m.xlsx",
                    "--report", workspace / "r.json"])
    assert r.returncode == 2
    # 一个源损坏, 其余继续
    cfg2 = dict(CFG)
    cfg2["sources"] = [{"file": "nope.xlsx", "source_name": "坏源", "column_map": {}}] + CFG["sources"][1:]
    (workspace / "cfg2.json").write_text(json.dumps(cfg2), encoding="utf-8")
    r2 = run_script("sheet-consolidate-reconcile", "consolidate.py",
                    ["--config", workspace / "cfg2.json", "--output", workspace / "m2.xlsx",
                    "--report", workspace / "r2.json"])
    assert r2.returncode == 0
    rep2 = json.loads((workspace / "r2.json").read_text(encoding="utf-8"))
    assert "坏源" in rep2["failed_sources"]
    assert rep2["match_stats"]["n_sources"] == 2
