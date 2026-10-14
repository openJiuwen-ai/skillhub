#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx-version-diff: 两个 .docx 的结构级对比。

用法:
  python docx_diff.py --old v1.docx --new v2.docx --report diff.json [--changes-doc changes.docx]

退出码: 0=完成(包括无差异) 2=输入错误
"""
import argparse
import difflib
import json
import sys

from docx import Document
from docx.oxml.ns import qn


def body_paragraphs(doc):
    """正文顶层段落(不含表格内), 按顺序。"""
    paras = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            # 顶层段落(不在表格内)
            paras.append(child)
    texts = []
    for p in paras:
        t = "".join(n.text or "" for n in p.iter(qn("w:t")))
        texts.append(t)
    return texts


def tables_of(doc):
    out = []
    for t in doc.tables:
        rows = []
        for row in t.rows:
            rows.append([c.text for c in row.cells])
        out.append(rows)
    return out


def headers_footers(doc):
    out = []
    for si, sec in enumerate(doc.sections):
        entry = {"section": si}
        try:
            entry["header"] = "\n".join(p.text for p in sec.header.paragraphs)
        except Exception:
            entry["header"] = ""
        try:
            entry["footer"] = "\n".join(p.text for p in sec.footer.paragraphs)
        except Exception:
            entry["footer"] = ""
        pg = sec.page_width, sec.page_height, sec.left_margin, sec.top_margin
        entry["page"] = [str(x) for x in pg]
        out.append(entry)
    return out


def diff_paragraphs(old, new):
    changes = []
    sm = difflib.SequenceMatcher(a=old, b=new, autojunk=False)
    n_mod = n_add = n_rem = 0
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace":
            # 逐段配对输出
            for k in range(max(i2 - i1, j2 - j1)):
                b = old[i1 + k] if i1 + k < i2 else None
                a = new[j1 + k] if j1 + k < j2 else None
                if b is not None and a is not None:
                    changes.append({"type": "paragraph", "change": "modified",
                                    "location": f"body paragraph {j1 + k}",
                                    "before": b, "after": a})
                    n_mod += 1
                elif b is not None:
                    changes.append({"type": "paragraph", "change": "removed",
                                    "location": f"old paragraph {i1 + k}",
                                    "before": b, "after": None})
                    n_rem += 1
                else:
                    changes.append({"type": "paragraph", "change": "added",
                                    "location": f"body paragraph {j1 + k}",
                                    "before": None, "after": a})
                    n_add += 1
        elif tag == "delete":
            for k in range(i1, i2):
                changes.append({"type": "paragraph", "change": "removed",
                                "location": f"old paragraph {k}",
                                "before": old[k], "after": None})
                n_rem += 1
        elif tag == "insert":
            for k in range(j1, j2):
                changes.append({"type": "paragraph", "change": "added",
                                "location": f"body paragraph {k}",
                                "before": None, "after": new[k]})
                n_add += 1
    return changes, n_mod, n_add, n_rem


def diff_tables(old, new):
    changes = []
    n_cell = 0
    n = min(len(old), len(new))
    for idx in range(n):
        to, tn = old[idx], new[idx]
        if len(to) != len(tn) or (to and tn and len(to[0]) != len(tn[0])):
            changes.append({"type": "table", "change": "structural",
                            "location": f"table {idx}",
                            "before": f"{len(to)}x{len(to[0]) if to else 0}",
                            "after": f"{len(tn)}x{len(tn[0]) if tn else 0}"})
            continue
        for r in range(len(to)):
            for c in range(len(to[r])):
                if to[r][c] != tn[r][c]:
                    changes.append({"type": "table_cell", "change": "modified",
                                    "location": f"table {idx} row {r} col {c}",
                                    "before": to[r][c], "after": tn[r][c]})
                    n_cell += 1
    for idx in range(n, len(old)):
        changes.append({"type": "table", "change": "removed", "location": f"table {idx}",
                        "before": f"{len(old[idx])} rows", "after": None})
    for idx in range(n, len(new)):
        changes.append({"type": "table", "change": "added", "location": f"table {idx}",
                        "before": None, "after": f"{len(new[idx])} rows"})
    return changes, n_cell


def diff_headers(old, new):
    changes = []
    n = 0
    for idx in range(max(len(old), len(new))):
        o = old[idx] if idx < len(old) else None
        nw = new[idx] if idx < len(new) else None
        if o is None or nw is None:
            changes.append({"type": "section", "change": "added" if o is None else "removed",
                            "location": f"section {idx}"})
            n += 1
            continue
        for key in ("header", "footer"):
            if o[key] != nw[key]:
                changes.append({"type": key, "change": "modified",
                                "location": f"section {idx} {key}",
                                "before": o[key], "after": nw[key]})
                n += 1
        if o["page"] != nw["page"]:
            changes.append({"type": "section", "change": "modified",
                            "location": f"section {idx} page setup",
                            "before": o["page"], "after": nw["page"]})
            n += 1
    return changes, n


def write_changes_docx(changes, path, old_name, new_name):
    from docx import Document as D
    doc = D()
    doc.add_heading("文档版本变更清单", level=1)
    doc.add_paragraph(f"旧版: {old_name}")
    doc.add_paragraph(f"新版: {new_name}")
    if not changes:
        doc.add_paragraph("两版内容一致，无变更。")
    for ch in changes:
        loc = ch.get("location", "")
        head = f"[{ch['type']}|{ch['change']}] {loc}"
        doc.add_heading(head, level=2)
        if ch.get("before") is not None:
            doc.add_paragraph(f"旧: {ch['before']}")
        if ch.get("after") is not None:
            doc.add_paragraph(f"新: {ch['after']}")
    doc.save(path)


def main():
    ap = argparse.ArgumentParser(description="Word两版结构级对比")
    ap.add_argument("--old", required=True)
    ap.add_argument("--new", required=True)
    ap.add_argument("--report")
    ap.add_argument("--changes-doc")
    args = ap.parse_args()

    try:
        old_doc = Document(args.old)
        new_doc = Document(args.new)
    except Exception as e:
        print(f"ERROR: cannot open input: {e}", file=sys.stderr)
        return 2

    changes = []
    summary = {}

    pc, pm, pa, pr = diff_paragraphs(body_paragraphs(old_doc), body_paragraphs(new_doc))
    changes.extend(pc)
    summary["paragraph_modified"] = pm
    summary["paragraph_added"] = pa
    summary["paragraph_removed"] = pr

    tc, tcell = diff_tables(tables_of(old_doc), tables_of(new_doc))
    changes.extend(tc)
    summary["table_cell_modified"] = tcell

    hc, hn = diff_headers(headers_footers(old_doc), headers_footers(new_doc))
    changes.extend(hc)
    summary["header_footer_changed"] = hn
    summary["section_changed"] = sum(1 for h in hc if h["type"] == "section")

    identical = not changes
    report = {
        "old": args.old,
        "new": args.new,
        "identical": identical,
        "summary": summary,
        "changes": changes,
    }
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
    if args.changes_doc:
        write_changes_docx(changes, args.changes_doc, args.old, args.new)
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
