#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""input-preflight-check: 清单三态核对 + 分析前数据范围校验。

checklist 模式：材料清单三态核对（ready/missing/needs_confirm），损坏文件不得因
文件名命中而判 ready。data 模式：必需列/重复表头/单位/数值列/真实日历日期范围/
最少行数/公式缓存校验；不可解析或缺失的日期/数值一律阻断，不静默通过。
verdict=block 时不产出分析结论。
"""
import argparse
import datetime as _dt
import json
import math
import os
import re
import sys

TEXT_EXTS = (".md", ".txt", ".csv")


def sniff_content(path):
    """返回 (可提取文本, 可读标志, 错误说明)。可读=False 表示文件存在但解析失败。"""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in TEXT_EXTS:
            with open(path, encoding="utf-8", errors="ignore") as f:
                return f.read(200000), True, None
        if ext == ".docx":
            from docx import Document
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs), True, None
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, data_only=True, read_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(max_row=50, values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            wb.close()
            return "\n".join(parts), True, None
    except Exception as e:
        return "", False, "%s: %s" % (type(e).__name__, e)
    return "", True, None  # 未知扩展名：不提取内容，也不判损坏


def parse_calendar_date(v):
    """真实日历日期解析。不存在的日期（如 2026-02-30）与无法解析的值返回 None。"""
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    if v is None or isinstance(v, bool):
        return None
    s = str(v).strip().replace("/", "-").replace(".", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if not m:
        return None
    try:
        return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _short(v, n=40):
    s = str(v)
    return s if len(s) <= n else s[:n] + "..."


def dump(path, obj):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    summary = {"verdict": obj["verdict"], "reasons": obj.get("reasons", [])[:5]}
    # 清单模式的业务完整性层结论一并暴露到 stdout（数据模式不带该层，保持旧输出）
    if "business_verdict" in obj:
        summary["business_verdict"] = obj["business_verdict"]
        summary["coverage_warnings"] = obj.get("coverage_warnings", [])
    print(json.dumps(summary, ensure_ascii=False))


_CN_NUM = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
_TEMPORAL_RE = re.compile(r"(?:最近|近)\s*([0-9]+|[一二三四五六七八九十两]+)\s*年")
_YEAR_RE = re.compile(r"(?:19|20)\d{2}")
_VERDICT_RANK = {"pass": 0, "confirm_needed": 1, "needs_confirmation": 1,
                 "block": 2}
_VERDICT_LABEL = {0: "pass", 1: "needs_confirmation", 2: "block"}
_CHECKLIST_NOTE = ("verdict为关键词匹配层结论（向后兼容保留）；business_verdict为"
                   "业务完整性层结论（跨年度覆盖核对）。有跨年度/近N年要求但未声明"
                   "具体年份或证据不足时 business_verdict=needs_confirmation/block，"
                   "不得据关键词层 pass 放行；block时不产出分析结论")


def _cn_int(tok):
    if tok.isdigit():
        return int(tok)
    if tok in _CN_NUM:
        return _CN_NUM[tok]
    if "十" in tok:  # 十一 / 二十 / 二十三 等
        parts = tok.split("十")
        tens = _CN_NUM.get(parts[0], 1) if parts[0] else 1
        ones = _CN_NUM.get(parts[1], 0) if len(parts) > 1 and parts[1] else 0
        return tens * 10 + ones
    return None


def detect_temporal_years(name):
    """检测名称中的"近N年/最近N年"跨年度要求，返回 N 或 None。

    只识别相对年限，不推断基准年份（"近三年"相对哪一年由人工核验）。
    """
    if not name:
        return None
    m = _TEMPORAL_RE.search(str(name))
    if not m:
        return None
    n = _cn_int(m.group(1))
    return n if n and n > 0 else None


def _normalize_years(raw):
    """把 required_years 归一化为 4 位年份字符串列表，返回 (years, invalid)。"""
    if not isinstance(raw, list):
        return [], [raw]
    years, invalid = [], []
    for y in raw:
        s = str(y).strip()
        if _YEAR_RE.fullmatch(s):
            years.append(s)
        else:
            invalid.append(y)
    return years, invalid


def _year_token(year):
    return re.compile(r"(?<!\d)" + re.escape(year) + r"(?!\d)")


def compute_coverage(item, matched):
    """业务完整性层：按配置的材料覆盖核对。返回 (coverage_state, coverage, warnings)。

    核对的是"每个要求年份是否配置了可核对的对应材料"，**不是**审计真实性验证。

    - 声明 required_years：逐年核对是否有"该项已命中的可读证据文件"，文件名含该
      年份（独立数字标记）才算证据；缺年 → incomplete。
    - 配置不自洽（非法年份 / 重复年份 / 与名称"近N年"数量不一致）→ needs_confirmation，
      不得因存在有效年份就放行。
    - 单文件名称并列多个要求年份（合并报告）默认不算多份证据 → needs_confirmation；
      仅在显式 `allow_merged_years: true` 时按可核对证据接受。
    - 未声明但名称含"近N年"：不猜基准年份 → needs_confirmation。
    - 两者皆无：None（普通清单，保持旧行为）。
    来源证据保留在 coverage["evidence"]。
    """
    name = item.get("name")
    raw = item.get("required_years")
    temporal_n = detect_temporal_years(name)
    warnings = []
    if raw is not None:
        years, invalid = _normalize_years(raw)
        if invalid:
            warnings.append("required_years含无法识别年份: %s（应为4位年份），"
                            "配置不可靠，无法据此核对覆盖" % [str(x) for x in invalid])
        if not years:
            warnings.append("清单项'%s'声明了required_years但无可识别年份，"
                            "无法核对覆盖，需人工核验" % name)
            return "needs_confirmation", {
                "required_years": None, "covered_years": [],
                "missing_years": [], "evidence": {}}, warnings
        # 配置自洽性：重复年份 / 与"近N年"数量不一致
        config_conflict = bool(invalid)
        if len(set(years)) != len(years):
            warnings.append("required_years存在重复年份: %s，无法核对覆盖，需人工核验"
                            % years)
            config_conflict = True
        if temporal_n and len(years) != temporal_n:
            warnings.append("清单项'%s'为'近%d年'但required_years为%d个年份，数量不一致，"
                            "无法据此判齐备，需人工核验"
                            % (name, temporal_n, len(years)))
            config_conflict = True
        allow_merged = bool(item.get("allow_merged_years", False))
        per_year = {}
        for y in years:
            pat = _year_token(y)
            per_year[y] = [mf for mf in matched
                           if pat.search(os.path.basename(mf))]

        def _req_years_in(fname):
            b = os.path.basename(fname)
            return [y for y in set(years) if _year_token(y).search(b)]

        single = {f for f in matched if len(_req_years_in(f)) == 1}
        missing = [y for y in years if not per_year[y]]
        merged_only = []
        if not allow_merged:
            for y in years:
                if per_year[y] and not any(f in single for f in per_year[y]):
                    merged_only.append(y)
        covered = [y for y in years
                   if per_year[y] and y not in missing and y not in merged_only]
        evidence = {y: per_year[y] for y in years if per_year[y]}
        coverage = {"required_years": years, "covered_years": covered,
                    "missing_years": missing, "evidence": evidence,
                    "merged_years_pending": merged_only}
        if config_conflict:
            return "needs_confirmation", coverage, warnings
        if missing:
            warnings.append("跨年度材料不齐：'%s'要求年份 %s，缺少 %s"
                            % (name, years, missing))
            return "incomplete", coverage, warnings
        if merged_only:
            warnings.append("清单项'%s'存在跨越多要求年份的单文件（名称并列 %s），"
                            "单文件名并列年份不等于多份证据；未配置 allow_merged_years，"
                            "不能静默通过，请拆分或显式允许合并并人工核对"
                            % (name, merged_only))
            return "needs_confirmation", coverage, warnings
        return "complete", coverage, warnings
    if temporal_n:
        warnings.append("清单项'%s'含'近%d年'要求但未声明具体年份(required_years)，"
                        "无法核对跨年度覆盖；不猜测基准年份，请提供 required_years "
                        "或人工核验" % (name, temporal_n))
        return "needs_confirmation", {
            "temporal_requirement_years": temporal_n, "required_years": None,
            "covered_years": [], "missing_years": [], "evidence": {}}, warnings
    return None, None, warnings


def mode_checklist(args):
    try:
        with open(args.checklist, encoding="utf-8") as f:
            cfg = json.load(f)
    except Exception as e:
        print("ERROR: bad checklist: %s" % e, file=sys.stderr)
        return 2
    if not isinstance(cfg, dict):
        print("ERROR: bad checklist: 顶层必须是对象", file=sys.stderr)
        return 2
    if not os.path.isdir(args.dir):
        print("ERROR: materials dir not found: %s" % args.dir, file=sys.stderr)
        return 2
    items = cfg.get("items")
    if not isinstance(items, list) or not items:
        dump(args.report, {"mode": "checklist", "verdict": "block",
                           "business_verdict": "block", "items": [],
                           "files_scanned": 0, "coverage_warnings": [],
                           "reasons": ["清单为空：没有任何条目"],
                           "note": _CHECKLIST_NOTE})
        return 0
    # 清单自身问题：缺 id / 重复 id / 条目无任何匹配模式
    ids = [i.get("id") for i in items]
    seen, dups = set(), []
    for i in ids:
        if i is None:
            continue
        if i in seen and i not in dups:
            dups.append(i)
        seen.add(i)
    config_reasons = []
    missing_ids = [i for i in items if i.get("id") is None]
    if missing_ids:
        config_reasons.append("清单项缺少id: %d 个" % len(missing_ids))
    if dups:
        config_reasons.append("清单项id重复: %s" % dups)
    no_pattern = [i.get("id") or i.get("name") for i in items
                  if not i.get("patterns") and not i.get("fuzzy")]
    if no_pattern:
        config_reasons.append("清单项未配置patterns/fuzzy: %s" % no_pattern)
    files = []
    for root, _dirs, names in os.walk(args.dir):
        for n in sorted(names):
            files.append(os.path.join(root, n))
    if not files:
        config_reasons.append("材料目录为空：未找到任何文件")
    cache = {}
    results = []
    reasons = list(config_reasons)
    for item in items:
        pats = item.get("patterns", [])
        fuzzy = item.get("fuzzy", [])
        required = item.get("required", True)
        matched, unreadable, fuzzy_matched = [], [], []
        for f in files:
            fname = os.path.basename(f)
            if f not in cache:
                cache[f] = sniff_content(f)
            content, readable, err = cache[f]
            rel = os.path.relpath(f, args.dir)
            if pats and all((p in fname) or (p in content) for p in pats):
                if readable:
                    matched.append(rel)
                else:
                    unreadable.append({"file": rel, "error": err})
            elif fuzzy and any(p in fname for p in fuzzy):
                fuzzy_matched.append(rel)
        if matched:
            state = "ready"
        elif unreadable:
            state = "needs_confirm"
            reasons.append("无法读取: %s（文件命中模式但不可读: %s）"
                           % (item.get("name"), [u["file"] for u in unreadable[:3]]))
        elif fuzzy_matched:
            state = "needs_confirm"
            reasons.append("需确认: %s（文件名近似命中但内容模式未验证: %s）"
                           % (item.get("name"), fuzzy_matched[:3]))
        else:
            state = "missing"
            if required:
                reasons.append("缺失: %s（无文件命中全部模式 %s）"
                               % (item.get("name"), pats))
        cov_state, cov_detail, cov_warnings = compute_coverage(item, matched)
        results.append({"id": item.get("id"), "name": item.get("name"),
                        "state": state, "matched_files": matched[:5],
                        "unreadable_files": unreadable[:5],
                        "fuzzy_files": fuzzy_matched[:5],
                        "required": required,
                        "coverage_state": cov_state, "coverage": cov_detail,
                        "coverage_warnings": cov_warnings})
    blocking = bool(dups) or bool(missing_ids) or bool(no_pattern)
    for r, item in zip(results, items):
        required = item.get("required", True)
        if required and r["state"] == "missing":
            blocking = True
        if required and r["state"] == "needs_confirm" and r["unreadable_files"]:
            blocking = True  # 必需材料仅命中损坏文件：不可放行
    confirm = any(r["state"] == "needs_confirm" for r in results)
    verdict = "block" if blocking else ("confirm_needed" if (confirm or reasons)
                                        else "pass")
    # 业务完整性层：绝不比关键词层更宽松；跨年度证据不足时收紧结论
    business_rank = _VERDICT_RANK[verdict]
    coverage_warnings = []
    for r, item in zip(results, items):
        coverage_warnings.extend(r.get("coverage_warnings", []))
        cs = r.get("coverage_state")
        required = item.get("required", True)
        if cs in ("incomplete", "needs_confirmation"):
            business_rank = max(business_rank, 2 if required else 1)
    business_verdict = _VERDICT_LABEL[business_rank]
    report = {"mode": "checklist", "verdict": verdict,
              "business_verdict": business_verdict, "items": results,
              "files_scanned": len(files), "reasons": reasons,
              "coverage_warnings": coverage_warnings, "note": _CHECKLIST_NOTE}
    dump(args.report, report)
    return 0


def mode_data(args):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    try:
        with open(args.requirements, encoding="utf-8") as f:
            req = json.load(f)
    except Exception as e:
        print("ERROR: bad requirements: %s" % e, file=sys.stderr)
        return 2
    if not isinstance(req, dict):
        print("ERROR: bad requirements: 顶层必须是对象", file=sys.stderr)
        return 2
    try:
        wb = load_workbook(args.data, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        sheet_name = ws.title
        wb.close()
    except Exception as e:
        print("ERROR: cannot read data: %s" % e, file=sys.stderr)
        return 2
    if not rows:
        dump(args.report, {"mode": "data", "verdict": "block", "reasons": ["空表"],
                           "note": "verdict=block时不产出分析结论"})
        return 0
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    data_rows = [r for r in rows[1:] if any(c is not None and str(c).strip()
                                            for c in r)]
    checks = []
    reasons = []
    blocking = False

    # 公式无缓存值：未读取到的范围不得静默通过
    uncached = []
    try:
        wb_raw = load_workbook(args.data, data_only=False)
        raw_rows = list(wb_raw.active.iter_rows(values_only=True))
        for ri, (raw, cached) in enumerate(zip(raw_rows, rows), start=1):
            for ci, (rv, cv) in enumerate(zip(raw, cached), start=1):
                if isinstance(rv, str) and rv.startswith("=") and cv is None:
                    uncached.append("%s%d" % (get_column_letter(ci), ri))
        wb_raw.close()
    except Exception:
        uncached = []
    if uncached:
        checks.append({"check": "formula_cache", "ok": False,
                       "count": len(uncached),
                       "formula_cells_without_cache": uncached[:20]})
        reasons.append("公式无缓存值: %d 个单元格无法读取数值" % len(uncached))
        blocking = True

    # requirements 未配置任何校验项：不得静默通过
    has_any_check = (bool(req.get("columns")) or bool(req.get("units"))
                     or bool(req.get("date_column"))
                     or bool(req.get("numeric_columns")) or "min_rows" in req)
    if not has_any_check:
        reasons.append("requirements为空：未配置任何校验项")
        blocking = True

    # 重复表头：列引用歧义
    non_empty = [h for h in header if h]
    dup_headers = sorted({h for h in non_empty if non_empty.count(h) > 1})
    checks.append({"check": "duplicate_headers", "ok": not dup_headers,
                   "duplicates": dup_headers})
    if dup_headers:
        reasons.append("表头重复: %s（列引用歧义）" % dup_headers)
        blocking = True

    # 必需列
    missing_cols = [c for c in req.get("columns", []) if c not in header]
    checks.append({"check": "required_columns", "ok": not missing_cols,
                   "missing": missing_cols})
    if missing_cols:
        reasons.append("必需列缺失: %s" % missing_cols)
        blocking = True

    # 单位口径
    for col, unit in (req.get("units") or {}).items():
        if col not in header:
            continue
        idx = header.index(col)
        vals = [r[idx] for r in data_rows[:20]
                if idx < len(r) and isinstance(r[idx], (int, float))
                and not isinstance(r[idx], bool)]
        header_has_unit = unit in header[idx]
        checks.append({"check": "unit", "column": col, "expected_unit": unit,
                       "header": header[idx], "numeric_values": len(vals),
                       "ok": header_has_unit or not vals,
                       "note": None if header_has_unit else
                       "表头未注明单位且列为数值，需确认口径"})
        if not header_has_unit and vals:
            reasons.append("单位待确认: %s 列要求单位'%s'但表头'%s'未注明"
                           % (col, unit, header[idx]))

    # 数值列：不可解析/缺失/非有限数一律阻断
    for col in req.get("numeric_columns") or []:
        ncheck = {"check": "numeric", "column": col}
        if col not in header:
            ncheck.update(ok=False, error="column_not_in_header")
            reasons.append("数值列缺失: %s" % col)
            blocking = True
        else:
            idx = header.index(col)
            bad, miss = [], []
            for ri, r in enumerate(data_rows, start=2):
                v = r[idx] if idx < len(r) else None
                if v is None or (isinstance(v, str) and not v.strip()):
                    miss.append(ri)
                elif isinstance(v, bool):
                    bad.append({"row": ri, "value": str(v)})
                elif isinstance(v, (int, float)):
                    if not math.isfinite(v):
                        bad.append({"row": ri, "value": str(v)})
                else:
                    try:
                        float(str(v).strip())
                    except ValueError:
                        bad.append({"row": ri, "value": _short(v)})
            ok = not bad and not miss
            ncheck.update(ok=ok, rows=len(data_rows),
                          unparseable_count=len(bad), unparseable=bad[:20],
                          missing_count=len(miss), missing_rows=miss[:20])
            if bad:
                reasons.append("数值不可解析: %s 列 %d 行" % (col, len(bad)))
            if miss:
                reasons.append("数值缺失: %s 列 %d 行" % (col, len(miss)))
            if not ok:
                blocking = True
        checks.append(ncheck)

    # 日期列：真实日历日期解析；不可解析/缺失一律阻断，越界按配置（默认阻断）
    dc = req.get("date_column")
    if dc:
        dr = req.get("date_range")
        dcheck = {"check": "date_range", "column": dc}
        if dc not in header:
            dcheck.update(ok=False, error="date_column_not_in_header")
            reasons.append("日期列缺失: %s" % dc)
            blocking = True
        else:
            lo = parse_calendar_date(dr[0]) if isinstance(dr, list) and len(dr) == 2 else None
            hi = parse_calendar_date(dr[1]) if isinstance(dr, list) and len(dr) == 2 else None
            if lo is None or hi is None or lo > hi:
                dcheck.update(ok=False, error="bad_date_range_config", expected=dr)
                reasons.append("date_range配置无效: %s" % dr)
                blocking = True
            else:
                idx = header.index(dc)
                out_of_range, unparseable, miss, parsed = [], [], [], 0
                for ri, r in enumerate(data_rows, start=2):
                    v = r[idx] if idx < len(r) else None
                    if v is None or (isinstance(v, str) and not v.strip()):
                        miss.append(ri)
                        continue
                    d = parse_calendar_date(v)
                    if d is None:
                        unparseable.append({"row": ri, "value": _short(v)})
                    else:
                        parsed += 1
                        if not (lo <= d <= hi):
                            out_of_range.append({"row": ri, "value": d.isoformat()})
                dcheck.update(expected=[lo.isoformat(), hi.isoformat()],
                              parsed_rows=parsed,
                              out_of_range_count=len(out_of_range),
                              out_of_range=out_of_range[:20],
                              unparseable_count=len(unparseable),
                              unparseable=unparseable[:20],
                              missing_count=len(miss), missing_rows=miss[:20],
                              ok=not (out_of_range or unparseable or miss))
                if unparseable:
                    reasons.append("日期不可解析: %d 行（值不是有效日历日期）"
                                   % len(unparseable))
                if miss:
                    reasons.append("日期缺失: %d 行" % len(miss))
                if out_of_range:
                    reasons.append("日期超出范围: %d 行不在 [%s, %s] 内"
                                   % (len(out_of_range), lo, hi))
                if unparseable or miss or (out_of_range
                                           and req.get("date_range_block", True)):
                    blocking = True
        checks.append(dcheck)

    # 最少行数
    min_rows = req.get("min_rows", 1)
    checks.append({"check": "min_rows", "ok": len(data_rows) >= min_rows,
                   "rows": len(data_rows), "min": min_rows})
    if len(data_rows) < min_rows:
        reasons.append("有效行数不足: %d < %d" % (len(data_rows), min_rows))
        blocking = True

    verdict = "block" if blocking else ("confirm_needed" if reasons else "pass")
    report = {"mode": "data", "verdict": verdict, "sheet": sheet_name,
              "checks": checks, "reasons": reasons, "rows": len(data_rows),
              "note": "verdict=block时不产出分析结论"}
    dump(args.report, report)
    return 0


def main():
    ap = argparse.ArgumentParser(description="材料齐全性与数据范围校验")
    sub = ap.add_subparsers(dest="mode", required=True)
    p1 = sub.add_parser("checklist")
    p1.add_argument("--checklist", required=True)
    p1.add_argument("--dir", required=True)
    p1.add_argument("--report", required=True)
    p1.set_defaults(func=mode_checklist)
    p2 = sub.add_parser("data")
    p2.add_argument("--data", required=True)
    p2.add_argument("--requirements", required=True)
    p2.add_argument("--report", required=True)
    p2.set_defaults(func=mode_data)
    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
