#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""evidence-digest: 多源材料证据台账（来源定位+口径冲突+问题覆盖）。

读取失败 / 未知类型 / 公式缺缓存 一律单列，绝不进入事实台账、
覆盖匹配或数字统计（冻结缺陷 benchmarks/evidence-failure 修复）。
关键词命中仅是检索候选，不构成已证实答案；冲突判定保留期间、
主体与单位，不同期间/主体的数值不判为同口径冲突。
"""
import argparse
import csv
import json
import os
import re
import sys

from docx import Document
from openpyxl import load_workbook

TEXT_EXTS = (".md", ".txt")
CSV_EXTS = (".csv",)
DOCX_EXTS = (".docx",)
XLSX_EXTS = (".xlsx",)
IMAGE_EXTS = (".png", ".jpg", ".jpeg", ".gif", ".bmp")
SUPPORTED_EXTS = TEXT_EXTS + CSV_EXTS + DOCX_EXTS + XLSX_EXTS + IMAGE_EXTS

NUM_RE = re.compile(r"\d+(?:\.\d+)?")
UNITS = r"%|万元|亿美元|万美元|亿元|元|万人|人|个|天|次|家|项|件|平方米|公里|吨"
CLAUSE_RE = re.compile(
    r"([^，。；;、\s]{2,14}?)\s*(?:为|是|:|：)?\s*(\d+(?:\.\d+)?)\s*(" + UNITS + r")")

YEAR_RE = re.compile(r"(20\d{2})\s*年")
QUARTER_RE = re.compile(r"Q([1-4])|([一二三四])季度|第([1-4])季度")
HALF_RE = re.compile(r"(上半|下半)年")
MONTH_RE = re.compile(r"(\d{1,2})月")
SCOPE_RE = re.compile(r"全年|本年度|当年度|年度")
SUBJECT_RE = re.compile(
    r"([\u4e00-\u9fff]{2,10}(?:公司|集团|子公司|部门|事业部|分行|项目组))"
    r"|(甲方|乙方|丙方|丁方|建设单位|施工单位|供应商|采购方|销售方|本公司|本部)")
_CN_Q = {"一": "Q1", "二": "Q2", "三": "Q3", "四": "Q4"}
# 期间口径粘滞字：出现在标签首部时剥离，避免"3月客户满意度"≠"客户满意度"
GLUE_CHARS = "年月日周季期间度上下本当该每各的"

def norm_label(label_part):
    """归一化口径标签：取数字前尾端汉语串（至多6字），剥离期间粘滞前缀。"""
    m = re.search(r"[\u4e00-\u9fff]+$", label_part)
    if m:
        seq = m.group(0)
        run = seq[-6:] if len(seq) > 6 else seq
    else:
        run = label_part[-6:]
    while len(run) > 2 and run[0] in GLUE_CHARS:
        run = run[1:]
    return run or label_part[-2:]


def _find_markers(text):
    """收集期间/主体标记及其区间。"""
    marks = []
    for m in YEAR_RE.finditer(text):
        marks.append((m.start(), m.end(), "year", m.group(1)))
    for m in QUARTER_RE.finditer(text):
        g = m.group(1) or m.group(2) or m.group(3)
        if g in _CN_Q:
            q = _CN_Q[g]
        elif g in ("1", "2", "3", "4"):
            q = "Q" + g
        else:
            q = g
        marks.append((m.start(), m.end(), "quarter", q))
    for m in HALF_RE.finditer(text):
        marks.append((m.start(), m.end(), "half", m.group(1)))
    for m in MONTH_RE.finditer(text):
        marks.append((m.start(), m.end(), "month", m.group(1) + "月"))
    for m in SCOPE_RE.finditer(text):
        marks.append((m.start(), m.end(), "scope", "全年"))
    for m in SUBJECT_RE.finditer(text):
        marks.append((m.start(), m.end(), "subject", m.group(1) or m.group(2)))
    return sorted(marks, key=lambda x: x[0])


def _cluster(marks, gap=4):
    """把相邻(gap字符内)的标记聚为一个口径簇。"""
    clusters = []
    for s, e, kind, val in marks:
        if clusters and s - clusters[-1][1] <= gap:
            clusters[-1][1] = e
            clusters[-1][2].setdefault(kind, val)
        else:
            clusters.append([s, e, {kind: val}])
    return clusters
def extract_md_lines(path):
    with open(path, encoding="utf-8-sig") as f:
        for i, line in enumerate(f, start=1):
            line = line.rstrip()
            if line.strip():
                yield f"行{i}", line


def extract_csv_rows(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.reader(f), start=1):
            cells = [c.strip() for c in row if c and c.strip()]
            if cells:
                yield f"行{i}", ": ".join(cells)


def extract_docx_paras(path):
    doc = Document(path)
    for i, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            yield f"段落{i}", p.text.strip()
    for ti, t in enumerate(doc.tables, start=1):
        for ri, row in enumerate(t.rows, start=1):
            for ci, c in enumerate(row.cells, start=1):
                if c.text.strip():
                    yield f"表{ti}行{ri}列{ci}", c.text.strip()


def extract_xlsx_cells(path):
    """返回 ([(location, text)], 公式缺缓存清单)。公式无缓存值单列，不入台账。"""
    wb_vals = load_workbook(path, data_only=True)
    wb_raw = None
    try:
        wb_raw = load_workbook(path, data_only=False)
    except Exception:
        wb_raw = None
    out, nocache = [], []
    try:
        for ws in wb_vals.worksheets:
            raw_map = {}
            if wb_raw is not None and ws.title in wb_raw.sheetnames:
                for row in wb_raw[ws.title].iter_rows():
                    for c in row:
                        if c.value is not None:
                            raw_map[c.coordinate] = c.value
            for row in ws.iter_rows():
                for c in row:
                    raw = raw_map.get(c.coordinate)
                    if (isinstance(raw, str) and raw.startswith("=")
                            and (c.value is None or not str(c.value).strip())):
                        nocache.append({"sheet": ws.title, "cell": c.coordinate,
                                        "formula": raw})
                cells = [(c.coordinate, c.value) for c in row
                         if c.value is not None and str(c.value).strip()]
                if not cells:
                    continue
                label = ""
                if len(cells) > 1 and not isinstance(cells[0][1], (int, float)):
                    label = str(cells[0][1]).strip()
                for coord, v in cells:
                    s = str(v).strip()
                    if label and coord != cells[0][0] and any(ch.isdigit() for ch in s):
                        out.append((f"{ws.title}!{coord}", f"{label}: {s}"))
                    else:
                        out.append((f"{ws.title}!{coord}", s))
    finally:
        wb_vals.close()
        if wb_raw is not None:
            wb_raw.close()
    return out, nocache


def scan_materials(materials_dir):
    """只收集真实证据；读取失败/未知类型/公式缺缓存/图片各自单列。"""
    items, vision = [], []
    unreadable, unsupported, nocache_all = [], [], []
    n_files = 0
    for root, _dirs, names in sorted(os.walk(materials_dir)):
        for n in sorted(names):
            p = os.path.join(root, n)
            rel = os.path.relpath(p, materials_dir)
            ext = os.path.splitext(n)[1].lower()
            n_files += 1
            if ext in IMAGE_EXTS:
                vision.append({"file": rel, "needs_vision": True,
                               "note": "图片内容需视觉识别后回填, 不编造"})
                continue
            if ext not in SUPPORTED_EXTS:
                unsupported.append({"file": rel, "ext": ext or "<无扩展名>",
                                    "reason": "未知文件类型，未读取，不计入台账"})
                continue
            try:
                if ext in TEXT_EXTS:
                    src = list(extract_md_lines(p))
                elif ext in CSV_EXTS:
                    src = list(extract_csv_rows(p))
                elif ext in DOCX_EXTS:
                    src = list(extract_docx_paras(p))
                else:
                    src, nocache = extract_xlsx_cells(p)
                    for nc in nocache:
                        nc["file"] = rel
                    nocache_all.extend(nocache)
                for loc, text in src:
                    items.append({"file": rel, "location": loc, "text": text,
                                  "type": ext.lstrip(".")})
            except Exception as e:
                unreadable.append({"file": rel, "error": f"{type(e).__name__}: {e}"})
    return (items, vision, unreadable, unsupported, nocache_all, n_files)


def nearest_period(text, pos):
    """条款位置最近的期间簇 -> (year, quarter, half, month, scope) 或 None。"""
    clusters = _cluster([m for m in _find_markers(text)
                         if m[2] != "subject"])
    if not clusters:
        return None
    best, best_d = None, None
    for s, e, comp in clusters:
        d = 0 if s <= pos <= e else (pos - e if pos > e else s - pos)
        if best_d is None or d < best_d:
            best, best_d = comp, d
    return (best.get("year"), best.get("quarter"), best.get("half"),
            best.get("month"), best.get("scope"))


def nearest_subject(text, pos):
    marks = [m for m in _find_markers(text) if m[2] == "subject"]
    if not marks:
        return None
    best, best_d = None, None
    for s, e, _kind, val in marks:
        d = 0 if s <= pos <= e else (pos - e if pos > e else s - pos)
        if best_d is None or d < best_d:
            best, best_d = val, d
    return best


def render_period(period):
    if not period:
        return None
    year, quarter, half, month, scope = period
    out = ((f"{year}年" if year else "") + (half or "") + (quarter or "")
           + (month or "") + (scope or ""))
    return out or None


def _has_sub_period(period):
    return period is not None and any(
        period[i] is not None for i in (1, 2, 3))


def periods_compatible(p1, p2):
    """分量逐一比较：双方均显式且不同才视为不同口径；未声明分量视为兼容。

    例外：全年口径 vs 显式季度/半年/月份属于不同期间，不兼容。
    """
    if p1 is None or p2 is None:
        return True
    for a, b in zip(p1, p2):
        if a is not None and b is not None and a != b:
            return False
    if p1[4] and _has_sub_period(p2):
        return False
    if p2[4] and _has_sub_period(p1):
        return False
    return True


def merge_caliber(p1, p2):
    if p1 is None:
        return p2
    if p2 is None:
        return p1
    return tuple(b if a is None else a for a, b in zip(p1, p2))
def build_ledger(items):
    ledger = []
    for i, it in enumerate(items, start=1):
        ledger.append({"id": f"E-{i:04d}", "text": it["text"],
                       "source_file": it["file"], "location": it["location"],
                       "type": it.get("type", ""),
                       "has_numbers": bool(NUM_RE.search(it["text"])),
                       "confidence": "sourced"})
    return ledger


def find_conflicts(ledger):
    """条款级 标签+数值+单位 冲突检测；口径=期间+主体+单位。

    同口径不同数值 -> conflicts；同标签+单位但期间/主体口径不同 -> 不判冲突，
    单列 cross_caliber_review 供人工核对（不把不同期间值定为同口径冲突）。
    """
    metrics = {}
    for e in ledger:
        if not e["has_numbers"]:
            continue
        for m in CLAUSE_RE.finditer(e["text"]):
            label = norm_label(m.group(1))
            value, unit = m.group(2), m.group(3)
            period = nearest_period(e["text"], m.start(2))
            subject = nearest_subject(e["text"], m.start(2))
            metrics.setdefault((label.strip(), unit), []).append(
                {"value": float(value), "id": e["id"], "file": e["source_file"],
                 "location": e["location"], "period": render_period(period),
                 "subject": subject, "unit": unit, "_period": period})
    conflicts, cross = [], []
    for (label, unit), hits in sorted(metrics.items()):
        if len({h["value"] for h in hits}) < 2 or len(hits) < 2:
            continue
        groups = []
        for h in hits:
            for g in groups:
                if all(m["subject"] == h["subject"]
                       and periods_compatible(m["_period"], h["_period"])
                       for m in g):
                    g.append(h)
                    break
            else:
                groups.append([h])
        for group in groups:
            if len({h["value"] for h in group}) > 1:
                caliber_p = group[0]["_period"]
                for h in group[1:]:
                    caliber_p = merge_caliber(caliber_p, h["_period"])
                subject = next((h["subject"] for h in group
                                if h["subject"] is not None), None)
                conflicts.append({
                    "metric": f"{label}（{unit}）",
                    "caliber": {"period": render_period(caliber_p),
                                "subject": subject, "unit": unit},
                    "entries": [{k: h[k] for k in
                                 ("value", "id", "file", "location",
                                  "period", "subject", "unit")} for h in group]})
        if len(groups) > 1:
            cross.append({
                "metric": f"{label}（{unit}）",
                "reason": "同一标签+单位但期间/主体口径不同，不判定为同口径冲突，需人工核对",
                "entries": [{k: h[k] for k in
                             ("value", "id", "file", "location",
                              "period", "subject", "unit")} for h in hits]})
    return conflicts, cross


def load_questions(path):
    qs = []
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            qs.append(re.sub(r"^[-*\d\.\、\)\s]+", "", line))
    return qs


Q_STOP_RE = re.compile(
    r"是多少|怎么样|怎样|如何|怎么|为什么|是什么|什么|哪些|哪个|几个|多少|是否"
    r"|有没有|能否|可否|告诉我|说明一下|讲一下|列一下|吗|呢|吧")


def keywords_of(question):
    """切分问句为检索关键词：去疑问套话，按同质 run（汉字/字母数字）切分。

    命中仅是检索候选，不构成已证实答案。
    """
    toks = []
    for part in Q_STOP_RE.split(question):
        toks.extend(re.findall(r"[\u4e00-\u9fff]{2,}|[A-Za-z0-9]{2,}", part))
    seen, out = set(), []
    for t in toks:
        if t not in seen:
            seen.add(t)
            out.append(t)
    return out


def coverage(questions, ledger):
    """关键词命中=检索候选（candidates_need_review），不是已证实答案。"""
    cov, open_qs = [], []
    for qi, q in enumerate(questions, start=1):
        kws = keywords_of(q)
        hits = [e["id"] for e in ledger if any(k in e["text"] for k in kws)]
        if hits:
            cov.append({"q_no": qi, "question": q,
                        "status": "candidates_need_review",
                        "confidence": "candidate",
                        "evidence_ids": hits[:20],
                        "note": "关键词命中仅为检索候选，不构成已证实答案，需人工核对原文"})
        else:
            open_qs.append({"q_no": qi, "question": q, "status": "no_evidence",
                            "confidence": "none", "evidence_ids": [],
                            "note": "可用材料文本中无关键词命中"
                                    "（读取失败/未知类型/公式缺缓存/图片不参与匹配）"})
    return cov, open_qs
def main():
    ap = argparse.ArgumentParser(description="多源证据台账")
    ap.add_argument("--materials", required=True)
    ap.add_argument("--questions", required=True)
    ap.add_argument("--outdir", required=True)
    args = ap.parse_args()
    if not os.path.isdir(args.materials):
        print("ERROR: materials dir not found", file=sys.stderr)
        return 2
    (items, vision, unreadable, unsupported, nocache, n_files) = \
        scan_materials(args.materials)
    ledger = build_ledger(items)
    conflicts, cross = find_conflicts(ledger)
    try:
        questions = load_questions(args.questions)
    except Exception as e:
        print(f"ERROR: cannot read questions: {e}", file=sys.stderr)
        return 2
    cov, open_qs = coverage(questions, ledger)

    warnings = []
    if n_files == 0:
        warnings.append("材料目录为空：未发现任何文件")
    if unsupported:
        warnings.append(f"{len(unsupported)} 个未知类型文件未读取，已单列")
    if unreadable:
        warnings.append(f"{len(unreadable)} 个文件读取失败，已单列，未计入台账/覆盖/统计")
    if nocache:
        warnings.append(f"{len(nocache)} 个公式单元格缺缓存值，已单列，未计入台账/统计")
    if vision:
        warnings.append(f"{len(vision)} 个图片来源待视觉识别回填，不编造内容")
    if n_files and not items and not vision:
        warnings.append("材料目录中没有任何可读取的受支持类型内容")

    os.makedirs(args.outdir, exist_ok=True)
    result = {
        "materials_dir": args.materials,
        "ledger": ledger,
        "image_sources_needing_vision": vision,
        "unreadable_files": unreadable,
        "unsupported_files": unsupported,
        "formula_cells_without_cache": nocache,
        "conflicts": conflicts,
        "cross_caliber_review": cross,
        "coverage": cov,
        "warnings": warnings,
        "stats": {
            "material_files": n_files,
            "evidence_items": len(ledger),
            "with_numbers": sum(1 for e in ledger if e["has_numbers"]),
            "conflicts": len(conflicts),
            "cross_caliber_review": len(cross),
            "questions": len(questions),
            "questions_candidates": len(cov),
            "questions_no_evidence": len(open_qs),
            "unreadable_files": len(unreadable),
            "unsupported_files": len(unsupported),
            "image_sources": len(vision),
            "formula_cells_without_cache": len(nocache),
        },
        "note": "台账每条含来源文件与定位且可回溯; 读取失败/未知类型/公式缺缓存/图片"
                "单列不进台账; 关键词命中仅为检索候选; 无来源推断不入台账",
    }
    with open(os.path.join(args.outdir, "evidence-ledger.json"), "w",
              encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    with open(os.path.join(args.outdir, "conflicts.json"), "w",
              encoding="utf-8") as f:
        json.dump({"conflicts": conflicts, "cross_caliber_review": cross,
                   "note": "同口径(期间+主体+单位)不同数值才判冲突;"
                           "期间/主体不同的并列值仅列review不判冲突"},
                  f, ensure_ascii=False, indent=2)
    lines = ["# 证据台账", ""]
    for e in ledger:
        lines.append(f"- **{e['id']}** `{e['source_file']}@{e['location']}`：{e['text']}")
    lines += ["", "## 冲突指标（同口径）", ""]
    for c in conflicts:
        lines.append(f"- {c['metric']}"
                     f"（期间：{c['caliber']['period'] or '未声明'}；"
                     f"主体：{c['caliber']['subject'] or '未声明'}）: " + "；".join(
                         f"{h['value']}（{h['file']}@{h['location']}，"
                         f"期间{h['period'] or '未声明'}）" for h in c["entries"]))
    lines += ["", "## 跨口径待核（不判冲突）", ""]
    for c in cross:
        lines.append(f"- {c['metric']}: " + "；".join(
            f"{h['value']}（{h['file']}@{h['location']}，"
            f"期间{h['period'] or '未声明'}，主体{h['subject'] or '未声明'}）"
            for h in c["entries"]))
    lines += ["", "## 读取失败（未计入台账）", ""]
    lines += [f"- {u['file']}: {u['error']}" for u in unreadable] or ["（无）"]
    lines += ["", "## 未知类型（未读取）", ""]
    lines += [f"- {u['file']}（{u['ext']}）" for u in unsupported] or ["（无）"]
    lines += ["", "## 公式缺缓存（未计入台账）", ""]
    lines += [f"- {n['file']} {n['sheet']}!{n['cell']}: {n['formula']}"
              for n in nocache] or ["（无）"]
    lines += ["", "## 待视觉识别来源", ""]
    lines += [f"- {v['file']}" for v in vision] or ["（无）"]
    lines += ["", "## 问题覆盖（候选需人工复核）", ""]
    lines += [f"- Q{c['q_no']} {c['question']}: 候选 {c['evidence_ids']}"
              for c in cov] or ["（无）"]
    lines += ["", "## 无证据问题", ""]
    lines += [f"- {o['question']}" for o in open_qs] or ["（无）"]
    if warnings:
        lines += ["", "## 提示", ""] + [f"- {w}" for w in warnings]
    with open(os.path.join(args.outdir, "evidence-ledger.md"), "w",
              encoding="utf-8") as f:
        f.write("\n".join(lines))
    with open(os.path.join(args.outdir, "open-questions.json"), "w",
              encoding="utf-8") as f:
        json.dump({"open_questions": open_qs}, f, ensure_ascii=False, indent=2)
    print(json.dumps(result["stats"], ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
