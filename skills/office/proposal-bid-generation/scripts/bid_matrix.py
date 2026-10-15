#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""proposal-bid-generation: 招标条款响应矩阵 + 投标草稿。

诚实规则（脚本保证，详见 SKILL.md）：
- 条款解析保留全部编号条款（含短条款，如“须盖章”）、续行与表格条款，并记录出处；
- 关键词命中只作为检索候选，不据此判定“满足”；
- 仅当资质文件“内容”中出现不含否定/过期表述的肯定陈述时才计 responded；
- 否定/过期陈述 → gap 并保留原文摘录；证据主体与 --company 不一致或无法核验
  → needs_confirmation；肯定证据与相反线索并存 → needs_confirmation（冲突待核）；
- 仅文件名命中（内容不可读或格式不支持提取）→ needs_confirmation，绝不 responded；
- 草稿不自动承诺资格满足，所有响应均要求人工复核。
"""
import argparse
import json
import os
import re
import sys

from docx import Document
from openpyxl import Workbook, load_workbook

# 泛化词：作为分隔符从条款文本中剔除，不构成“特定资质/能力”证据关键词。
# 纯数字 token（如 500/1000）同样不作为证据关键词，金额类要求不靠数字命中放行。
GENERIC_WORDS = {
    "万元", "亿元", "项目", "公司", "本公司", "我公司", "我司", "本司",
    "投标", "投标人", "招标", "招标人", "采购", "供应商", "要求", "提供",
    "具备", "具有", "有效", "相关", "以下", "上述", "乙方", "甲方", "承诺",
    "保证", "内容", "文件", "材料", "说明", "情况", "其他", "所有", "响应",
    "满足", "需要", "进行", "完成", "中标", "合同", "报价", "预算", "金额",
    "不得", "必须", "应当", "可以", "以及", "对于", "根据", "按照", "意见",
    "有关", "数量", "标准", "范围", "方式", "期限", "证书", "认证", "超过",
    "低于", "高于", "达到", "并且", "或者", "同时", "投标文件", "招标文件",
    "服务期", "质量", "服务",
}
# 否定表述（出现在含关键词的句子中 → 否定证据）。排除“没有发现问题”类反例。
NEG_RE = re.compile(
    r"未取得|尚未|暂未|不曾|并未|从未|不具备|不持有|未持有|未获得|未通过"
    r"|不得宣称|无法提供|不满足|没有(?!发现|问题|异常|遗漏|变化)")
# 过期/无效表述
EXPIRY_PATTERNS = ("已过期", "过期", "已失效", "失效", "吊销", "撤销", "注销",
                   "已到期", "有效期已过", "有效期届满")
# 主体识别：公司/机构名称（用于证据主体一致性核验）
COMPANY_RE = re.compile(r"[\u4e00-\u9fa5A-Za-z0-9]{2,24}(?:有限公司|股份有限公司|集团公司|责任公司|事务所|研究院|集团)")
SELF_REFS = ("本公司", "我公司", "我司", "本司", "贵公司", "该公司", "投标人", "供应商", "招标人", "采购人")
# 条款编号行：第X条 / 一、 / 1. / （一）；编号后紧跟数字（如 3.5）不视为条款
CLAUSE_RE = re.compile(
    r"^\s*(?:第[一二三四五六七八九十百\d]+条"
    r"|[（(][一二三四五六七八九十\d]+[)）]"
    r"|[一二三四五六七八九十\d]+[、\.．)）:：](?!\d)"
    r")\s*(.*)$")
HEADER_TOKENS = ("序号", "条款", "要求", "编号", "内容", "项目名称", "说明")
REASON_ZH = {
    "positive_evidence": "内容级肯定证据",
    "negative_evidence": "存在否定陈述",
    "expired_or_invalid": "证据过期或无效",
    "conflicting_evidence": "证据冲突待核",
    "subject_mismatch": "证据主体与指定公司不一致",
    "subject_unverified": "证据主体无法核验",
    "candidate_only": "仅检索到候选（文件名或弱命中），内容未证实",
    "no_evidence": "未检索到证据",
    "no_specific_keyword": "无可检索的特定关键词",
}


def _read_text(path, limit=200000):
    """utf-8 优先、gbk 回退读取文本；都失败时以替换字符读出。"""
    for enc in ("utf-8", "gbk"):
        try:
            with open(path, encoding=enc) as f:
                return f.read(limit)
        except UnicodeDecodeError:
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read(limit)


def _docx_parts(path):
    """docx 段落与表格文本（表格逐行拼接，保留全部内容）。"""
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return parts


def keywords(text):
    """提取特定资质/能力关键词。

    泛化词先作为分隔符剔除，剩余片段取中文 2-3 字整词、4 字以上滑窗 4-gram
    （≤6 字整词一并保留）；字母数字 token（≥3 位且含字母，如 ISO9001/PMP）
    保留为强关键词；纯数字（如 500/1000）不作为证据关键词。
    """
    seg = text
    for g in sorted(GENERIC_WORDS, key=len, reverse=True):
        seg = seg.replace(g, "|")
    kws = []

    def _add(k):
        if k and k not in kws and not any(len(g) >= 2 and g in k for g in GENERIC_WORDS):
            kws.append(k)

    for part in re.split(r"[|\s]+", seg):
        for tok in re.findall(r"[\u4e00-\u9fa5]{2,}", part):
            if len(tok) >= 4:
                for i in range(len(tok) - 3):
                    _add(tok[i:i + 4])
                if len(tok) <= 6:
                    _add(tok)
            else:
                _add(tok)
        for t in re.findall(r"[A-Za-z0-9]{3,}", part):
            if any(c.isalpha() for c in t):
                tl = t.lower()
                if tl not in kws:
                    kws.append(tl)
    return kws[:40]


def _strong_kws(kws):
    """强关键词：字母数字标识（如 iso9001/pmp）。存在时肯定证据须命中其一。"""
    return [k for k in kws if re.fullmatch(r"[a-z0-9]+", k)]


def _is_header_row(texts):
    """表格首行是否为表头（全短单元格且含序号/条款/要求等表头词）。"""
    return all(len(t) <= 10 for t in texts) and any(
        any(h in t for h in HEADER_TOKENS) for t in texts)


def _clean_body(body):
    return body.strip(" ：:、.．\u3000")


def parse_tender(path):
    """解析招标文件 → (clauses, mode, meta)。

    编号模式：编号行（第X条/一、/1./（一））为条款起始，无编号的后续行并入
    上一条款（续行），首次编号行之前的内容计为前言跳过；无编号时按段落回退，
    每个非空段落各成一条（虚拟编号）。docx 表格逐行成条（首行表头跳过）。
    短条款（如“须盖章”）不做长度过滤，全部保留并记录出处。
    """
    tables = []
    if path.lower().endswith(".docx"):
        doc = Document(path)
        units = [(f"docx:段落{i}", p.text.strip())
                 for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
        for ti, t in enumerate(doc.tables, 1):
            rows = []
            for row in t.rows:
                rows.append([c.text.strip() for c in row.cells])
            tables.append((ti, rows))
    else:
        lines = _read_text(path).splitlines()
        units = [(f"txt:行{i}", ln.strip())
                 for i, ln in enumerate(lines, 1) if ln.strip()]
    units = [(src, t) for src, t in units if t]

    clauses = []
    preamble = 0
    mode = None
    numbered = [t for _, t in units if CLAUSE_RE.match(t)]
    if units and len(numbered) >= 2 and len(numbered) * 3 >= len(units):
        mode = "numbered"
        cur = None
        for src, t in units:
            m = CLAUSE_RE.match(t)
            if m:
                if cur is not None:
                    clauses.append(cur)
                body = _clean_body(m.group(1))
                cur = {"text": body, "source": src, "continued": not body}
            elif cur is None:
                preamble += 1
            else:
                cur["text"] = (cur["text"] + " " + t).strip()
                cur["continued"] = True
        if cur is not None:
            clauses.append(cur)
    elif units:
        mode = "paragraph_fallback"
        for src, t in units:
            clauses.append({"text": t, "source": src, "continued": False})

    table_clauses = []
    for ti, rows in tables:
        for ri, cells in enumerate(rows, 1):
            texts = [c for c in cells if c]
            if not texts:
                continue
            if ri == 1 and _is_header_row(texts):
                continue
            table_clauses.append({"text": "；".join(texts),
                                  "source": f"docx:表{ti}行{ri}",
                                  "continued": False})
    if table_clauses:
        mode = (mode + "+table") if mode else "table"
        clauses.extend(table_clauses)

    if not clauses:
        raise ValueError("未从招标文件解析到任何条款（空文件或无有效内容）")
    for i, c in enumerate(clauses, 1):
        c["no"] = f"R{i:02d}"
    meta = {"preamble": preamble, "table_clauses": len(table_clauses),
            "continued": sum(1 for c in clauses if c.get("continued"))}
    return clauses, mode, meta


def load_quals(quals_dir):
    """加载资质材料：.md/.txt/.docx/.xlsx 提取内容；其余格式或读取失败只记文件名。

    返回 [{file, name, content, read, error}]；read=False 的文件只能做文件名
    级候选，绝不作为 responded 证据。
    """
    quals = []
    for root, _dirs, names in sorted(os.walk(quals_dir)):
        for n in sorted(names):
            p = os.path.join(root, n)
            ext = os.path.splitext(n)[1].lower()
            q = {"file": os.path.relpath(p, quals_dir).replace("\\", "/"),
                 "name": n, "content": "", "read": False, "error": ""}
            try:
                if ext in (".md", ".txt"):
                    q["content"] = _read_text(p)
                    q["read"] = True
                elif ext == ".docx":
                    q["content"] = "\n".join(_docx_parts(p))
                    q["read"] = True
                elif ext == ".xlsx":
                    wb = load_workbook(p, read_only=True, data_only=True)
                    parts = []
                    for ws in wb.worksheets:
                        for row in ws.iter_rows(values_only=True):
                            parts.append(" ".join(
                                str(v) for v in row if v is not None))
                    wb.close()
                    q["content"] = "\n".join(parts)
                    q["read"] = True
                else:
                    q["error"] = "unsupported_format"
            except Exception as e:  # 读取失败：仅文件名候选，不中断批处理
                q["error"] = f"{type(e).__name__}: {e}"
            quals.append(q)
    return quals


def _is_self_ref(name):
    return name in SELF_REFS or (name and name[0] in "本我贵该")


def _same_company(name, company):
    return name == company or company in name or name in company


def _sentence_kind(sentence, company):
    """判定含关键词句子的证据性质：negative/expired/subject/positive。"""
    if NEG_RE.search(sentence):
        return "negative", None
    if any(p in sentence for p in EXPIRY_PATTERNS):
        return "expired", None
    names = [m for m in COMPANY_RE.findall(sentence) if not _is_self_ref(m)]
    if names:
        if company:
            if all(_same_company(n, company) for n in names):
                return "positive", None
            return "subject", "subject_mismatch"
        return "subject", "subject_unverified"
    return "positive", None


def _dedupe_evidence(items):
    out, seen = [], set()
    for it in items:
        if it["file"] not in seen:
            seen.add(it["file"])
            out.append({"file": it["file"], "sentence": it["sentence"][:80],
                        "term": it["term"]})
        if len(out) >= 3:
            break
    return out


def match_clause(clause, quals, company):
    """单条款响应判定。

    关键词只用于检索候选；状态由句子级证据决定：
    - 强关键词（字母数字标识）存在时，肯定证据必须命中其一，防止近邻资质误配；
    - 无任何有效证据/否定/过期 → gap；主体待核/仅文件名命中/证据冲突
      → needs_confirmation；仅有内容级肯定证据 → responded。
    """
    kws = keywords(clause["text"])
    res = {"no": clause["no"], "text": clause["text"],
           "source": clause["source"], "continued": bool(clause.get("continued")),
           "status": "gap", "reason": "no_specific_keyword", "keywords": kws[:8],
           "evidence": [], "candidates": [], "warnings": []}
    if not kws:
        return res
    res["reason"] = "no_evidence"
    strong = _strong_kws(kws)
    positives, negatives, expired, subjects = [], [], [], []
    subject_reason = "subject_unverified"
    for q in quals:
        name_l = q["name"].lower()
        name_hits = [k for k in kws if k in name_l]
        content_l = q["content"].lower()
        content_terms = [k for k in kws if k in content_l]
        if name_hits or content_terms:
            res["candidates"].append({
                "file": q["file"], "read": q["read"],
                "matched": (name_hits + [t for t in content_terms
                                         if t not in name_hits])[:5]})
        for s in re.split(r"[。；;！!？?\n\r\t]+", q["content"]):
            if not s.strip():
                continue
            s_l = s.lower()
            hit_terms = [k for k in kws if k in s_l]
            if not hit_terms:
                continue
            if strong and not any(t in strong for t in hit_terms):
                continue  # 弱命中仅作候选，不参与证据/否定判定
            kind, detail = _sentence_kind(s, company)
            item = {"file": q["file"], "sentence": s.strip(),
                    "term": hit_terms[0], "kind": kind}
            if kind == "negative":
                negatives.append(item)
            elif kind == "expired":
                expired.append(item)
            elif kind == "subject":
                subjects.append(item)
                if detail:
                    subject_reason = detail
            else:
                positives.append(item)

    if positives and (negatives or expired or subjects):
        res["status"], res["reason"] = "needs_confirmation", "conflicting_evidence"
    elif positives:
        res["status"], res["reason"] = "responded", "positive_evidence"
        res["evidence"] = _dedupe_evidence(positives)
    elif negatives:
        res["status"], res["reason"] = "gap", "negative_evidence"
    elif expired:
        res["status"], res["reason"] = "gap", "expired_or_invalid"
    elif subjects:
        res["status"], res["reason"] = "needs_confirmation", subject_reason
    elif res["candidates"]:
        res["status"], res["reason"] = "needs_confirmation", "candidate_only"
    res["warnings"] = (negatives + expired + subjects)[:4]
    return res


def build_xlsx(results, outdir):
    wb = Workbook()
    ws = wb.active
    ws.title = "响应矩阵"
    ws.append(["条款编号", "招标要求", "响应状态", "证据文件", "匹配词", "原因/说明", "条款出处"])
    for r in results:
        files = "；".join(e["file"] for e in r["evidence"]) or \
            "；".join(c["file"] for c in r["candidates"])
        reason = REASON_ZH.get(r["reason"], r["reason"])
        for w in r["warnings"]:
            reason += f"｜{w['file']}：「{w['sentence'][:40]}」"
        ws.append([r["no"], r["text"], r["status"], files,
                   "、".join(r["keywords"][:6]), reason, r["source"]])
    wb.save(os.path.join(outdir, "response-matrix.xlsx"))


def build_draft(results, outdir, company):
    doc = Document()
    doc.add_heading("投标响应文件（草稿）", level=1)
    doc.add_paragraph(
        "本稿由招标条款响应矩阵生成。状态含义：responded=已检索到内容级肯定证据"
        "（仍需人工复核有效期、主体与适用范围）；needs_confirmation=仅检索到候选"
        "或证据待核验，未证实；gap=无有效证据或存在否定/过期陈述。本草稿不自动"
        "承诺任何资格满足，商务内容（报价、方案细节）由投标方补充。")
    if company:
        doc.add_paragraph(f"指定投标主体：{company}（证据主体一致性按此核验）。")
    for r in results:
        title = r["text"] if len(r["text"]) <= 40 else r["text"][:40] + "…"
        doc.add_heading(f"{r['no']} {title}", level=2)
        p = doc.add_paragraph()
        if r["status"] == "responded":
            p.add_run(f"响应说明（{REASON_ZH[r['reason']]}）：检索到以下支撑材料，"
                      "请复核其有效性、有效期与主体后确认响应程度：")
            for e in r["evidence"]:
                doc.add_paragraph(f"— {e['file']}：「{e['sentence']}」",
                                  style="List Bullet")
        elif r["status"] == "needs_confirmation":
            p.add_run(f"【待确认-{REASON_ZH[r['reason']]}】未取得可核验的肯定"
                      "证据。").bold = True
            cands = "；".join(c["file"] for c in r["candidates"]) or "（无）"
            doc.add_paragraph(f"检索候选：{cands}。请人工核验文件内容后再确定响应。")
        else:
            p.add_run("【待补充-缺口】本条款暂无有效资质材料支撑。").bold = True
            for w in r["warnings"]:
                if w["kind"] in ("negative", "expired"):
                    label = "过期/无效" if w["kind"] == "expired" else "否定"
                    doc.add_paragraph(
                        f"注意：材料中存在{label}陈述（{w['file']}："
                        f"「{w['sentence'][:60]}」），不得据此宣称满足。")
    doc.add_paragraph("以上由脚本生成，不构成对资格满足的承诺；最终响应以经复核"
                      "的有效证明文件为准。")
    doc.save(os.path.join(outdir, "bid-draft.docx"))


def main():
    ap = argparse.ArgumentParser(description="招标条款响应矩阵与投标草稿")
    ap.add_argument("--tender", required=True, help="招标文件（.docx/.md/.txt）")
    ap.add_argument("--quals", required=True, help="资质材料目录")
    ap.add_argument("--outdir", required=True, help="输出目录")
    ap.add_argument("--company", default=None,
                    help="投标主体名称（可选）：用于证据主体一致性核验")
    args = ap.parse_args()
    if not os.path.isfile(args.tender):
        print(f"ERROR: 招标文件不存在或不可读: {args.tender}", file=sys.stderr)
        return 2
    if not os.path.isdir(args.quals):
        print(f"ERROR: 资质材料目录不存在: {args.quals}", file=sys.stderr)
        return 2
    try:
        clauses, mode, meta = parse_tender(args.tender)
    except Exception as e:
        print(f"ERROR: cannot parse tender: {e}", file=sys.stderr)
        return 2
    quals = load_quals(args.quals)
    results = [match_clause(c, quals, args.company) for c in clauses]

    counts = {
        "clauses": len(results),
        "responded": sum(1 for r in results if r["status"] == "responded"),
        "needs_confirmation": sum(1 for r in results if r["status"] == "needs_confirmation"),
        "gaps": sum(1 for r in results if r["status"] == "gap"),
    }
    os.makedirs(args.outdir, exist_ok=True)
    build_xlsx(results, args.outdir)
    build_draft(results, args.outdir, args.company)

    matrix = {"tender": args.tender, "parse_mode": mode, "company": args.company,
              "keyword_role": "candidate_search_only", "counts": counts,
              "clauses": results}
    with open(os.path.join(args.outdir, "response-matrix.json"), "w",
              encoding="utf-8") as f:
        json.dump(matrix, f, ensure_ascii=False, indent=2)

    gaps = [r for r in results if r["status"] == "gap"]
    needs = [r for r in results if r["status"] == "needs_confirmation"]
    with open(os.path.join(args.outdir, "gaps.json"), "w", encoding="utf-8") as f:
        json.dump({
            "gaps": [{"no": r["no"], "text": r["text"], "source": r["source"],
                      "reason": r["reason"],
                      "negative_evidence": [w for w in r["warnings"]
                                            if w["kind"] == "negative"],
                      "expired_evidence": [w for w in r["warnings"]
                                           if w["kind"] == "expired"]}
                     for r in gaps],
            "needs_confirmation": [{"no": r["no"], "text": r["text"],
                                    "source": r["source"], "reason": r["reason"],
                                    "candidates": [c["file"] for c in r["candidates"]]}
                                   for r in needs],
            "total": len(results),
            "note": "缺口条款无有效证据或存在否定/过期陈述，不编造资质；"
                    "needs_confirmation 需人工核验后再定响应"},
            f, ensure_ascii=False, indent=2)

    report = {"tender": args.tender, "parse_mode": mode, **counts,
              "quals_files": len(quals),
              "quals_unreadable": sum(1 for q in quals
                                      if not q["read"] and q["error"] != "unsupported_format"),
              "quals_unsupported": sum(1 for q in quals
                                       if q["error"] == "unsupported_format"),
              "preamble_skipped": meta["preamble"],
              "continued_clauses": meta["continued"],
              "table_clauses": meta["table_clauses"],
              "company": args.company,
              "keyword_role": "candidate_search_only"}
    with open(os.path.join(args.outdir, "report.json"), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
