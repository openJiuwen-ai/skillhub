#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""doc-translate-layout: 文档翻译并保持版式（extract / apply / verify）。

用法:
  python translate_layout.py extract --input src.docx --worksheet ws.json
  python translate_layout.py apply --input src.docx --worksheet ws.json \
      --translations tr.json --output out.docx [--glossary g.json] \
      [--report r.json] [--overflow-ratio 2.0]
  python translate_layout.py verify --input src.docx --worksheet ws.json \
      --output out.docx [--translations tr.json] [--report vr.json] \
      [--overflow-ratio 2.0]

安全约束:
- 工作表与源文件强绑定（source_sha256）：apply/verify 拒绝源已变更或无绑定的
  旧工作表，防止把译文回填到错误对象（benchmarks/translation-complex 冻结缺陷）。
- 段落 id 在 shape 范围内唯一稳定（pptx 含 shape 标识），同页不同 shape 不会
  共享 id 导致译文互相覆盖。
- 回填完整替换段落全部 run 文本，多 run 混合格式合并到首 run 时显式记录
  format_degradations，不用成功退出码掩盖。

退出码: 0=成功 2=输入/校验错误
"""
import argparse
import hashlib
import json
import os
import sys

WS_FORMAT_VERSION = 2


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def iter_docx_runs(p):
    """段落内可回填的 w:r（跳过删除修订中的 run）。"""
    from docx.oxml.ns import qn
    return [r for r in p.iter(qn("w:r"))
            if r.getparent() is not None and r.getparent().tag != qn("w:del")
            and r.findall(qn("w:t"))]


def runs_text(runs):
    from docx.oxml.ns import qn
    return "".join("".join(t.text or "" for t in r.findall(qn("w:t"))) for r in runs)


def docx_paragraphs(doc):
    """正文(含表格内段落) + 页眉页脚；id 顺序稳定，el 供回填。"""
    from docx.oxml.ns import qn
    items = []
    idx = 0
    for p in doc.element.body.iter(qn("w:p")):
        text = runs_text(iter_docx_runs(p))
        if not text.strip():
            continue
        node, in_table = p.getparent(), False
        while node is not None:
            if node.tag == qn("w:tbl"):
                in_table = True
                break
            node = node.getparent()
        style = ""
        pPr = p.find(qn("w:pPr"))
        if pPr is not None:
            ps = pPr.find(qn("w:pStyle"))
            if ps is not None:
                style = ps.get(qn("w:val")) or ""
        items.append({"id": f"p-{idx:04d}", "text": text, "style": style,
                      "location": "table" if in_table else "body", "el": p})
        idx += 1
    for si, sec in enumerate(doc.sections):
        for kind in ("header", "footer"):
            try:
                el = getattr(sec, kind)._element
            except Exception:
                continue
            if el is None:
                continue
            for p in el.iter(qn("w:p")):
                text = runs_text([r for r in p.iter(qn("w:r"))
                                  if r.findall(qn("w:t"))])
                if not text.strip():
                    continue
                items.append({"id": f"{kind[0]}-{si}-{idx:04d}", "text": text,
                              "style": "", "location": f"section{si} {kind}",
                              "el": p})
                idx += 1
    return items


def pptx_paragraphs(prs):
    """每段 id 含 slide+shape 标识：同页不同 shape 不共享段落 id。"""
    items = []
    seen = {}
    for si, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            try:
                shid = shape.shape_id
            except Exception:
                shid = 0
            for pi, para in enumerate(shape.text_frame.paragraphs):
                text = "".join(r.text or "" for r in para.runs)
                if not text.strip():
                    continue
                pid = f"s{si:02d}-sp{shid}-p{pi:03d}"
                if pid in seen:
                    seen[pid] += 1
                    pid = f"{pid}-x{seen[pid]}"
                else:
                    seen[pid] = 0
                items.append({"id": pid, "text": text, "style": "",
                              "location": f"slide {si}", "para": para})
    return items


def extract_items(path):
    """打开文档并抽取段落项；返回 (items, container)。"""
    if path.lower().endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(path)
        return pptx_paragraphs(prs), prs
    from docx import Document
    doc = Document(path)
    return docx_paragraphs(doc), doc


def set_docx_para_text(item, new_text):
    """完整替换段落文本：清空全部 run 的 w:t，译文写入第一个 run。

    返回替换前有文本的 run 数（>1 表示混合字符格式被合并为首 run 格式，
    调用方必须显式记录降级）；段落无可用 run 时返回 None。
    """
    from docx.oxml.ns import qn
    runs = iter_docx_runs(item["el"])
    if not runs:
        return None
    nonempty = 0
    for r in runs:
        ts = r.findall(qn("w:t"))
        if "".join(t.text or "" for t in ts):
            nonempty += 1
        for t in ts:
            t.text = ""
    first = runs[0].findall(qn("w:t"))[0]
    first.text = new_text
    first.set(qn("xml:space"), "preserve")
    return nonempty


def set_pptx_para_text(item, new_text):
    """完整替换 pptx 段落文本；返回语义同 set_docx_para_text。"""
    para = item["para"]
    if not para.runs:
        return None
    nonempty = sum(1 for r in para.runs if (r.text or ""))
    para.runs[0].text = new_text
    for r in para.runs[1:]:
        r.text = ""
    return nonempty


def shrink_docx_para_font(item, min_pt=9.0):
    """对承载译文的 run 有显式字号时降 1pt；返回是否实际缩放。"""
    from docx.oxml.ns import qn
    runs = iter_docx_runs(item["el"])
    if not runs:
        return False
    rPr = runs[0].find(qn("w:rPr"))
    if rPr is None:
        return False
    sz = rPr.find(qn("w:sz"))
    if sz is None:
        return False
    pt = int(sz.get(qn("w:val"))) / 2.0
    if pt > min_pt:
        sz.set(qn("w:val"), str(int(max(min_pt, pt - 1) * 2)))
        return True
    return False


def parse_translations(raw):
    """支持 {id: 译文} 或 [{id, translation}] 两种形态；非法返回 None。"""
    if isinstance(raw, dict):
        return dict(raw)
    if isinstance(raw, list) and all(isinstance(x, dict) and "id" in x for x in raw):
        return {x["id"]: x.get("translation", "") for x in raw}
    return None


def validate_binding(input_path, ws):
    """工作表与源文件绑定校验；返回 None(通过) 或错误信息。"""
    if not isinstance(ws, dict) or not isinstance(ws.get("items"), list):
        return "worksheet format invalid: expect an object with items[]"
    if not ws.get("source_sha256"):
        return ("worksheet has no source binding (source_sha256 missing); re-run "
                "extract on the current source file before applying translations")
    actual = file_sha256(input_path)
    if actual != ws["source_sha256"]:
        return (f"source binding mismatch: worksheet was extracted from a different "
                f"file (worksheet sha256 {ws['source_sha256'][:12]}..., input "
                f"sha256 {actual[:12]}...). Re-run extract; refuse to apply stale "
                f"positional translations to a changed document")
    return None


def check_items_match(ws_items, items):
    """重抽取项与工作表逐项核对（id + 原文）；返回差异描述列表。"""
    diffs = []
    cur = {}
    for it in items:
        cur.setdefault(it["id"], it)
    if len(cur) != len(items):
        diffs.append("duplicate ids in current file extraction")
    if len(ws_items) != len(items):
        diffs.append(f"item count differs: worksheet={len(ws_items)} "
                     f"current={len(items)}")
    for w in ws_items:
        c = cur.get(w.get("id"))
        if c is None:
            diffs.append(f"worksheet id {w.get('id')!r} not found in current file")
        elif (w.get("text") or "") != c["text"]:
            diffs.append(f"id {w['id']}: source text differs from worksheet")
    return diffs


def cmd_extract(args):
    try:
        items, _ = extract_items(args.input)
    except Exception as e:
        print(f"ERROR: cannot read input {args.input}: {e}", file=sys.stderr)
        return 2
    worksheet = {
        "format_version": WS_FORMAT_VERSION,
        "source": args.input,
        "source_sha256": file_sha256(args.input),
        "source_size": os.path.getsize(args.input),
        "item_count": len(items),
        "items": [{"id": it["id"], "text": it["text"],
                   "style": it["style"], "location": it["location"]}
                  for it in items],
    }
    try:
        with open(args.worksheet, "w", encoding="utf-8") as f:
            json.dump(worksheet, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"ERROR: cannot write worksheet {args.worksheet}: {e}", file=sys.stderr)
        return 2
    print(json.dumps({"items": len(items), "worksheet": args.worksheet,
                      "source_sha256": worksheet["source_sha256"][:12] + "..."},
                     ensure_ascii=False))
    return 0


def cmd_apply(args):
    try:
        with open(args.worksheet, encoding="utf-8") as f:
            ws = json.load(f)
        with open(args.translations, encoding="utf-8") as f:
            trans_map = parse_translations(json.load(f))
        glossary = {}
        if args.glossary:
            with open(args.glossary, encoding="utf-8") as f:
                glossary = json.load(f)
    except Exception as e:
        print(f"ERROR: bad input files: {e}", file=sys.stderr)
        return 2
    if trans_map is None:
        print("ERROR: translations must be {id: text} or [{id, translation}]",
              file=sys.stderr)
        return 2
    err = validate_binding(args.input, ws)
    if err:
        print(f"ERROR: {err}", file=sys.stderr)
        return 2
    try:
        items, container = extract_items(args.input)
    except Exception as e:
        print(f"ERROR: cannot read input {args.input}: {e}", file=sys.stderr)
        return 2
    diffs = check_items_match(ws["items"], items)
    if diffs:
        print(f"ERROR: worksheet does not match current file structure: "
              f"{diffs[:5]}", file=sys.stderr)
        return 2
    known = {it["id"] for it in items}
    unknown = sorted(tid for tid in trans_map if tid not in known)
    if unknown:
        print(f"ERROR: translations reference ids not in worksheet: {unknown[:5]}; "
              f"refusing to apply (translations appear to be for a different "
              f"document)", file=sys.stderr)
        return 2

    is_pptx = args.input.lower().endswith(".pptx")
    translated = 0
    untranslated = []
    fallback = []
    glossary_hits = []
    glossary_misses = []
    format_degradations = []
    for it in items:
        tid = it["id"]
        new_text = trans_map.get(tid)
        if not new_text:
            untranslated.append(tid)
            continue
        for term, tgt in glossary.items():
            if term in it["text"]:
                if tgt in new_text:
                    glossary_hits.append({"id": tid, "term": term})
                else:
                    glossary_misses.append({"id": tid, "term": term,
                                            "expected": tgt})
        nonempty = (set_pptx_para_text(it, new_text) if is_pptx
                    else set_docx_para_text(it, new_text))
        if nonempty is None:
            untranslated.append(tid)
            continue
        translated += 1
        if nonempty > 1:
            format_degradations.append(
                {"id": tid, "runs_with_text_before": nonempty,
                 "note": "mixed run formatting collapsed into first run's format"})
        if len(new_text) > args.overflow_ratio * max(1, len(it["text"])):
            if is_pptx:
                action = "needs_attention"
            else:
                action = ("font_shrunk" if shrink_docx_para_font(it)
                          else "needs_attention")
            fallback.append({"id": tid, "src_len": len(it["text"]),
                             "tgt_len": len(new_text), "action": action})
    try:
        container.save(args.output)
    except Exception as e:
        print(f"ERROR: cannot save output {args.output}: {e}", file=sys.stderr)
        return 2

    report = {
        "input": args.input,
        "output": args.output,
        "source_binding": {"source_sha256": ws["source_sha256"], "matched": True},
        "items_total": len(items),
        "translated": translated,
        "untranslated_ids": untranslated,
        "glossary_hits": glossary_hits,
        "glossary_misses": glossary_misses,
        "fallback_list": fallback,
        "format_degradations": format_degradations,
        "overflow_ratio_threshold": args.overflow_ratio,
        "mt_engine": "external (caller-provided translations)",
        "note": ("结构保持为确定性验证范围; 译文质量未验证(由提供方负责), "
                 "见 SKILL.md 诚实边界; 回读检查请运行 verify"),
    }
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
    print(json.dumps({"translated": translated, "untranslated": len(untranslated),
                      "fallback": len(fallback),
                      "format_degradations": len(format_degradations),
                      "glossary_misses": len(glossary_misses)}, ensure_ascii=False))
    return 0


def cmd_verify(args):
    """结构回读验证：输出与源逐项核对 id/style/location 与译文/原文，长文本风险。"""
    try:
        with open(args.worksheet, encoding="utf-8") as f:
            ws = json.load(f)
        trans_map = {}
        if args.translations:
            with open(args.translations, encoding="utf-8") as f:
                trans_map = parse_translations(json.load(f))
    except Exception as e:
        print(f"ERROR: bad input files: {e}", file=sys.stderr)
        return 2
    if trans_map is None:
        print("ERROR: translations must be {id: text} or [{id, translation}]",
              file=sys.stderr)
        return 2
    err = validate_binding(args.input, ws)
    if err:
        print(f"ERROR: {err}", file=sys.stderr)
        return 2
    try:
        src_items, _ = extract_items(args.input)
        out_items, _ = extract_items(args.output)
    except Exception as e:
        print(f"ERROR: cannot read documents: {e}", file=sys.stderr)
        return 2
    diffs = check_items_match(ws["items"], src_items)
    if diffs:
        print(f"ERROR: worksheet does not match source file structure: "
              f"{diffs[:5]}", file=sys.stderr)
        return 2

    structure_problems = []
    text_problems = []
    if len(src_items) != len(out_items):
        structure_problems.append(f"item count changed: source={len(src_items)} "
                                  f"output={len(out_items)}")
    text_checks = []
    long_text_risks = []
    for s, o in zip(src_items, out_items):
        if s["id"] != o["id"]:
            structure_problems.append(f"id sequence changed: {s['id']} -> {o['id']}")
        elif s["style"] != o["style"] or s["location"] != o["location"]:
            structure_problems.append(f"style/location changed for {s['id']}")
        t = trans_map.get(s["id"])
        expected = t or s["text"]
        ok = o["text"] == expected
        text_checks.append({"id": s["id"],
                            "expected_kind": "translation" if t else "original",
                            "ok": ok})
        if not ok:
            text_problems.append(f"id {s['id']}: output text != expected")
        if t and len(t) > args.overflow_ratio * max(1, len(s["text"])):
            long_text_risks.append({"id": s["id"], "src_len": len(s["text"]),
                                    "tgt_len": len(t)})
    result = {
        "input": args.input,
        "output": args.output,
        "source_binding": {"source_sha256": ws["source_sha256"], "matched": True},
        "items_total": len(src_items),
        "structure_preserved": not structure_problems,
        "structure_problems": structure_problems,
        "texts_verified": len(text_checks),
        "text_failures": [c["id"] for c in text_checks if not c["ok"]],
        "long_text_risks": long_text_risks,
        "verified": not (structure_problems or text_problems),
    }
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    for p in structure_problems + text_problems:
        print(f"VERIFY FAILED: {p}", file=sys.stderr)
    print(json.dumps({"structure_preserved": result["structure_preserved"],
                      "text_failures": len(result["text_failures"]),
                      "long_text_risks": len(long_text_risks),
                      "verified": result["verified"]}, ensure_ascii=False))
    return 0 if result["verified"] else 2


def main():
    ap = argparse.ArgumentParser(
        description="文档翻译保持版式（extract/apply/verify，工作表与源强绑定）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p1 = sub.add_parser("extract", help="抽取待译工作表（含源绑定 sha256）")
    p1.add_argument("--input", required=True)
    p1.add_argument("--worksheet", required=True)
    p1.set_defaults(func=cmd_extract)
    p2 = sub.add_parser("apply", help="校验源绑定后完整回填译文")
    p2.add_argument("--input", required=True)
    p2.add_argument("--worksheet", required=True)
    p2.add_argument("--translations", required=True)
    p2.add_argument("--output", required=True)
    p2.add_argument("--glossary")
    p2.add_argument("--report")
    p2.add_argument("--overflow-ratio", type=float, default=2.0)
    p2.set_defaults(func=cmd_apply)
    p3 = sub.add_parser("verify", help="输出结构/译文回读验证与长文本风险清单")
    p3.add_argument("--input", required=True)
    p3.add_argument("--worksheet", required=True)
    p3.add_argument("--output", required=True)
    p3.add_argument("--translations")
    p3.add_argument("--report")
    p3.add_argument("--overflow-ratio", type=float, default=2.0)
    p3.set_defaults(func=cmd_verify)
    args = ap.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
