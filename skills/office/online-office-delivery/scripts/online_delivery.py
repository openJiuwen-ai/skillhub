#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""online-office-delivery: 在线办公文档交付与回读（依赖宿主授权 CLI）。

命令语法不凭猜测：执行前先运行 CLI 自身 `--help` 核对必需 shortcut 是否存在，
核对不过即 blocked，绝不拼造不存在的子命令。交付成功必须同时满足：
真实导入完成（ready/job_status 闭环）+ CLI 返回的真实 URL + 内容回读一致 + 
可编辑校验（金丝雀写入回读）通过；缺任何一环即 failed/blocked，绝不伪造链接。

用法:
  python online_delivery.py --plan plan.json --outdir out/ [--cli lark-cli]

退出码: 0=delivered(全部文件导入+URL+内容回读+可编辑校验全通过)
        2=输入错误(计划非法/空清单/本地文件缺失/不支持的目标或格式)
        3=blocked(无授权CLI/命令语法无法从CLI帮助核实/CLI未授权)
        4=failed(导入/回读/权限校验失败或内容不一致)
"""
import argparse
import csv
import hashlib
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import uuid

CLI_CANDIDATES = ["lark-cli"]
ENV_VAR_NAMES = ["LARK_TOKEN", "FEISHU_TOKEN", "LARK_APP_ID", "FEISHU_APP_ID"]
REQUIRED_SHORTCUTS = ["+workbook-import", "+workbook-info", "+csv-get",
                      "+csv-put", "+sheet-create", "+sheet-delete"]
TARGET_SPECS = {
    "lark_sheet": {
        "subcommand": "sheets",
        "extensions": [".xlsx", ".csv"],
        "ext_note": ".xls 无可信本地值读取器（无法生成回读基准），请先转为 .xlsx",
    },
}
UNSUPPORTED_TARGET_NOTE = (
    "当前仅 lark_sheet 的导入/回读命令语法经过官方文档核对；"
    "docx/pptx/bitable 等目标的导入命令未获核实，不猜测子命令，拒绝执行")
CHECK_SHEET_TITLE = "__delivery_check__"
SAMPLE_MAX_SHEETS = 3
SAMPLE_MAX_ROWS = 5
SAMPLE_MAX_COLS = 5
IMPORT_TIMEOUT = 180
CALL_TIMEOUT = 60
NEXT_COMMAND_MAX_RETRIES = 3
AUTH_ERROR_PATTERNS = [
    "unauthorized", "401", "403", "invalid token", "authentication",
    "credential", "not logged in", "no valid token", "need_user_authorization",
    "token_missing", "missing_scope", "登录", "授权失败", "未授权",
]
# +csv-get 默认给每逻辑行加 "[row=N] " 前缀（annotated_csv）；解析前必须去前缀。
ROW_PREFIX_RE = re.compile(r"(?m)^\[row=\d+\] ?")
# JSON 信封机器解析前静音 CLI 升级/skills 提示，避免 _notice 混入 stdout。
QUIET_ENV = {
    "LARKSUITE_CLI_NO_UPDATE_NOTIFIER": "1",
    "LARKSUITE_CLI_NO_SKILLS_NOTIFIER": "1",
}


class BlockedError(Exception):
    def __init__(self, reason, missing=None):
        super().__init__(reason)
        self.reason = reason
        self.missing = missing or {}


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_envelope(stdout):
    """解析 CLI 输出的 {ok, identity, data, ...} 信封；解析不出返回 None。"""
    text = (stdout or "").strip()
    if not text:
        return None
    try:
        obj = json.loads(text)
        if isinstance(obj, dict):
            return obj
    except ValueError:
        pass
    start = text.find("{")
    while start != -1:
        depth = 0
        for i in range(start, len(text)):
            if text[i] == "{":
                depth += 1
            elif text[i] == "}":
                depth -= 1
                if depth == 0:
                    try:
                        obj = json.loads(text[start:i + 1])
                        if isinstance(obj, dict):
                            return obj
                    except ValueError:
                        break
        start = text.find("{", start + 1)
    return None


def strip_row_prefix(text):
    """去掉 +csv-get annotated_csv 每逻辑行前的 "[row=N] " 前缀。"""
    return ROW_PREFIX_RE.sub("", text or "")


def csv_text_to_matrix(text):
    """annotated_csv/csv 文本 → 二维值矩阵（先去 [row=N] 前缀）。"""
    return [row for row in csv.reader(strip_row_prefix(text).splitlines())]


def parse_error_envelope(text):
    """解析 CLI 错误信封（stderr）：返回 error 内的权限/修复字段；解析不出返回 None。"""
    env = parse_envelope(text)
    if isinstance(env, dict) and env.get("ok") is False and isinstance(
            env.get("error"), dict):
        err = env["error"]
        out = {k: err.get(k) for k in
               ("type", "subtype", "code", "message", "hint",
                "missing_scopes", "console_url")}
        return {k: v for k, v in out.items() if v is not None}
    return None


def deep_find(obj, key):
    """在嵌套 dict 中查找第一个名为 key 的字段值。"""
    if isinstance(obj, dict):
        if key in obj and obj[key] is not None:
            return obj[key]
        for v in obj.values():
            found = deep_find(v, key)
            if found is not None:
                return found
    elif isinstance(obj, list):
        for v in obj:
            found = deep_find(v, key)
            if found is not None:
                return found
    return None


def extract_url(envelope, stdout):
    """从信封或原始输出提取真实 URL；只提取，绝不构造。"""
    if envelope:
        url = deep_find(envelope, "url")
        if isinstance(url, str) and url.startswith(("http://", "https://")):
            return url, "envelope"
    for m in re.finditer(r"https?://[^\s\"'<>]+", stdout or ""):
        return m.group(0), "stdout_regex"
    return None, None


def looks_like_auth_error(text):
    low = (text or "").lower()
    return any(p in low for p in AUTH_ERROR_PATTERNS)


def split_command(cmd):
    """next_command 续查命令拆分（CLI 返回的字符串或列表）。"""
    if isinstance(cmd, list):
        return [str(x) for x in cmd]
    parts = shlex.split(cmd, posix=(os.name != "nt"))
    if os.name == "nt":
        parts = [p[1:-1] if len(p) >= 2 and p[0] == p[-1] and p[0] in "\"'"
                 else p for p in parts]
    return parts


class CliRunner:
    """执行宿主 CLI 并留全量调用审计日志。

    identity 非空时给 sheets 操作显式追加 `--as <identity>`（跳过 --help 探测），
    避免 CLI 自动在 bot/user 间隐式选择身份。
    """

    def __init__(self, exe, invocation_log, identity=None):
        self.exe = exe
        self.log = invocation_log
        self.identity = identity

    def _with_identity(self, argv):
        if (self.identity and "sheets" in argv and "--help" not in argv
                and "--as" not in argv):
            return argv + ["--as", self.identity]
        return argv

    def run(self, args, cwd, timeout=CALL_TIMEOUT):
        return self.run_argv([self.exe] + [str(a) for a in args], cwd,
                             timeout=timeout)

    def run_argv(self, argv, cwd, timeout=CALL_TIMEOUT):
        """执行完整命令（next_command 续查）；首 token 按需经 PATH 解析。"""
        argv = self._with_identity([str(a) for a in argv])
        if not os.path.isfile(argv[0]):
            resolved = shutil.which(argv[0])
            if resolved:
                argv = [resolved] + argv[1:]
        env = dict(os.environ)
        env.update(QUIET_ENV)
        try:
            proc = subprocess.run(
                argv, capture_output=True, text=True, encoding="utf-8",
                errors="replace", cwd=cwd, timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            self.log.append({"argv": argv, "cwd": cwd, "timeout": timeout})
            return None
        self.log.append({
            "argv": argv, "cwd": cwd, "returncode": proc.returncode,
            "stdout_excerpt": (proc.stdout or "")[:2000],
            "stderr_excerpt": (proc.stderr or "")[:500],
        })
        return proc


def load_local_sheets(path):
    """本地回读基准：xlsx 用 openpyxl 读 sheet 名与抽样值；csv 用标准库。

    返回 [(sheet_name_or_None, [[v, ...], ...])]，抽样不超过
    SAMPLE_MAX_SHEETS 张子表、每表前 SAMPLE_MAX_ROWS 行、前 SAMPLE_MAX_COLS 列。
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".csv":
        rows = []
        with open(path, encoding="utf-8-sig", newline="") as f:
            for i, row in enumerate(csv.reader(f)):
                if i >= SAMPLE_MAX_ROWS:
                    break
                rows.append(row[:SAMPLE_MAX_COLS])
        return [(None, rows)]
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for name in wb.sheetnames[:SAMPLE_MAX_SHEETS]:
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= SAMPLE_MAX_ROWS:
                    break
                rows.append(list(row)[:SAMPLE_MAX_COLS])
            sheets.append((name, rows))
        return sheets
    finally:
        wb.close()


def cells_equal(local, remote):
    """单元格级比对：本地数值型按数值容差比对，文本按字面比对。

    文本不走数值化（"001" 与 "1" 视为不一致，捕获导入端静默数值化失真）。
    """
    if local is None:
        local = ""
    if remote is None:
        remote = ""
    if isinstance(local, bool):
        return str(local).lower() == str(remote).strip().lower()
    if isinstance(local, (int, float)):
        try:
            return abs(float(local) - float(remote)) <= 1e-9
        except (TypeError, ValueError):
            return False
    return str(local).strip() == str(remote).strip()


def sample_dims(rows):
    nrows = len(rows)
    ncols = max((len(r) for r in rows), default=0)
    return max(nrows, 0), max(ncols, 0 if nrows == 0 else 1)


def a1_range(nrows, ncols):
    def col_name(n):
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s
    if nrows <= 0 or ncols <= 0:
        return None
    return "A1:%s%d" % (col_name(ncols), nrows)


def extract_matrix(envelope, stdout):
    """从 +csv-get 信封提取二维值矩阵。

    真实 lark-cli 的主入口是 data.annotated_csv（含 "[row=N] " 前缀），
    去前缀后解析；同时容忍 data.values / data.csv / data.cells / 纯 CSV 回退。
    """
    if envelope:
        data = envelope.get("data") if isinstance(envelope.get("data"), dict) else {}
        for holder in (data, envelope):
            values = holder.get("values")
            if isinstance(values, list) and all(
                    isinstance(r, list) for r in values):
                return values
            for key in ("annotated_csv", "csv"):
                csv_text = holder.get(key)
                if isinstance(csv_text, str) and csv_text.strip():
                    return csv_text_to_matrix(csv_text)
            cells = holder.get("cells")
            if isinstance(cells, list) and cells and isinstance(cells[0], list):
                return cells
    text = (stdout or "").strip()
    if text and not text.startswith("{"):
        try:
            rows = csv_text_to_matrix(text)
            if rows and any(any(c.strip() for c in r) for r in rows):
                return rows
        except csv.Error:
            pass
    return None


def run_import(runner, local_path, name, folder_token, identity=None):
    """执行 +workbook-import；ready/job_status 闭环，timed_out 按 next_command 续查。

    要求交付使用的身份（identity）与 CLI 信封 identity 一致，避免以错误身份
    （如 bot 而非 user）静默落在错误账号下。
    """
    cwd = os.path.dirname(os.path.abspath(local_path))
    args = ["sheets", "+workbook-import",
            "--file", os.path.basename(local_path)]
    if folder_token:
        args += ["--folder-token", folder_token]
    if name:
        args += ["--name", name]
    record = {"argv_display": " ".join(["<cli>", "sheets", "+workbook-import",
                                        "--file", os.path.basename(local_path)]
                                       + (["--name", name] if name else [])),
              "attempts": 0, "followed_next_command": False,
              "identity_requested": identity}
    proc = runner.run(args, cwd=cwd, timeout=IMPORT_TIMEOUT)
    record["attempts"] = 1
    if proc is None:
        record["ok"] = False
        record["error"] = "import 命令超时（%ds）" % IMPORT_TIMEOUT
        return None, record
    envelope = parse_envelope(proc.stdout)
    err_text = (proc.stderr or "") + (proc.stdout or "")
    if proc.returncode != 0 or not envelope or envelope.get("ok") is not True:
        record["ok"] = False
        record["returncode"] = proc.returncode
        record["error"] = (err_text.strip()[:400] or "CLI 无输出")
        if looks_like_auth_error(err_text):
            missing = {"cli_authorization":
                       "CLI 已安装但未通过授权校验：%s" % err_text.strip()[:200]}
            err_env = (parse_error_envelope(proc.stderr)
                       or parse_error_envelope(proc.stdout))
            if err_env:
                missing["error"] = err_env
                if err_env.get("missing_scopes"):
                    missing["missing_scopes"] = err_env["missing_scopes"]
                if err_env.get("console_url"):
                    missing["console_url"] = err_env["console_url"]
            raise BlockedError(
                "CLI 存在但调用被拒绝（疑似未授权）", missing)
        return None, record
    data = envelope.get("data") if isinstance(envelope.get("data"), dict) else {}
    record["identity"] = envelope.get("identity")
    retries = 0
    while (data.get("timed_out") is True and retries < NEXT_COMMAND_MAX_RETRIES):
        next_cmd = data.get("next_command")
        if not next_cmd:
            break
        retries += 1
        record["attempts"] += 1
        record["followed_next_command"] = True
        proc = runner.run_argv(split_command(next_cmd), cwd=cwd,
                               timeout=IMPORT_TIMEOUT)
        if proc is None:
            record["ok"] = False
            record["error"] = "next_command 续查超时"
            return None, record
        envelope = parse_envelope(proc.stdout)
        if proc.returncode != 0 or not envelope or envelope.get("ok") is not True:
            record["ok"] = False
            record["error"] = ((proc.stderr or "") + (proc.stdout or "")).strip()[:400]
            return None, record
        data = envelope.get("data") if isinstance(envelope.get("data"), dict) else {}
        record["identity"] = envelope.get("identity")
    record["timed_out_retries"] = retries
    url, url_source = extract_url(envelope, proc.stdout)
    token = deep_find(envelope, "token")
    ready = data.get("ready")
    job_status = data.get("job_status")
    record["url_source"] = url_source
    record["ready"] = ready
    record["job_status"] = job_status
    problems = []
    if not url:
        problems.append("CLI 成功返回但未给出 URL（拒绝无链接交付，不伪造链接）")
    if ready is not True:
        problems.append("导入未就绪（ready=%r）" % (ready,))
    if job_status not in (0, "0", None):
        problems.append("导入任务状态异常（job_status=%r）" % (job_status,))
    if job_status is None and ready is not True:
        problems.append("信封缺少 ready/job_status 完成信号")
    env_identity = record.get("identity")
    if identity and env_identity and env_identity != identity:
        problems.append("身份不符：请求 %r，CLI 实际 %r（拒绝在错误身份账号下交付）"
                        % (identity, env_identity))
    if problems:
        record["ok"] = False
        record["error"] = "; ".join(problems)
        return None, record
    record["ok"] = True
    return {"url": url, "token": token if isinstance(token, str) else None,
            "envelope_excerpt": json.dumps(envelope, ensure_ascii=False)[:800],
            "import": record}, record


def readback_content(runner, url, local_path):
    """内容回读：+workbook-info 核对子表集合，+csv-get 抽样比对本地基准值。"""
    result = {"verified": False, "reasons": [], "sampled": []}
    proc = runner.run(["sheets", "+workbook-info", "--url", url], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        result["reasons"].append("+workbook-info 调用失败")
        return result
    envelope = parse_envelope(proc.stdout)
    sheets_info = None
    if envelope:
        for holder in (envelope.get("data"), envelope):
            if isinstance(holder, dict) and isinstance(holder.get("sheets"), list):
                sheets_info = holder["sheets"]
                break
    if not sheets_info:
        result["reasons"].append("无法从 +workbook-info 解析子表清单")
        return result
    online_titles = []
    online_ids = {}
    for s in sheets_info:
        if not isinstance(s, dict):
            continue
        title = s.get("title") or s.get("sheet_name")
        if title:
            online_titles.append(title)
            if s.get("sheet_id"):
                online_ids[title] = s["sheet_id"]
    result["sheets_online"] = online_titles
    try:
        local_sheets = load_local_sheets(local_path)
    except Exception as e:
        result["reasons"].append("无法生成本地回读基准（%s）" % e)
        return result
    local_names = [n for n, _ in local_sheets if n is not None]
    if local_names:
        result["sheets_local"] = local_names
        missing = [n for n in local_names if n not in online_titles]
        extra = [t for t in online_titles if t not in local_names
                 and t != CHECK_SHEET_TITLE]
        if missing:
            result["reasons"].append("线上缺少子表: %s" % ", ".join(missing))
        if extra:
            result["reasons"].append("线上多出未知子表: %s" % ", ".join(extra))
        if missing or extra:
            return result
    mismatches_total = 0
    for sheet_name, local_rows in local_sheets:
        nrows, ncols = sample_dims(local_rows)
        rng = a1_range(nrows, ncols)
        if not rng:
            result["sampled"].append({"sheet": sheet_name or online_titles[0],
                                      "range": None, "cells": 0,
                                      "mismatches": [], "note": "本地无数据可抽样"})
            continue
        target_name = sheet_name or (online_titles[0] if online_titles else None)
        if not target_name:
            result["reasons"].append("无法确定回读目标子表")
            return result
        locate = (["--sheet-id", online_ids[target_name]]
                  if target_name in online_ids
                  else ["--sheet-name", target_name])
        args = ["sheets", "+csv-get", "--url", url] + locate + ["--range", rng]
        proc = runner.run(args, cwd=os.getcwd())
        if proc is None or proc.returncode != 0:
            result["reasons"].append("+csv-get 调用失败（%s）" % target_name)
            return result
        matrix = extract_matrix(parse_envelope(proc.stdout), proc.stdout)
        if matrix is None:
            result["reasons"].append("无法解析 +csv-get 返回值（%s!%s）"
                                     % (target_name, rng))
            return result
        mismatches = []
        for r in range(nrows):
            for c in range(ncols):
                lv = local_rows[r][c] if c < len(local_rows[r]) else ""
                rv = matrix[r][c] if r < len(matrix) and c < len(matrix[r]) else ""
                if not cells_equal(lv, rv):
                    col = chr(65 + c) if c < 26 else "C%d" % c
                    mismatches.append({"cell": "%s%d" % (col, r + 1),
                                       "local": lv if not isinstance(lv, str) else lv[:80],
                                       "online": str(rv)[:80]})
        mismatches_total += len(mismatches)
        result["sampled"].append({
            "sheet": target_name, "range": rng,
            "cells": nrows * ncols,
            "mismatches": mismatches[:10], "mismatch_count": len(mismatches)})
        if mismatches:
            result["reasons"].append(
                "回读不一致（%s!%s 共%d处，首处 %s 本地=%r 线上=%r）"
                % (target_name, rng, len(mismatches), mismatches[0]["cell"],
                   mismatches[0]["local"], mismatches[0]["online"]))
    result["verified"] = not result["reasons"]
    return result


def verify_edit_permission(runner, url, keep_check_sheet, warnings):
    """可编辑校验：新建检查子表→金丝雀写入→回读比对→（默认）删除检查子表。

    只有真实写入并被回读命中才判可编辑；任何一环失败即 not verified。
    """
    marker = "delivery-check-%s" % uuid.uuid4().hex[:12]
    check = {"method": "canary-write-readback", "sheet": CHECK_SHEET_TITLE,
             "marker_prefix": "delivery-check-", "verified": False}
    proc = runner.run(["sheets", "+sheet-create", "--url", url,
                       "--title", CHECK_SHEET_TITLE], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "+sheet-create 调用失败（无法创建检查子表）"
        return check
    envelope = parse_envelope(proc.stdout)
    sheet_id = deep_find(envelope, "sheet_id")
    if not isinstance(sheet_id, str) or not sheet_id:
        check["error"] = "未从 +sheet-create 返回中取得 sheet_id"
        return check
    check["sheet_id"] = sheet_id
    locate = ["--sheet-id", sheet_id]
    proc = runner.run(["sheets", "+csv-put", "--url", url] + locate +
                      ["--start-cell", "A1", "--csv", marker], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "金丝雀写入失败（+csv-put）：%s" % (
            ((proc.stderr or "") + (proc.stdout or "")).strip()[:200]
            if proc else "超时")
        return check
    proc = runner.run(["sheets", "+csv-get", "--url", url] + locate +
                      ["--range", "A1"], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "金丝雀回读失败（+csv-get）"
        return check
    matrix = extract_matrix(parse_envelope(proc.stdout), proc.stdout)
    value = None
    if matrix and isinstance(matrix[0], list) and matrix[0]:
        value = matrix[0][0]
    if value != marker:
        check["error"] = "金丝雀回读不一致（写入 %r，读回 %r）" % (marker, value)
        return check
    check["write_verified"] = True
    check["readback_matched"] = True
    check["verified"] = True
    check["cleanup"] = {"attempted": False, "ok": None}
    if not keep_check_sheet:
        check["cleanup"]["attempted"] = True
        proc = runner.run(["sheets", "+sheet-delete", "--url", url,
                           "--sheet-id", sheet_id, "--yes"], cwd=os.getcwd())
        ok = proc is not None and proc.returncode == 0
        check["cleanup"]["ok"] = ok
        if not ok:
            warnings.append(
                "检查子表 %s（sheet_id=%s）删除失败，保留在交付物中；"
                "可编辑校验已通过，不影响交付判定" % (CHECK_SHEET_TITLE, sheet_id))
    return check


def probe_cli(cli_name, invocation_log):
    """探测授权 CLI：PATH 解析 + 用 CLI 自身 --help 核对命令语法。

    返回 (cli_path, syntax_verification)；不读取任何环境变量的值。
    """
    cli_path = shutil.which(cli_name)
    if not cli_path:
        return None, None
    runner = CliRunner(cli_path, invocation_log)
    top = runner.run(["--help"], cwd=os.getcwd())
    sub = runner.run(["sheets", "--help"], cwd=os.getcwd())
    version = runner.run(["--version"], cwd=os.getcwd())
    verification = {
        "source": "CLI 自身 --help 输出（不猜测子命令）",
        "cli_resolved": cli_path,
        "cli_version": ((version.stdout or "").strip()[:100]
                        if version is not None else None),
        "top_level_help_ok": bool(top is not None and top.returncode == 0),
        "commands_verified": {},
    }
    help_text = ""
    if top is not None:
        help_text += (top.stdout or "")
    if sub is not None:
        help_text += "\n" + (sub.stdout or "")
        verification["subcommand_help_ok"] = sub.returncode == 0
    for shortcut in REQUIRED_SHORTCUTS:
        name = shortcut.lstrip("+")
        found = shortcut in help_text
        if not found:
            found = bool(re.search(r"(?m)^\s{0,6}%s\b" % re.escape(name),
                                   help_text or ""))
        verification["commands_verified"][shortcut] = found
    verification["all_required_verified"] = all(
        verification["commands_verified"].values())
    verification["help_excerpt"] = help_text[:2000]
    return cli_path, verification


def validate_plan(plan, plan_path):
    errors = []
    if not isinstance(plan, dict):
        return ["计划不是 JSON 对象"]
    target = plan.get("target") or {}
    if not isinstance(target, dict):
        errors.append("target 不是对象")
        target = {}
    ttype = target.get("type", "lark_sheet")
    if ttype not in TARGET_SPECS:
        errors.append("不支持的目标类型 %r：%s" % (ttype, UNSUPPORTED_TARGET_NOTE))
        return errors, None
    spec = TARGET_SPECS[ttype]
    files = plan.get("files")
    if not isinstance(files, list) or not files:
        errors.append("files 缺失或为空清单：没有待交付文件不得判 delivered")
        files = []
    seen = set()
    for i, f in enumerate(files):
        if not isinstance(f, dict) or not f.get("local"):
            errors.append("files[%d] 缺少 local" % i)
            continue
        local = f["local"]
        if not os.path.isabs(local):
            local = os.path.join(os.path.dirname(os.path.abspath(plan_path)),
                                 local)
        if local in seen:
            errors.append("files[%d] 重复文件 %s" % (i, local))
            seen.add(local)
        if not os.path.exists(local):
            errors.append("files[%d] 本地文件不存在: %s" % (i, local))
            continue
        if not os.path.isfile(local):
            errors.append("files[%d] 不是普通文件: %s" % (i, local))
            continue
        ext = os.path.splitext(local)[1].lower()
        if ext not in spec["extensions"]:
            errors.append("files[%d] 扩展名 %s 不受支持（%s 目标支持 %s；%s）"
                          % (i, ext, ttype, "/".join(spec["extensions"]),
                             spec["ext_note"]))
    return errors, ttype


def write_report(outdir, report):
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, "delivery-report.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    return path


def main():
    ap = argparse.ArgumentParser(description="在线办公文档交付与回读")
    ap.add_argument("--plan", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--cli", default=None,
                    help="宿主授权 CLI 可执行名（默认依次探测 %s）"
                         % "/".join(CLI_CANDIDATES))
    ap.add_argument("--as", dest="identity", default=None,
                    choices=["user", "bot"],
                    help="交付使用的显式身份（透传 CLI --as）；计划里的 "
                         "identity 优先级低于本参数")
    ap.add_argument("--profile", default=None,
                    help="仅 dingtalk_sheet 后端使用：透传 dws --profile")
    args = ap.parse_args()
    invocation_log = []
    warnings = []
    report = {
        "skill": "online-office-delivery",
        "status": None,
        "plan_path": os.path.abspath(args.plan),
        "probe": {"cli_candidates": [args.cli] if args.cli else CLI_CANDIDATES,
                  "env_vars_checked": ENV_VAR_NAMES,
                  "env_vars_present": [v for v in ENV_VAR_NAMES
                                       if v in os.environ]},
        "cli_invocation_log": invocation_log,
        "warnings": warnings,
        "notes": ("delivered 仅表示：经探测到的宿主授权 CLI 完成导入、取得 CLI "
                  "返回的真实 URL、内容回读与可编辑校验通过。链接一律取自 CLI "
                  "输出，绝不伪造；无授权 CLI 或语法无法核实时显式 blocked。"),
    }
    try:
        with open(args.plan, encoding="utf-8") as f:
            plan = json.load(f)
    except Exception as e:
        report["status"] = "input_error"
        report["errors"] = ["计划文件无法读取/解析: %s" % e]
        write_report(args.outdir, report)
        print(json.dumps({"status": "input_error",
                          "error": report["errors"][0]}, ensure_ascii=False))
        return 2
    # 可选钉钉后端：仅当计划显式选择 dingtalk_sheet 时分派到独立小模块；
    # lark_sheet（默认）路径完全不变，钉钉/dws 不是本技能的必需依赖。
    if (isinstance(plan, dict) and isinstance(plan.get("target"), dict)
            and plan["target"].get("type") == "dingtalk_sheet"):
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import dingtalk_delivery
        return dingtalk_delivery.run(args.plan, args.outdir, cli=args.cli,
                                     profile=getattr(args, "profile", None))
    errors, ttype = validate_plan(plan, args.plan)
    if isinstance(plan, dict) and plan.get("readback") is False:
        warnings.append("readback=false 被忽略：交付必须完成内容回读校验")
    if errors:
        report["status"] = "input_error"
        report["errors"] = errors
        write_report(args.outdir, report)
        print(json.dumps({"status": "input_error", "errors": errors},
                         ensure_ascii=False))
        return 2
    cli_names = [args.cli] if args.cli else CLI_CANDIDATES
    cli_path = None
    verification = None
    for name in cli_names:
        cli_path, verification = probe_cli(name, invocation_log)
        if cli_path:
            report["probe"]["cli_name"] = name
            report["probe"]["syntax_verification"] = verification
            break
    if not cli_path:
        report["status"] = "blocked"
        report["missing"] = {
            "clis": cli_names,
            "reason": "PATH 中未找到已授权的在线办公 CLI（如 lark-cli）；"
                      "技能不内置凭据，不以本地文件或截图冒充在线交付",
        }
        write_report(args.outdir, report)
        print(json.dumps({"status": "blocked",
                          "missing": report["missing"]["reason"]},
                         ensure_ascii=False))
        return 3
    if not report["probe"]["syntax_verification"]["all_required_verified"]:
        missing_cmds = [c for c, ok in
                        report["probe"]["syntax_verification"]
                        ["commands_verified"].items() if not ok]
        report["status"] = "blocked"
        report["missing"] = {
            "commands_from_help": missing_cmds,
            "reason": "必需命令未出现在 CLI 自身 --help 中，无法核实语法，"
                      "拒绝猜测子命令执行",
        }
        write_report(args.outdir, report)
        print(json.dumps({"status": "blocked",
                          "missing": report["missing"]["reason"]},
                         ensure_ascii=False))
        return 3
    identity = args.identity or (plan.get("identity") if isinstance(plan, dict)
                                 else None)
    if identity not in (None, "user", "bot"):
        warnings.append("计划 identity=%r 非法，已忽略（仅支持 user/bot）"
                        % identity)
        identity = None
    report["identity_requested"] = identity
    runner = CliRunner(cli_path, invocation_log, identity=identity)
    target = plan.get("target") or {}
    folder_token = target.get("folder_token") or ""
    keep = bool(plan.get("keep_check_sheet"))
    results = []
    all_ok = True
    for f in plan["files"]:
        local = f["local"]
        if not os.path.isabs(local):
            local = os.path.join(
                os.path.dirname(os.path.abspath(args.plan)), local)
        entry = {"local": local,
                 "name": f.get("name") or os.path.splitext(
                     os.path.basename(local))[0],
                 "local_sha256": file_sha256(local)}
        try:
            imported, import_rec = run_import(runner, local, entry["name"],
                                              folder_token, identity)
        except BlockedError as e:
            report["status"] = "blocked"
            report["missing"] = e.missing
            entry["status"] = "failed"
            entry["reasons"] = [e.reason]
            results.append(entry)
            report["results"] = results
            write_report(args.outdir, report)
            print(json.dumps({"status": "blocked", "missing": e.missing},
                             ensure_ascii=False))
            return 3
        entry["import"] = import_rec
        if not imported:
            entry["status"] = "failed"
            entry["reasons"] = [import_rec.get("error", "导入失败")]
            all_ok = False
            results.append(entry)
            continue
        entry["url"] = imported["url"]
        entry["token"] = imported["token"]
        entry["readback"] = readback_content(runner, imported["url"], local)
        entry["permission_check"] = verify_edit_permission(
            runner, imported["url"], keep, warnings)
        reasons = []
        if not entry["readback"]["verified"]:
            reasons += entry["readback"]["reasons"]
        if not entry["permission_check"]["verified"]:
            reasons.append("可编辑校验未通过：%s"
                           % entry["permission_check"].get("error", ""))
        entry["reasons"] = reasons
        entry["status"] = "delivered" if not reasons else "failed"
        if reasons:
            all_ok = False
        results.append(entry)
    report["results"] = results
    report["status"] = "delivered" if all_ok else "failed"
    write_report(args.outdir, report)
    print(json.dumps({"status": report["status"],
                      "files": [{"local": r["local"], "status": r["status"],
                                 "url": r.get("url")} for r in results]},
                     ensure_ascii=False))
    return 0 if all_ok else 4


if __name__ == "__main__":
    sys.exit(main())
