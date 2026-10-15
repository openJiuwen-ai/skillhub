#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""dingtalk_delivery: online-office-delivery 的钉钉（dws sheet）可选后端。

独立小模块，仅在计划 target.type == "dingtalk_sheet" 时由 online_delivery.py
分派调用，或直接作为 CLI 运行；不改变 lark_sheet 契约与默认路径，也不把
dingtalk/dws 作为其它办公技能的必需依赖。

命令语法不凭猜测：执行前先用官方 DingTalk Workspace CLI 自身 `--help`
（`dws --help` / `dws sheet --help` / `dws sheet range --help`）核对要用到的
命令是否存在，核对不过即 blocked，绝不拼造不存在的子命令。

本后端的命令事实仅来自官方 dws v1.0.64 的 --help / schema / 包内 reference：
  - 导入：`dws sheet import create --file <路径> (--folder-token|--workspace) [--name]`
    成功返回 {success:true, taskId, documentUrl, documentName, documentType,
    nodeId}；轮询超时返回 {success:false, timed_out:true, taskId, status,
    next_command}，退出码仍为 0，需按 next_command 续查（≤3 次）。
  - 列子表：`dws sheet +list-sheets --node <nodeId>`（严格，缺集合即失败）。
  - 回读：`dws sheet +read --node --sheet-id --range`（严格，截断即失败）
    或 `dws sheet csv-get`（带 "[row=N] " 前缀的 CSV）。
  - 新建/删除工作表：`dws sheet new --node --name` /
    `dws sheet delete-sheet --node --sheet-id --yes`（删除为 user_required）。
  - 写入：`dws sheet csv-put --node --sheet-id --start-cell --csv`。

交付成功必须同时满足：导入完成 + CLI 返回的真实 documentUrl + 内容回读
一致（严格未截断）+ 可编辑校验（金丝雀写入回读）通过；缺任何一环即
failed/blocked，绝不伪造链接、不以本地文件冒充在线交付。
"""
import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import uuid

CLI_CANDIDATES = ["dws"]
# 只核对本管线实际使用的 sheet 命令（子表集合/严格读/CSV读写/工作表增删/区域）。
REQUIRED_SHEET_COMMANDS = [
    "import", "new", "delete-sheet", "+list-sheets", "+read",
    "csv-get", "csv-put", "range",
]
REQUIRED_RANGE_COMMANDS = ["read", "update"]
TARGET_TYPE = "dingtalk_sheet"
# 官方 sheet import 只支持 xlsx/xls，且要求 --folder-token 或 --workspace 至少一个；
# 本地回读基准需可信读取器，openpyxl 不读 .xls，故本后端只接受 .xlsx（.xls 请先转存）。
SUPPORTED_EXTENSIONS = [".xlsx"]
EXT_NOTE = ("钉钉 sheet import 支持 xlsx/xls，但本后端回读基准依赖 openpyxl（不读 .xls），"
            "且 import 不支持 csv；请先将 .xls/.csv 转为 .xlsx")
UNSUPPORTED_TARGET_NOTE = (
    "dingtalk_delivery 仅处理 target.type=dingtalk_sheet")
CHECK_SHEET_TITLE = "__delivery_check__"
SAMPLE_MAX_SHEETS = 3
SAMPLE_MAX_ROWS = 5
SAMPLE_MAX_COLS = 5
IMPORT_TIMEOUT = 360
CALL_TIMEOUT = 60
NEXT_COMMAND_MAX_RETRIES = 3
# dws 未授权/未登录时的提示（access token 缺失、未登录等）。
AUTH_ERROR_PATTERNS = [
    "access token", "resolve access token", "not logged in", "未登录",
    "unauthorized", "401", "403", "invalid token", "authentication",
    "credential", "token", "授权", "未授权",
]
ORG_DENIED_PATTERNS = [
    "组织未授权", "需要组织主管理员", "cli访问", "cli access",
    "organization", "admin",
]
ADMIN_URL = "https://open-dev.dingtalk.com/fe/old#/developerSettings"
# dws 用双信封：raw 命令 {success,result,...}；shortcut {ok,outcome,result/error}；
# import create 为扁平 {success, documentUrl, nodeId, ...}；错误 {error:{...}} 走 stderr。
ROW_PREFIX_RE = re.compile(r"(?m)^\[row=\d+\] ?")


class BlockedError(Exception):
    def __init__(self, reason, missing=None):
        super().__init__(reason)
        self.reason = reason
        self.missing = missing or {}


def file_sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def parse_envelope(stdout):
    """解析 dws JSON 信封；解析不出返回 None。"""
    text = (stdout or "").strip()
    if not text:
        return None
    try:
        obj = json.loads(text)
        if isinstance(obj, dict):
            return obj
    except ValueError:
        pass
    start = text.find("{")
    while start != -1:
        depth = 0
        for i in range(start, len(text)):
            if text[i] == "{":
                depth += 1
            elif text[i] == "}":
                depth -= 1
                if depth == 0:
                    try:
                        obj = json.loads(text[start:i + 1])
                        if isinstance(obj, dict):
                            return obj
                    except ValueError:
                        break
        start = text.find("{", start + 1)
    return None


def envelope_ok(env):
    """dws 成功/失败判定：success/ok/outcome 三种形态；未知返回 None。"""
    if not isinstance(env, dict):
        return None
    if env.get("error"):
        return False
    if env.get("success") is True or env.get("ok") is True:
        return True
    if env.get("outcome") == "success":
        return True
    if (env.get("success") is False or env.get("ok") is False
            or env.get("outcome") == "failure"):
        return False
    return None


def envelope_error(env):
    if isinstance(env, dict) and isinstance(env.get("error"), dict):
        err = env["error"]
        out = {k: err.get(k) for k in
               ("type", "subtype", "category", "code", "message", "hint",
                "operation", "stage")}
        return {k: v for k, v in out.items() if v is not None}
    return {}


def extract_data(env):
    """返回承载业务字段的对象：优先 env['result']（dict），否则 env 本身。"""
    if isinstance(env, dict) and isinstance(env.get("result"), dict):
        return env["result"]
    return env if isinstance(env, dict) else {}


def deep_find(obj, key):
    if isinstance(obj, dict):
        if key in obj and obj[key] is not None:
            return obj[key]
        for v in obj.values():
            found = deep_find(v, key)
            if found is not None:
                return found
    elif isinstance(obj, list):
        for v in obj:
            found = deep_find(v, key)
            if found is not None:
                return found
    return None


def extract_url(env, stdout):
    """只提取真实 URL，绝不构造：优先 documentUrl，其次任意 url 字段/输出文本。"""
    for key in ("documentUrl", "url"):
        val = deep_find(env, key)
        if isinstance(val, str) and val.startswith(("http://", "https://")):
            return val, "envelope"
    for m in re.finditer(r"https?://[^\s\"'<>]+", stdout or ""):
        return m.group(0), "stdout_regex"
    return None, None


def looks_like_auth_error(text):
    low = (text or "").lower()
    return any(p.lower() in low for p in AUTH_ERROR_PATTERNS)


def looks_like_org_denied(text):
    low = (text or "").lower()
    return any(p.lower() in low for p in ORG_DENIED_PATTERNS)


def strip_row_prefix(text):
    return ROW_PREFIX_RE.sub("", text or "")


def csv_text_to_matrix(text):
    return [row for row in csv.reader(strip_row_prefix(text).splitlines())]


def split_command(cmd):
    """next_command 续查命令拆分（字符串或列表）。"""
    if isinstance(cmd, list):
        return [str(x) for x in cmd]
    parts = re.findall(r'"[^"]*"|\'[^\']*\'|\S+', cmd or "")
    out = []
    for p in parts:
        if len(p) >= 2 and p[0] == p[-1] and p[0] in "\"'":
            p = p[1:-1]
        out.append(p)
    return out


class DwsRunner:
    """执行官方 dws 并留全量调用审计日志。

    profile 非空时给业务命令显式追加 `--profile <profile>`，避免依赖隐式默认账号；
    机器可解析输出显式追加 `-f json`；`--help` 探测不追加。
    """

    def __init__(self, exe, invocation_log, profile=None):
        self.exe = exe
        self.log = invocation_log
        self.profile = profile

    def _decorate(self, argv):
        if "--help" in argv or "-h" in argv:
            return argv
        out = list(argv)
        if self.profile and "--profile" not in out:
            out += ["--profile", self.profile]
        if "-f" not in out and "--format" not in out:
            out += ["-f", "json"]
        return out

    def run(self, args, cwd, timeout=CALL_TIMEOUT):
        return self.run_argv([self.exe] + [str(a) for a in args], cwd,
                             timeout=timeout)

    def run_argv(self, argv, cwd, timeout=CALL_TIMEOUT):
        argv = [str(a) for a in argv]
        # CLI 返回的 next_command 首 token 可能是 "dws"/"dws.exe"：替换为本进程使用的
        # 可执行文件，避免解析到另一个 dws。
        if argv and (os.path.basename(argv[0]).lower() in ("dws", "dws.exe")
                     or argv[0] == self.exe):
            argv[0] = self.exe
        elif not os.path.isfile(argv[0]):
            resolved = shutil.which(argv[0])
            if resolved:
                argv = [resolved] + argv[1:]
        argv = self._decorate(argv)
        env = dict(os.environ)
        env["PYTHONUTF8"] = "1"
        try:
            proc = subprocess.run(
                argv, capture_output=True, text=True, encoding="utf-8",
                errors="replace", cwd=cwd, timeout=timeout, env=env)
        except subprocess.TimeoutExpired:
            self.log.append({"argv": argv, "cwd": cwd, "timeout": timeout})
            return None
        self.log.append({
            "argv": argv, "cwd": cwd, "returncode": proc.returncode,
            "stdout_excerpt": (proc.stdout or "")[:2000],
            "stderr_excerpt": (proc.stderr or "")[:500],
        })
        return proc


def load_local_sheets(path):
    """本地回读基准：xlsx 用 openpyxl 读取子表名与前 5×5 抽样值。"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for name in wb.sheetnames[:SAMPLE_MAX_SHEETS]:
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= SAMPLE_MAX_ROWS:
                    break
                rows.append(list(row)[:SAMPLE_MAX_COLS])
            sheets.append((name, rows))
        return sheets
    finally:
        wb.close()


def cells_equal(local, remote):
    """数值按容差、文本按字面；"001" != 1（捕获导入端静默数值化失真）。"""
    if local is None:
        local = ""
    if remote is None:
        remote = ""
    if isinstance(local, bool):
        return str(local).lower() == str(remote).strip().lower()
    if isinstance(local, (int, float)):
        try:
            return abs(float(local) - float(remote)) <= 1e-9
        except (TypeError, ValueError):
            return False
    return str(local).strip() == str(remote).strip()


def sample_dims(rows):
    nrows = len(rows)
    ncols = max((len(r) for r in rows), default=0)
    return max(nrows, 0), max(ncols, 0 if nrows == 0 else 1)


def a1_range(nrows, ncols):
    def col_name(n):
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s
    if nrows <= 0 or ncols <= 0:
        return None
    return "A1:%s%d" % (col_name(ncols), nrows)


def cli_error_text(proc):
    if proc is None:
        return ""
    return ((proc.stderr or "") + "\n" + (proc.stdout or "")).strip()


def classify_blocked(proc, context):
    """把 CLI 失败的输出归类为 blocked（未登录/组织未授权）或 None（普通失败）。"""
    text = cli_error_text(proc)
    if not text:
        return None
    env = parse_envelope(proc.stderr) or parse_envelope(proc.stdout)
    err = envelope_error(env) if env else {}
    if looks_like_org_denied(text):
        return {
            "cli_authorization": {
                "reason": "组织级 CLI 访问被拒绝（需组织主管理员开启）",
                "admin_url": ADMIN_URL,
                "message": text[:400],
                "error": err,
            }
        }
    if looks_like_auth_error(text):
        return {
            "cli_authorization": {
                "reason": "dws 未通过授权校验（未登录/缺 access token）",
                "message": text[:400],
                "error": err,
            }
        }
    return None


def probe_cli(cli_name, invocation_log, profile=None):
    """探测 dws：PATH 解析 + 用 --help 核对实际使用的 sheet 命令语法。"""
    cli_path = cli_name if cli_name and os.path.isfile(cli_name) else shutil.which(
        cli_name or "dws")
    if not cli_path:
        return None, None
    runner = DwsRunner(cli_path, invocation_log, profile=profile)
    top = runner.run(["--help"], cwd=os.getcwd())
    sheet = runner.run(["sheet", "--help"], cwd=os.getcwd())
    rng = runner.run(["sheet", "range", "--help"], cwd=os.getcwd())
    version = runner.run(["--version"], cwd=os.getcwd())
    verification = {
        "source": "dws 自身 --help 输出（不猜测子命令）",
        "cli_resolved": cli_path,
        "cli_version": ((version.stdout or "").strip()[:100]
                        if version is not None else None),
        "top_level_help_ok": bool(top is not None and top.returncode == 0),
        "sheet_help_ok": bool(sheet is not None and sheet.returncode == 0),
        "range_help_ok": bool(rng is not None and rng.returncode == 0),
        "commands_verified": {},
        "range_commands_verified": {},
    }
    sheet_text = sheet.stdout or "" if sheet is not None else ""
    range_text = rng.stdout or "" if rng is not None else ""
    for cmd in REQUIRED_SHEET_COMMANDS:
        verification["commands_verified"][cmd] = _help_has(sheet_text, cmd)
    for cmd in REQUIRED_RANGE_COMMANDS:
        verification["range_commands_verified"][cmd] = _help_has(range_text, cmd)
    verification["all_required_verified"] = (
        all(verification["commands_verified"].values())
        and all(verification["range_commands_verified"].values()))
    verification["help_excerpt"] = sheet_text[:1500]
    return cli_path, verification


def _help_has(text, token):
    if not text:
        return False
    if re.search(r"(?m)^\s{0,6}%s(\s|$)" % re.escape(token), text):
        return True
    return token in text


def validate_plan(plan, plan_path):
    errors = []
    if not isinstance(plan, dict):
        return ["计划不是 JSON 对象"]
    target = plan.get("target")
    if not isinstance(target, dict):
        return ["target 缺失或不是对象"]
    if target.get("type") != TARGET_TYPE:
        return ["%s（target.type=%r）" % (UNSUPPORTED_TARGET_NOTE,
                                          target.get("type"))]
    folder_token = target.get("folder_token")
    workspace = target.get("workspace")
    if not folder_token and not workspace:
        errors.append(
            "target 必须提供 folder_token 或 workspace 之一：官方 sheet import "
            "要求导入到指定文件夹/知识库（--folder-token 与 --workspace 至少一个），"
            "缺少目标位置时拒绝执行，不猜测默认位置")
    files = plan.get("files")
    if not isinstance(files, list) or not files:
        errors.append("files 缺失或为空清单：没有待交付文件不得判 delivered")
        files = []
    seen = set()
    for i, f in enumerate(files):
        if not isinstance(f, dict) or not f.get("local"):
            errors.append("files[%d] 缺少 local" % i)
            continue
        local = f["local"]
        if not os.path.isabs(local):
            local = os.path.join(os.path.dirname(os.path.abspath(plan_path)),
                                 local)
        if local in seen:
            errors.append("files[%d] 重复文件 %s" % (i, local))
        seen.add(local)
        if not os.path.exists(local):
            errors.append("files[%d] 本地文件不存在: %s" % (i, local))
            continue
        if not os.path.isfile(local):
            errors.append("files[%d] 不是普通文件: %s" % (i, local))
            continue
        ext = os.path.splitext(local)[1].lower()
        if ext not in SUPPORTED_EXTENSIONS:
            errors.append("files[%d] 扩展名 %s 不受支持（%s 目标支持 %s；%s）"
                          % (i, ext, TARGET_TYPE, "/".join(SUPPORTED_EXTENSIONS),
                             EXT_NOTE))
    return errors


def run_import(runner, local_path, name, target):
    """执行 `sheet import create`；timed_out 按 next_command 续查（≤3 次）。

    ready/job_status 不适用于钉钉（其信封为 success/documentUrl/nodeId），
    改为要求 success=true 且取得真实 documentUrl。
    """
    cwd = os.path.dirname(os.path.abspath(local_path))
    args = ["sheet", "import", "create",
            "--file", os.path.basename(local_path)]
    if target.get("folder_token"):
        args += ["--folder-token", target["folder_token"]]
    if target.get("workspace"):
        args += ["--workspace", target["workspace"]]
    if name:
        args += ["--name", name]
    record = {"argv_display": " ".join(
        ["<dws>", "sheet", "import", "create", "--file",
         os.path.basename(local_path)] + (["--name", name] if name else [])),
        "attempts": 0, "followed_next_command": False}
    proc = runner.run(args, cwd=cwd, timeout=IMPORT_TIMEOUT)
    record["attempts"] = 1
    if proc is None:
        record["ok"] = False
        record["error"] = "import 命令超时（%ds）" % IMPORT_TIMEOUT
        return None, record
    blocked = classify_blocked(proc, "import")
    if blocked:
        raise BlockedError("dws 调用被拒绝（未授权/组织未开启 CLI 访问）", blocked)
    env = parse_envelope(proc.stdout)
    # timed_out 信封为 {success:false, timed_out:true, next_command}（退出码 0），
    # 属于业务态异步未完成，需先于 success 判定处理。
    timed_out = deep_find(env, "timed_out") is True
    if proc.returncode != 0 or env is None or (
            not timed_out and envelope_ok(env) is not True):
        record["ok"] = False
        record["returncode"] = proc.returncode
        record["error"] = (cli_error_text(proc)[:400] or "CLI 无输出")
        return None, record
    retries = 0
    while (deep_find(env, "timed_out") is True
           and retries < NEXT_COMMAND_MAX_RETRIES):
        next_cmd = deep_find(env, "next_command")
        if not next_cmd:
            break
        retries += 1
        record["attempts"] += 1
        record["followed_next_command"] = True
        proc = runner.run_argv(split_command(next_cmd), cwd=cwd,
                               timeout=IMPORT_TIMEOUT)
        if proc is None:
            record["ok"] = False
            record["error"] = "next_command 续查超时"
            return None, record
        blocked = classify_blocked(proc, "import_get")
        if blocked:
            raise BlockedError("dws 续查被拒绝（未授权/组织未开启 CLI 访问）", blocked)
        env = parse_envelope(proc.stdout)
        if (proc.returncode != 0 or env is None
                or (deep_find(env, "timed_out") is not True
                    and envelope_ok(env) is not True)):
            record["ok"] = False
            record["error"] = cli_error_text(proc)[:400]
            return None, record
    record["timed_out_retries"] = retries
    if deep_find(env, "timed_out") is True:
        record["ok"] = False
        record["error"] = ("导入持续 timed_out（续查 %d 次未收敛），"
                           "保留 taskId 供人工兜底" % retries)
        return None, record
    url, url_source = extract_url(env, proc.stdout)
    node_id = deep_find(env, "nodeId")
    task_id = deep_find(env, "taskId")
    record["url_source"] = url_source
    record["task_id"] = task_id
    record["node_id"] = node_id
    problems = []
    if not url:
        problems.append("CLI 成功返回但未给出 documentUrl（拒绝无链接交付，不伪造链接）")
    if not isinstance(node_id, str) or not node_id:
        problems.append("CLI 成功返回但未给出 nodeId（后续回读/校验无文档标识）")
    if problems:
        record["ok"] = False
        record["error"] = "; ".join(problems)
        return None, record
    record["ok"] = True
    return {"url": url, "node_id": node_id, "task_id": task_id,
            "envelope_excerpt": json.dumps(env, ensure_ascii=False)[:800],
            "import": record}, record


def _extract_sheets(env):
    for holder in (extract_data(env), env):
        if isinstance(holder, dict) and isinstance(holder.get("sheets"), list):
            return holder["sheets"]
    return None


def _extract_cells_matrix(env):
    """严格读取 (+read) 的 cells → 值矩阵。"""
    for holder in (extract_data(env), env):
        if not isinstance(holder, dict):
            continue
        cells = holder.get("cells")
        if not isinstance(cells, list) or not cells:
            continue
        matrix = []
        for row in cells:
            if not isinstance(row, list):
                return None
            matrix.append([(c.get("value") if isinstance(c, dict) else c)
                           for c in row])
        return matrix
    return None


def _csv_from_envelope(env):
    for holder in (extract_data(env), env):
        if isinstance(holder, dict):
            for key in ("csv", "annotated_csv"):
                if isinstance(holder.get(key), str) and holder[key].strip():
                    return holder[key]
    return None


def list_sheets(runner, node_id):
    proc = runner.run(["sheet", "+list-sheets", "--node", node_id],
                      cwd=os.getcwd())
    if proc is None:
        return None, "调用超时"
    blocked = classify_blocked(proc, "list_sheets")
    if blocked:
        raise BlockedError("dws 列子表被拒绝（未授权/组织未开启 CLI 访问）", blocked)
    env = parse_envelope(proc.stdout)
    if proc.returncode != 0 or env is None or envelope_ok(env) is not True:
        return None, (cli_error_text(proc)[:300] or "CLI 无输出")
    sheets = _extract_sheets(env)
    if sheets is None:
        return None, "无法从 +list-sheets 解析子表集合（拒绝把未知结构当空结果）"
    return sheets, None


def readback_content(runner, node_id, local_path):
    result = {"verified": False, "reasons": [], "sampled": []}
    sheets_info, err = list_sheets(runner, node_id)
    if sheets_info is None:
        result["reasons"].append("+list-sheets 调用失败：%s" % err)
        return result
    online_titles = []
    online_ids = {}
    for s in sheets_info:
        if not isinstance(s, dict):
            continue
        title = s.get("title") or s.get("name") or s.get("sheetName")
        sid = s.get("sheetId") or s.get("sheet_id")
        if title:
            online_titles.append(title)
            if sid:
                online_ids[title] = sid
    result["sheets_online"] = online_titles
    try:
        local_sheets = load_local_sheets(local_path)
    except Exception as e:
        result["reasons"].append("无法生成本地回读基准（%s）" % e)
        return result
    local_names = [n for n, _ in local_sheets if n is not None]
    if local_names:
        result["sheets_local"] = local_names
        missing = [n for n in local_names if n not in online_titles]
        extra = [t for t in online_titles if t not in local_names
                 and t != CHECK_SHEET_TITLE]
        if missing:
            result["reasons"].append("线上缺少子表: %s" % ", ".join(missing))
        if extra:
            result["reasons"].append("线上多出未知子表: %s" % ", ".join(extra))
        if missing or extra:
            return result
    for sheet_name, local_rows in local_sheets:
        nrows, ncols = sample_dims(local_rows)
        rng = a1_range(nrows, ncols)
        target_name = sheet_name or (online_titles[0] if online_titles else None)
        if not rng:
            result["sampled"].append({"sheet": target_name, "range": None,
                                      "cells": 0, "mismatches": [],
                                      "note": "本地无数据可抽样"})
            continue
        if not target_name:
            result["reasons"].append("无法确定回读目标子表")
            return result
        locate = (["--sheet-id", online_ids[target_name]]
                  if target_name in online_ids else ["--sheet-id", target_name])
        proc = runner.run(["sheet", "+read", "--node", node_id] + locate +
                          ["--range", rng, "--value-render-option", "raw_value"],
                          cwd=os.getcwd())
        matrix = None
        strict_meta = None
        if proc is not None and proc.returncode == 0:
            env = parse_envelope(proc.stdout)
            if env is not None and envelope_ok(env) is True:
                # 严格读取：截断结果失败关闭，绝不把部分结果当完整。
                complete = deep_find(env, "complete")
                has_more = deep_find(env, "hasMore")
                reasons = deep_find(env, "truncationReasons")
                strict_meta = {"complete": complete, "hasMore": has_more,
                               "truncationReasons": reasons}
                if (has_more is True or (reasons not in (None, [], ()))
                        or complete is False):
                    result["reasons"].append(
                        "+read 结果为截断/不完整（%s），按严格契约失败关闭"
                        % strict_meta)
                    return result
                matrix = _extract_cells_matrix(env)
        if matrix is None:
            # 回退 csv-get（带 [row=N] 前缀的 CSV）。
            proc = runner.run(
                ["sheet", "csv-get", "--node", node_id] + locate +
                ["--range", rng, "--value-render-option", "raw_value"],
                cwd=os.getcwd())
            if proc is None or proc.returncode != 0:
                result["reasons"].append("+read 与 csv-get 均无法读取（%s）"
                                         % target_name)
                return result
            env = parse_envelope(proc.stdout)
            csv_text = _csv_from_envelope(env) if env is not None else None
            if csv_text is None:
                result["reasons"].append("无法解析回读返回值（%s!%s）"
                                         % (target_name, rng))
                return result
            matrix = csv_text_to_matrix(csv_text)
            strict_meta = {"source": "csv_get"}
        mismatches = []
        for r in range(nrows):
            for c in range(ncols):
                lv = local_rows[r][c] if c < len(local_rows[r]) else ""
                rv = matrix[r][c] if r < len(matrix) and c < len(matrix[r]) else ""
                if not cells_equal(lv, rv):
                    col = chr(65 + c) if c < 26 else "C%d" % c
                    mismatches.append({"cell": "%s%d" % (col, r + 1),
                                       "local": lv if not isinstance(lv, str)
                                       else lv[:80],
                                       "online": str(rv)[:80]})
        result["sampled"].append({
            "sheet": target_name, "range": rng, "cells": nrows * ncols,
            "source": (strict_meta or {}).get("source", "read_strict"),
            "mismatches": mismatches[:10], "mismatch_count": len(mismatches)})
        if mismatches:
            result["reasons"].append(
                "回读不一致（%s!%s 共%d处，首处 %s 本地=%r 线上=%r）"
                % (target_name, rng, len(mismatches), mismatches[0]["cell"],
                   mismatches[0]["local"], mismatches[0]["online"]))
    result["verified"] = not result["reasons"]
    return result


def verify_edit_permission(runner, node_id, keep_check_sheet, warnings):
    """可编辑校验：新建检查子表→金丝雀写入→回读比对→（默认）删除。"""
    marker = "delivery-check-%s" % uuid.uuid4().hex[:12]
    check = {"method": "canary-write-readback", "sheet": CHECK_SHEET_TITLE,
             "marker_prefix": "delivery-check-", "verified": False}
    proc = runner.run(["sheet", "new", "--node", node_id,
                       "--name", CHECK_SHEET_TITLE], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "sheet new 调用失败（无法创建检查子表）"
        return check
    env = parse_envelope(proc.stdout)
    sheet_id = deep_find(env, "sheetId") or deep_find(env, "sheet_id")
    if not isinstance(sheet_id, str) or not sheet_id:
        # new 未回传 sheetId 时用 +list-sheets 精确取。
        sheets, _ = list_sheets(runner, node_id)
        if sheets:
            for s in sheets:
                if isinstance(s, dict) and (
                        s.get("title") == CHECK_SHEET_TITLE):
                    sheet_id = s.get("sheetId") or s.get("sheet_id")
                    break
    if not isinstance(sheet_id, str) or not sheet_id:
        check["error"] = "未取得检查子表 sheetId"
        return check
    check["sheet_id"] = sheet_id
    locate = ["--sheet-id", sheet_id]
    proc = runner.run(["sheet", "csv-put", "--node", node_id] + locate +
                      ["--start-cell", "A1", "--csv", marker], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "金丝雀写入失败（csv-put）：%s" % (
            cli_error_text(proc)[:200] or "超时")
        return check
    proc = runner.run(["sheet", "csv-get", "--node", node_id] + locate +
                      ["--range", "A1"], cwd=os.getcwd())
    if proc is None or proc.returncode != 0:
        check["error"] = "金丝雀回读失败（csv-get）"
        return check
    env = parse_envelope(proc.stdout)
    csv_text = _csv_from_envelope(env) if env is not None else None
    value = None
    if csv_text is not None:
        matrix = csv_text_to_matrix(csv_text)
        if matrix and isinstance(matrix[0], list) and matrix[0]:
            value = matrix[0][0]
    if value != marker:
        check["error"] = "金丝雀回读不一致（写入 %r，读回 %r）" % (marker, value)
        return check
    check["write_verified"] = True
    check["readback_matched"] = True
    check["verified"] = True
    check["cleanup"] = {"attempted": False, "ok": None}
    if not keep_check_sheet:
        check["cleanup"]["attempted"] = True
        # delete-sheet 为 user_required：仅删除本轮自建的检查子表，显式 --yes。
        proc = runner.run(["sheet", "delete-sheet", "--node", node_id,
                           "--sheet-id", sheet_id, "--yes"], cwd=os.getcwd())
        ok = proc is not None and proc.returncode == 0
        check["cleanup"]["ok"] = ok
        if not ok:
            warnings.append(
                "检查子表 %s（sheetId=%s）删除失败，保留在交付物中；"
                "可编辑校验已通过，不影响交付判定"
                % (CHECK_SHEET_TITLE, sheet_id))
    return check


def check_auth_status(runner):
    """只读查看非秘密授权状态：authenticated 布尔与 message 文本。"""
    proc = runner.run(["auth", "status", "--readonly", "--format", "json"],
                      cwd=os.getcwd())
    if proc is None:
        return {"checked": True, "ok": None, "note": "auth status 调用超时"}
    env = parse_envelope(proc.stdout)
    out = {"checked": True, "returncode": proc.returncode}
    if isinstance(env, dict):
        out["authenticated"] = env.get("authenticated")
        out["message"] = env.get("message")
        out["ok"] = envelope_ok(env)
    else:
        out["ok"] = None
        out["message"] = cli_error_text(proc)[:200]
    return out


def write_report(outdir, report):
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, "delivery-report.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    return path


def _finish(report, outdir, code, stdout_obj):
    write_report(outdir, report)
    print(json.dumps(stdout_obj, ensure_ascii=False))
    return code


def run(plan_path, outdir, cli=None, profile=None):
    invocation_log = []
    warnings = []
    report = {
        "skill": "online-office-delivery",
        "backend": "dingtalk_sheet",
        "status": None,
        "plan_path": os.path.abspath(plan_path),
        "probe": {"cli_candidates": [cli] if cli else CLI_CANDIDATES},
        "cli_invocation_log": invocation_log,
        "warnings": warnings,
        "notes": ("dingtalk_sheet 后端：delivered 仅表示经探测到的官方 dws CLI 完成 "
                  "sheet import create、取得真实 documentUrl、严格回读一致与金丝雀"
                  "可编辑校验通过。链接一律取自 CLI 输出，绝不伪造；未授权/组织未"
                  "开启 CLI 访问或语法无法核实时显式 blocked（非能力缺失）。"),
    }
    try:
        with open(plan_path, encoding="utf-8") as f:
            plan = json.load(f)
    except Exception as e:
        report["status"] = "input_error"
        report["errors"] = ["计划文件无法读取/解析: %s" % e]
        return _finish(report, outdir, 2, {"status": "input_error",
                                           "error": report["errors"][0]})
    errors = validate_plan(plan, plan_path)
    if isinstance(plan, dict) and plan.get("readback") is False:
        warnings.append("readback=false 被忽略：交付必须完成内容回读校验")
    if errors:
        report["status"] = "input_error"
        report["errors"] = errors
        return _finish(report, outdir, 2, {"status": "input_error",
                                           "errors": errors})
    target = plan["target"]
    cli_names = [cli] if cli else CLI_CANDIDATES
    cli_path = None
    verification = None
    for name in cli_names:
        cli_path, verification = probe_cli(name, invocation_log, profile=profile)
        if cli_path:
            report["probe"]["cli_name"] = name
            report["probe"]["syntax_verification"] = verification
            break
    if not cli_path:
        report["status"] = "blocked"
        report["missing"] = {
            "clis": cli_names,
            "reason": "PATH 中未找到官方钉钉 CLI dws（可经 --cli 指定官方二进制）；"
                      "技能不内置凭据，不以本地文件或截图冒充在线交付",
        }
        return _finish(report, outdir, 3, {"status": "blocked",
                                           "missing": report["missing"]["reason"]})
    if not verification["all_required_verified"]:
        missing_cmds = [c for c, ok in verification["commands_verified"].items()
                        if not ok]
        missing_cmds += ["range %s" % c for c, ok in
                         verification["range_commands_verified"].items()
                         if not ok]
        report["status"] = "blocked"
        report["missing"] = {
            "commands_from_help": missing_cmds,
            "reason": "必需命令未出现在 dws 自身 --help 中，无法核实语法，"
                      "拒绝猜测子命令执行",
        }
        return _finish(report, outdir, 3, {"status": "blocked",
                                           "missing": report["missing"]["reason"]})
    runner = DwsRunner(cli_path, invocation_log, profile=profile)
    report["auth_status"] = check_auth_status(runner)
    if report["auth_status"].get("authenticated") is False:
        report["status"] = "blocked"
        report["missing"] = {
            "cli_authorization": {
                "reason": "dws 未登录（auth status authenticated=false）；"
                          "需要先由用户完成授权（登录/组织开启 CLI 访问）",
                "auth_status": report["auth_status"],
                "admin_url": ADMIN_URL,
            },
        }
        return _finish(report, outdir, 3, {"status": "blocked",
                                           "missing": report["missing"]})
    folder_token = target.get("folder_token") or ""
    keep = bool(plan.get("keep_check_sheet"))
    results = []
    all_ok = True
    for f in plan["files"]:
        local = f["local"]
        if not os.path.isabs(local):
            local = os.path.join(os.path.dirname(os.path.abspath(plan_path)),
                                 local)
        entry = {"local": local,
                 "name": f.get("name") or os.path.splitext(
                     os.path.basename(local))[0],
                 "local_sha256": file_sha256(local)}
        try:
            imported, import_rec = run_import(runner, local, entry["name"],
                                              target)
        except BlockedError as e:
            report["status"] = "blocked"
            report["missing"] = e.missing
            entry["status"] = "failed"
            entry["reasons"] = [e.reason]
            results.append(entry)
            report["results"] = results
            return _finish(report, outdir, 3, {"status": "blocked",
                                               "missing": e.missing})
        entry["import"] = import_rec
        if not imported:
            entry["status"] = "failed"
            entry["reasons"] = [import_rec.get("error", "导入失败")]
            all_ok = False
            results.append(entry)
            continue
        entry["url"] = imported["url"]
        entry["node_id"] = imported["node_id"]
        entry["readback"] = readback_content(runner, imported["node_id"], local)
        entry["permission_check"] = verify_edit_permission(
            runner, imported["node_id"], keep, warnings)
        reasons = []
        if not entry["readback"]["verified"]:
            reasons += entry["readback"]["reasons"]
        if not entry["permission_check"]["verified"]:
            reasons.append("可编辑校验未通过：%s"
                           % entry["permission_check"].get("error", ""))
        entry["reasons"] = reasons
        entry["status"] = "delivered" if not reasons else "failed"
        if reasons:
            all_ok = False
        results.append(entry)
    report["results"] = results
    report["status"] = "delivered" if all_ok else "failed"
    return _finish(report, outdir, 0 if all_ok else 4, {
        "status": report["status"],
        "backend": "dingtalk_sheet",
        "files": [{"local": r["local"], "status": r["status"],
                   "url": r.get("url")} for r in results]})


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="online-office-delivery 钉钉（dws sheet）可选后端")
    ap.add_argument("--plan", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--cli", default=None,
                    help="官方 dws 可执行文件路径（默认探测 PATH 中的 dws）")
    ap.add_argument("--profile", default=None,
                    help="传给 dws 的组织/账号 profile（透传 --profile）")
    args = ap.parse_args(argv)
    return run(args.plan, args.outdir, cli=args.cli, profile=args.profile)


if __name__ == "__main__":
    sys.exit(main())
