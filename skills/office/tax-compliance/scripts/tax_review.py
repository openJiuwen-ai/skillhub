#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""tax-compliance: 申报材料归集 + 适用税种规则分析（可溯源）。

门控原则：
- 规则仅在业务期间(period)/管辖地(jurisdiction)/纳税人类型(taxpayer_type)等事实齐备、
  且业务期间落在规则生效/失效窗口内时才可输出 applicable；
- 事实缺失、期间不可解析、阶段性优惠条件未确认 → needs_confirmation 并列出缺失项；
- 期间在规则窗口外 → not_applicable_for_period（历史/现行规则分轨）；
- 金额用 Decimal 解析；无法解析/空值/无缓存公式/非有限数记入 data_errors，
  绝不静默按 0 归总；单位固定标注“元（人民币，材料口径）”；
- 材料目录为空或全部不可读 → materials_status=not_ready，全部规则降为 needs_confirmation；
- 永不输出“合规通过”结论，不代申报、不过账。
退出码：0=分析完成且无阻断性数据问题；1=存在数据错误或材料未就绪（报告仍写出）；2=用法/输入错误。
"""
import argparse
import csv
import json
import math
import os
import re
import sys
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook

RULES_PATH = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "references", "tax-rules-cn.json")
DISCLAIMER = ("本分析基于材料归集与公开法规条目工程摘录，不构成专业税务意见；"
              "适用以主管税务机关解释及现行有效法规文本为准。本工具不代为申报、不汇算、不过账。")
UNIT_LABEL = "元（人民币，材料口径）"

FACT_LABELS = {
    "period": "业务期间(period，支持YYYY/YYYYQn/YYYY-MM/YYYY-MM-DD)",
    "jurisdiction": "管辖地/纳税地点(jurisdiction)",
    "taxpayer_type": "纳税人类型(taxpayer_type：一般纳税人/小规模纳税人)",
    "activities": "业务活动(activities)",
    "entity_type": "主体身份(entity_type，如居民企业)",
    "filing_period_type": "计税期间口径(filing_period_type：月/季/按次)",
    "sales_computation_basis": "销售额口径(sales_computation_basis)",
    "sales_by_category": "应税交易类别构成(sales_by_category，区分不动产类与其他)",
    "preference_election": "优惠选择声明(preference_election)",
    "cross_border_service_details": "跨境服务交易条件(cross_border_service_details：service/buyer/consumed_overseas)",
    "export_documents": "出口单证(export_documents：customs_declaration/export_invoice/forex_receipt/retention_documents)",
}
KNOWN_TAXPAYER_TYPES = {"一般纳税人", "小规模纳税人"}

_RE_Q = re.compile(r"^(\d{4})[Qq]([1-4])$")
_RE_M = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")
_RE_D = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])-(0[1-9]|[12]\d|3[01])$")
_RE_Y = re.compile(r"^(\d{4})$")
_CLEAN_AMOUNT = re.compile(r"[,\s￥¥]")


class AmountError(ValueError):
    pass


def _end_of_month(year, month):
    if month == 12:
        return date(year, 12, 31)
    return date.fromordinal(date(year, month + 1, 1).toordinal() - 1)


def parse_period(raw):
    """把 business.period 解析为 {start,end}（date）；缺失或不可解析返回 None。"""
    if raw is None:
        return None
    s = str(raw).strip()
    try:
        m = _RE_Q.match(s)
        if m:
            y, q = int(m.group(1)), int(m.group(2))
            start = date(y, 3 * q - 2, 1)
            end = _end_of_month(y, 3 * q)
            return {"start": start, "end": end}
        m = _RE_M.match(s)
        if m:
            y, mo = int(m.group(1)), int(m.group(2))
            return {"start": date(y, mo, 1), "end": _end_of_month(y, mo)}
        m = _RE_D.match(s)
        if m:
            d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
            return {"start": d, "end": d}
        m = _RE_Y.match(s)
        if m:
            y = int(m.group(1))
            return {"start": date(y, 1, 1), "end": date(y, 12, 31)}
    except ValueError:
        return None
    return None


def parse_amount(raw):
    """解析金额为 Decimal；失败抛 AmountError（带中文原因），绝不静默返回 0。"""
    if raw is None:
        raise AmountError("空值（无缓存值的公式单元格或未填写）")
    if isinstance(raw, float) and not math.isfinite(raw):
        raise AmountError("非有限数值")
    if isinstance(raw, bool):
        raise AmountError(f"无法解析为金额：{raw!r}")
    if isinstance(raw, (int, float)):
        d = Decimal(str(raw))
    else:
        s = _CLEAN_AMOUNT.sub("", str(raw).replace("元", ""))
        if s == "":
            raise AmountError("金额为空")
        try:
            d = Decimal(s)
        except InvalidOperation:
            raise AmountError(f"无法解析为金额：{str(raw)!r}")
    if not d.is_finite():
        raise AmountError("非有限数值（NaN/Infinity）")
    return d


def read_table_rows(path):
    """读取表格为 (rows, error)。error 非 None 表示读取失败（文件损坏/编码错误等）。"""
    try:
        if path.lower().endswith(".csv"):
            with open(path, encoding="utf-8-sig", newline="") as f:
                return list(csv.reader(f)), None
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows, None
    except Exception as e:  # noqa: BLE001 - 读取失败必须上报而非吞掉
        return None, f"读取失败：{type(e).__name__}: {e}"


def _header_of(rows):
    return [str(c).strip() if c is not None else "" for c in (rows[0] if rows else [])]


def classify_file(path, rel, data_errors):
    """返回 (分类, stats)。读取失败记入 data_errors。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs)
            if "合同" in text[:500]:
                return "合同", {"paragraphs": len(doc.paragraphs)}
            return "其他文档", {"paragraphs": len(doc.paragraphs)}
        except Exception as e:  # noqa: BLE001
            data_errors.append({"file": rel, "row": None, "column": None,
                                "value": None, "reason": f"读取失败：{type(e).__name__}: {e}"})
            return "读取失败", {}
    if ext in (".xlsx", ".csv"):
        rows, err = read_table_rows(path)
        if err:
            data_errors.append({"file": rel, "row": None, "column": None,
                                "value": None, "reason": err})
            return "读取失败", {}
        header = _header_of(rows)
        flat = " ".join(header)
        if "发票" in flat or ("税额" in flat and "金额" in flat):
            return "发票台账", gather_invoice_stats(rows, rel, data_errors)
        if "日期" in flat and "金额" in flat:
            return "收支流水", gather_flow_stats(rows, rel, data_errors)
        return "其他表格", {"rows": max(len(rows) - 1, 0)}
    return "其他", {}


def _find_cols(header, keywords, exclude=()):
    for i, h in enumerate(header):
        if any(k in h for k in keywords) and not any(e in h for e in exclude):
            return i
    return None


def _iter_data_rows(rows):
    """跳过全空行；产出 (excel行号, row)。行号从2起（1为表头）。"""
    for idx, r in enumerate(rows[1:], start=2):
        if r and not all(c is None or str(c).strip() == "" for c in r):
            yield idx, r


def gather_invoice_stats(rows, rel, data_errors):
    """发票台账金额/税额合计（Decimal）。无法解析的值记入 data_errors，绝不按0归总。"""
    header = _header_of(rows)
    amount_col = _find_cols(header, ["金额"], exclude=("税额",))
    tax_col = _find_cols(header, ["税额"])
    amount_total = tax_total = Decimal("0")
    n = err_cells = 0
    for rowno, r in _iter_data_rows(rows):
        n += 1
        for col, total_name in ((amount_col, "amount"), (tax_col, "tax")):
            if col is None or col >= len(r):
                continue
            raw = r[col]
            try:
                val = parse_amount(raw)
            except AmountError as e:
                err_cells += 1
                data_errors.append({
                    "file": rel, "row": rowno, "column": header[col] if col < len(header) else col,
                    "value": str(raw) if raw is not None else None, "reason": str(e)})
                continue
            if total_name == "amount":
                amount_total += val
            else:
                tax_total += val
    q = lambda d: d.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)  # noqa: E731
    return {"invoice_rows": n, "error_cells": err_cells,
            "amount_total": float(q(amount_total)), "amount_total_exact": str(q(amount_total)),
            "tax_total": float(q(tax_total)), "tax_total_exact": str(q(tax_total)),
            "unit": UNIT_LABEL}


def gather_flow_stats(rows, rel, data_errors):
    """收支流水行数与金额可解析性检查（不归总收支，方向口径不明）。"""
    header = _header_of(rows)
    amount_col = _find_cols(header, ["金额"])
    n = err_cells = 0
    for rowno, r in _iter_data_rows(rows):
        n += 1
        if amount_col is not None and amount_col < len(r):
            try:
                parse_amount(r[amount_col])
            except AmountError as e:
                err_cells += 1
                data_errors.append({
                    "file": rel, "row": rowno,
                    "column": header[amount_col] if amount_col < len(header) else amount_col,
                    "value": str(r[amount_col]) if r[amount_col] is not None else None,
                    "reason": str(e)})
    return {"flow_rows": n, "amount_error_cells": err_cells, "unit": UNIT_LABEL}


def build_fact_state(business, period_info):
    """判定各事实是否具备。未知/缺失一律 False，不得据此输出 applicable。"""
    facts = {}
    facts["period"] = period_info is not None
    facts["jurisdiction"] = bool(str(business.get("jurisdiction") or "").strip())
    facts["taxpayer_type"] = business.get("taxpayer_type") in KNOWN_TAXPAYER_TYPES
    acts = business.get("activities")
    facts["activities"] = isinstance(acts, list) and len(acts) > 0
    et = str(business.get("entity_type") or "").strip()
    facts["entity_type"] = any(k in et for k in ("企业", "公司", "单位")) and bool(et)
    facts["filing_period_type"] = business.get("filing_period_type") in {"月", "季", "按次"}
    facts["sales_computation_basis"] = bool(str(business.get("sales_computation_basis") or "").strip())
    sbc = business.get("sales_by_category")
    facts["sales_by_category"] = isinstance(sbc, dict) and len(sbc) > 0
    pe = business.get("preference_election")
    facts["preference_election"] = bool(pe) and isinstance(pe, (dict, str))
    cbd = business.get("cross_border_service_details")
    facts["cross_border_service_details"] = (
        isinstance(cbd, dict) and bool(str(cbd.get("service") or "").strip())
        and bool(str(cbd.get("buyer") or "").strip())
        and cbd.get("consumed_overseas") is not None)
    ed = business.get("export_documents")
    facts["export_documents"] = isinstance(ed, dict) and all(
        ed.get(k) for k in ("customs_declaration", "export_invoice",
                            "forex_receipt", "retention_documents"))
    return facts


def select_rules(rules, business):
    """按 taxpayer_type/activities 粗筛；事实缺失时不筛除，交给事实核查标注。"""
    ttype = business.get("taxpayer_type")
    ttype_known = ttype in KNOWN_TAXPAYER_TYPES
    acts = set(business.get("activities") or [])
    selected = []
    for r in rules:
        cond = r.get("applies_when", {})
        types = cond.get("taxpayer_type", [])
        acts_req = cond.get("activities", [])
        if ttype_known and types and ttype not in types:
            continue
        if acts and acts_req and not acts & set(acts_req):
            continue
        selected.append(r)
    return selected


def evaluate_rule(rule, facts, period_info, mflags, materials_not_ready):
    src = rule["source"]
    entry = {
        "id": rule["id"], "tax_type": rule["tax_type"], "regime": rule.get("regime", "current"),
        "summary": rule["summary"],
        "source": {"regulation": src["regulation"], "document_number": src.get("document_number"),
                   "article": src["article"], "url": src.get("url"),
                   "url_verified": src.get("url_verified", False),
                   "effective_from": src.get("effective_from"),
                   "effective_to": src.get("effective_to"),
                   "checked_on": src.get("checked_on")},
        "scope": rule["scope"],
    }
    for opt in ("url_note", "supersedes", "cross_refs"):
        if src.get(opt):
            entry["source"][opt] = src[opt]
    missing, notes = [], []
    period_flagged = False
    if period_info is None:
        missing.append(FACT_LABELS["period"])
        notes.append("业务期间缺失或不可解析，无法进行法规有效期匹配；不得确认适用")
        period_flagged = True
    else:
        eff_from = src.get("effective_from")
        eff_to = src.get("effective_to")
        if eff_from and period_info["end"] < date.fromisoformat(eff_from):
            entry["status"] = "not_applicable_for_period"
            notes.append(f"业务期间早于本规则生效日（{eff_from}）；该期间应适用当时有效的规则"
                         "（见规则库 historical 条目或总局法规库检索当时文本）")
            entry["missing_facts"], entry["notes"] = missing, notes
            return entry
        if eff_to and period_info["start"] > date.fromisoformat(eff_to):
            entry["status"] = "not_applicable_for_period"
            notes.append(f"业务期间晚于本规则失效日（{eff_to}）；现行政策见规则库 current 条目")
            entry["missing_facts"], entry["notes"] = missing, notes
            return entry
    for f in rule.get("required_facts", []):
        if f == "period" and period_flagged:
            continue
        if not facts.get(f):
            missing.append(FACT_LABELS.get(f, f))
    rid = rule["id"]
    if rid == "INV-COMPLIANCE" and not mflags["has_invoice"]:
        missing.append("发票台账材料（未发现发票类材料，发票合规核对范围待确认）")
    if rid.startswith("VAT-CROSS") and not mflags["has_contract"]:
        missing.append("合同等跨境业务支撑材料（未发现合同文件）")
    if rid == "CIT-25" and not mflags["has_flow"]:
        missing.append("收支流水材料（收入口径待确认）")
    if materials_not_ready:
        missing.append("申报材料（材料目录为空或全部不可读，材料归集未就绪）")
    if rule.get("is_preference"):
        entry["conditions"] = rule.get("conditions_to_confirm", [])
        notes.append("阶段性优惠：与法定征收率/税率分条列示；上列条件全部确认前不得默认享受")
    entry["status"] = "needs_confirmation" if missing else "applicable"
    entry["missing_facts"] = missing
    entry["notes"] = notes
    return entry


def main():
    ap = argparse.ArgumentParser(description="税务合规材料归集与规则分析")
    ap.add_argument("--materials", required=True, help="材料目录")
    ap.add_argument("--business", required=True, help="业务背景 json")
    ap.add_argument("--outdir", required=True, help="输出目录")
    args = ap.parse_args()
    try:
        with open(args.business, encoding="utf-8") as f:
            business = json.load(f)
    except Exception as e:
        print(f"ERROR: bad business json: {e}", file=sys.stderr)
        return 2
    try:
        with open(RULES_PATH, encoding="utf-8") as f:
            rules = json.load(f).get("rules", [])
    except Exception as e:
        print(f"ERROR: bad rules library: {e}", file=sys.stderr)
        return 2
    if not os.path.isdir(args.materials):
        print(f"ERROR: materials dir not found: {args.materials}", file=sys.stderr)
        return 2

    files = []
    for root, _dirs, names in sorted(os.walk(args.materials)):
        for n in sorted(names):
            if not n.startswith("~$") and not n.startswith("."):
                files.append(os.path.join(root, n))
    data_errors = []
    ledger = []
    for f in files:
        rel = os.path.relpath(f, args.materials)
        cls, stats = classify_file(f, rel, data_errors)
        ledger.append({"file": rel, "class": cls, "stats": stats})
    read_failures = sum(1 for x in ledger if x["class"] == "读取失败")
    materials_status = "not_ready" if (not files or read_failures == len(files)) else "ready"

    period_info = parse_period(business.get("period"))
    facts = build_fact_state(business, period_info)
    mflags = {"has_contract": any(x["class"] == "合同" for x in ledger),
              "has_invoice": any(x["class"] == "发票台账" for x in ledger),
              "has_flow": any(x["class"] == "收支流水" for x in ledger)}
    selected = select_rules(rules, business)
    analysis = [evaluate_rule(r, facts, period_info, mflags,
                             materials_status == "not_ready") for r in selected]
    confirmations = []
    for a in analysis:
        for mf in a.get("missing_facts", []):
            confirmations.append({"rule_id": a["id"], "missing_fact": mf,
                                  "detail": f"[{a['id']}] {mf}"})

    os.makedirs(args.outdir, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "材料归集"
    ws.append(["文件", "分类", "关键统计"])
    for x in ledger:
        ws.append([x["file"], x["class"], json.dumps(x["stats"], ensure_ascii=False)])
    if data_errors:
        ws2 = wb.create_sheet("数据错误")
        ws2.append(["文件", "行", "列", "值", "原因"])
        for e in data_errors:
            ws2.append([e["file"], e["row"], e["column"],
                        e["value"] if e["value"] is not None else "", e["reason"]])
    wb.save(os.path.join(args.outdir, "materials-ledger.xlsx"))

    result = {
        "business": business,
        "business_facts": {
            "period_raw": business.get("period"),
            "period_start": period_info["start"].isoformat() if period_info else None,
            "period_end": period_info["end"].isoformat() if period_info else None,
            "jurisdiction": business.get("jurisdiction"),
            "taxpayer_type": business.get("taxpayer_type"),
            "entity_type": business.get("entity_type"),
            "activities": business.get("activities", []),
        },
        "materials_status": materials_status,
        "materials": ledger,
        "materials_summary": {
            "contracts": sum(1 for x in ledger if x["class"] == "合同"),
            "invoice_files": sum(1 for x in ledger if x["class"] == "发票台账"),
            "flow_files": sum(1 for x in ledger if x["class"] == "收支流水"),
            "read_failures": read_failures},
        "data_errors": data_errors,
        "rules": analysis,
        "confirmations": confirmations,
        "disclaimer": DISCLAIMER,
    }
    with open(os.path.join(args.outdir, "rule-analysis.json"), "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    with open(os.path.join(args.outdir, "confirmations.json"), "w", encoding="utf-8") as f:
        json.dump({"confirmations": confirmations}, f, ensure_ascii=False, indent=2)

    md = ["# 税务合规材料归集与规则分析", "",
          f"- 纳税人类型：{business.get('taxpayer_type')}",
          f"- 业务活动：{'、'.join(business.get('activities') or [])}",
          f"- 业务期间：{business.get('period')}"
          f"（{result['business_facts']['period_start']}~{result['business_facts']['period_end']}）"
          if period_info else f"- 业务期间：{business.get('period')}（缺失或不可解析）",
          f"- 管辖地：{business.get('jurisdiction')}",
          f"- 材料状态：{materials_status}（合同 {result['materials_summary']['contracts']} 份，"
          f"发票台账 {result['materials_summary']['invoice_files']} 份，"
          f"流水 {result['materials_summary']['flow_files']} 份，"
          f"读取失败 {read_failures} 份）", ""]
    if data_errors:
        md += ["## 数据错误（不按0静默归总）", ""]
        md += [f"- {e['file']} 行{e['row']} 列[{e['column']}] 值{e['value']!r}：{e['reason']}"
               for e in data_errors]
        md.append("")
    md += ["## 适用规则（逐条溯源）", ""]
    for a in analysis:
        md.append(f"### {a['id']} {a['tax_type']}（{a['status']}）")
        md.append(f"- 要点：{a['summary']}")
        md.append(f"- 来源：{a['source']['regulation']} {a['source']['article']}"
                   f"（{a['source']['document_number']}）生效 {a['source']['effective_from']}"
                   f" 至 {a['source']['effective_to'] or '—'}，核查 {a['source']['checked_on']}")
        if a["source"].get("url"):
            md.append(f"- 来源URL：{a['source']['url']}"
                       f"（{'已核验' if a['source']['url_verified'] else '未核验，见url_note'}）")
        md.append(f"- 适用范围：{a['scope']}")
        if a.get("conditions"):
            md.append(f"- 须确认条件：{'；'.join(a['conditions'])}")
        if a.get("missing_facts"):
            md.append(f"- 缺失事实：{'；'.join(a['missing_facts'])}")
        for n in a.get("notes", []):
            md.append(f"- 说明：{n}")
        md.append("")
    md += ["## 待确认项", ""]
    md += [f"- {c['detail']}" for c in confirmations] or ["（无）"]
    md += ["", "---", DISCLAIMER]
    with open(os.path.join(args.outdir, "rule-analysis.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(md))

    summary = {"files": len(ledger), "materials_status": materials_status,
               "data_errors": len(data_errors), "rules_selected": len(analysis),
               "rules_applicable": sum(1 for a in analysis if a["status"] == "applicable"),
               "confirmations": len(confirmations)}
    print(json.dumps(summary, ensure_ascii=False))
    if data_errors or materials_status != "ready":
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())



