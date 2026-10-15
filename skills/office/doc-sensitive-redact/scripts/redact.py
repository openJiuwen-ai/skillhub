#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""doc-sensitive-redact: DOCX/XLSX/PPTX/PDF 敏感信息打码（双轨识别 + 底层删除 + 交付级回读验证）。

用法:
  python redact.py --input <file_or_dir> --outdir out/ [--report manifest.json]
                   [--types id_card,phone,bank_card,email]

处理范围（按格式）:
  DOCX/PPTX: 直接在包内所有 XML 部件上做 run 级/文本节点级/属性级替换
             （正文、表格、页眉页脚、删除修订 w:delText、批注、脚注尾注、
              PPT 备注、幻灯片/版式/母版、图表文本、docProps、rels 等），
             未改动的部件按原始字节复制。
  XLSX: openpyxl 处理字符串/数值单元格、单元格批注、文档属性；
        保存后再对包内剩余 XML 部件做底层清洗，并保留（清洗后的）线程批注。
  PDF: 页面文本 redaction 底层删除 + 元数据/XMP/书签/批注信息/链接目标
       (URI/启动文件等字符串字段)/表单域值清洗 + 文件附件批注内容清洗
       (zip/Office 附件逐条目重写并保留无关内容, 附件包内 OOXML 部件按
        段落 run 拼接做跨 run 清洗, 无法可靠重写的二进制
        明文/加密/超深嵌套显式阻断) + 含敏感值的文档级附件移除;
       保存后对页面文本/链接/批注/附件(zip 递归展开)/元数据/XMP/书签
       做全类型回读验证, 表单/含图像页等未核验范围以 needs_review
       标注且不宣称 pdf_removed_and_verified。

安全策略:
  - 每个输出文件都做底层回读验证（逐部件搜索已识别的敏感值，
    二进制部件按原始字节搜索；PDF 额外对全部启用类型重新识别，
    覆盖链接目标与附件内容等新表面，而非仅搜索此前已识别的值）；
    任何残留 => 该文件标记 failed 并删除输出，
    不做笼统成功。无法清洗的位置（如二进制部件内的明文、不可靠
    重写的附件内容）会被显式阻断。
  - 打码清单不包含任何原始敏感值。

退出码: 0=全部完成且通过回读验证（可含 needs_review） 2=输入错误或任一文件失败/被阻断
"""
import argparse
import io
import json
import math
import os
import posixpath
import re
import sys
import zipfile

RE_ID = re.compile(r"(?<![0-9Xx])[1-9]\d{5}(?:18|19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3}[\dXx](?![0-9Xx])")
RE_ID_15 = re.compile(r"(?<!\d)[1-9]\d{5}\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3}(?!\d)")
RE_PHONE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
RE_BANK = re.compile(r"(?<!\d)\d{13,19}(?!\d)")
RE_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+")
ID_WEIGHTS = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
ID_CHECK = "10X98765432"
ALL_TYPES = ("id_card", "phone", "bank_card", "email")

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
SS_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
W_P, W_R = "{%s}p" % W_NS, "{%s}r" % W_NS
W_T, W_DELTEXT = "{%s}t" % W_NS, "{%s}delText" % W_NS
A_P, A_R, A_FLD, A_T = "{%s}p" % A_NS, "{%s}r" % A_NS, "{%s}fld" % A_NS, "{%s}t" % A_NS
SS_V = "{%s}v" % SS_NS
SS_SI, SS_R, SS_T = "{%s}si" % SS_NS, "{%s}r" % SS_NS, "{%s}t" % SS_NS
PARA_SPECS = {
    "docx": (W_P, (W_R,), (W_T, W_DELTEXT)),
    "pptx": (A_P, (A_R, A_FLD), (A_T,)),
}
# 附件包内 OOXML 部件 => 按格式选择结构化跨 run 清洗规格。
# word/ppt 走段落 run 拼接; xl/sharedStrings 的富文本 <r><t> 同样可能
# 跨 run 拆分敏感值, 以 <si> 为"段落"做同等拼接处理。
ATTACH_PART_SPECS = (
    ("word/", PARA_SPECS["docx"]),
    ("ppt/", PARA_SPECS["pptx"]),
    ("xl/sharedstrings.xml", (SS_SI, (SS_R,), (SS_T,))),
)


def _attach_para_spec(name):
    """附件包内条目名 => OOXML 结构化清洗规格(无匹配返回 None)。"""
    if not name:
        return None
    low = name.lower()
    for prefix, spec in ATTACH_PART_SPECS:
        if low.startswith(prefix):
            return spec
    return None
THREADED_CT = "application/vnd.ms-excel.threadedcomments+xml"
THREADED_REL_TYPE = "http://schemas.microsoft.com/office/2017/10/relationships/threadedComment"
DROP_WARN_RE = re.compile(r"(chart|oleobject|activex|control|customxml|querytable|connection)", re.I)


class VerificationLeak(Exception):
    def __init__(self, parts):
        self.parts = parts
        super().__init__(
            "verification_leak: sensitive value remains accessible in parts: "
            + ", ".join(parts[:10]))


class UncleanableAttachment(Exception):
    """附件内容含敏感值但格式无法可靠重写（二进制明文/加密/超深嵌套），
    必须显式阻断交付而不是笼统宣称安全。"""


def id_checksum_ok(v):
    try:
        s = sum(int(v[i]) * ID_WEIGHTS[i] for i in range(17))
        return ID_CHECK[s % 11] == v[17].upper()
    except (ValueError, IndexError):
        return False


def luhn_ok(v):
    digits = [int(c) for c in v]
    total = 0
    for i, d in enumerate(reversed(digits)):
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def find_hits(text, types=None):
    """返回 (hits, needs_review)。hits: [{type, value, start, end}]，按位置倒序。

    types 为 None 时识别全部类型；否则只保留指定类型（review 项随 id_card 开关）。
    """
    hits = []
    review = []
    taken = []
    want_id = types is None or "id_card" in types
    if want_id:
        for m in RE_ID.finditer(text):
            v = m.group(0)
            if id_checksum_ok(v):
                hits.append({"type": "id_card", "value": v, "start": m.start(), "end": m.end()})
            else:
                review.append({"type": "id_card_invalid_checksum", "value": v,
                               "start": m.start(), "end": m.end()})
            taken.append((m.start(), m.end()))
        for m in RE_ID_15.finditer(text):
            if any(s <= m.start() < e for s, e in taken):
                continue
            review.append({"type": "legacy_15_digit", "value": m.group(0),
                           "start": m.start(), "end": m.end()})
            taken.append((m.start(), m.end()))
    for m in RE_PHONE.finditer(text):
        if any(s <= m.start() < e for s, e in taken):
            continue
        taken.append((m.start(), m.end()))
        if types is not None and "phone" not in types:
            continue
        hits.append({"type": "phone", "value": m.group(0), "start": m.start(), "end": m.end()})
    for m in RE_BANK.finditer(text):
        if any(s <= m.start() < e for s, e in taken):
            continue
        taken.append((m.start(), m.end()))
        if types is not None and "bank_card" not in types:
            continue
        if luhn_ok(m.group(0)):
            hits.append({"type": "bank_card", "value": m.group(0),
                         "start": m.start(), "end": m.end()})
    for m in RE_EMAIL.finditer(text):
        if any(s <= m.start() < e for s, e in taken):
            continue
        if types is not None and "email" not in types:
            continue
        hits.append({"type": "email", "value": m.group(0), "start": m.start(), "end": m.end()})
    hits.sort(key=lambda h: h["start"], reverse=True)
    review.sort(key=lambda h: h["start"], reverse=True)
    return hits, review


def mask_for(value):
    return "■" * min(len(value), 12)


def _replace_in(text, hits):
    out = []
    last = 0
    for h in sorted(hits, key=lambda x: x["start"]):
        out.append(text[last:h["start"]])
        out.append(mask_for(h["value"]))
        last = h["end"]
    out.append(text[last:])
    return "".join(out)


def _keep_space(el, text_tags):
    if el.tag in text_tags and (el.text != (el.text or "").strip() or el.text == ""):
        el.set(XML_SPACE, "preserve")


def _direct_runs(p, para_tag, run_tags):
    """段落直属 run（排除嵌套段落如文本框内的段落，避免跨结构拼接与重复处理）。"""
    runs = []
    for r in p.iter():
        if r.tag not in run_tags:
            continue
        anc = r.getparent()
        nested = False
        while anc is not None and anc is not p:
            if anc.tag == para_tag:
                nested = True
                break
            anc = anc.getparent()
        if not nested:
            runs.append(r)
    return runs


def _splice_runs(run_nodes, text_tags, hits):
    """按命中（位置倒序、互不重叠）在 run 序列上做跨 run 完整替换。"""
    for h in hits:
        start, end = h["start"], h["end"]
        mask = mask_for(h["value"])
        pos = 0
        first_done = False
        for r in run_nodes:
            ts = [c for c in r if c.tag in text_tags]
            if not ts:
                continue
            rt = "".join(t.text or "" for t in ts)
            rs, re_ = pos, pos + len(rt)
            if re_ <= start or rs >= end:
                pos = re_
                continue
            prefix = rt[:max(0, start - rs)]
            suffix = rt[max(0, min(len(rt), end - rs)):]
            new_text = prefix + (mask if not first_done else "") + suffix
            first_done = True
            for t in ts:
                t.text = ""
            ts[0].text = new_text
            _keep_space(ts[0], text_tags)
            pos = re_


def scrub_tree(root, para_spec, types, loc_prefix="", skip_v=False):
    """对一棵 XML 树做三级清洗: 段落 run 级 / 文本节点级 / 属性级。

    返回 (hits, review, modified)；hits/review 均为清洗前在原文上识别到的内容。
    """
    hits = []
    review = []
    modified = False
    text_tags_all = para_spec[2] if para_spec is not None else (W_T, W_DELTEXT, A_T)
    if para_spec is not None:
        para_tag, run_tags, text_tags = para_spec
        for pi, p in enumerate(root.iter(para_tag)):
            run_nodes = _direct_runs(p, para_tag, run_tags)
            ts_all = []
            parts = []
            for r in run_nodes:
                ts = [c for c in r if c.tag in text_tags]
                if ts:
                    ts_all.append(ts)
                    parts.append("".join(t.text or "" for t in ts))
            if not ts_all:
                continue
            full = "".join(parts)
            hh, rr = find_hits(full, types)
            review.extend(rr)
            if hh:
                for h in hh:
                    h["location"] = "%s#p%d" % (loc_prefix, pi)
                hits.extend(hh)
                _splice_runs(run_nodes, text_tags, hh)
                modified = True
    run_tags_all = para_spec[1] if para_spec is not None else ()
    for el in root.iter():
        if skip_v and el.tag == SS_V:
            continue
        if el.tag in text_tags_all and run_tags_all:
            parent = el.getparent()
            if parent is not None and parent.tag in run_tags_all:
                continue
        txt = el.text
        if isinstance(txt, str) and txt:
            hh, rr = find_hits(txt, types)
            review.extend(rr)
            if hh:
                el.text = _replace_in(txt, hh)
                _keep_space(el, (W_T, W_DELTEXT, A_T))
                for h in hh:
                    h["location"] = loc_prefix
                hits.extend(hh)
                modified = True
    for el in root.iter():
        for k, v in list(el.attrib.items()):
            if isinstance(v, str) and v:
                hh, rr = find_hits(v, types)
                review.extend(rr)
                if hh:
                    el.set(k, _replace_in(v, hh))
                    for h in hh:
                        h["location"] = loc_prefix + "(attribute)"
                    hits.extend(hh)
                    modified = True
    return hits, review, modified

# ---------- XML 基础 ----------

def parse_xml(data):
    from lxml import etree
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    try:
        return etree.fromstring(data)
    except Exception:
        return None


def serialize_xml(root):
    from lxml import etree
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _xml_text_blob(data):
    """全文视图: 所有文本节点 + 属性值，各项之间以换行分隔（防止跨节点误拼接）。"""
    root = parse_xml(data)
    if root is None:
        return None
    parts = []
    for el in root.iter():
        if el.text:
            parts.append(el.text)
        for v in el.attrib.values():
            parts.append(v)
    return "\n".join(parts)


XML_SUFFIXES = (".xml", ".rels", ".vml")


def detect_in_part(name, data, types):
    """底层识别: XML 类部件用全文视图，其余按字节 latin-1 识别。"""
    vals = []
    if name.lower().endswith(XML_SUFFIXES):
        blob = _xml_text_blob(data)
        if blob is not None:
            hits, _ = find_hits(blob, types)
            vals.extend(h["value"] for h in hits)
            return vals
    hits, _ = find_hits(data.decode("latin-1"), types)
    vals.extend(h["value"] for h in hits)
    return vals


def _record(result, hits, review):
    for h in hits:
        item = result["redactions"].setdefault(h["type"], {"count": 0, "locations": []})
        item["count"] += 1
        item["locations"].append(h.get("location", ""))
    for h in review:
        result["needs_review"][h["type"]] = result["needs_review"].get(h["type"], 0) + 1


def _finalize(result):
    return {
        "redactions": [{"type": k, "count": v["count"], "locations": v["locations"]}
                       for k, v in result["redactions"].items()],
        "needs_review": [{"type": k, "count": v} for k, v in result["needs_review"].items()],
    }


def verify_ooxml(path, values):
    """底层回读验证: 输出包内任何部件（含二进制）不得再出现已识别的敏感值。"""
    if not values:
        return
    uniq = set(values)
    leaks = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            data = z.read(name)
            hit = any(v.encode("utf-8") in data for v in uniq)
            if not hit and name.lower().endswith(XML_SUFFIXES):
                blob = _xml_text_blob(data)
                if blob is not None:
                    hit = any(v in blob for v in uniq)
            if hit:
                leaks.append(name)
    if leaks:
        raise VerificationLeak(leaks)


def _write_zip(path, entries):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            z.writestr(name, data)


# ---------- DOCX / PPTX ----------

def process_ooxml(in_path, out_path, kind, types):
    spec = PARA_SPECS[kind]
    with zipfile.ZipFile(in_path) as z:
        entries = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    result = {"redactions": {}, "needs_review": {}}
    detected = []
    new_entries = []
    for name, data in entries:
        if name.lower().endswith(XML_SUFFIXES):
            root = parse_xml(data)
            if root is not None:
                hits, review, modified = scrub_tree(root, spec, types, loc_prefix=name)
                _record(result, hits, review)
                detected.extend(h["value"] for h in hits)
                if modified:
                    data = serialize_xml(root)
        else:
            detected.extend(detect_in_part(name, data, types))
        new_entries.append((name, data))
    if any("/embeddings/" in n for n, _ in entries):
        result["needs_review"]["embedded_object_needs_manual_review"] = 1
    _write_zip(out_path, new_entries)
    verify_ooxml(out_path, detected)
    return _finalize(result), detected


# ---------- XLSX ----------

def _partname_from_rel(base_part, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _sheet_part_map(data):
    """workbook.xml 中 sheet 名称 -> 部件名。"""
    wb = parse_xml(data.get("xl/workbook.xml", b""))
    rels = parse_xml(data.get("xl/_rels/workbook.xml.rels", b""))
    if wb is None or rels is None:
        return {}
    rel_map = {}
    for rel in rels:
        if rel.tag == "{%s}Relationship" % REL_NS:
            rel_map[rel.get("Id")] = rel.get("Target")
    out = {}
    for sh in wb.iter("{%s}sheet" % SS_NS):
        target = rel_map.get(sh.get("{%s}id" % R_NS))
        if target:
            out[sh.get("name")] = _partname_from_rel("xl/workbook.xml", target)
    return out


def _threaded_map(data):
    """sheet 名称 -> 线程批注部件名列表。"""
    out = {}
    sheets = _sheet_part_map(data)
    for name in data:
        m = re.match(r"xl/worksheets/_rels/(sheet\d+)\.xml\.rels$", name)
        if not m:
            continue
        rels = parse_xml(data[name])
        if rels is None:
            continue
        sheet_part = "xl/worksheets/%s.xml" % m.group(1)
        targets = []
        for rel in rels:
            if rel.tag == "{%s}Relationship" % REL_NS \
                    and (rel.get("Type") or "").endswith("/threadedComment"):
                targets.append(_partname_from_rel(sheet_part, rel.get("Target")))
        for sheet_name, part in sheets.items():
            if part == sheet_part and targets:
                out.setdefault(sheet_name, []).extend(targets)
    return out


def _ensure_override(ct_root, partname, content_type):
    pn = "/" + partname
    for ov in ct_root.iter("{%s}Override" % CT_NS):
        if ov.get("PartName") == pn:
            return
    from lxml import etree
    ov = etree.SubElement(ct_root, "{%s}Override" % CT_NS)
    ov.set("PartName", pn)
    ov.set("ContentType", content_type)


def _ensure_sheet_rel(entries_dict, sheet_part, threaded_part):
    from lxml import etree
    rels_part = posixpath.join(posixpath.dirname(sheet_part),
                               "_rels", posixpath.basename(sheet_part) + ".rels")
    target = posixpath.relpath(threaded_part, posixpath.dirname(sheet_part))
    root = None
    if rels_part in entries_dict:
        root = parse_xml(entries_dict[rels_part])
    if root is None:
        root = etree.fromstring(
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
    existing = {rel.get("Id") for rel in root}
    n = 1
    while "rIdTC%d" % n in existing:
        n += 1
    rel = etree.SubElement(root, "{%s}Relationship" % REL_NS)
    rel.set("Id", "rIdTC%d" % n)
    rel.set("Type", THREADED_REL_TYPE)
    rel.set("Target", target)
    entries_dict[rels_part] = serialize_xml(root)
    return True


def process_xlsx(in_path, out_path, types):
    from openpyxl import load_workbook
    with zipfile.ZipFile(in_path) as z:
        orig = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    orig_data = dict(orig)
    result = {"redactions": {}, "needs_review": {}}
    detected = []
    for name, data in orig:
        detected.extend(detect_in_part(name, data, types))
    threaded = _threaded_map(orig_data)

    wb = load_workbook(in_path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v:
                    hits, _ = find_hits(v, types)
                    if hits:
                        cell.value = _replace_in(v, hits)
                        for h in hits:
                            detected.append(h["value"])
                        _record(result, [dict(h, location="%s!%s" % (ws.title, cell.coordinate))], [])
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    if isinstance(v, float):
                        if not math.isfinite(v) or v != int(v):
                            continue
                        s = str(int(v))
                    else:
                        s = str(v)
                    if len(s) < 11:
                        continue
                    hits, review = find_hits(s, types)
                    _record(result, [], review)
                    if hits:
                        cell.value = _replace_in(s, hits)
                        for h in hits:
                            detected.append(h["value"])
                        _record(result, [dict(h, location="%s!%s" % (ws.title, cell.coordinate))], [])
                cmt = cell.comment
                if cmt is not None:
                    if cmt.text:
                        hits, _ = find_hits(cmt.text, types)
                        if hits:
                            cmt.text = _replace_in(cmt.text, hits)
                            for h in hits:
                                detected.append(h["value"])
                            _record(result, [dict(h, location="%s!%s(comment)" % (ws.title, cell.coordinate))], [])
                    if cmt.author:
                        hits, _ = find_hits(cmt.author, types)
                        if hits:
                            cmt.author = _replace_in(cmt.author, hits)
                            for h in hits:
                                detected.append(h["value"])
                            _record(result, [dict(h, location="%s!%s(comment-author)" % (ws.title, cell.coordinate))], [])
    for prop in ("creator", "lastModifiedBy", "title", "subject", "description",
                 "keywords", "identifier", "language", "category", "contentStatus"):
        val = getattr(wb.properties, prop, None)
        if isinstance(val, str) and val:
            hits, _ = find_hits(val, types)
            if hits:
                setattr(wb.properties, prop, _replace_in(val, hits))
                for h in hits:
                    detected.append(h["value"])
                _record(result, [dict(h, location="docProps")], [])
    wb.save(out_path)

    with zipfile.ZipFile(out_path) as z:
        cur = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    entries_dict = dict(cur)
    for name, data in cur:
        if not name.lower().endswith(XML_SUFFIXES):
            continue
        root = parse_xml(data)
        if root is None:
            continue
        hits, review, modified = scrub_tree(
            root, None, types, loc_prefix=name,
            skip_v=name.startswith("xl/worksheets/"))
        _record(result, hits, review)
        detected.extend(h["value"] for h in hits)
        if modified:
            entries_dict[name] = serialize_xml(root)

    preserved = 0
    out_sheets = _sheet_part_map(entries_dict)
    ct = parse_xml(entries_dict.get("[Content_Types].xml", b""))
    for sheet_name, parts in threaded.items():
        sheet_part = out_sheets.get(sheet_name)
        if not sheet_part:
            continue
        for tp in parts:
            if tp in entries_dict:
                continue
            root = parse_xml(orig_data.get(tp, b""))
            if root is None:
                continue
            hits, review, _m = scrub_tree(root, None, types, loc_prefix=tp)
            _record(result, hits, review)
            detected.extend(h["value"] for h in hits)
            entries_dict[tp] = serialize_xml(root)
            if ct is not None:
                _ensure_override(ct, tp, THREADED_CT)
            if _ensure_sheet_rel(entries_dict, sheet_part, tp):
                preserved += 1
    if ct is not None:
        entries_dict["[Content_Types].xml"] = serialize_xml(ct)
    _write_zip(out_path, list(entries_dict.items()))

    with zipfile.ZipFile(out_path) as z:
        out_names = set(z.namelist())
    dropped = sorted(
        n for n, _ in orig
        if n not in out_names and DROP_WARN_RE.search(n)
        and not n.startswith("xl/threadedComments/"))
    fin = _finalize(result)
    if dropped:
        fin["dropped_unmodeled_parts"] = dropped
    if preserved:
        fin["threaded_comments_preserved"] = preserved
    verify_ooxml(out_path, detected)
    return fin, detected

# ---------- PDF ----------

ZIP_HEADS = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")
MAX_NEST_DEPTH = 4


def _is_zip_data(data):
    return data[:4] in ZIP_HEADS


def _text_blob_clean(data, types):
    """文本视图清洗: 返回 (bytes, hits); 无法按文本处理时返回 None。"""
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = None
    if text is not None:
        hits, _ = find_hits(text, types)
        if hits:
            return _replace_in(text, hits).encode("utf-8"), hits
        return data, []
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        codec = "utf-16-le" if data[:2] == b"\xff\xfe" else "utf-16-be"
        try:
            text = data[2:].decode(codec)
        except UnicodeDecodeError:
            text = None
        if text is not None:
            hits, _ = find_hits(text, types)
            if hits:
                return data[:2] + _replace_in(text, hits).encode(codec), hits
            return data, []
    return None


def _ooxml_part_text(name, data):
    """附件包内 OOXML 部件的结构化文本视图: 逐段拼接 run 文本,
    使跨 run 拆分的敏感值(如 split@ + example.test)可被识别;
    部件名/XML 不匹配时返回 None。"""
    spec = _attach_para_spec(name)
    if spec is None:
        return None
    root = parse_xml(data)
    if root is None:
        return None
    para_tag, run_tags, text_tags = spec
    parts = []
    for p in root.iter(para_tag):
        parts.append("".join(
            t.text or "" for r in _direct_runs(p, para_tag, run_tags)
            for t in r if t.tag in text_tags))
    return "\n".join(parts)


def _blob_views(data, depth=0, name=None):
    """字节内容的可读文本视图: utf-8 / 带 BOM 的 utf-16 / latin-1;
    zip 容器递归展开条目(深度受限); OOXML 条目额外给出跨 run 结构化视图。"""
    views = []
    if name:
        structured = _ooxml_part_text(name, data)
        if structured is not None:
            views.append(structured)
    if _is_zip_data(data):
        if depth >= MAX_NEST_DEPTH:
            views.append(data.decode("latin-1", errors="ignore"))
            return views
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for n in z.namelist():
                    views.extend(_blob_views(z.read(n), depth + 1, n))
        except Exception:
            views.append(data.decode("latin-1", errors="ignore"))
        return views
    try:
        views.append(data.decode("utf-8"))
        return views
    except UnicodeDecodeError:
        pass
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        codec = "utf-16-le" if data[:2] == b"\xff\xfe" else "utf-16-be"
        try:
            views.append(data[2:].decode(codec))
            return views
        except UnicodeDecodeError:
            pass
    views.append(data.decode("latin-1", errors="ignore"))
    return views


def _clean_ooxml_part(name, data, types, where):
    """结构化清洗包内 OOXML 部件: 跨 run 拆分的敏感值按段落拼接识别并替换,
    保留其余段落/结构。非目标部件或无法解析返回 None。"""
    spec = _attach_para_spec(name)
    if spec is None:
        return None
    root = parse_xml(data)
    if root is None:
        return None
    hits, _review, modified = scrub_tree(root, spec, types, loc_prefix=where)
    if not modified:
        return data, []
    return serialize_xml(root), hits


def _clean_attachment_blob(blob, types, where, depth=0, name=None):
    """清洗附件内容, 返回 (cleaned_bytes, all_hits)。

    zip 容器逐条目递归重写(保留无关条目), OOXML 部件走结构化跨 run 清洗;
    文本视图做全类型替换; 二进制明文命中或超深嵌套/加密压缩包 =>
    抛 UncleanableAttachment 阻断整份交付。
    """
    if _is_zip_data(blob):
        if depth >= MAX_NEST_DEPTH:
            raise UncleanableAttachment(
                "%s: nested archive beyond supported depth" % where)
        try:
            with zipfile.ZipFile(io.BytesIO(blob)) as zin:
                items = [(n, zin.read(n)) for n in zin.namelist()]
        except Exception as e:
            raise UncleanableAttachment(
                "%s: unreadable or encrypted archive (%s)"
                % (where, type(e).__name__))
        out = io.BytesIO()
        all_hits = []
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for n, data in items:
                cleaned, hits = _clean_attachment_blob(
                    data, types, "%s:%s" % (where, n), depth + 1, n)
                zout.writestr(n, cleaned)
                all_hits.extend(hits)
        return out.getvalue(), all_hits
    sres = _clean_ooxml_part(name, blob, types, where)
    if sres is not None and sres[1]:
        return sres
    res = _text_blob_clean(blob, types)
    if res is not None:
        return res
    hits, _ = find_hits(blob.decode("latin-1"), types)
    if hits:
        raise UncleanableAttachment(
            "%s: sensitive values in non-rewritable binary content" % where)
    return blob, []


# ---------------------------------------------------------------------------
# PDF 宽松栈辅助: 字体编解码与内容流算子级删除。
# 来源: 本仓库 skills/office/pdf-nl-edit/scripts/pdf_edit.py(本仓库自有实现,
# 随根 LICENSE 分发)。为保持本技能自包含, 复制其字体解码/内容流分词与命中
# 切除逻辑, 未修改 pdf-nl-edit; 复用目的正是避免重复实现复杂字体缺陷。
# ---------------------------------------------------------------------------
try:
    from pdfminer.fontmetrics import FONT_METRICS as _STD_FONT_METRICS
except Exception:  # pragma: no cover - pdfminer.six 随 pdfplumber 提供
    _STD_FONT_METRICS = {}

_NON_TEXT_STANDARD_FONTS = ("Symbol", "ZapfDingbats")


def _standard_char_width(basefont, ch):
    name = (basefont or "").lstrip("/")
    if "+" in name:
        name = name.split("+", 1)[1]
    if not name or name in _NON_TEXT_STANDARD_FONTS:
        return None
    entry = _STD_FONT_METRICS.get(name)
    if not entry:
        return None
    w = entry[1].get(ch)
    return float(w) if w is not None else None


def _is_wide(ch):
    o = ord(ch)
    return (
        0x1100 <= o <= 0x115F
        or 0x2E80 <= o <= 0xA4CF
        or 0xAC00 <= o <= 0xD7A3
        or 0xF900 <= o <= 0xFAFF
        or 0xFE30 <= o <= 0xFE4F
        or 0xFF00 <= o <= 0xFF60
        or 0xFFE0 <= o <= 0xFFE6
    )


_WS = b"\x00\x09\x0a\x0c\x0d\x20"
_DELIM = b"()<>[]{}/%"


def _parse_tounicode(data):
    mapping = {}
    for blk in re.findall(rb"beginbfchar(.*?)endbfchar", data, re.S):
        for src, dst in re.findall(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]*)>", blk):
            try:
                mapping[int(src, 16)] = bytes.fromhex(dst.decode()).decode(
                    "utf-16-be", "replace")
            except ValueError:
                continue
    for blk in re.findall(rb"beginbfrange(.*?)endbfrange", data, re.S):
        for lo, hi, dst in re.findall(
                rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", blk):
            lo_i, hi_i, dv = int(lo, 16), int(hi, 16), int(dst, 16)
            for k in range(lo_i, hi_i + 1):
                mapping[k] = chr(dv + (k - lo_i))
    return mapping


class _FontCodec:
    """把字体资源归一化为 2 字节(CID)或 1 字节(简单字体)编解码。"""

    def __init__(self, name, font_dict):
        self.name = name
        self.subtype = str(font_dict.get("/Subtype", ""))
        self.basefont = str(font_dict.get("/BaseFont", ""))
        encoding = str(font_dict.get("/Encoding", ""))
        self.is_cid = self.subtype == "/Type0"
        self.two_byte = self.is_cid
        self.tounicode = {}
        tu = font_dict.get("/ToUnicode")
        if tu is not None:
            try:
                self.tounicode = _parse_tounicode(tu.get_object().get_data())
            except Exception:
                self.tounicode = {}
        self.predefined_unicode = False
        if self.is_cid:
            self.predefined_unicode = encoding.startswith("/Uni")
            if not self.predefined_unicode and encoding.startswith("/Identity"):
                self.predefined_unicode = True
        else:
            self.winansi = encoding in ("/WinAnsiEncoding", "/StandardEncoding", "")
            self.macroman = encoding == "/MacRomanEncoding"
        self.widths = {}
        self.cid_widths = {}
        self.dw = 1000.0
        self.missing_width = None
        self._rev = None
        try:
            if self.is_cid:
                self._load_cid_widths(font_dict)
            else:
                self._load_simple_widths(font_dict)
        except Exception:
            self.widths = {}
            self.cid_widths = {}

    def _load_simple_widths(self, font_dict):
        widths = font_dict.get("/Widths")
        first = font_dict.get("/FirstChar")
        if widths is not None and first is not None:
            try:
                arr = widths.get_object()
                base = int(first)
                for k in range(len(arr)):
                    try:
                        self.widths[base + k] = float(arr[k])
                    except (TypeError, ValueError, IndexError):
                        pass
            except Exception:
                self.widths = {}
        fd = font_dict.get("/FontDescriptor")
        if fd is not None:
            try:
                mw = fd.get_object().get("/MissingWidth")
                if mw is not None:
                    self.missing_width = float(mw)
            except Exception:
                pass

    def _load_cid_widths(self, font_dict):
        desc = font_dict.get("/DescendantFonts")
        d0 = None
        if desc is not None:
            try:
                d0 = desc.get_object()[0].get_object()
            except Exception:
                d0 = None
        if d0 is None:
            return
        dw = d0.get("/DW")
        if dw is not None:
            try:
                self.dw = float(dw)
            except (TypeError, ValueError):
                pass
        warr = d0.get("/W")
        if warr is None:
            return
        try:
            arr = warr.get_object()
        except Exception:
            return
        i, n = 0, len(arr)
        while i < n:
            try:
                first = int(arr[i])
            except (TypeError, ValueError):
                i += 1
                continue
            i += 1
            if i >= n:
                break
            item = arr[i]
            if hasattr(item, "get_object"):
                item = item.get_object()
            if isinstance(item, (list, tuple)) or type(item).__name__ == "ArrayObject":
                for k, v in enumerate(item):
                    try:
                        self.cid_widths[first + k] = float(v)
                    except (TypeError, ValueError):
                        pass
                i += 1
            else:
                i += 1
                if i >= n:
                    break
                try:
                    last = int(item)
                    wv = float(arr[i])
                except (TypeError, ValueError):
                    i += 1
                    continue
                i += 1
                for c in range(first, last + 1):
                    self.cid_widths[c] = wv

    def code_for(self, ch):
        if self.two_byte:
            if self.predefined_unicode:
                cp = ord(ch)
                return cp if cp <= 0xFFFF else None
            if self.tounicode:
                if self._rev is None:
                    self._rev = {v: k for k, v in self.tounicode.items()}
                return self._rev.get(ch)
            return None
        try:
            return ch.encode("cp1252")[0]
        except (UnicodeEncodeError, IndexError):
            return None

    def width_for_code(self, code, ch=None):
        if self.is_cid:
            return self.cid_widths.get(code, self.dw)
        if code in self.widths:
            return self.widths[code]
        w = _standard_char_width(self.basefont, ch) if ch is not None else None
        if w is not None:
            return w
        if self.missing_width is not None:
            return self.missing_width
        return 500.0

    def decode(self, raw):
        out = []
        if self.two_byte:
            codes = [int.from_bytes(raw[i:i + 2], "big")
                     for i in range(0, len(raw) - 1, 2)]
            if self.tounicode:
                for c in codes:
                    out.append(self.tounicode.get(c, "\ufffd"))
                return "".join(out)
            return raw.decode("utf-16-be", "replace")
        else:
            for b in raw:
                if self.tounicode and b in self.tounicode:
                    out.append(self.tounicode[b])
                elif self.macroman:
                    out.append(bytes([b]).decode("mac-roman", "replace"))
                else:
                    out.append(bytes([b]).decode("cp1252", "replace"))
            return "".join(out)


class _Tok:
    __slots__ = ("kind", "value", "start", "end", "items")

    def __init__(self, kind, value, start, end, items=None):
        self.kind = kind
        self.value = value
        self.start = start
        self.end = end
        self.items = items


def _scan_literal(data, i):
    depth = 1
    j = i + 1
    n = len(data)
    out = bytearray()
    while j < n:
        c = data[j]
        if c == 0x5C:
            if j + 1 < n:
                nxt = data[j + 1]
                simple = {0x6E: 10, 0x72: 13, 0x74: 9, 0x62: 8, 0x66: 12,
                          0x28: 0x28, 0x29: 0x29, 0x5C: 0x5C}
                if nxt in simple:
                    out.append(simple[nxt])
                    j += 2
                    continue
                if 0x30 <= nxt <= 0x37:
                    oct_digits = []
                    k = j + 1
                    while k < n and len(oct_digits) < 3 and 0x30 <= data[k] <= 0x37:
                        oct_digits.append(data[k])
                        k += 1
                    out.append(int(bytes(oct_digits), 8) & 0xFF)
                    j = k
                    continue
                if nxt in (0x0A, 0x0D):
                    j += 2
                    if nxt == 0x0D and j < n and data[j] == 0x0A:
                        j += 1
                    continue
                out.append(nxt)
                j += 2
                continue
            j += 1
            continue
        if c == 0x28:
            depth += 1
            out.append(c)
            j += 1
            continue
        if c == 0x29:
            depth -= 1
            if depth == 0:
                return bytes(out), j + 1
            out.append(c)
            j += 1
            continue
        out.append(c)
        j += 1
    return bytes(out), j


def _scan_hex(data, i):
    j = data.find(b">", i)
    n = len(data)
    if j < 0:
        j = n
    raw = re.sub(rb"[^0-9A-Fa-f]", b"", data[i + 1:j])
    if len(raw) % 2:
        raw += b"0"
    return bytes.fromhex(raw.decode()), (j + 1 if j < n else n)


def _skip_balanced(data, i, opener, closer):
    depth = 0
    j = i
    n = len(data)
    while j < n:
        if data[j:j + len(opener)] == opener:
            depth += 1
            j += len(opener)
            continue
        if data[j:j + len(closer)] == closer:
            depth -= 1
            j += len(closer)
            if depth == 0:
                return j
            continue
        j += 1
    return n


def _next_token(data, i):
    n = len(data)
    while i < n and data[i] in _WS:
        i += 1
    if i >= n:
        return None, i
    c = data[i]
    if c == 0x25:  # %
        j = data.find(b"\n", i)
        return _next_token(data, n if j < 0 else j + 1)
    if c == 0x2F:  # /
        j = i + 1
        while j < n and data[j] not in _WS and data[j] not in _DELIM:
            j += 1
        return _Tok("name", data[i + 1:j], i, j), j
    if c == 0x28:  # (
        raw, j = _scan_literal(data, i)
        return _Tok("str", raw, i, j), j
    if c == 0x3C:  # <
        if data[i + 1:i + 2] == b"<":
            j = _skip_balanced(data, i, b"<<", b">>")
            return _Tok("dict", data[i:j], i, j), j
        raw, j = _scan_hex(data, i)
        return _Tok("str", raw, i, j), j
    if c == 0x5B:  # [
        items = []
        j = i + 1
        while True:
            t, j = _next_token(data, j)
            if t is None:
                break
            if t.kind == "op" and t.value == b"]":
                return _Tok("array", None, i, j, items), j
            items.append(t)
        return _Tok("array", None, i, i + 1, items), i + 1
    if c in (0x2B, 0x2D, 0x2E) or 0x30 <= c <= 0x39:  # number
        j = i
        while j < n and (data[j] in b"+-.0123456789"):
            j += 1
        return _Tok("num", data[i:j], i, j), j
    if c == 0x5D:  # ]
        return _Tok("op", b"]", i, i + 1), i + 1
    j = i
    while j < n and data[j] not in _WS and data[j] not in _DELIM:
        j += 1
    if j == i:
        return _Tok("op", data[i:i + 1], i, i + 1), i + 1
    return _Tok("op", data[i:j], i, j), j


def _skip_inline_image(data, i):
    m = re.search(rb"\sEI(?=[\s\[\]<>()/%]|\Z)", data[i:])
    if not m:
        return len(data)
    return i + m.end()


def _tokenize(data):
    toks = []
    i = 0
    n = len(data)
    while i < n:
        t, i = _next_token(data, i)
        if t is None:
            break
        toks.append(t)
        if t.kind == "op" and t.value == b"ID":
            i = _skip_inline_image(data, i)
    return toks


def _decode_op_text(font, tok):
    if tok.kind == "str":
        return font.decode(tok.value)
    if tok.kind == "array":
        return "".join(font.decode(it.value) for it in tok.items
                       if it.kind == "str")
    return ""


def _decode_with_codes(font, raw):
    out = []
    if font.two_byte:
        for i in range(0, len(raw) - 1, 2):
            code = int.from_bytes(raw[i:i + 2], "big")
            if font.tounicode:
                ch = font.tounicode.get(code, "\ufffd")
            else:
                ch = raw[i:i + 2].decode("utf-16-be", "replace")
            out.append((ch, code))
    else:
        for b in raw:
            if font.tounicode and b in font.tounicode:
                ch = font.tounicode[b]
            elif font.macroman:
                ch = bytes([b]).decode("mac-roman", "replace")
            else:
                ch = bytes([b]).decode("cp1252", "replace")
            out.append((ch, b))
    return out


def _split_occurrences(pairs, find):
    text = "".join(ch for ch, _ in pairs)
    segments = []
    idx = 0
    while True:
        pos = text.find(find, idx)
        if pos < 0:
            segments.append(("keep", pairs[idx:]))
            break
        segments.append(("keep", pairs[idx:pos]))
        segments.append(("drop", pairs[pos:pos + len(find)]))
        idx = pos + len(find)
    return segments


class _UnsupportedEdit(Exception):
    """无法可靠手术的内容结构(显式失败, 不猜改)。"""


def _codes_to_hex(font, seg):
    if font.two_byte:
        return b"".join(code.to_bytes(2, "big") for _, code in seg).hex().encode()
    return bytes(code & 0xFF for _, code in seg).hex().encode()


def _segments_to_dump(font, segments):
    chunks = []
    has_keep = False
    for kind, seg in segments:
        if kind == "keep":
            if seg:
                chunks.append(b"<" + _codes_to_hex(font, seg) + b">")
                has_keep = True
        else:
            if seg:
                total = sum(font.width_for_code(code, ch) for ch, code in seg)
                kern = -int(round(total))
                if kern:
                    chunks.append(b" " + str(kern).encode())
    return b" ".join(chunks), has_keep


def _rewrite_show_token(font, show_tok, find):
    if show_tok.kind == "str":
        segments = _split_occurrences(_decode_with_codes(font, show_tok.value), find)
        dumps, has_keep = _segments_to_dump(font, segments)
        if has_keep:
            return b"[" + dumps + b"]TJ"
        return None

    items = show_tok.items
    texts = []
    pos_map = []
    for idx, it in enumerate(items):
        if it.kind == "str":
            d = font.decode(it.value)
            texts.append(d)
            pos_map.extend([idx] * len(d))
    full = "".join(texts)
    start = 0
    while True:
        p = full.find(find, start)
        if p < 0:
            break
        if pos_map[p] != pos_map[p + len(find) - 1]:
            raise _UnsupportedEdit(
                "find string spans multiple strings inside one TJ array; "
                "refusing to guess at kerning")
        start = p + len(find)
    parts = []
    has_keep = False
    for it in items:
        if it.kind == "str":
            seg = _split_occurrences(_decode_with_codes(font, it.value), find)
            chunk, keep = _segments_to_dump(font, seg)
            has_keep = has_keep or keep
            parts.append(chunk)
        elif it.kind == "num":
            parts.append(it.value)
        else:
            parts.append(b"")
    if has_keep:
        body = b" ".join(p for p in parts if p)
        return b"[" + body + b"]TJ"
    return None


def _edit_content(content, fonts, find):
    """在内容流中删除 `find` 的字形码, 返回 (新内容字节, 命中列表)。"""
    if not find:
        return content, []
    toks = _tokenize(content)
    edits = []
    hits = []
    operands = []
    cur_font = None
    cur_size = 0.0
    for t in toks:
        if t.kind != "op":
            operands.append(t)
            continue
        op = t.value
        if op == b"Tf":
            if len(operands) >= 2 and operands[-2].kind == "name":
                cur_font = fonts.get(operands[-2].value.decode("latin-1"))
                try:
                    cur_size = float(operands[-1].value)
                except (TypeError, ValueError):
                    cur_size = 0.0
            operands = []
            continue
        show_tok = None
        if op in (b"Tj", b"'", b'"') and operands and operands[-1].kind == "str":
            show_tok = operands[-1]
        elif op == b"TJ" and operands and operands[-1].kind == "array":
            show_tok = operands[-1]
        if show_tok is None or cur_font is None:
            operands = []
            continue
        decoded = _decode_op_text(cur_font, show_tok)
        occ = decoded.count(find)
        if occ:
            span_start = operands[0].start
            span_end = t.end
            replacement = _rewrite_show_token(cur_font, show_tok, find)
            if replacement is None:
                edits.append((span_start, span_end, b""))
            else:
                edits.append((show_tok.start, span_end, replacement))
            hits.extend([{"font": cur_font, "size": cur_size}] * occ)
        operands = []
    out = bytearray(content)
    for start, end, repl in sorted(edits, key=lambda e: e[0], reverse=True):
        out[start:end] = repl
    return bytes(out), hits


def _iter_fonts(resources):
    fonts = resources.get("/Font") if resources else None
    if fonts is None:
        return {}
    if hasattr(fonts, "get_object"):
        fonts = fonts.get_object()
    if not isinstance(fonts, dict):
        return {}
    out = {}
    for key, ref in fonts.items():
        bare = key[1:] if str(key).startswith("/") else str(key)
        out[bare] = ref
    return out


class _PdfLibs:
    pypdf = None
    pdfplumber = None


def _pdf_libs():
    if _PdfLibs.pypdf is None:
        import pypdf as _pypdf
        import pdfplumber as _pdfplumber
        _PdfLibs.pypdf = _pypdf
        _PdfLibs.pdfplumber = _pdfplumber
    return _PdfLibs.pypdf, _PdfLibs.pdfplumber


def _pdf_string(value):
    pypdf, _ = _pdf_libs()
    return pypdf.generic.create_string_object(value)


def _pdf_fonts(page_obj):
    try:
        resources = page_obj.get("/Resources")
        if resources is not None:
            resources = resources.get_object()
    except Exception:
        resources = None
    fonts = {}
    for name, ref in _iter_fonts(resources).items():
        try:
            fdict = ref.get_object() if hasattr(ref, "get_object") else ref
            fonts[name] = _FontCodec(name, fdict)
        except Exception:
            continue
    return fonts


def _redact_content_stream(content, fonts, values):
    """按值删除内容流中的敏感字形码(实际删除, 非白框覆盖)。"""
    total = 0
    cur = content
    for v in sorted(set(values), key=len, reverse=True):
        try:
            cur, hits = _edit_content(cur, fonts, v)
        except _UnsupportedEdit:
            continue
        total += len(hits)
    return cur, total


def _filespec_blob(fs):
    ef = fs.get("/EF")
    if ef is None:
        return None
    ef = ef.get_object()
    for key in ("/F", "/UF"):
        if key in ef:
            try:
                return ef[key].get_object().get_data()
            except Exception:
                continue
    return None


def _set_filespec_blob(fs, data):
    ef = fs.get("/EF")
    if ef is None:
        return
    ef = ef.get_object()
    for key in ("/F", "/UF"):
        if key in ef:
            try:
                ef[key].get_object().set_data(data)
                return
            except Exception:
                continue


def _clean_pdf_attachment_annot(annot, result, detected, types, where):
    """文件附件批注: 嵌入内容与文件名含敏感值时清洗重写并保留无关内容;
    无法可靠清洗时抛 UncleanableAttachment(整文件阻断, 不笼统成功)。"""
    pypdf, _ = _pdf_libs()
    fs = annot.get("/FS")
    if fs is None:
        return
    fs = fs.get_object()
    blob = _filespec_blob(fs)
    cleaned = None
    hits_all = []
    if blob is not None:
        cleaned, hits_all = _clean_attachment_blob(blob, types, where)
    fhits = []
    for key in ("/F", "/UF", "/Desc"):
        v = fs.get(key)
        if isinstance(v, str) and v:
            hh, _ = find_hits(v, types)
            if hh:
                fs[pypdf.generic.NameObject(key)] = _pdf_string(_replace_in(v, hh))
                fhits.extend(hh)
    if cleaned is not None and cleaned != blob:
        _set_filespec_blob(fs, cleaned)
    if not hits_all and not fhits:
        return
    for h in hits_all:
        detected.append(h["value"])
    _record(result, [dict(h, location=where) for h in hits_all], [])
    for h in fhits:
        detected.append(h["value"])
    _record(result, [dict(h, location=where + ":filename") for h in fhits], [])


def _annot_list(page_obj):
    annots = page_obj.get("/Annots")
    if annots is None:
        return []
    try:
        return [a.get_object() for a in annots]
    except Exception:
        return []


def _res_of(obj):
    try:
        res = obj.get("/Resources")
        if res is not None:
            res = res.get_object()
    except Exception:
        res = None
    return res


def _fonts_from_resources(resources):
    fonts = {}
    for name, ref in _iter_fonts(resources).items():
        try:
            fdict = ref.get_object() if hasattr(ref, "get_object") else ref
            fonts[name] = _FontCodec(name, fdict)
        except Exception:
            continue
    return fonts


def _page_content(page_obj):
    try:
        return page_obj.get_contents().get_data()
    except Exception:
        return b""


def _content_decoded_text(content, fonts):
    """按算子顺序解码内容流中的文本(分离算子以换行避免跨算子误拼接)。"""
    toks = _tokenize(content)
    lines = []
    operands = []
    cur_font = None
    for t in toks:
        if t.kind != "op":
            operands.append(t)
            continue
        op = t.value
        if op == b"Tf":
            if len(operands) >= 2 and operands[-2].kind == "name":
                cur_font = fonts.get(operands[-2].value.decode("latin-1"))
            operands = []
            continue
        show = None
        if op in (b"Tj", b"'", b'"') and operands and operands[-1].kind == "str":
            show = operands[-1]
        elif op == b"TJ" and operands and operands[-1].kind == "array":
            show = operands[-1]
        if show is not None and cur_font is not None:
            lines.append(_decode_op_text(cur_font, show))
        operands = []
    return "\n".join(lines)


def _pdf_page_texts(path):
    _, pdfplumber = _pdf_libs()
    texts = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            texts.append(page.extract_text() or "")
    return texts


def _stream_decoded_text(stream_obj, resources):
    try:
        content = stream_obj.get_data()
    except Exception:
        return ""
    return _content_decoded_text(content, _fonts_from_resources(resources))


def _redact_stream_obj(stream_obj, values, seen):
    """递归清洗内容流(Form XObject/外观流): 删除命中字形。"""
    try:
        content = stream_obj.get_data()
    except Exception:
        return
    res = _res_of(stream_obj)
    new, _ = _redact_content_stream(content, _fonts_from_resources(res), values)
    if new != content:
        try:
            stream_obj.set_data(new)
        except Exception:
            pass
    _redact_form_xobjects(res, values, seen)


def _redact_form_xobjects(resources, values, seen):
    if resources is None:
        return
    xo = resources.get("/XObject")
    if xo is None:
        return
    try:
        xo = xo.get_object()
    except Exception:
        return
    for _, ref in list(xo.items()):
        try:
            obj = ref.get_object() if hasattr(ref, "get_object") else ref
        except Exception:
            continue
        if not hasattr(obj, "get_data") or str(obj.get("/Subtype")) != "/Form":
            continue
        if id(obj) in seen:
            continue
        seen.add(id(obj))
        _redact_stream_obj(obj, values, seen)


def _redact_ap_streams(annot, values, seen):
    ap = annot.get("/AP")
    if ap is None:
        return
    try:
        ap = ap.get_object()
    except Exception:
        return

    def walk(v):
        try:
            v = v.get_object() if hasattr(v, "get_object") else v
        except Exception:
            return
        if hasattr(v, "get_data"):
            if id(v) in seen:
                return
            seen.add(id(v))
            _redact_stream_obj(v, values, seen)
        elif hasattr(v, "items"):
            for _, sub in list(v.items()):
                walk(sub)

    for _, sub in list(ap.items()):
        walk(sub)


def _scrub_link(annot, types, result, detected, pno):
    pypdf, _ = _pdf_libs()
    if str(annot.get("/Subtype")) != "/Link":
        return
    where = "link:page %d" % pno

    def clean(container, key):
        v = container.get(key)
        if not isinstance(v, str) or not v:
            return
        hh, _ = find_hits(v, types)
        if not hh:
            return
        container[pypdf.generic.NameObject(key)] = _pdf_string(_replace_in(v, hh))
        for h in hh:
            detected.append(h["value"])
        _record(result, [dict(h, location=where) for h in hh], [])

    action = annot.get("/A")
    if action is not None:
        action = action.get_object()
        for key in ("/URI", "/N"):
            clean(action, key)
        fval = action.get("/F")
        if isinstance(fval, str):
            clean(action, "/F")
        elif fval is not None and hasattr(fval, "get_object"):
            fs = fval.get_object()
            if hasattr(fs, "get"):
                for fk in ("/F", "/UF"):
                    clean(fs, fk)
        clean(action, "/D")
    clean(annot, "/Dest")


def _scrub_annot_strings(annot, keys, types, result, detected, where_prefix):
    pypdf, _ = _pdf_libs()
    for key in keys:
        v = annot.get(key)
        if isinstance(v, str) and v:
            hh, _ = find_hits(v, types)
            if hh:
                annot[pypdf.generic.NameObject(key)] = _pdf_string(_replace_in(v, hh))
                for h in hh:
                    detected.append(h["value"])
                _record(result, [dict(h, location="%s:%s" % (where_prefix, key))
                                 for h in hh], [])


def _scrub_widget(annot, types, result, detected, where, seen):
    """清洗逻辑值 /V,/DV 及外观流 /AP, 避免只改 V 而保留旧值外观的假成功。"""
    pypdf, _ = _pdf_libs()
    hit_values = []
    for key in ("/V", "/DV"):
        v = annot.get(key)
        if isinstance(v, str) and v:
            hh, _ = find_hits(v, types)
            if hh:
                annot[pypdf.generic.NameObject(key)] = _pdf_string(_replace_in(v, hh))
                for h in hh:
                    detected.append(h["value"])
                    hit_values.append(h["value"])
                _record(result, [dict(h, location=where) for h in hh], [])
    if hit_values:
        _redact_ap_streams(annot, hit_values, seen)


def _outline_titles(root_obj):
    out = []
    root = root_obj.get("/Outlines")
    if root is None:
        return out
    seen = set()
    stack = [root]
    while stack:
        try:
            node = stack.pop().get_object()
        except Exception:
            continue
        if id(node) in seen:
            continue
        seen.add(id(node))
        title = node.get("/Title")
        if isinstance(title, str):
            out.append(title)
        for key in ("/First", "/Next"):
            if node.get(key) is not None:
                stack.append(node.get(key))
    return out


def _scrub_outlines(root_obj, types, result, detected):
    pypdf, _ = _pdf_libs()
    root = root_obj.get("/Outlines")
    if root is None:
        return
    seen = set()
    stack = [root]
    while stack:
        try:
            node = stack.pop().get_object()
        except Exception:
            continue
        if id(node) in seen:
            continue
        seen.add(id(node))
        title = node.get("/Title")
        if isinstance(title, str) and title:
            hh, _ = find_hits(title, types)
            if hh:
                node[pypdf.generic.NameObject("/Title")] = _pdf_string(
                    _replace_in(title, hh))
                for h in hh:
                    detected.append(h["value"])
                _record(result, [dict(h, location="toc") for h in hh], [])
        for key in ("/First", "/Next"):
            if node.get(key) is not None:
                stack.append(node.get(key))


def _embedded_file_scopes(root_obj):
    scopes = []
    names = root_obj.get("/Names")
    if names is None:
        return scopes
    ef = names.get_object().get("/EmbeddedFiles")
    if ef is None:
        return scopes
    stack = [ef]
    seen = set()
    while stack:
        try:
            node = stack.pop().get_object()
        except Exception:
            continue
        if id(node) in seen:
            continue
        seen.add(id(node))
        pairs = node.get("/Names")
        if pairs is not None:
            i = 0
            while i + 1 < len(pairs):
                name, spec = pairs[i], pairs[i + 1]
                if isinstance(name, str):
                    scopes.append(name)
                try:
                    spec = spec.get_object()
                except Exception:
                    spec = None
                if spec is not None and hasattr(spec, "get"):
                    for key in ("/F", "/UF", "/Desc"):
                        v = spec.get(key)
                        if isinstance(v, str):
                            scopes.append(v)
                    blob = _filespec_blob(spec)
                    if blob is not None:
                        scopes.extend(_blob_views(
                            blob, 0, spec.get("/UF") or spec.get("/F")))
                i += 2
        for kid in (node.get("/Kids") or []):
            stack.append(kid)
    return scopes


def _scrub_embedded_files(root_obj, types, detected):
    pypdf, _ = _pdf_libs()
    from pypdf.generic import ArrayObject
    names = root_obj.get("/Names")
    if names is None:
        return 0
    ef = names.get_object().get("/EmbeddedFiles")
    if ef is None:
        return 0
    removed = 0
    seen = set()
    stack = [ef]
    while stack:
        try:
            node = stack.pop().get_object()
        except Exception:
            continue
        if id(node) in seen or not hasattr(node, "get"):
            continue
        seen.add(id(node))
        pairs = node.get("/Names")
        if pairs is not None:
            keep = []
            i = 0
            while i + 1 < len(pairs):
                name, spec = pairs[i], pairs[i + 1]
                sensitive = False
                if isinstance(name, str):
                    hh, _ = find_hits(name, types)
                    if hh:
                        sensitive = True
                        for h in hh:
                            detected.append(h["value"])
                try:
                    spec_obj = spec.get_object()
                except Exception:
                    spec_obj = None
                if spec_obj is not None and hasattr(spec_obj, "get"):
                    for key in ("/F", "/UF", "/Desc"):
                        v = spec_obj.get(key)
                        if isinstance(v, str):
                            hh, _ = find_hits(v, types)
                            if hh:
                                sensitive = True
                                for h in hh:
                                    detected.append(h["value"])
                    blob = _filespec_blob(spec_obj)
                    if blob is not None:
                        for view in _blob_views(
                                blob, 0, spec_obj.get("/UF") or spec_obj.get("/F")):
                            hh, _ = find_hits(view, types)
                            if hh:
                                sensitive = True
                                for h in hh:
                                    detected.append(h["value"])
                                break
                if sensitive:
                    removed += 1
                else:
                    keep.append(name)
                    keep.append(spec)
                i += 2
            node[pypdf.generic.NameObject("/Names")] = ArrayObject(keep)
        for kid in (node.get("/Kids") or []):
            stack.append(kid)
    return removed


def _page_has_image(page_obj):
    try:
        res = _res_of(page_obj)
        if res is None:
            return False
        xo = res.get("/XObject")
        if xo is None:
            return False
        xo = xo.get_object()
        for _, ref in xo.items():
            obj = ref.get_object() if hasattr(ref, "get_object") else ref
            if str(obj.get("/Subtype")) == "/Image":
                return True
    except Exception:
        return True
    return False


def _sweep_orphans(writer):
    """将 writer 中不可达的间接对象替换为 NullObject。

    pypdf 从 reader clone 后, 被替换/移除的旧内容流、旧附件、旧元数据对象
    可能仍留在对象表中作为孤立对象写回, 从而残留敏感字节。此步从 root/info
    出发做可达性分析, 断开孤立对象, 确保不残留旧值。"""
    pypdf, _ = _pdf_libs()
    from pypdf.generic import (
        IndirectObject, DictionaryObject, ArrayObject, NullObject)
    reachable = set()

    def visit(obj):
        if obj is None or isinstance(obj, (str, bytes, int, float, bool)):
            return
        if isinstance(obj, IndirectObject):
            oid = getattr(obj, "idnum", None)
            if oid in reachable:
                return
            if oid is not None:
                reachable.add(oid)
            try:
                visit(obj.get_object())
            except Exception:
                return
            return
        ref = getattr(obj, "indirect_reference", None)
        if ref is not None:
            rid = getattr(ref, "idnum", None)
            if rid is not None:
                reachable.add(rid)
        if isinstance(obj, DictionaryObject):
            for v in obj.values():
                visit(v)
        elif isinstance(obj, ArrayObject):
            for v in obj:
                visit(v)

    try:
        visit(writer.root_object)
    except Exception:
        pass
    for root in (getattr(writer, "_info", None), getattr(writer, "_encryption", None)):
        try:
            visit(root)
        except Exception:
            pass

    objs = writer._objects
    if isinstance(objs, list):
        for idx in range(len(objs)):
            obj = objs[idx]
            if obj is None:
                continue
            ref = getattr(obj, "indirect_reference", None)
            oid = getattr(ref, "idnum", None) if ref is not None else None
            if oid is None:
                oid = idx + 1
            if oid not in reachable:
                objs[idx] = NullObject()
    elif isinstance(objs, dict):
        for k in list(objs.keys()):
            if k not in reachable:
                objs[k] = NullObject()


def _collect_xobject_text(resources, out, seen):
    if resources is None:
        return
    xo = resources.get("/XObject")
    if xo is None:
        return
    try:
        xo = xo.get_object()
    except Exception:
        return
    for _, ref in list(xo.items()):
        try:
            obj = ref.get_object() if hasattr(ref, "get_object") else ref
        except Exception:
            continue
        if not hasattr(obj, "get_data") or id(obj) in seen:
            continue
        seen.add(id(obj))
        if str(obj.get("/Subtype")) == "/Form":
            out.append(_stream_decoded_text(obj, _res_of(obj)))
            _collect_xobject_text(_res_of(obj), out, seen)


def _ap_texts(annot):
    out = []
    ap = annot.get("/AP")
    if ap is None:
        return out
    try:
        ap = ap.get_object()
    except Exception:
        return out

    def walk(v):
        try:
            v = v.get_object() if hasattr(v, "get_object") else v
        except Exception:
            return
        if hasattr(v, "get_data"):
            out.append(_stream_decoded_text(v, _res_of(v)))
        elif hasattr(v, "items"):
            for _, sub in list(v.items()):
                walk(sub)

    for _, sub in list(ap.items()):
        walk(sub)
    return out


def _pdf_scopes(path, types):
    """回读视图: 页面正文(pdfplumber 与内容流解码双视图)/Form XObject/
    链接目标/表单 V,DV 与外观 AP/批注信息/附件(zip 递归)/元数据/XMP/书签。"""
    pypdf, pdfplumber = _pdf_libs()
    scopes = []
    reader = pypdf.PdfReader(path)
    root = reader.trailer["/Root"]
    seen_xo = set()
    for page in reader.pages:
        scopes.append(_stream_decoded_text(page.get_contents(), _res_of(page)))
        _collect_xobject_text(_res_of(page), scopes, seen_xo)
        for annot in _annot_list(page):
            st = str(annot.get("/Subtype"))
            if st == "/Link":
                action = annot.get("/A")
                if action is not None:
                    action = action.get_object()
                    for key in ("/URI", "/F", "/N", "/D"):
                        v = action.get(key)
                        if isinstance(v, str):
                            scopes.append(v)
                dest = annot.get("/Dest")
                if isinstance(dest, str):
                    scopes.append(dest)
            for key in ("/Contents", "/T", "/Subj"):
                v = annot.get(key)
                if isinstance(v, str):
                    scopes.append(v)
            if st == "/Widget":
                for key in ("/V", "/DV"):
                    v = annot.get(key)
                    if isinstance(v, str):
                        scopes.append(v)
                scopes.extend(_ap_texts(annot))
            if st == "/FileAttachment":
                fs = annot.get("/FS")
                if fs is not None:
                    fs = fs.get_object()
                    for key in ("/F", "/UF", "/Desc"):
                        v = fs.get(key)
                        if isinstance(v, str):
                            scopes.append(v)
                    blob = _filespec_blob(fs)
                    if blob is not None:
                        scopes.extend(_blob_views(
                            blob, 0, fs.get("/UF") or fs.get("/F")))
    md = reader.metadata or {}
    scopes.extend(str(v) for v in md.values() if isinstance(v, str))
    if root.get("/Metadata") is not None:
        try:
            scopes.append(root["/Metadata"].get_object().get_data().decode(
                "utf-8", "replace"))
        except Exception:
            pass
    scopes.extend(_outline_titles(root))
    scopes.extend(_embedded_file_scopes(root))
    # 关键输出回读: pdfplumber 页面文本抽取失败必须 fail closed,
    # 不能吞掉异常后仍宣称完整核验(空白页正常返回 "" 不触发)。
    try:
        with pdfplumber.open(path) as pdf:
            for pg in pdf.pages:
                scopes.append(pg.extract_text() or "")
    except Exception as e:
        raise RuntimeError(
            "pdf_output_text_extraction_failed: %s" % type(e).__name__)
    return scopes


def process_pdf(in_path, out_path, types):
    """PDF 宽松栈实现(pypdf BSD + pdfplumber MIT, 无 PyMuPDF):
    页面正文按内容流算子级删除(含 Form XObject 递归), 元数据/XMP/书签/
    批注/链接目标/表单 V,DV 与外观 AP/文件附件(zip/Office 递归清洗)/
    文档级附件一并清洗; 保存后双视图(内容流解码 + pdfplumber)回读验证。
    """
    pypdf, _ = _pdf_libs()
    reader = pypdf.PdfReader(in_path)
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise RuntimeError("encrypted PDF cannot be reliably rewritten")
    writer = pypdf.PdfWriter(clone_from=reader)
    result = {"redactions": {}, "needs_review": {}}
    detected = []
    removed_emb = 0
    seen_xobj = set()
    seen_ap = set()

    in_text_ok = True
    try:
        texts = _pdf_page_texts(in_path)
    except Exception:
        texts = []
        in_text_ok = False

    for pno, page in enumerate(writer.pages, start=1):
        text = texts[pno - 1] if pno - 1 < len(texts) else ""
        hits, review = find_hits(text, types)
        _record(result, [], review)
        values = [h["value"] for h in hits]
        res = _res_of(page)
        extra = _content_decoded_text(_page_content(page), _fonts_from_resources(res))
        ehits, ereview = find_hits(extra, types)
        _record(result, [], ereview)
        for h in ehits:
            if h["value"] not in values:
                values.append(h["value"])
        if values:
            try:
                content = _page_content(page)
                new, _n = _redact_content_stream(
                    content, _fonts_from_resources(res), values)
                if new != content:
                    stream = pypdf.generic.DecodedStreamObject()
                    stream.set_data(new)
                    page.replace_contents(stream)
            except Exception:
                pass
            try:
                _redact_form_xobjects(res, values, seen_xobj)
            except Exception:
                pass
        for h in hits:
            detected.append(h["value"])
        if hits:
            _record(result, [dict(h, location="page %d" % pno) for h in hits], [])

        for annot in _annot_list(page):
            st = str(annot.get("/Subtype"))
            _scrub_link(annot, types, result, detected, pno)
            _scrub_annot_strings(annot, ("/Contents", "/T", "/Subj"),
                                 types, result, detected, "annotation")
            if st == "/Widget":
                _scrub_widget(annot, types, result, detected,
                              "form_field:page %d" % pno, seen_ap)
            if st == "/FileAttachment":
                _clean_pdf_attachment_annot(
                    annot, result, detected, types,
                    "file_attachment:page %d" % pno)

    root = writer.root_object

    md = dict(writer.metadata or {})
    md_hits = False
    for k, v in md.items():
        if isinstance(v, str) and v:
            hh, _ = find_hits(v, types)
            if hh:
                md[k] = _replace_in(v, hh)
                md_hits = True
                for h in hh:
                    detected.append(h["value"])
                _record(result, [dict(h, location="metadata:" + k) for h in hh], [])
    if md_hits:
        writer.add_metadata(md)

    if root.get("/Metadata") is not None:
        try:
            mobj = root["/Metadata"].get_object()
            xmp_text = mobj.get_data().decode("utf-8")
        except Exception:
            xmp_text = None
        if xmp_text:
            hh, _ = find_hits(xmp_text, types)
            if hh:
                mobj.set_data(_replace_in(xmp_text, hh).encode("utf-8"))
                for h in hh:
                    detected.append(h["value"])
                _record(result, [dict(h, location="xmp_metadata") for h in hh], [])

    _scrub_outlines(root, types, result, detected)
    removed_emb += _scrub_embedded_files(root, types, detected)

    if root.get("/AcroForm") is not None:
        af = root["/AcroForm"].get_object()
        for f in (af.get("/Fields") or []):
            try:
                fo = f.get_object()
            except Exception:
                continue
            _scrub_widget(fo, types, result, detected, "form_field", seen_ap)

    unverified = {}
    if not in_text_ok:
        result["needs_review"]["text_extraction_unverified_needs_manual"] = 1
        unverified["input_page_text_extraction"] = \
            "输入页面文本抽取不可用, 未据文本识别/核验"
    if root.get("/AcroForm") is not None:
        result["needs_review"]["form_fields_present_needs_manual_review"] = 1
        unverified["form_fields"] = "表单结构/脚本未逐项核验"
    img_pages = []
    for pno, page in enumerate(writer.pages, start=1):
        if _page_has_image(page):
            img_pages.append(pno)
    if img_pages:
        result["needs_review"]["images_present_not_ocr_scanned"] = len(img_pages)
        unverified["image_pages_not_ocr_scanned"] = img_pages

    _sweep_orphans(writer)
    with open(out_path, "wb") as fh:
        writer.write(fh)

    uniq = set(detected)
    leaks = []
    for scope in _pdf_scopes(out_path, types):
        hh, _ = find_hits(scope, types)
        if hh:
            leaks.append("pdf surface (%s)" % hh[0]["type"])
            break
        if any(v in scope for v in uniq):
            leaks.append("pdf content/metadata/attachment")
            break
    raw = open(out_path, "rb").read()
    for v in uniq:
        if v.encode("utf-8") in raw:
            leaks.append("pdf raw bytes")
            break
    if leaks:
        raise VerificationLeak(sorted(set(leaks)))

    fin = _finalize(result)
    if unverified:
        fin["pdf_removed_and_verified"] = False
        fin["unverified_surfaces"] = unverified
    else:
        fin["pdf_removed_and_verified"] = True
    if removed_emb:
        fin["removed_embedded_files"] = removed_emb
    return fin, detected


# ---------- main ----------

def _output_name(f, input_root, is_dir):
    ext = os.path.splitext(f)[1].lower()
    if is_dir:
        rel = os.path.relpath(f, input_root)
        rel_no_ext, _ = os.path.splitext(rel)
        flat = re.sub(r'[\\/:*?"<>|]+', "__", rel_no_ext)
        return "%s_REDACTED%s" % (flat, ext)
    stem = os.path.splitext(os.path.basename(f))[0]
    return "%s_REDACTED%s" % (stem, ext)


def main():
    ap = argparse.ArgumentParser(description="敏感信息打码(DOCX/XLSX/PPTX/PDF)")
    ap.add_argument("--input", required=True)
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--report")
    ap.add_argument("--types", help="逗号分隔: id_card,phone,bank_card,email")
    args = ap.parse_args()

    types = None
    if args.types:
        types = {t.strip() for t in args.types.split(",") if t.strip()}
        bad = types - set(ALL_TYPES)
        if bad:
            print("ERROR: unknown types: %s" % ",".join(sorted(bad)), file=sys.stderr)
            return 2

    is_dir = os.path.isdir(args.input)
    if not is_dir and not os.path.isfile(args.input):
        print("ERROR: input not found", file=sys.stderr)
        return 2
    outdir_abs = os.path.normcase(os.path.abspath(args.outdir))
    files = []
    if is_dir:
        for root, dirs, names in os.walk(args.input):
            dirs[:] = [d for d in dirs
                       if os.path.normcase(os.path.abspath(os.path.join(root, d))) != outdir_abs]
            for n in sorted(names):
                if n.lower().endswith((".docx", ".xlsx", ".pptx", ".pdf")) and not n.startswith("~$"):
                    p = os.path.join(root, n)
                    if os.path.normcase(os.path.abspath(p)).startswith(outdir_abs + os.sep):
                        continue
                    files.append(p)
    else:
        if not args.input.lower().endswith((".docx", ".xlsx", ".pptx", ".pdf")):
            print("ERROR: unsupported file type", file=sys.stderr)
            return 2
        files = [args.input]
    if not files:
        print("ERROR: no processable files", file=sys.stderr)
        return 2

    os.makedirs(args.outdir, exist_ok=True)
    results = []
    all_values = []
    used_outputs = {}
    for f in files:
        ext = os.path.splitext(f)[1].lower()
        out_name = _output_name(f, args.input, is_dir)
        out_path = os.path.join(args.outdir, out_name)
        key = os.path.normcase(os.path.abspath(out_path))
        if key in used_outputs:
            results.append({"file": f, "output": None, "failed": True,
                            "error": "output_name_collision: 另一输入已映射到相同输出名 " + out_name,
                            "redactions": [], "needs_review": []})
            continue
        used_outputs[key] = True
        try:
            if ext == ".docx":
                r, vals = process_ooxml(f, out_path, "docx", types)
            elif ext == ".pptx":
                r, vals = process_ooxml(f, out_path, "pptx", types)
            elif ext == ".xlsx":
                r, vals = process_xlsx(f, out_path, types)
            else:
                r, vals = process_pdf(f, out_path, types)
            all_values.extend(vals)
            r["file"] = f
            r["output"] = out_path
            r["failed"] = False
            results.append(r)
        except Exception as e:
            if os.path.exists(out_path):
                try:
                    os.remove(out_path)
                except OSError:
                    pass
            results.append({"file": f, "output": None, "failed": True,
                            "error": "%s: %s" % (type(e).__name__, e),
                            "redactions": [], "needs_review": []})

    manifest = {"files": results,
                "summary": {"files_processed": len([r for r in results if not r["failed"]]),
                            "files_failed": len([r for r in results if r["failed"]]),
                            "total_redactions": sum(
                                sum(x["count"] for x in r.get("redactions", []))
                                for r in results)}}
    report_path = args.report or os.path.join(args.outdir, "redact-manifest.json")
    raw = json.dumps(manifest, ensure_ascii=False, indent=2)
    for v in sorted(set(all_values), key=len, reverse=True):
        raw = raw.replace(v, mask_for(v))
    with open(report_path, "w", encoding="utf-8") as fh:
        fh.write(raw)
    print(json.dumps(manifest["summary"], ensure_ascii=False))
    if manifest["summary"]["files_failed"]:
        print("ERROR: %d file(s) failed or blocked" % manifest["summary"]["files_failed"],
              file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())
