#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""hr-material-analysis: 薪资考勤比对 + 名单汇总 + HR简报。

原创实现（office-delivery-glm53 独立实现，无第三方代码复用）。
运行时依赖：openpyxl（MIT）、python-docx（MIT）。

核心防静默规则：
- 表内工号冲突（同工号多行且字段值不一致）不挑行、不丢行：冲突人员显式标注，
  不计入部门汇总，明细（来源文件+行号+字段值）单列于「重复工号」sheet。
- 提供的每张表必须有工号列（工号/员工编号/员工号/编号），否则报错退出，
  绝不以 0 人静默成功。
- 未提供考勤表时不做缺勤结论（不默认满勤，不写已核查 0 天）。
- 金额/天数为空、不可解析、非有限数或公式无缓存值时显式列为数据质量问题，
  不按 0 计入合计。
"""
import argparse
import json
import math
import os
import re
import sys

from docx import Document
from openpyxl import Workbook, load_workbook

TOTAL_RE = re.compile(r"^(合计|总计)")
SENSITIVE_HEADERS = ("身份证", "银行卡", "手机", "电话")
ID_KEYWORDS = ("工号", "员工编号", "员工号")
NUM_FIELDS = ("应发", "实发", "应出勤", "实出勤")
TABLE_LABELS = {"payroll": "薪资表", "attendance": "考勤表", "roster": "花名册"}


def is_id_header(h):
    return any(k in h for k in ID_KEYWORDS) or h == "编号"


def find_cols(header):
    idx = {}
    for i, h in enumerate(header):
        h = str(h).strip() if h is not None else ""
        if not h:
            continue
        if is_id_header(h) and "工号" not in idx:
            idx["工号"] = i
        elif "姓名" in h and "姓名" not in idx:
            idx["姓名"] = i
        elif "部门" in h and "部门" not in idx:
            idx["部门"] = i
        elif "应发" in h and "应发" not in idx:
            idx["应发"] = i
        elif "实发" in h and "实发" not in idx:
            idx["实发"] = i
        elif "应出勤" in h and "应出勤" not in idx:
            idx["应出勤"] = i
        elif "实出勤" in h and "实出勤" not in idx:
            idx["实出勤"] = i
        else:
            for s in SENSITIVE_HEADERS:
                if s in h:
                    idx.setdefault(f"敏感_{h}", i)
    return idx


def parse_numeric(raw, formula_raw):
    """数值列解析状态：ok/unparseable/non_finite/uncached_formula/missing。"""
    is_formula = isinstance(formula_raw, str) and formula_raw.startswith("=")
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        if is_formula:
            return {"raw": None, "num": None, "status": "uncached_formula"}
        return {"raw": None, "num": None, "status": "missing"}
    if isinstance(raw, bool):
        return {"raw": raw, "num": None, "status": "unparseable"}
    if isinstance(raw, (int, float)):
        try:
            v = float(raw)
        except OverflowError:
            return {"raw": raw, "num": None, "status": "non_finite"}
        if math.isfinite(v):
            return {"raw": raw, "num": v, "status": "ok"}
        return {"raw": raw, "num": None, "status": "non_finite"}
    s = str(raw).strip().replace(",", "")
    try:
        v = float(s)
    except ValueError:
        st = "uncached_formula" if is_formula else "unparseable"
        return {"raw": raw, "num": None, "status": st}
    if math.isfinite(v):
        return {"raw": raw, "num": v, "status": "ok"}
    return {"raw": raw, "num": None, "status": "non_finite"}


def display_num(f):
    """数值字段展示：有限数→数值；不可解析→原样；非有限→标注；其余空。"""
    if f is None:
        return ""
    st = f["status"]
    if st == "ok":
        return f["num"]
    if st == "unparseable":
        return str(f["raw"])
    if st == "non_finite":
        return f"非有限数({f['raw']})"
    return ""


def mask_value(v):
    s = str(v)
    return "■" * min(len(s), 12)


def row_summary(row, sensitive_idx):
    """未关联行内容摘要（跳过敏感列，防明文泄漏）。"""
    parts = []
    for i, c in enumerate(row):
        if c is None or str(c).strip() == "" or i in sensitive_idx:
            continue
        parts.append(str(c).strip())
    return " | ".join(parts[:6])


def _sens_sig(rec):
    return tuple(sorted((k, str(v)) for k, v in rec["sensitive"].items()))


def read_sheet(path):
    """读取 active sheet，两次加载区分公式与缓存值。返回 (idx, header, data, dropped, totals)。"""
    wb_v = load_workbook(path, data_only=True)
    wb_f = load_workbook(path, data_only=False)
    rows_v = [list(r) for r in wb_v.active.iter_rows(values_only=True)]
    rows_f = [list(r) for r in wb_f.active.iter_rows(values_only=True)]
    wb_v.close()
    wb_f.close()
    if not rows_v:
        return {}, [], [], [], []
    header = [str(c).strip() if c is not None else "" for c in rows_v[0]]
    idx = find_cols(header)
    sensitive_idx = {i for h, i in idx.items() if h.startswith("敏感_")}
    data, dropped, totals = [], [], []
    for rn in range(2, len(rows_v) + 1):
        r = rows_v[rn - 1]
        rf = rows_f[rn - 1] if rn - 1 < len(rows_f) else []
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue
        emp = str(r[idx["工号"]]).strip() if "工号" in idx and idx["工号"] < len(r) \
            and r[idx["工号"]] is not None else ""
        name = str(r[idx["姓名"]]).strip() if "姓名" in idx and idx["姓名"] < len(r) \
            and r[idx["姓名"]] is not None else ""
        if TOTAL_RE.match(emp) or TOTAL_RE.match(name):
            totals.append(rn)
            continue
        if not emp:
            dropped.append({"row": rn, "summary": row_summary(r, sensitive_idx)})
            continue
        rec = {"工号": emp, "姓名": name, "row": rn, "fields": {}, "sensitive": {}}
        if "部门" in idx and idx["部门"] < len(r) and r[idx["部门"]] is not None:
            rec["部门"] = str(r[idx["部门"]]).strip()
        else:
            rec["部门"] = ""
        for k in NUM_FIELDS:
            if k in idx and idx[k] < len(r):
                fr = rf[idx[k]] if idx[k] < len(rf) else None
                rec["fields"][k] = parse_numeric(r[idx[k]], fr)
        for h, i in idx.items():
            if h.startswith("敏感_") and i < len(r) and r[i] is not None:
                rec["sensitive"][h[len("敏感_"):]] = r[i]
        data.append(rec)
    return idx, header, data, dropped, totals


def _sig(rec):
    """行内容签名（含敏感字段，用于重复工号一致性判定）。"""
    parts = [rec.get("姓名", ""), rec.get("部门", ""), _sens_sig(rec)]
    for k in NUM_FIELDS:
        f = rec["fields"].get(k)
        parts.append(None if f is None else (str(f["raw"]), f["status"]))
    return tuple(parts)


def merge_table(records, table, filename):
    """表内工号分组：一致重复去重计入；冲突不挑行、显式标注并排除出汇总。"""
    groups = {}
    for r in records:
        groups.setdefault(r["工号"], []).append(r)
    by_id, dups = {}, []
    for emp, rows in groups.items():
        if len(rows) == 1:
            by_id[emp] = rows[0]
            continue
        dup = {"table": table, "file": filename, "工号": emp,
               "rows": [r["row"] for r in rows], "field_values": {}}
        for r in rows:
            dup["field_values"].setdefault("姓名", []).append(r.get("姓名", ""))
            dup["field_values"].setdefault("部门", []).append(r.get("部门", ""))
            for k in NUM_FIELDS:
                f = r["fields"].get(k)
                dup["field_values"].setdefault(k, []).append(
                    None if f is None else f["raw"])
        if len({_sig(r) for r in rows}) == 1:
            dup["consistent"] = True
            keep = dict(rows[0])
            keep["dup_consistent"] = True
            by_id[emp] = keep
            dups.append(dup)
            continue
        dup["consistent"] = False
        dup["sensitive_conflict"] = len({_sens_sig(r) for r in rows}) > 1
        merged = {"工号": emp, "row": rows[0]["row"], "rows": dup["rows"],
                  "fields": {}, "conflicted": True, "conflict_fields": [],
                  "姓名": "", "部门": "", "sensitive": rows[0]["sensitive"],
                  "dup_consistent": False}
        for k in ("姓名", "部门"):
            vals = [r.get(k, "") for r in rows]
            if len(set(vals)) == 1:
                merged[k] = vals[0]
            else:
                merged["conflict_fields"].append(k)
        for k in NUM_FIELDS:
            vals = [r["fields"].get(k) for r in rows]
            present = [v for v in vals if v is not None]
            if present and len(present) == len(vals) and \
                    all(v["raw"] == present[0]["raw"] and v["status"] == present[0]["status"]
                        for v in present):
                merged["fields"][k] = present[0]
            elif present:
                merged["conflict_fields"].append(k)
        by_id[emp] = merged
        dups.append(dup)
    return by_id, dups


def conflict_detail(emp, dups):
    """冲突明细字符串：来源表[文件]行号 + 冲突字段及各值（敏感字段只说不一致）。"""
    parts = []
    for d in dups:
        if d["工号"] != emp or d["consistent"]:
            continue
        fv = []
        for k in ("姓名", "部门") + NUM_FIELDS:
            vals = d["field_values"].get(k)
            if not vals or all(v is None for v in vals) or len(set(map(str, vals))) == 1:
                continue
            fv.append(f"{k}:" + "/".join(str(v) for v in vals))
        if d.get("sensitive_conflict"):
            fv.append("敏感字段:不一致(已脱敏)")
        parts.append(f"{TABLE_LABELS[d['table']]}[{d['file']}]行"
                     + "、".join(str(x) for x in d["rows"])
                     + ("：" + "、".join(fv) if fv else ""))
    return "；".join(parts)


def build_compare(tables, opts):
    """跨表按工号关联，生成逐人明细、部门汇总与总计。"""
    all_ids = set()
    for by_id in tables.values():
        all_ids.update(by_id.keys())
    people_detail, dept = [], {}
    totals = {"people": 0, "in_summary": 0, "excluded_conflict": 0,
              "anomalies": 0, "absence_sum": 0.0, "absence_checked": 0,
              "pay_sum": 0.0, "dq_issues": 0}
    for emp in sorted(all_ids):
        p = tables.get("payroll", {}).get(emp)
        a = tables.get("attendance", {}).get(emp)
        r = tables.get("roster", {}).get(emp)
        recs = {"payroll": p, "attendance": a, "roster": r}
        flags, dq = [], []
        conflicted = any(x and x.get("conflicted") for x in recs.values())
        if conflicted:
            flags.append("工号冲突")
        if any(x and x.get("dup_consistent") for x in recs.values()):
            flags.append("工号重复(一致)")
        names = {t: x["姓名"] for t, x in recs.items() if x and x.get("姓名")}
        if len(set(names.values())) > 1:
            flags.append("姓名不一致")
        if opts["attendance_provided"] and p and not a:
            flags.append("缺考勤记录")
        if opts["pay_provided"] and a and not p:
            flags.append("缺薪资记录")
        for t, x in recs.items():
            if not x or x.get("conflicted"):
                continue
            label = TABLE_LABELS[t]
            for k, f in x["fields"].items():
                st = f["status"]
                if st == "ok":
                    continue
                rn = x.get("row")
                if st == "missing":
                    dq.append(f"{label}{k}为空(行{rn})")
                elif st == "unparseable":
                    dq.append(f"{label}{k}不可解析:'{f['raw']}'(行{rn})")
                elif st == "non_finite":
                    dq.append(f"{label}{k}非有限数(行{rn})")
                elif st == "uncached_formula":
                    dq.append(f"{label}{k}公式无缓存值(行{rn})")
        absence = None
        if a and not a.get("conflicted"):
            f1, f2 = a["fields"].get("应出勤"), a["fields"].get("实出勤")
            if f1 and f2 and f1["status"] == "ok" and f2["status"] == "ok":
                absence = round(f1["num"] - f2["num"], 2)
                totals["absence_checked"] += 1
                if absence > 0:
                    totals["absence_sum"] += absence
        in_summary = not conflicted

        def dept_of(x):
            # 冲突记录的字段级置空已在 merge_table 完成（一致字段保留公共值）
            return (x.get("部门") if x else "") or ""

        if in_summary:
            totals["in_summary"] += 1
            d = dept_of(r) or dept_of(p) or "未分配"
            dept.setdefault(d, {"count": 0, "pay": 0.0})
            dept[d]["count"] += 1
            pf = (p or {}).get("fields", {}).get("应发")
            if pf and pf["status"] == "ok":
                dept[d]["pay"] += pf["num"]
                totals["pay_sum"] += pf["num"]
        else:
            totals["excluded_conflict"] += 1
        if flags:
            totals["anomalies"] += 1
        totals["people"] += 1
        totals["dq_issues"] += len(dq)
        people_detail.append({
            "工号": emp, "names": names, "flags": flags, "dq": dq,
            "in_summary": in_summary, "absence": absence,
            "部门": dept_of(r) or dept_of(p),
            "应出勤": display_num(a["fields"].get("应出勤")) if a else "",
            "实出勤": display_num(a["fields"].get("实出勤")) if a else "",
            "应发": display_num(p["fields"].get("应发")) if p else "",
            "实发": display_num(p["fields"].get("实发")) if p else "",
            "敏感字段": dict(p["sensitive"]) if p else {},
        })
    return people_detail, dept, totals


def write_compare_xlsx(path, people_detail, dups, dropped_rows, mask_on):
    wb = Workbook()
    ws = wb.active
    ws.title = "人员比对"
    headers = ["工号", "姓名(薪资)", "姓名(考勤)", "部门", "应出勤", "实出勤",
               "缺勤天数", "应发", "实发", "异常", "数据质量", "冲突明细"]
    sens_cols = sorted({k for pd_ in people_detail for k in pd_["敏感字段"]})
    headers += [f"{c}(脱敏)" for c in sens_cols] if mask_on else sens_cols
    ws.append(headers)
    for pd_ in people_detail:
        row = [pd_["工号"], pd_["names"].get("payroll", ""),
               pd_["names"].get("attendance", ""), pd_["部门"],
               pd_["应出勤"], pd_["实出勤"],
               pd_["absence"] if pd_["absence"] is not None else "",
               pd_["应发"], pd_["实发"], "、".join(pd_["flags"]),
               "；".join(pd_["dq"]), conflict_detail(pd_["工号"], dups)]
        for c in sens_cols:
            v = pd_["敏感字段"].get(c)
            row.append(mask_value(v) if (mask_on and v is not None)
                       else (v if v is not None else ""))
        ws.append(row)
    ws2 = wb.create_sheet("重复工号")
    ws2.append(["来源文件", "表", "工号", "出现行号", "判定", "字段值明细"])
    for d in dups:
        fv = []
        for k in ("姓名", "部门") + NUM_FIELDS:
            vals = d["field_values"].get(k)
            if not vals or all(v is None for v in vals):
                continue
            fv.append(f"{k}=" + "/".join(str(v) for v in vals))
        if d.get("sensitive_conflict"):
            fv.append("敏感字段=不一致(已脱敏)")
        ws2.append([d["file"], TABLE_LABELS[d["table"]], d["工号"],
                    "、".join(str(x) for x in d["rows"]),
                    "内容一致(已去重计入)" if d["consistent"] else "内容冲突(未计入汇总)",
                    "；".join(fv)])
    ws3 = wb.create_sheet("未关联行")
    ws3.append(["来源文件", "表", "行号", "原因", "内容摘要(敏感列已省略)"])
    for dr in dropped_rows:
        ws3.append([dr["file"], TABLE_LABELS[dr["table"]], dr["row"],
                    dr["reason"], dr["summary"]])
    wb.save(path)


def write_summary_xlsx(path, dept, totals, opts):
    wb = Workbook()
    ws = wb.active
    ws.title = "部门汇总"
    ws.append(["部门", "人数", "应发合计"])
    for dname, v in sorted(dept.items()):
        ws.append([dname, v["count"],
                   round(v["pay"], 2) if opts["pay_available"] else None])
    ws5 = wb.create_sheet("汇总说明")
    ws5.append(["项目", "说明"])
    ws5.append(["汇总范围", f"无工号冲突人员 {totals['in_summary']} 人；"
                f"{totals['excluded_conflict']} 人因工号冲突未计入"
                f"（见compare-report.xlsx「重复工号」，含来源文件与行号）"])
    if opts["pay_available"]:
        ws5.append(["应发合计口径", "仅计入应发为有限数值且无工号冲突的人员；"
                    "不可解析/非有限/无缓存值/为空字段未计入"])
    else:
        ws5.append(["应发合计口径", "薪资表未提供或缺少应发列，应发合计不可计算（留空，不按0计）"])
    if not opts["attendance_provided"]:
        ws5.append(["考勤口径", "未提供考勤表，缺勤情况未核查（不默认满勤，不写0天）"])
    elif not opts["absence_computable"]:
        ws5.append(["考勤口径", "考勤表缺少应出勤/实出勤列，缺勤天数不可计算"])
    else:
        ws5.append(["考勤口径", f"缺勤=应出勤-实出勤，基于 {totals['absence_checked']} 名"
                    f"考勤可核查人员，合计 {round(totals['absence_sum'], 2)} 天"])
    wb.save(path)


def write_draft(path, totals, opts, dropped_rows, n_consistent_dup=0):
    doc = Document()
    doc.add_heading("人力材料分析简报（草稿）", level=1)
    if totals["excluded_conflict"]:
        doc.add_paragraph(
            f"本期识别人员 {totals['people']} 人；部门汇总基于 {totals['in_summary']} 名"
            f"无工号冲突人员，{totals['excluded_conflict']} 人因工号冲突未计入汇总"
            f"（明细见 compare-report.xlsx「重复工号」，含来源文件与行号，不挑行、不合并）。")
    elif n_consistent_dup:
        doc.add_paragraph(f"本期识别人员 {totals['people']} 人，部门汇总基于全部识别人员；"
                          f"检测到内容一致的重复工号 {n_consistent_dup} 组（已去重计入，"
                          f"明细见 compare-report.xlsx「重复工号」），未发现内容冲突的工号重复。")
    else:
        doc.add_paragraph(f"本期识别人员 {totals['people']} 人，部门汇总基于全部识别人员，"
                          f"未发现表内工号重复。")
    if not opts["attendance_provided"]:
        doc.add_paragraph("未提供考勤表，本期未核查缺勤情况（既不默认满勤，也不默认缺勤）。")
    elif not opts["absence_computable"]:
        doc.add_paragraph("考勤表缺少应出勤/实出勤列，缺勤天数不可计算，本期未核查缺勤。")
    else:
        doc.add_paragraph(f"考勤比对合计缺勤 {totals['absence_sum']:.2f} 天"
                          f"（基于 {totals['absence_checked']} 名考勤可核查人员，"
                          f"缺勤=应出勤-实出勤，可手工复算）。")
    doc.add_paragraph(f"异常人员 {totals['anomalies']} 人（姓名不一致/缺记录/工号冲突/"
                      f"工号重复等，明细见 compare-report.xlsx「人员比对」异常列）。")
    if totals["dq_issues"]:
        doc.add_paragraph(f"数据质量问题 {totals['dq_issues']} 项（金额/天数为空、不可解析、"
                          f"非有限数或公式无缓存值），相关字段未计入合计，明细见 hr-report.json。")
    if dropped_rows:
        doc.add_paragraph(f"另有 {len(dropped_rows)} 行因工号为空无法关联，已单列于"
                          f" compare-report.xlsx「未关联行」，未计入任何汇总（不静默丢行）。")
    doc.add_paragraph("敏感字段（身份证/银行卡/手机号）在输出中默认脱敏，明细见 mask-manifest.json。")
    doc.save(path)


def main():
    ap = argparse.ArgumentParser(description="HR材料比对分析")
    ap.add_argument("--payroll")
    ap.add_argument("--attendance")
    ap.add_argument("--roster")
    ap.add_argument("--outdir", required=True)
    ap.add_argument("--no-mask", action="store_true")
    args = ap.parse_args()
    sources = [("payroll", args.payroll), ("attendance", args.attendance),
               ("roster", args.roster)]
    provided = [(t, f) for t, f in sources if f]
    if not provided:
        print("ERROR: need at least one of --payroll/--attendance/--roster", file=sys.stderr)
        return 2
    tables, table_idx, dups = {}, {}, []
    dropped_rows, skipped_totals, table_files = [], [], {}
    try:
        for table, path in provided:
            idx, header, data, dropped, totals_rows = read_sheet(path)
            if "工号" not in idx:
                hdr = "、".join(h for h in header if h) or "(空)"
                print(f"ERROR: {os.path.basename(path)} 缺少必需标识列(工号)，"
                      f"实际表头: {hdr}；无法可靠关联人员，拒绝以0人静默成功",
                      file=sys.stderr)
                return 2
            table_idx[table] = idx
            table_files[table] = os.path.basename(path)
            base = os.path.basename(path)
            for dr in dropped:
                dropped_rows.append({"file": base, "table": table, "row": dr["row"],
                                     "reason": "工号为空", "summary": dr["summary"]})
            for rn in totals_rows:
                skipped_totals.append({"file": base, "table": table, "row": rn,
                                       "reason": "合计/总计行"})
            by_id, ddups = merge_table(data, table, base)
            tables[table] = by_id
            dups.extend(ddups)
    except Exception as e:
        print(f"ERROR: cannot read input: {e}", file=sys.stderr)
        return 2
    opts = {
        "attendance_provided": args.attendance is not None,
        "pay_provided": args.payroll is not None,
        "pay_available": "payroll" in table_idx and "应发" in table_idx["payroll"],
        "absence_computable": "attendance" in table_idx and
        "应出勤" in table_idx["attendance"] and "实出勤" in table_idx["attendance"],
    }
    people_detail, dept, totals = build_compare(tables, opts)
    mask_on = not args.no_mask
    os.makedirs(args.outdir, exist_ok=True)
    write_compare_xlsx(os.path.join(args.outdir, "compare-report.xlsx"),
                       people_detail, dups, dropped_rows, mask_on)
    write_summary_xlsx(os.path.join(args.outdir, "roster-summary.xlsx"),
                       dept, totals, opts)
    write_draft(os.path.join(args.outdir, "draft.docx"), totals, opts, dropped_rows,
                sum(1 for d in dups if d["consistent"]))
    sens_cols = sorted({k for pd_ in people_detail for k in pd_["敏感字段"]})
    conflict_groups = [d for d in dups if not d["consistent"]]
    report = {
        "people_total": totals["people"],
        "people_in_summary": totals["in_summary"],
        "people_excluded_conflict": totals["excluded_conflict"],
        "attendance_provided": opts["attendance_provided"],
        "absence_computable": opts["absence_computable"],
        "total_absence": (round(totals["absence_sum"], 2)
                          if opts["attendance_provided"] and opts["absence_computable"]
                          else None),
        "people_absence_checked": totals["absence_checked"] if opts["absence_computable"] else 0,
        "pay_available": opts["pay_available"],
        "pay_sum": round(totals["pay_sum"], 2) if opts["pay_available"] else None,
        "anomalies": totals["anomalies"],
        "data_quality_issue_count": totals["dq_issues"],
        "duplicate_groups": len(dups),
        "conflict_groups": len(conflict_groups),
        "consistent_duplicate_groups": len(dups) - len(conflict_groups),
        "conflicts": conflict_groups,
        "consistent_duplicates": [d for d in dups if d["consistent"]],
        "dropped_rows": dropped_rows,
        "skipped_total_rows": skipped_totals,
        "dept_summary": {d: {"count": v["count"],
                             "应发合计": round(v["pay"], 2) if opts["pay_available"] else None}
                         for d, v in dept.items()},
        "people_detail": [{"工号": pd_["工号"], "names": pd_["names"],
                           "flags": pd_["flags"], "data_quality": pd_["dq"],
                           "in_summary": pd_["in_summary"],
                           "absence": pd_["absence"]} for pd_ in people_detail],
        "input_files": table_files,
    }
    with open(os.path.join(args.outdir, "hr-report.json"), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    with open(os.path.join(args.outdir, "mask-manifest.json"), "w", encoding="utf-8") as f:
        json.dump({"masked": mask_on,
                   "fields": list(SENSITIVE_HEADERS) if mask_on else [],
                   "masked_columns": [f"{c}(脱敏)" for c in sens_cols] if mask_on else []},
                  f, ensure_ascii=False, indent=2)
    print(json.dumps({"people": report["people_total"],
                      "people_in_summary": report["people_in_summary"],
                      "people_excluded_conflict": report["people_excluded_conflict"],
                      "anomalies": report["anomalies"],
                      "attendance_provided": report["attendance_provided"],
                      "total_absence": report["total_absence"],
                      "pay_sum": report["pay_sum"],
                      "duplicate_groups": report["duplicate_groups"],
                      "conflict_groups": report["conflict_groups"],
                      "data_quality_issues": report["data_quality_issue_count"]},
                     ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
