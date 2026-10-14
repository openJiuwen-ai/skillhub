#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""sheet-consolidate-reconcile: 多源表格合并、清洗与对账。

用法:
  python consolidate.py --config cfg.json --output merged.xlsx --report r.json

退出码: 0=成功(可含差异) 2=配置/输入错误(全部源失败)
"""
import argparse
import csv
import json
import math
import os
import re
import sys

from openpyxl import Workbook, load_workbook

TOTAL_ROW_RE = re.compile(r"^(合计|总计|小计|累计)\s*$")


def _display_num(v):
    try:
        if v == int(v):
            return int(v)
    except (ValueError, OverflowError):
        pass
    return v


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        f = float(v)
    else:
        s = str(v).strip().replace(",", "").replace("，", "")
        if not s:
            return None
        try:
            f = float(s)
        except ValueError:
            return None
    # 非有限值(inf/nan/infinity 等文本)不作为数值参与对账，按原值保留走待核
    return f if math.isfinite(f) else None


def _cell(rowvals, col_idx, cmap, name):
    std = cmap.get(name, name)
    j = col_idx.get(std)
    if rowvals is None or j is None or j >= len(rowvals):
        return None
    return rowvals[j]


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def read_source(spec, pk, metrics, base_dir):
    """读取一个源 -> (rows, issues)。

    rows: [{pk, row_num, metrics{m: {raw, num, formula?}}}]
    issues: 字符串=致命错误(该源failed); 元组=(kind, row, payload)，
            kind in {total_row_skipped, formula_cached, formula_no_cache}
    """
    path = spec["file"]
    if not os.path.isabs(path):
        path = os.path.join(base_dir, path)
    issues = []
    cmap = spec.get("column_map", {}) or {}
    data_rows = []  # [(行号, 值元组)]
    f_map = {}      # 行号 -> 公式元组（仅xlsx，来自 data_only=False 遍）
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            table = list(csv.reader(f))
        header_idx = None
        header = None
        target_pk = cmap.get(pk, pk)
        for i, r in enumerate(table[:20]):
            vals = [c.strip() for c in r]
            if target_pk in vals:
                header_idx = i
                header = vals
                break
        if header_idx is None:
            return [], [f"header not found (need column {target_pk!r})"]
        col_idx = {name: j for j, name in enumerate(header)}
        for rn, r in enumerate(table[header_idx + 1:], header_idx + 2):
            if not r or all(not c.strip() for c in r):
                continue
            data_rows.append((rn, r))
    else:
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[spec.get("sheet")] if spec.get("sheet") else wb.active
        header_idx = None
        header = None
        target_pk = cmap.get(pk, pk)
        for i, r in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True)):
            vals = [str(c).strip() if c is not None else "" for c in r]
            if target_pk in vals:
                header_idx = i
                header = vals
                break
        if header_idx is None:
            wb.close()
            return [], [f"header not found (need column {target_pk!r})"]
        col_idx = {name: j for j, name in enumerate(header) if name}
        for rn, r in enumerate(ws.iter_rows(min_row=header_idx + 2, values_only=True), header_idx + 2):
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            data_rows.append((rn, r))
        wb.close()
        # 公式识别遍：data_only=False 时公式单元格返回 "=..." 字符串
        try:
            wbf = load_workbook(path, data_only=False, read_only=True)
            wsf = wbf[spec.get("sheet")] if spec.get("sheet") else wbf.active
            for rn, fr in enumerate(wsf.iter_rows(min_row=header_idx + 2, values_only=True), header_idx + 2):
                if fr is None or all(c is None or str(c).strip() == "" for c in fr):
                    continue
                f_map[rn] = fr
            wbf.close()
        except Exception:
            f_map = {}
        # 值遍全空但公式遍非空 => 整行公式且无任何缓存值
        seen = {rn for rn, _ in data_rows}
        for rn, fr in f_map.items():
            if rn not in seen:
                fs = [c for c in fr if _is_formula(c)]
                if fs:
                    issues.append(("formula_no_cache", rn,
                                   {"pk": "", "metric": "(整行)",
                                    "formula": "; ".join(str(x) for x in fs[:3])}))

    out = []
    for rn, r in data_rows:
        fr = f_map.get(rn)
        pk_val = _cell(r, col_idx, cmap, pk)
        pk_val = str(pk_val).strip() if pk_val is not None else ""
        if TOTAL_ROW_RE.match(pk_val):
            issues.append(("total_row_skipped", rn, pk_val))
            continue
        if not pk_val:
            f_pk = _cell(fr, col_idx, cmap, pk)
            if _is_formula(f_pk):
                issues.append(("formula_no_cache", rn,
                               {"pk": "", "metric": pk, "formula": f_pk}))
            continue
        rec = {"pk": pk_val, "row_num": rn, "metrics": {}}
        for m in metrics:
            v = _cell(r, col_idx, cmap, m)
            f = _cell(fr, col_idx, cmap, m)
            num = to_num(v)
            entry = {"raw": v, "num": num}
            if _is_formula(f):
                entry["formula"] = f
                if num is None:
                    issues.append(("formula_no_cache", rn,
                                   {"pk": pk_val, "metric": m, "formula": f}))
                else:
                    issues.append(("formula_cached", rn, m))
            rec["metrics"][m] = entry
        out.append(rec)
    return out, issues


def main():
    ap = argparse.ArgumentParser(description="多源表格合并对账")
    ap.add_argument("--config", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--report", required=True)
    args = ap.parse_args()
    try:
        with open(args.config, encoding="utf-8") as f:
            cfg = json.load(f)
    except Exception as e:
        print(f"ERROR: bad config: {e}", file=sys.stderr)
        return 2
    pk = cfg.get("primary_key")
    metrics = cfg.get("metrics", [])
    sources = cfg.get("sources", [])
    base_dir = os.path.dirname(os.path.abspath(args.config))
    if not pk or not metrics or not sources:
        print("ERROR: config needs primary_key/metrics/sources", file=sys.stderr)
        return 2

    per_source = {}
    failed = {}
    skipped_total = []
    formula_no_cache = []
    formula_cached = {}
    for spec in sources:
        name = spec.get("source_name") or spec["file"]
        try:
            rows, issues = read_source(spec, pk, metrics, base_dir)
        except Exception as e:
            rows, issues = [], [f"read error: {e}"]
        hard = [i for i in issues if isinstance(i, str)]
        if hard:
            failed[name] = hard
            per_source[name] = []
        else:
            per_source[name] = rows
            for kind, rn, payload in issues:
                if kind == "total_row_skipped":
                    skipped_total.append({"source": name, "row": rn, "pk_value": payload})
                elif kind == "formula_cached":
                    formula_cached[name] = formula_cached.get(name, 0) + 1
                elif kind == "formula_no_cache":
                    entry = {"source": name, "row": rn}
                    entry.update(payload)
                    formula_no_cache.append(entry)

    if all(not v for v in per_source.values()):
        print(f"ERROR: all sources failed: {failed}", file=sys.stderr)
        return 2

    names = [s.get("source_name") or s["file"] for s in sources]
    all_pks = []
    for n in names:
        for rec in per_source[n]:
            if rec["pk"] not in all_pks:
                all_pks.append(rec["pk"])
    all_pks.sort()

    # 重复主键（同源）
    duplicates = []
    for n in names:
        seen = {}
        for rec in per_source[n]:
            seen.setdefault(rec["pk"], []).append(rec)
        for p, recs in seen.items():
            if len(recs) > 1:
                duplicates.append({"source": n, "pk": p, "count": len(recs),
                                   "rows": [r["row_num"] for r in recs]})

    # 逐主键逐指标对账
    mismatches = []
    pending = []
    for p in all_pks:
        for m in metrics:
            vals = {}
            for n in names:
                nums = []
                for r in per_source[n]:
                    if r["pk"] != p:
                        continue
                    e = r["metrics"][m]
                    if e["num"] is not None:
                        nums.append(e["num"])
                    elif e.get("formula") is not None:
                        # 无缓存公式：显式列为待核，不猜不补0，不参与数值比较
                        pending.append({"pk": p, "metric": m, "source": n,
                                        "row": r["row_num"], "formula": e["formula"]})
                    elif e.get("raw") is not None:
                        pending.append({"pk": p, "metric": m, "source": n,
                                        "row": r["row_num"], "raw": str(e["raw"])})
                if nums:
                    vals[n] = nums
            if not vals:
                continue
            distinct = set()
            for n, nums in vals.items():
                for x in nums:
                    distinct.add(round(x, 6))
            if len(distinct) > 1 or any(len(v) > 1 for v in vals.values()):
                mismatches.append({"pk": p, "metric": m,
                                   "values": {n: v for n, v in vals.items()}})

    n_present_all = 0
    for p in all_pks:
        if all(any(r["pk"] == p for r in per_source[n]) for n in names if per_source[n]):
            n_present_all += 1
    active = [n for n in names if per_source[n]]
    source_missing = {n: [p for p in all_pks
                          if not any(r["pk"] == p for r in per_source[n])]
                      for n in active}

    n_unique = len(all_pks)
    n_mismatch_pk = len({mm["pk"] for mm in mismatches})
    match_rate = round(n_present_all / n_unique, 4) if n_unique else 0.0

    # 写 merged.xlsx
    pending_keys = {(e["pk"], e["metric"], e["source"]) for e in pending}
    wb = Workbook()
    ws = wb.active
    ws.title = "合并汇总"
    header = [pk] + [f"{m}@{n}" for m in metrics for n in names] + ["数据一致", "缺失来源"]
    ws.append(header)
    for p in all_pks:
        row = [p]
        all_consistent = True
        has_pending = False
        missing = []
        for m in metrics:
            vals = []
            for n in names:
                nums = [r["metrics"][m]["num"] for r in per_source[n]
                        if r["pk"] == p and r["metrics"][m]["num"] is not None]
                if nums:
                    v = nums[0]
                    vals.append(_display_num(v))
                    if len(nums) > 1:
                        all_consistent = False
                elif (p, m, n) in pending_keys:
                    has_pending = True
                    if any(r["metrics"][m].get("formula") for r in per_source[n] if r["pk"] == p):
                        vals.append("(公式未算)")
                    else:
                        vals.append("")
                else:
                    vals.append("")
            # 一致性判断
            nums_flat = []
            for n in names:
                for r in per_source[n]:
                    if r["pk"] == p and r["metrics"][m]["num"] is not None:
                        nums_flat.append(r["metrics"][m]["num"])
            if len(set(round(x, 6) for x in nums_flat)) > 1:
                all_consistent = False
            row.extend(vals)
        for n in names:
            if per_source[n] and not any(r["pk"] == p for r in per_source[n]):
                missing.append(n)
        if not all_consistent:
            row.append("否")
        elif has_pending:
            row.append("待核")
        else:
            row.append("是")
        row.append("、".join(missing))
        ws.append(row)
    # 源明细
    ws2 = wb.create_sheet("源明细")
    ws2.append(["来源", pk, "行号"] + metrics)
    for n in names:
        for rec in per_source[n]:
            cells = []
            for m in metrics:
                e = rec["metrics"][m]
                if e.get("formula") is not None and e["num"] is None:
                    cells.append(f"(公式未算){e['formula']}")
                else:
                    cells.append(e["raw"])
            ws2.append([n, rec["pk"], rec["row_num"]] + cells)
    # 重复主键
    ws3 = wb.create_sheet("重复主键")
    ws3.append(["来源", pk, "出现次数", "行号"])
    for d in duplicates:
        ws3.append([d["source"], d["pk"], d["count"], "、".join(str(x) for x in d["rows"])])
    try:
        wb.save(args.output)
    except Exception as e:
        print(f"ERROR: cannot save xlsx: {e}", file=sys.stderr)
        return 2

    report = {
        "output": args.output,
        "match_stats": {
            "n_sources": len(active),
            "n_unique_pk": n_unique,
            "n_present_all": n_present_all,
            "n_single_source": sum(1 for p in all_pks
                                   if sum(1 for n in active
                                          if any(r["pk"] == p for r in per_source[n])) == 1),
            "n_mismatch_pk": n_mismatch_pk,
            "match_rate": match_rate,
        },
        "mismatches": mismatches,
        "duplicates": duplicates,
        "source_missing": {k: v for k, v in source_missing.items() if v},
        "skipped_total_rows": skipped_total,
        "pending_values": pending,
        "formula_handling": {
            "note": "xlsx 公式单元格取来源文件保存时写入的缓存值（静态读值，本工具不重算公式）；无缓存公式的值不参与数值对账，不猜测不补0，逐条列于 no_cache",
            "cached_cells_per_source": formula_cached,
            "no_cache": formula_no_cache,
        },
        "failed_sources": failed,
    }
    with open(args.report, "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2, default=str)
    print(json.dumps({"match_stats": report["match_stats"],
                      "mismatch_count": len(mismatches),
                      "duplicate_groups": len(duplicates),
                      "formula_no_cache": len(formula_no_cache),
                      "pending": len(pending)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
