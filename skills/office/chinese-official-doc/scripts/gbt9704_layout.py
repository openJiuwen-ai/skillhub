#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""chinese-official-doc: GB/T 9704 版式 docx 排版与要素齐全性机检。

用法:
  排版: python gbt9704_layout.py --draft 公文.txt --output 公文.docx [--report r.json]
  核查: python gbt9704_layout.py --check 公文.docx --report check.json

草稿格式(UTF-8 文本):
  首行 "标题：xxx" (或 "# xxx")
  "主送：xxx" (可选)
  空行分隔的正文段落 (以 一、/二、 开头的段落用黑体一级标题样式)
  末尾 "落款：xxx" 与 "日期：xxxx年xx月xx日" (可选)

默认版式参数(见 references/format-gbt9704.md):
  A4; 页边距上37/下35/左28/右26mm; 主标题 方正小标宋简体 二号(22pt) 居中 不加粗 行距固定31磅;
  主送 仿宋_GB2312 三号(16pt) 顶格; 正文 仿宋_GB2312 三号 首行缩进2字符 行距固定28磅;
  落款/成文日期 右对齐(右空2字)。

退出码: 0=成功 2=输入错误
"""
import argparse
import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

FANGSONG = "仿宋_GB2312"
XIAOBIAOSONG = "方正小标宋简体"
HEITI = "黑体"
KAITI = "楷体_GB2312"


def set_cjk_font(run, eastasia):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), eastasia)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def parse_draft(path):
    with open(path, encoding="utf-8") as f:
        raw = f.read()
    lines = [ln.strip() for ln in raw.splitlines()]
    data = {"title": None, "recipient": None, "body": [], "signoff": None, "date": None}
    for ln in lines:
        if not ln:
            continue
        if ln.startswith("# ") and data["title"] is None:
            data["title"] = ln[2:].strip()
        elif re.match(r"^标题[：:]", ln) and data["title"] is None:
            data["title"] = re.sub(r"^标题[：:]\s*", "", ln)
        elif re.match(r"^主送[：:]", ln) and data["recipient"] is None:
            data["recipient"] = re.sub(r"^主送[：:]\s*", "", ln)
        elif re.match(r"^落款[：:]", ln) and data["signoff"] is None:
            data["signoff"] = re.sub(r"^落款[：:]\s*", "", ln)
        elif re.match(r"^日期[：:]", ln) and data["date"] is None:
            data["date"] = re.sub(r"^日期[：:]\s*", "", ln)
        else:
            data["body"].append(ln)
    return data


def add_para(doc, text, *, size, eastasia, align=None, indent_chars=0.0,
             right_indent_chars=0.0, line_spacing_pt=28, bold=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    else:
        pf.first_line_indent = Pt(0)
    if right_indent_chars:
        pf.right_indent = Pt(size * right_indent_chars)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_spacing_pt)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    set_cjk_font(r, eastasia)
    return p


def layout(draft, output_path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(37)
    sec.bottom_margin = Mm(35)
    sec.left_margin = Mm(28)
    sec.right_margin = Mm(26)

    if draft["title"]:
        add_para(doc, draft["title"], size=22, eastasia=XIAOBIAOSONG,
                 align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing_pt=31, bold=False)
    if draft["recipient"]:
        add_para(doc, draft["recipient"], size=16, eastasia=FANGSONG,
                 line_spacing_pt=28)
    for para in draft["body"]:
        if re.match(r"^[一二三四五六七八九十百]+、", para):
            add_para(doc, para, size=16, eastasia=HEITI, bold=False, line_spacing_pt=28)
        elif re.match(r"^（[一二三四五六七八九十]+）", para):
            add_para(doc, para, size=16, eastasia=KAITI, bold=False, line_spacing_pt=28)
        else:
            add_para(doc, para, size=16, eastasia=FANGSONG, indent_chars=2,
                     line_spacing_pt=28)
    if draft["signoff"]:
        add_para(doc, draft["signoff"], size=16, eastasia=FANGSONG,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent_chars=2, line_spacing_pt=28)
    if draft["date"]:
        add_para(doc, draft["date"], size=16, eastasia=FANGSONG,
                 align=WD_ALIGN_PARAGRAPH.RIGHT, right_indent_chars=2, line_spacing_pt=28)
    doc.save(output_path)


def check_docx(path):
    """机检 GB/T 9704 主要元素齐全性。"""
    doc = Document(path)
    sec = doc.sections[0]
    checks = []

    def add(item, ok, detail=None):
        c = {"item": item, "ok": bool(ok)}
        if detail is not None:
            c["detail"] = detail
        checks.append(c)

    add("paper_a4", abs(sec.page_width.mm - 210) < 1 and abs(sec.page_height.mm - 297) < 1,
        f"{sec.page_width.mm:.0f}x{sec.page_height.mm:.0f}mm")
    add("margin_top_37mm", abs(sec.top_margin.mm - 37) < 1, f"{sec.top_margin.mm:.1f}")
    add("margin_bottom_35mm", abs(sec.bottom_margin.mm - 35) < 1, f"{sec.bottom_margin.mm:.1f}")
    add("margin_left_28mm", abs(sec.left_margin.mm - 28) < 1, f"{sec.left_margin.mm:.1f}")
    add("margin_right_26mm", abs(sec.right_margin.mm - 26) < 1, f"{sec.right_margin.mm:.1f}")

    paras = [p for p in doc.paragraphs]
    nonempty = [p for p in paras if p.text.strip()]
    add("title_exists", bool(nonempty))
    if nonempty:
        t = nonempty[0]
        add("title_centered", t.alignment == WD_ALIGN_PARAGRAPH.CENTER, str(t.alignment))
        sz = t.runs[0].font.size if t.runs else None
        add("title_22pt_erhao", sz is not None and abs(sz.pt - 22) < 0.6,
            f"{sz.pt if sz else None}")
        ea = None
        if t.runs:
            rPr = t.runs[0]._element.find(qn("w:rPr"))
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    ea = rf.get(qn("w:eastAsia"))
        add("title_xiaobiaosong", ea is not None and ("小标宋" in ea or "宋" in ea), ea)
    body = [p for p in nonempty[1:]
            if p.alignment != WD_ALIGN_PARAGRAPH.RIGHT
            and p.paragraph_format.first_line_indent is not None
            and p.paragraph_format.first_line_indent.pt > 10]
    add("body_exists", len(body) > 0)
    if body:
        b = body[0]
        sz = b.runs[0].font.size if b.runs else None
        add("body_16pt_sanhaofangsong", sz is not None and abs(sz.pt - 16) < 0.6,
            f"{sz.pt if sz else None}")
        fli = b.paragraph_format.first_line_indent
        add("body_first_line_indent_2chars",
            fli is not None and abs(fli.pt - 32) < 2, f"{fli.pt if fli else None}")
        ea = None
        if b.runs:
            rPr = b.runs[0]._element.find(qn("w:rPr"))
            if rPr is not None:
                rf = rPr.find(qn("w:rFonts"))
                if rf is not None:
                    ea = rf.get(qn("w:eastAsia"))
        add("body_fangsong", ea is not None and "仿宋" in ea, ea)
    right_aligned = [p for p in paras if p.alignment == WD_ALIGN_PARAGRAPH.RIGHT]
    add("signoff_right_aligned", len(right_aligned) >= 2,
        f"右对齐段数={len(right_aligned)}(落款+日期)")
    return {"file": path, "checks": checks,
            "elements_complete": all(c["ok"] for c in checks),
            "missing": [c["item"] for c in checks if not c["ok"]]}


def main():
    ap = argparse.ArgumentParser(description="GB/T 9704 公文排版与要素核查")
    ap.add_argument("--draft")
    ap.add_argument("--output")
    ap.add_argument("--check")
    ap.add_argument("--report")
    args = ap.parse_args()

    if args.check:
        try:
            result = check_docx(args.check)
        except Exception as e:
            print(f"ERROR: cannot check: {e}", file=sys.stderr)
            return 2
        if args.report:
            with open(args.report, "w", encoding="utf-8") as f:
                json.dump(result, f, ensure_ascii=False, indent=2)
        print(json.dumps(result, ensure_ascii=False))
        return 0

    if not (args.draft and args.output):
        print("ERROR: --draft and --output required", file=sys.stderr)
        return 2
    try:
        draft = parse_draft(args.draft)
    except Exception as e:
        print(f"ERROR: cannot read draft: {e}", file=sys.stderr)
        return 2
    if not (draft["title"] or draft["body"]):
        print("ERROR: draft has no title and no body", file=sys.stderr)
        return 2
    layout(draft, args.output)
    result = check_docx(args.output)
    result["draft_elements"] = {k: bool(v) if not isinstance(v, list) else len(v)
                                for k, v in draft.items()}
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
    print(json.dumps(result, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
